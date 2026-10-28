#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# ==========================================================
#  VERSION 1.36 - 27 de octubre de 2026 (centro MX)
#    Actas: catalogo de PERSONAS (se dan de alta una vez, desde el acta, y despues solo se eligen; opcion en blanco); entrega PARCIAL:
#    al cerrar con pendientes el acta queda ACTA-X-1 cerrada y los pendientes pasan a ACTA-X-2 (borrador), y asi sucesivamente;
#    inventario de pendientes (pantalla, CSV, imprimir); conformes plegados, botones de agregar arriba, observaciones al final;
#    lista separada en "en proceso" y "realizadas"; vista de CONSULTA (solo lectura) para supervisores.
#  VERSION 1.35 - 27 de octubre de 2026 (centro MX)
#    Actas: "Agregar actividades" ya deja elegir de CUALQUIER zona/bloque (filtros por bloque, area, proveedor); lo que esta en otro
#    borrador se pasa y el borrador que queda vacio se borra solo; boton "Actividad nueva" (Torre B, nace por validar); PDF con area.
#    Plano: clic en el encabezado de zona (bloque - area) acerca el plano a esa zona.
#  VERSION 1.34 - 26 de octubre de 2026 (centro MX)
#    Actas: el PDF ya no falla con muchos pendientes; filtros por proveedor/area/especialidad/resultado (pantalla y
#    PDF parcial de un solo proveedor); acta de VARIAS areas; "Quitar" ahora DESACTIVA (reversible); boton para marcar
#    conformes los visibles; todas las actividades de obra quedan con "requiere pruebas" = SI (una sola vez) y las nuevas nacen asi.
#    (27-oct) Actas: lo que esta en OTRO borrador ya se puede pasar al acta nueva (antes salia tenue y no dejaba marcar);
#    en actividades generalizadas (ej. 4 tomas en una sola) el boton + sube tambien el total por probar.
#    (27-oct) Actas: boton "Generar actas automaticas" (un borrador por area, sin duplicar actividades, sin fecha hasta abrirla,
#    supervision Hector Longoria); paginas HTML sin cache del navegador.
#  VERSION 1.33 - 26 de octubre de 2026 (centro MX)
#    Carga masiva (Torre A y Torre B): casilla "Cargar por validar" (nacen como propuesta, no suman en ningun total
#    hasta que las apruebes en Validacion) y columna opcional "Fecha inicio". Sin casilla, la carga funciona igual que antes.
#  VERSION 1.29 - 22 de octubre de 2026 (centro MX)
#    Melissa (fase 1): asistente de consulta solo para admin. Boton flotante en las pantallas, estatus por area/bloque,
#    por avance y avance general, con datos reales de la torre abierta. Codigo en melissa.py y static/melissa.*; las
#    preguntas libres usan Claude solo si existe ANTHROPIC_API_KEY. Solo agrega la tabla melissa_log; no cambia nada mas.
#  VERSION 1.25 - 16 de octubre de 2026 (centro MX)
#    Dashboard para directivos (/dashboard): nuevo rol "directivo" solo lectura con notas. Avance global y por
#    bloque/area, semaforo por tendencia de 7 dias, plano pintado por avance, actas firmadas y avisos. Selector
#    Torre A / Torre B. El directivo NO ve proveedores ni actividades y no puede cambiar nada. Solo tablas nuevas.
#  VERSION 1.24 - 16 de octubre de 2026 (centro MX)
#    Torre A y Torre B: Actas de entrega-recepcion y pruebas de funcionamiento (menu Reportes > Actas). Un acta
#    junta las actividades que tu eliges de un bloque o area, lleva cuantos elementos se probaron (contactos, tomas
#    de gases, alarmas...), quien los probo y el resultado, y sale en un PDF con folio y firmas. Independiente del
#    avance. Folio: Torre A ACTA-01.01-001; Torre B ACTA-B-001. Solo agrega 2 tablas nuevas; no cambia lo existente.
#  VERSION 1.23 - 15 de octubre de 2026 (centro MX)
#    Torre A: cada lado de un area dibujada puede llevar una o varias especialidades (con metros y m2 de muro
#    segun la altura); totales por especialidad; altura promedio de muro editable. Ajustes de v1.22
#    (selector de bloque sin preseleccion, tipos fijos, aviso de "Ver zona"). Torre B no cambia.
#  VERSION 1.21 - 14 de octubre de 2026 (centro MX)
#    Torre A ya NO tiene pantallas aparte: es la misma pantalla principal y el mismo Plano de obra de
#    siempre, con el selector "Proyecto" junto al logo HAP. Catalogos fijos (20 bloques, 85 areas, 15
#    proveedores numerados, especialidades de Torre B), codigo BLOQUE.AREA.PROVEEDOR.ACTIVIDAD (1.1.1.1),
#    fecha de entrega propia, carga masiva por Excel y, en el Plano de obra, bloques/areas por piezas,
#    calibrar escala, medidas (m / m2) y filtro "Ver zona". Torre B no cambia.
#  VERSION 1.20 - 14 de octubre de 2026 (centro MX)
#    Torre A: contorno de los 20 bloques (pestana Bloques), varias piezas por bloque o area,
#    filtro por zona (bloque) y boton "Escala de la hoja".
#  VERSION 1.18 - 14 de octubre de 2026 (centro MX)
#    Torre A: capas del plano (puntos, lineas, zonas) ligadas a proveedor y actividad;
#    escala calibrada y medidas aproximadas (m lineales / m2); rectangulo de 2 clics.
#  VERSION 1.17 - 14 de octubre de 2026 (centro MX)
#    Torre A: plano con zonas por area (/torre-a/plano); desde ahi se ven y se capturan actividades.
#  VERSION 1.16 - 14 de octubre de 2026 (centro MX)
#    Torre A: proveedores de Torre B ya numerados; alta y renombre de bloques/areas del catalogo.
#  VERSION 1.15 - Torre A: actividades de 5 niveles.
#    (v1.15 detalle:)
#    Torre A: proveedores numerados y actividades de 5 niveles
#    (bloque.area.proveedor.actividad.detalle) con avance en los detalles (solo admin).
#  VERSION 1.14 - Torre A: catalogo de espacios (20 bloques, 85 areas), pagina /torre-a.
#  VERSION 1.13 - 14 de octubre de 2026 (centro MX)
#  v1.13 (1) Al corregir el avance como gestor tambien se alinea el avance
#       declarado (ACT-0430/0432: Validar ya no regresa el valor viejo).
#       (2) Base y plano separados por proyecto (Torre B / Torre A); Torre A
#       solo la ve el admin. Torre B queda igual.
#  v1.1 columnas departamento/mundo/tipo_interno en CREATE TABLE
#  v1.2 modales faltantes en portal.html
#  v1.3 auditoria: 4 pantallas muertas por botones inexistentes
#  v1.4 CORRECCION: el blindaje de la 1.3 rompio la sintaxis de
#       app.js y validacion.js. Rehecho con metodo seguro.
#       Probado en navegador real: 0 errores de JavaScript.
#  v1.5 NUEVO: consulta de dias pasados (/historico). Ver como estaba
#       la obra en una fecha, comparar dos fechas, por responsable y
#       por bloque, en los dos mundos. Reconstruye el pasado desde el
#       historial cuando no hay foto. Foto semanal AUTOMATICA.
#  v1.6 LISTA PARA INTERNET: base en disco persistente (DATA_DIR),
#       puerto y clave desde el servidor, cookies seguras en HTTPS,
#       arranque compatible con gunicorn. Ver GUIA_PUBLICAR_EN_INTERNET.md
#  v1.7 SIEMBRA AUTOMATICA: en un servidor nuevo con disco vacio, copia
#       semilla/obra_inicial.db (955 partidas) la primera vez. Ya no hay
#       que subir la base a mano.
#       VISTA DE CELULAR corregida: los botones del encabezado ya no se
#       salen de la pantalla; tablero, filtros y tabla adaptados.
#  v1.8 CONTRASEÑAS SEGURAS (10 de octubre de 2026): se guardan con
#       generate_password_hash (pbkdf2 + sal por usuario). Las antiguas
#       (SHA-256) siguen funcionando y se convierten solas al iniciar sesión.
#       Se quitó la contraseña fija del admin inicial: ahora viene de
#       ADMIN_CLAVE_INICIAL o se genera al azar (solo en servidor nuevo).
#       La semilla ya no trae usuarios.
#  v1.9 PLANO: al llegar desde el 📍 de una actividad se resaltan TODOS sus
#       pines (borde amarillo, 3 pulsos y etiqueta) y se atenúa el resto hasta
#       tocar una parte vacía. Nuevo botón "📍 Ver en plano" en el portal de
#       proveedores. Se agrega .gitignore para no subir respaldos.
#  v1.10 PLANO DEL PROVEEDOR: el panel ahora muestra "Lo mío" (pendientes en
#       orden de prioridad, toca y se resalta en el plano), "Quién trabaja en mis
#       áreas" (avance promedio de las otras empresas en sus áreas) y avisos de
#       fechas clave. Capas y filtros pasan a "Más opciones". Ruta nueva de solo
#       lectura /api/portal/mis_areas; fechas y empresas en PRIORIDAD_PLANO.
#  v1.11 Panel del proveedor: "Lo mío" y "Quién trabaja en mis áreas" ahora son
#       secciones plegables; se despliegan al tocarlas.
#  v1.12 "Lo mío" incluye también las actividades SIN pin, con filtro
#       Todas / En plano / Sin ubicar.
# ==========================================================
"""
PLATAFORMA DE CONTROL DE OBRA — Unidad Quirúrgica HAP
Servidor local. Base de datos SQLite en ./data/obra.db
Entrega meta: 31 de octubre de 2026.

Uso:
    python app.py
Luego abre en tu navegador:  http://localhost:5000
"""

import os
import re
import sqlite3
import datetime
import json
import hashlib
import hmac
import secrets
import base64
from functools import wraps
from flask import Flask, request, jsonify, render_template, g, send_file, Response, session, redirect, has_request_context, after_this_request
from werkzeug.security import generate_password_hash, check_password_hash

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
# En internet (Render/Railway) la base debe vivir en un DISCO PERSISTENTE,
# si no se borra en cada actualizacion. Se define con la variable DATA_DIR.
DATA_DIR = os.environ.get("DATA_DIR") or os.path.join(BASE_DIR, "data")
os.makedirs(DATA_DIR, exist_ok=True)
DB_PATH = os.path.join(DATA_DIR, "obra.db")
# Torre A: base y plano propios en el mismo disco. Torre B (obra.db) no se toca.
DB_PATH_TORRE_A = os.path.join(DATA_DIR, "obra_torre_a.db")
PROYECTOS = ("torre_b", "torre_a")
FECHA_ENTREGA = "2026-10-31"

app = Flask(__name__)


@app.after_request
def _sin_cache_en_paginas(resp):
    """Blindaje contra rebotes: nunca dejar que el navegador ni ningun
    proxy/CDN guarde en cache las paginas HTML. Asi jamas se puede quedar
    pegada una version vieja del panel o del portal."""
    if resp.mimetype == "text/html":
        resp.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        resp.headers["Pragma"] = "no-cache"
    return resp

# clave para las sesiones; se guarda en data/ para que no cambie entre reinicios
_secret_file = os.path.join(DATA_DIR, ".secret")
def _cargar_secret():
    try:
        if os.path.exists(_secret_file):
            return open(_secret_file).read().strip()
        os.makedirs(os.path.dirname(_secret_file), exist_ok=True)
        s = secrets.token_hex(32)
        open(_secret_file, "w").write(s)
        return s
    except Exception:
        return "hap-clave-temporal-cambiar"
app.secret_key = os.environ.get("SECRET_KEY") or _cargar_secret()

# --- Notificaciones push (celular/navegador) para el admin, sin depender
# de ningun servicio externo (correo, WhatsApp, etc.) ---
VAPID_PRIVATE_FILE = os.path.join(DATA_DIR, "vapid_private.pem")
VAPID_PUBLIC_FILE = os.path.join(DATA_DIR, "vapid_public.txt")
VAPID_CLAIM_EMAIL = os.environ.get("VAPID_CONTACT_EMAIL", "mailto:soporte@controlobrahap.local")

def _vapid_keys():
    """Genera (una sola vez) o carga el par de llaves VAPID que permiten
    mandar notificaciones push directo al navegador/celular. La llave
    privada se guarda en disco igual que .secret, para que NO cambie entre
    reinicios -- si cambiara, todos los que ya activaron avisos dejarian
    de recibirlos y tendrian que activarlos de nuevo."""
    try:
        if os.path.exists(VAPID_PRIVATE_FILE) and os.path.exists(VAPID_PUBLIC_FILE):
            return VAPID_PRIVATE_FILE, open(VAPID_PUBLIC_FILE).read().strip()
        from py_vapid import Vapid02
        from cryptography.hazmat.primitives.serialization import Encoding, PublicFormat
        v = Vapid02()
        v.generate_keys()
        os.makedirs(DATA_DIR, exist_ok=True)
        with open(VAPID_PRIVATE_FILE, "wb") as f:
            f.write(v.private_pem())
        pub_raw = v.public_key.public_bytes(Encoding.X962, PublicFormat.UncompressedPoint)
        pub_b64 = base64.urlsafe_b64encode(pub_raw).rstrip(b"=").decode()
        with open(VAPID_PUBLIC_FILE, "w") as f:
            f.write(pub_b64)
        return VAPID_PRIVATE_FILE, pub_b64
    except Exception as e:
        print("  Aviso: no se pudieron preparar las llaves de notificaciones push:", e)
        return None, None

def enviar_push_admins(titulo, cuerpo, url="/validacion"):
    """Manda una notificacion push a todos los celulares/navegadores que el
    admin haya activado. Nunca debe tronar la accion que lo disparo (guardar
    un avance, una propuesta, etc.) -- es un extra, no algo de lo que
    dependa el flujo principal."""
    try:
        from pywebpush import webpush, WebPushException
    except ImportError:
        return
    priv_path, pub_b64 = _vapid_keys()
    if not priv_path:
        return
    try:
        db = get_db_usuarios()
        subs = db.execute("SELECT * FROM push_subscriptions").fetchall()
    except Exception:
        return
    if not subs:
        return
    payload = json.dumps({"titulo": titulo, "cuerpo": cuerpo, "url": url})
    for s in subs:
        subscription_info = {
            "endpoint": s["endpoint"],
            "keys": {"p256dh": s["p256dh"], "auth": s["auth"]},
        }
        try:
            webpush(
                subscription_info=subscription_info,
                data=payload,
                vapid_private_key=priv_path,
                vapid_claims={"sub": VAPID_CLAIM_EMAIL},
            )
        except WebPushException as e:
            status = getattr(getattr(e, "response", None), "status_code", None)
            if status in (404, 410):
                try:
                    db.execute("DELETE FROM push_subscriptions WHERE id=?", (s["id"],))
                    db.commit()
                except Exception:
                    pass
        except Exception:
            pass

# --- Seguridad de la sesion cuando corre en internet (HTTPS) ---
EN_INTERNET = bool(os.environ.get("EN_INTERNET"))
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=EN_INTERNET,
    PERMANENT_SESSION_LIFETIME=datetime.timedelta(hours=12),
)


# ----------------------------------------------------------------------------
# Contraseñas
# Formato nuevo (v1.8): generate_password_hash de Werkzeug (pbkdf2 con sal
# propia por usuario). Formato antiguo: SHA-256 de 64 caracteres, con o sin
# SAL_CLAVES. Las contraseñas antiguas siguen funcionando y se convierten al
# formato nuevo automáticamente la próxima vez que el usuario inicia sesión.
# ----------------------------------------------------------------------------
SAL = os.environ.get("SAL_CLAVES", "")
# 600,000 iteraciones: recomendación OWASP para pbkdf2-sha256, sin cargar de más
# el CPU del servidor (1 worker en Render).
METODO_HASH_CLAVE = "pbkdf2:sha256:600000"


def crear_hash_clave(txt):
    """Devuelve el hash seguro (formato nuevo) de una contraseña."""
    return generate_password_hash(txt or "", method=METODO_HASH_CLAVE)


def es_hash_antiguo(guardada):
    """True si la contraseña guardada está en el formato SHA-256 antiguo."""
    if not guardada or len(guardada) != 64:
        return False
    return all(c in "0123456789abcdef" for c in guardada.lower())


def _hash_antiguo(txt, sal):
    """Calcula el hash del formato antiguo, solo para verificar contraseñas viejas."""
    return hashlib.sha256((sal + (txt or "")).encode()).hexdigest()


def verificar_clave(guardada, txt):
    """Compara la contraseña escrita contra la guardada, en formato nuevo o antiguo.
    Nunca lanza error: ante cualquier problema responde False (acceso negado)."""
    if not guardada or txt is None:
        return False
    try:
        if es_hash_antiguo(guardada):
            candidatos = [_hash_antiguo(txt, "")]
            if SAL:
                candidatos.append(_hash_antiguo(txt, SAL))
            return any(hmac.compare_digest(guardada.lower(), c) for c in candidatos)
        return check_password_hash(guardada, txt)
    except Exception as e:
        print("  Aviso: no se pudo verificar una contraseña:", e)
        return False


def actualizar_hash_si_antiguo(db, usuario_id, guardada, txt):
    """Si la contraseña estaba en formato antiguo, la guarda en el formato nuevo.
    Se llama solo después de verificarla. Si falla, no impide el inicio de sesión."""
    if not es_hash_antiguo(guardada):
        return
    try:
        # "AND clave=?" evita pisar una contraseña que el admin cambió en ese mismo instante
        db.execute("UPDATE usuarios SET clave=? WHERE id=? AND clave=?",
                   (crear_hash_clave(txt), usuario_id, guardada))
        db.commit()
    except Exception as e:
        try:
            db.rollback()
        except Exception:
            pass
        print("  Aviso: no se pudo actualizar el formato de la contraseña:", e)


# Hash de relleno: cuando el usuario no existe se verifica contra él, para que la
# respuesta tarde lo mismo y no delate qué usuarios existen.
_HASH_RELLENO = crear_hash_clave(secrets.token_hex(16))


def contar_claves_antiguas(con):
    """Cuántos usuarios siguen con contraseña en formato antiguo (para los logs)."""
    try:
        filas = con.execute("SELECT clave FROM usuarios").fetchall()
        return sum(1 for f in filas if es_hash_antiguo(f[0]))
    except Exception:
        return -1


def usuario_actual():
    return session.get("usuario"), session.get("rol"), session.get("proveedor")


def col_duenio():
    """Devuelve ('departamento', mundo) o ('proveedor', mundo) según el mundo del usuario."""
    mundo = session.get("mundo", "obra")
    col = "departamento" if mundo == "interno" else "proveedor"
    return col, mundo


def requiere_login(fn):
    @wraps(fn)
    def envoltura(*a, **k):
        if "usuario" not in session:
            return jsonify({"error": "no_login"}), 401
        # candado: si al usuario lo desactivaron DESPUÉS de que ya había iniciado
        # sesión, su sesión abierta no debe seguir funcionando. Sin esto, un
        # usuario desactivado que ya estaba adentro podía seguir usando la
        # plataforma hasta que él mismo cerrara sesión.
        db = get_db_usuarios()
        u = db.execute("SELECT activo FROM usuarios WHERE usuario=?", (session["usuario"],)).fetchone()
        if u is not None and "activo" in u.keys() and u["activo"] == 0:
            session.clear()
            return jsonify({"error": "cuenta_desactivada"}), 401
        return fn(*a, **k)
    return envoltura


def requiere_admin(fn):
    @wraps(fn)
    def envoltura(*a, **k):
        if session.get("rol") != "admin":
            return jsonify({"error": "solo_admin"}), 403
        return fn(*a, **k)
    return envoltura


# Roles que pueden GESTIONAR la obra (validar, asignar, editar, dependencias):
# el admin y los dos supervisores. NO incluye tocar usuarios, respaldos ni borrar.
ROLES_GESTORES = ("admin", "supervisor", "supervisor_obra", "supervisor_depto")

def es_gestor(rol):
    return rol in ROLES_GESTORES

def requiere_gestor(fn):
    """Deja pasar a admin y supervisores. Para validar/asignar/editar/dependencias."""
    @wraps(fn)
    def envoltura(*a, **k):
        if session.get("rol") not in ROLES_GESTORES:
            return jsonify({"error": "solo_gestores"}), 403
        return fn(*a, **k)
    return envoltura

# ----------------------------------------------------------------------------
# Base de datos
# ----------------------------------------------------------------------------
import unicodedata


def sin_acentos(txt):
    """Quita acentos y pasa a minúsculas, para buscar sin que estorbe la tilde."""
    if txt is None:
        return ""
    t = unicodedata.normalize("NFD", str(txt))
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return t.lower()


def proyecto_actual():
    """Proyecto de la sesion: 'torre_a' SOLO si quien pregunta es admin y lo eligio.
    Cualquier otro caso (proveedor, supervisor, sin sesion, arranque) es 'torre_b'.
    Este es el candado unico: aunque una cookie traiga torre_a, un no-admin cae en Torre B."""
    try:
        if has_request_context() and session.get("rol") == "admin" and session.get("proyecto") == "torre_a":
            return "torre_a"
    except Exception:
        pass
    return "torre_b"


def ruta_db(proyecto=None):
    """Archivo de base de datos del proyecto (por defecto, el de la sesion)."""
    proyecto = proyecto or proyecto_actual()
    return DB_PATH_TORRE_A if proyecto == "torre_a" else DB_PATH


def ruta_salida_temporal(nombre):
    """Archivo temporal UNICO por peticion para exportaciones (Excel/PDF), que se borra
    al terminar la respuesta. Antes todos escribian al mismo nombre fijo y, con varios
    hilos, una exportacion podia mezclarse con otra (ahora tambien importa por Torre A)."""
    carpeta = os.path.join(BASE_DIR, "data")
    os.makedirs(carpeta, exist_ok=True)
    ruta = os.path.join(carpeta, f"{secrets.token_hex(6)}_{nombre}")

    @after_this_request
    def _borrar(resp):
        try:
            os.remove(ruta)
        except OSError:
            pass
        return resp
    return ruta


def get_db_usuarios():
    """Conexion SIEMPRE a la base principal (Torre B): ahi viven los usuarios, las
    contrasenas y las suscripciones de avisos. Torre A no tiene usuarios propios;
    asi un admin en modo Torre A sigue autenticandose y administrando usuarios bien."""
    if "db_usuarios" not in g:
        g.db_usuarios = sqlite3.connect(DB_PATH, timeout=10.0)
        g.db_usuarios.row_factory = sqlite3.Row
        g.db_usuarios.execute("PRAGMA journal_mode = WAL")
        g.db_usuarios.execute("PRAGMA busy_timeout = 5000")
        g.db_usuarios.execute("PRAGMA foreign_keys = ON")
    return g.db_usuarios


def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(ruta_db(), timeout=10.0)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA journal_mode = WAL")
        g.db.execute("PRAGMA busy_timeout = 5000")
        g.db.execute("PRAGMA foreign_keys = ON")
        # función propia para buscar ignorando acentos y mayúsculas
        g.db.create_function("sinac", 1, sin_acentos)
    return g.db


@app.teardown_appcontext
def close_db(exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()
    dbu = g.pop("db_usuarios", None)
    if dbu is not None:
        dbu.close()


def generar_respaldo_bd(ruta=None):
    """Genera una copia de seguridad timestamped en data/backups/ (mantiene los últimos 30).
    Sin parametro respalda la base del proyecto de la sesion (Torre B si no hay sesion)."""
    ruta = ruta or ruta_db()
    es_torre_a = (ruta == DB_PATH_TORRE_A)
    prefijo = "obra_torre_a_backup_" if es_torre_a else "obra_backup_"
    try:
        backup_dir = os.path.join(BASE_DIR, "data", "backups")
        os.makedirs(backup_dir, exist_ok=True)
        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        dest = os.path.join(backup_dir, f"{prefijo}{ts}.db")
        if os.path.exists(ruta):
            src_con = sqlite3.connect(ruta)
            dst_con = sqlite3.connect(dest)
            src_con.backup(dst_con)
            dst_con.close()
            src_con.close()
            # Limpiar respaldos viejos (mantener los últimos 30)
            archivos = sorted([os.path.join(backup_dir, f) for f in os.listdir(backup_dir) if f.startswith(prefijo) and f.endswith(".db")])
            if len(archivos) > 30:
                for f in archivos[:-30]:
                    try: os.remove(f)
                    except Exception: pass
            return dest
    except Exception as e:
        print(f"Error generando respaldo automático: {e}")
    return None


PROVEEDORES_INICIALES_TORRE_A = (
    "Aluoferta", "Longoria", "Andrés Colunga", "Juan Carlos Lozano", "ACSA", "Óscar Crespo",
    "Murosol", "Ana Rada", "Daniel Contreras", "G2", "Marquinox", "Luis Carlos López",
    "Muñoz", "Julio Sojo",
)  # numeros 2 a 15; CEBSA es el 1

ARCHIVO_CATALOGO_TORRE_A = os.path.join(BASE_DIR, "semilla", "torre_a_catalogo.json")


def crear_tablas_actas(con):
    """Actas de entrega-recepcion y pruebas. Se crean en AMBAS bases (Torre A y Torre B);
    solo agrega tablas nuevas (IF NOT EXISTS), no toca nada de lo que ya existe."""
    con.execute("""
        CREATE TABLE IF NOT EXISTS ta_actas (
            id            INTEGER PRIMARY KEY AUTOINCREMENT,
            folio         TEXT NOT NULL UNIQUE,
            bloque        TEXT NOT NULL,
            area          TEXT,
            titulo        TEXT,
            estado        TEXT NOT NULL DEFAULT 'borrador',
            fecha         TEXT,
            recibe_obra   TEXT,
            recibe_area   TEXT,
            proveedor_rep TEXT,
            supervision   TEXT,
            observaciones TEXT,
            creada        TEXT,
            creada_por    TEXT,
            cerrada       TEXT,
            cerrada_por   TEXT,
            eliminada     INTEGER NOT NULL DEFAULT 0
        )""")
    con.execute("""
        CREATE TABLE IF NOT EXISTS ta_acta_items (
            id           INTEGER PRIMARY KEY AUTOINCREMENT,
            acta_id      INTEGER NOT NULL,
            actividad_id INTEGER NOT NULL,
            unidad       TEXT NOT NULL DEFAULT 'piezas',
            cant_total   INTEGER NOT NULL DEFAULT 1,
            cant_ok      INTEGER NOT NULL DEFAULT 0,
            resultado    TEXT NOT NULL DEFAULT 'pendiente',
            nota         TEXT,
            probado_por  TEXT,
            fecha_prueba TEXT,
            orden        INTEGER NOT NULL DEFAULT 0,
            UNIQUE (acta_id, actividad_id)
        )""")
    # columnas nuevas (si la tabla ya existía): areas_json = varias áreas; activo = 0 -> desactivado del acta (reversible)
    if "areas_json" not in [r[1] for r in con.execute("PRAGMA table_info(ta_actas)")]:
        con.execute("ALTER TABLE ta_actas ADD COLUMN areas_json TEXT")
    if "activo" not in [r[1] for r in con.execute("PRAGMA table_info(ta_acta_items)")]:
        con.execute("ALTER TABLE ta_acta_items ADD COLUMN activo INTEGER NOT NULL DEFAULT 1")
    # v1.36: entregas parciales encadenadas (ACTA-B-010-1 -> ACTA-B-010-2 ...)
    cols = [r[1] for r in con.execute("PRAGMA table_info(ta_actas)")]
    for c in ("origen_id", "parcial_n", "continua_id"):
        if c not in cols:
            con.execute("ALTER TABLE ta_actas ADD COLUMN %s INTEGER" % c)
    # v1.36: catalogo de personas (quien recibe, quien prueba, representante, supervision). Se da de alta una vez y se elige.
    con.execute("""
        CREATE TABLE IF NOT EXISTS ta_personas (
            id        INTEGER PRIMARY KEY AUTOINCREMENT,
            nombre    TEXT NOT NULL,
            funcion   TEXT NOT NULL DEFAULT 'interno',
            empresa   TEXT,
            activo    INTEGER NOT NULL DEFAULT 1,
            creada    TEXT,
            creada_por TEXT
        )""")
    if con.execute("SELECT COUNT(*) FROM ta_personas").fetchone()[0] == 0:
        # primera vez: se llena con los nombres que ya se habian escrito a mano en las actas
        vistos = {}
        def _poner(nombre, funcion):
            n = " ".join(str(nombre or "").split())[:120]
            if len(n) >= 2 and n.casefold() not in vistos:
                vistos[n.casefold()] = (n, funcion)
        _poner("Héctor Longoria", "supervisor")
        for campo, fn in (("supervision", "supervisor"), ("recibe_obra", "interno"), ("recibe_area", "interno"), ("proveedor_rep", "proveedor")):
            for r in con.execute("SELECT DISTINCT %s FROM ta_actas WHERE eliminada=0 AND %s IS NOT NULL" % (campo, campo)).fetchall():
                _poner(r[0], fn)
        for r in con.execute("SELECT DISTINCT probado_por FROM ta_acta_items WHERE probado_por IS NOT NULL").fetchall():
            _poner(r[0], "interno")
        ahora = datetime.datetime.now().isoformat(timespec="seconds")
        for n, fn in vistos.values():
            con.execute("INSERT INTO ta_personas (nombre,funcion,creada,creada_por) VALUES (?,?,?,?)", (n, fn, ahora, "migración"))


def crear_tablas_dashboard(con):
    """Notas y avisos del dashboard para directivos. Se crean en AMBAS bases (solo tablas nuevas)."""
    con.execute("""
        CREATE TABLE IF NOT EXISTS dir_notas (
            id        INTEGER PRIMARY KEY AUTOINCREMENT,
            texto     TEXT NOT NULL,
            bloque    TEXT,
            area      TEXT,
            aviso     INTEGER NOT NULL DEFAULT 0,
            autor     TEXT,
            usuario   TEXT,
            rol       TEXT,
            creada    TEXT,
            eliminada INTEGER NOT NULL DEFAULT 0
        )""")


SH_NORMAS_INICIALES = [
    ("NOM-001-STPS-2008", "Edificios, locales, instalaciones y áreas en los centros de trabajo · condiciones de seguridad"),
    ("NOM-009-STPS-2011", "Condiciones de seguridad para realizar trabajos en altura"),
    ("NOM-017-STPS-2008", "Equipo de protección personal · selección, uso y manejo"),
    ("NOM-026-STPS-2008", "Colores y señales de seguridad e higiene, e identificación de riesgos"),
    ("NOM-031-STPS-2011", "Construcción · condiciones de seguridad y salud en el trabajo"),
]
# (categoria, punto de revision, [(hallazgo precargado, tipo, [normas])])
SH_CATALOGO_INICIAL = [
    ("EPP y señalización", "Casco, lentes y calzado de seguridad en uso", [
        ("Personal sin casco de seguridad", "acto", ["NOM-017-STPS-2008"]),
        ("Personal sin lentes de seguridad", "acto", ["NOM-017-STPS-2008"]),
        ("Personal sin calzado de seguridad", "acto", ["NOM-017-STPS-2008"]),
        ("Personal sin guantes para la tarea que realiza", "acto", ["NOM-017-STPS-2008"]),
    ]),
    ("EPP y señalización", "Protección en trabajos en altura", [
        ("Trabajo en altura sin arnés o línea de vida", "acto", ["NOM-009-STPS-2011", "NOM-031-STPS-2011"]),
        ("Andamio o plataforma sin barandal o mal armado", "condicion", ["NOM-009-STPS-2011", "NOM-031-STPS-2011"]),
    ]),
    ("EPP y señalización", "Señalización de áreas de riesgo", [
        ("Falta señalización del área de riesgo", "condicion", ["NOM-026-STPS-2008"]),
        ("Señal dañada, ilegible o mal colocada", "condicion", ["NOM-026-STPS-2008"]),
    ]),
    ("Orden y limpieza", "Pasillos y accesos libres", [
        ("Pasillo o acceso obstruido con material o escombro", "condicion", ["NOM-031-STPS-2011", "NOM-001-STPS-2008"]),
        ("Ruta de evacuación o salida bloqueada", "condicion", ["NOM-001-STPS-2008"]),
    ]),
    ("Orden y limpieza", "Material y herramienta en orden", [
        ("Material mal acopiado o apilado", "condicion", ["NOM-031-STPS-2011"]),
        ("Herramienta o cables tirados en zona de paso", "acto", ["NOM-031-STPS-2011"]),
    ]),
    ("Orden y limpieza", "Residuos y escombro retirados", [
        ("Escombro o residuos acumulados", "condicion", ["NOM-031-STPS-2011"]),
        ("Derrame o piso resbaloso sin señalar", "condicion", ["NOM-001-STPS-2008"]),
    ]),
]


def crear_tablas_sh(con):
    """Seguridad e Higiene: visitas, hallazgos y fotos. Se crean en AMBAS bases (Torre A y Torre B);
    solo tablas nuevas (IF NOT EXISTS), no toca nada de lo que ya existe."""
    con.execute("""
        CREATE TABLE IF NOT EXISTS sh_visitas (
            id         INTEGER PRIMARY KEY AUTOINCREMENT,
            folio      TEXT NOT NULL UNIQUE,
            estado     TEXT NOT NULL DEFAULT 'en_curso',
            inicio     TEXT,
            fin        TEXT,
            iniciada_por TEXT,
            cerrada_por  TEXT,
            notas      TEXT
        )""")
    # La revision del checklist es POR LUGAR (bloque o bloque|area). Si la tabla ya existia sin esa
    # columna (version anterior), se reconstruye conservando los datos con lugar = '' (sin lugar).
    cols = [r[1] for r in con.execute("PRAGMA table_info(sh_checklist)")]
    if cols and "lugar" not in cols:
        con.execute("ALTER TABLE sh_checklist RENAME TO sh_checklist_vieja")
    con.execute("""
        CREATE TABLE IF NOT EXISTS sh_checklist (
            visita_id INTEGER NOT NULL,
            punto_id  INTEGER NOT NULL,
            lugar     TEXT NOT NULL DEFAULT '',
            resultado TEXT NOT NULL,
            por       TEXT,
            cuando    TEXT,
            PRIMARY KEY (visita_id, punto_id, lugar)
        )""")
    if cols and "lugar" not in cols:
        con.execute("INSERT INTO sh_checklist (visita_id,punto_id,lugar,resultado,por,cuando) "
                    "SELECT visita_id,punto_id,'',resultado,por,cuando FROM sh_checklist_vieja")
        con.execute("DROP TABLE sh_checklist_vieja")
    con.execute("""
        CREATE TABLE IF NOT EXISTS sh_hallazgos (
            id              INTEGER PRIMARY KEY AUTOINCREMENT,
            visita_id       INTEGER,
            punto_id        INTEGER,
            categoria       TEXT,
            punto_texto     TEXT,
            tipo            TEXT NOT NULL,
            texto           TEXT NOT NULL,
            normas          TEXT,
            bloque          TEXT,
            area            TEXT,
            proveedor       TEXT,
            fecha_hallazgo  TEXT,
            fecha_captura   TEXT,
            retroactivo     INTEGER NOT NULL DEFAULT 0,
            evidencia_ref   TEXT,
            fecha_compromiso TEXT,
            estado          TEXT NOT NULL DEFAULT 'registrado',
            capturado_por   TEXT,
            val_mtto_por    TEXT, val_mtto_at TEXT, val_mtto_nota TEXT,
            val_arq_por     TEXT, val_arq_at  TEXT, val_arq_nota  TEXT,
            devuelto_por    TEXT, devuelto_at TEXT, devuelto_nota TEXT,
            enviado_at      TEXT,
            resp_texto      TEXT, resp_compromiso TEXT, resp_por TEXT, resp_at TEXT,
            cerrado_por     TEXT, cerrado_at TEXT, cierre_nota TEXT,
            eliminada       INTEGER NOT NULL DEFAULT 0
        )""")
    con.execute("""
        CREATE TABLE IF NOT EXISTS sh_fotos (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            hallazgo_id INTEGER NOT NULL,
            fase        TEXT NOT NULL DEFAULT 'hallazgo',
            mime        TEXT NOT NULL DEFAULT 'image/jpeg',
            datos       BLOB NOT NULL,
            subida_por  TEXT,
            creada      TEXT
        )""")
    con.execute("CREATE INDEX IF NOT EXISTS idx_sh_fotos_h ON sh_fotos(hallazgo_id)")
    # columna agregada despues (accion propuesta por el inspector); aditiva para bases que ya existen
    if "accion_propuesta" not in [r[1] for r in con.execute("PRAGMA table_info(sh_hallazgos)")]:
        con.execute("ALTER TABLE sh_hallazgos ADD COLUMN accion_propuesta TEXT")
    con.execute("CREATE INDEX IF NOT EXISTS idx_sh_hall_estado ON sh_hallazgos(estado, proveedor)")


def crear_catalogo_sh(con):
    """Catalogo de Seguridad e Higiene y equipo. Vive SOLO en la base principal (la misma que los
    usuarios) para que sea uno solo para las dos torres. Se siembra una unica vez."""
    con.execute("""
        CREATE TABLE IF NOT EXISTS sh_roles (
            usuario TEXT PRIMARY KEY,
            rol     TEXT NOT NULL
        )""")
    con.execute("""
        CREATE TABLE IF NOT EXISTS sh_normas (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            clave       TEXT NOT NULL UNIQUE,
            descripcion TEXT,
            activa      INTEGER NOT NULL DEFAULT 1
        )""")
    con.execute("""
        CREATE TABLE IF NOT EXISTS sh_puntos (
            id        INTEGER PRIMARY KEY AUTOINCREMENT,
            categoria TEXT NOT NULL,
            texto     TEXT NOT NULL,
            orden     INTEGER NOT NULL DEFAULT 0,
            activo    INTEGER NOT NULL DEFAULT 1
        )""")
    con.execute("""
        CREATE TABLE IF NOT EXISTS sh_opciones (
            id       INTEGER PRIMARY KEY AUTOINCREMENT,
            punto_id INTEGER NOT NULL,
            texto    TEXT NOT NULL,
            tipo     TEXT NOT NULL DEFAULT 'condicion',
            normas   TEXT,
            orden    INTEGER NOT NULL DEFAULT 0,
            activo   INTEGER NOT NULL DEFAULT 1
        )""")
    con.execute("CREATE TABLE IF NOT EXISTS config (clave TEXT PRIMARY KEY, valor TEXT)")
    ya = con.execute("SELECT valor FROM config WHERE clave='sh_sembrado'").fetchone()
    if ya:
        return
    for clave, desc in SH_NORMAS_INICIALES:
        con.execute("INSERT OR IGNORE INTO sh_normas (clave,descripcion,activa) VALUES (?,?,1)", (clave, desc))
    orden_p = 0
    for categoria, punto, opciones in SH_CATALOGO_INICIAL:
        orden_p += 1
        cur = con.execute("INSERT INTO sh_puntos (categoria,texto,orden,activo) VALUES (?,?,?,1)",
                          (categoria, punto, orden_p))
        for n, (txt, tipo, normas) in enumerate(opciones, 1):
            con.execute("INSERT INTO sh_opciones (punto_id,texto,tipo,normas,orden,activo) VALUES (?,?,?,?,?,1)",
                        (cur.lastrowid, txt, tipo, ",".join(normas), n))
    con.execute("INSERT OR REPLACE INTO config (clave,valor) VALUES ('sh_sembrado','1')")


def crear_catalogo_torre_a(con):
    """Catalogo de espacios de Torre A (bloques y areas). Solo existe en la base de Torre A.
    La clave se arma sola: bloque.area (1.1, 19.10). Se siembra una unica vez desde
    semilla/torre_a_catalogo.json y despues ya no se vuelve a tocar."""
    con.execute("""
        CREATE TABLE IF NOT EXISTS ta_bloques (
            num    INTEGER PRIMARY KEY,
            nombre TEXT NOT NULL
        )""")
    con.execute("""
        CREATE TABLE IF NOT EXISTS ta_areas (
            bloque INTEGER NOT NULL,
            num    INTEGER NOT NULL,
            nombre TEXT NOT NULL,
            PRIMARY KEY (bloque, num)
        )""")
    # Proveedores de Torre A: el numero (nivel 3 del codigo) es fijo y nunca se reutiliza.
    con.execute("""
        CREATE TABLE IF NOT EXISTS ta_proveedores (
            num    INTEGER PRIMARY KEY AUTOINCREMENT,
            nombre TEXT NOT NULL UNIQUE
        )""")
    con.execute("INSERT OR IGNORE INTO ta_proveedores (num,nombre) VALUES (1,'CEBSA')")
    # Los mismos proveedores de obra de Torre B, numerados una sola vez (por volumen de actividades).
    # Despues, cualquier proveedor nuevo se da de alta desde la pantalla y toma el siguiente numero.
    if con.execute("PRAGMA user_version").fetchone()[0] < 1:
        for num, nombre in enumerate(PROVEEDORES_INICIALES_TORRE_A, start=2):
            if con.execute("SELECT 1 FROM ta_proveedores WHERE lower(nombre)=lower(?)", (nombre,)).fetchone():
                continue
            if con.execute("SELECT 1 FROM ta_proveedores WHERE num=?", (num,)).fetchone():
                con.execute("INSERT INTO ta_proveedores (nombre) VALUES (?)", (nombre,))
            else:
                con.execute("INSERT INTO ta_proveedores (num,nombre) VALUES (?,?)", (num, nombre))
        con.execute("PRAGMA user_version = 1")
    # v1.21: el responsable de Torre A puede ser un departamento (mundo interno): se numera igual
    # pero no sale en la lista de proveedores.
    cols_prov = [r[1] for r in con.execute("PRAGMA table_info(ta_proveedores)").fetchall()]
    if "interno" not in cols_prov:
        con.execute("ALTER TABLE ta_proveedores ADD COLUMN interno INTEGER NOT NULL DEFAULT 0")
    # v1.21: especialidades de Torre A (se siembran una vez con las de Torre B; despues se agregan a mano).
    con.execute("""
        CREATE TABLE IF NOT EXISTS ta_especialidades (
            nombre TEXT PRIMARY KEY
        )""")
    if con.execute("PRAGMA user_version").fetchone()[0] < 2:
        try:
            principal = sqlite3.connect("file:%s?mode=ro" % DB_PATH, uri=True, timeout=5.0)
            giros = [r[0] for r in principal.execute(
                "SELECT DISTINCT giro FROM actividades WHERE giro IS NOT NULL AND giro<>'' "
                "AND (eliminada IS NULL OR eliminada=0) AND COALESCE(mundo,'obra')='obra' ORDER BY giro")]
            principal.close()
            for g in giros:
                con.execute("INSERT OR IGNORE INTO ta_especialidades (nombre) VALUES (?)", (g,))
            con.execute("PRAGMA user_version = 2")
        except Exception as e:   # sin base principal todavia: se reintenta en el siguiente arranque
            print("Aviso: no se pudieron copiar las especialidades de Torre B:", e)
    # v1.23: lados de las areas dibujadas (cada lado puede llevar varias especialidades) y altura de muro por area.
    con.execute("""
        CREATE TABLE IF NOT EXISTS ta_lados (
            trazo_id INTEGER NOT NULL,
            lado     INTEGER NOT NULL,
            giro     TEXT NOT NULL,
            PRIMARY KEY (trazo_id, lado, giro)
        )""")
    con.execute("""
        CREATE TABLE IF NOT EXISTS ta_alturas (
            trazo_id INTEGER PRIMARY KEY,
            altura   REAL NOT NULL
        )""")
    # v1.21: el codigo de actividad (1.1.1.1) no se puede repetir en Torre A.
    try:
        con.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_actividades_codigo_ta ON actividades(codigo)")
    except sqlite3.Error as e:
        print("Aviso: no se pudo crear el indice unico de codigos de Torre A:", e)
    # Actividades: nivel 4 (detalle=0) = actividad general; nivel 5 (detalle>=1) = detalle.
    # Codigo = bloque.area.proveedor.actividad[.detalle]; las columnas numericas sirven para ordenar.
    con.execute("""
        CREATE TABLE IF NOT EXISTS ta_actividades (
            id        INTEGER PRIMARY KEY AUTOINCREMENT,
            bloque    INTEGER NOT NULL,
            area      INTEGER NOT NULL,
            proveedor INTEGER NOT NULL,
            actividad INTEGER NOT NULL,
            detalle   INTEGER NOT NULL DEFAULT 0,
            nombre    TEXT NOT NULL,
            avance    INTEGER NOT NULL DEFAULT 0,
            eliminada INTEGER NOT NULL DEFAULT 0,
            creada    TEXT,
            UNIQUE (bloque, area, proveedor, actividad, detalle)
        )""")
    # Zonas dibujadas en el plano: un poligono por area (x,y entre 0 y 1, relativos a la imagen).
    # v1.20: un area puede tener varias piezas dibujadas (ya no hay UNIQUE por area).
    con.execute("""
        CREATE TABLE IF NOT EXISTS ta_zonas (
            id      INTEGER PRIMARY KEY AUTOINCREMENT,
            bloque  INTEGER NOT NULL,
            area    INTEGER NOT NULL,
            puntos  TEXT NOT NULL,
            creada  TEXT
        )""")
    fila_z = con.execute("SELECT sql FROM sqlite_master WHERE type='table' AND name='ta_zonas'").fetchone()
    if fila_z and fila_z[0] and "UNIQUE" in fila_z[0].upper():
        # base creada con v1.17-v1.19: se reconstruye la tabla sin UNIQUE conservando las zonas
        con.execute("ALTER TABLE ta_zonas RENAME TO ta_zonas_vieja")
        con.execute("""
            CREATE TABLE ta_zonas (
                id      INTEGER PRIMARY KEY AUTOINCREMENT,
                bloque  INTEGER NOT NULL,
                area    INTEGER NOT NULL,
                puntos  TEXT NOT NULL,
                creada  TEXT
            )""")
        con.execute("INSERT INTO ta_zonas (id,bloque,area,puntos,creada) SELECT id,bloque,area,puntos,creada FROM ta_zonas_vieja")
        con.execute("DROP TABLE ta_zonas_vieja")
    # Contorno de cada bloque (tambien por piezas).
    con.execute("""
        CREATE TABLE IF NOT EXISTS ta_bloque_zonas (
            id      INTEGER PRIMARY KEY AUTOINCREMENT,
            bloque  INTEGER NOT NULL,
            puntos  TEXT NOT NULL,
            creada  TEXT
        )""")
    # Capas del plano (como en AutoCAD): cada capa agrupa elementos del mismo tipo
    # (puntos: contactos/apagadores/luminarias; lineas: muros/ductos; zonas: plafon) y dice
    # que hace el proveedor en ella (proveedor + actividad ligada + descripcion).
    con.execute("""
        CREATE TABLE IF NOT EXISTS ta_capas (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            nombre      TEXT NOT NULL UNIQUE,
            tipo        TEXT NOT NULL CHECK (tipo IN ('punto','linea','zona')),
            color       TEXT NOT NULL DEFAULT '#2563eb',
            proveedor   INTEGER,
            actividad_id INTEGER,
            descripcion TEXT,
            creada      TEXT
        )""")
    con.execute("""
        CREATE TABLE IF NOT EXISTS ta_elementos (
            id      INTEGER PRIMARY KEY AUTOINCREMENT,
            capa_id INTEGER NOT NULL,
            puntos  TEXT NOT NULL,
            nota    TEXT,
            creada  TEXT
        )""")
    crear_tablas_actas(con)
    con.execute("""
        CREATE TABLE IF NOT EXISTS ta_historial (
            id      INTEGER PRIMARY KEY AUTOINCREMENT,
            act_id  INTEGER NOT NULL,
            antes   INTEGER,
            despues INTEGER,
            fecha   TEXT
        )""")
    if con.execute("SELECT COUNT(*) FROM ta_bloques").fetchone()[0] > 0:
        return
    try:
        with open(ARCHIVO_CATALOGO_TORRE_A, encoding="utf-8") as f:
            datos = json.load(f)
        for b in datos.get("bloques", []):
            con.execute("INSERT INTO ta_bloques (num,nombre) VALUES (?,?)", (b["num"], b["nombre"]))
            for a in b.get("areas", []):
                con.execute("INSERT INTO ta_areas (bloque,num,nombre) VALUES (?,?,?)",
                            (b["num"], a["num"], a["nombre"]))
        print("  Catalogo de Torre A sembrado:", len(datos.get("bloques", [])), "bloques")
    except Exception as e:
        print("  (no se pudo sembrar el catalogo de Torre A:", e, ")")


def init_db(ruta=None):
    ruta = ruta or DB_PATH
    es_principal = (ruta == DB_PATH)
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    con = sqlite3.connect(ruta, timeout=10.0)
    con.execute("PRAGMA journal_mode = WAL")
    con.execute("PRAGMA busy_timeout = 5000")
    con.executescript(
        """
    CREATE TABLE IF NOT EXISTS actividades (
        id             INTEGER PRIMARY KEY AUTOINCREMENT,
        codigo         TEXT,
        bloque         TEXT,
        area           TEXT,
        giro           TEXT,
        proveedor      TEXT,
        departamento   TEXT,
        mundo          TEXT DEFAULT 'obra',
        tipo_interno   TEXT,
        partida        TEXT,
        tipo           TEXT,
        tipo_partida   TEXT DEFAULT 'Construcción',
        aplica         TEXT DEFAULT 'SÍ',
        avance         INTEGER DEFAULT 0,
        f_inicio       TEXT,
        f_fin          TEXT,
        duracion_dias  INTEGER,
        estatus        TEXT DEFAULT 'Pendiente',
        depende_de     INTEGER,
        definido       TEXT DEFAULT 'NO',
        causa_retraso  TEXT,
        nota_proveedor TEXT,
        causa_por      TEXT,
        causa_fecha    TEXT,
        -- Portal de proveedores: origen y estado de validación
        origen         TEXT DEFAULT 'oficial',   -- 'oficial' (la diste de alta tú) | 'propuesta' (la agregó el proveedor)
        estado_val     TEXT DEFAULT 'validado',  -- 'validado' | 'propuesta' (espera tu firma) | 'rechazada'
        avance_decl    INTEGER,                  -- lo que el proveedor DICE que lleva (declarado)
        avance_decl_por TEXT,                     -- quién lo declaró
        avance_decl_fecha TEXT,                   -- cuándo lo declaró
        definido_por   TEXT,                      -- 'plano' | 'adicional' | 'comentario' (para que el proveedor se proteja)
        creado_por     TEXT,                      -- usuario que la creó (si es propuesta)
        reconocida     TEXT DEFAULT 'NO',         -- 'SÍ' cuando el proveedor acepta que es su trabajo
        reconocida_por TEXT,
        reconocida_fecha TEXT,
        no_reconocida_nota TEXT,                   -- si el proveedor dice "esto no es mío", su comentario
        fuera_zona     INTEGER DEFAULT 0,
        rechazo_motivo TEXT,
        rechazado_por  TEXT,
        rechazado_fecha TEXT,
        avance_decl_rechazado INTEGER DEFAULT 0,
        notas          TEXT,
        actualizado    TEXT,
        FOREIGN KEY (depende_de) REFERENCES actividades(id) ON DELETE SET NULL
    );

    -- Modelo de Dependencias por área (auto-liberación y forzado con historial)
    CREATE TABLE IF NOT EXISTS dependencias (
        id                 INTEGER PRIMARY KEY AUTOINCREMENT,
        area               TEXT NOT NULL,
        tipo_sucesor       TEXT NOT NULL,
        tipos_predecesores TEXT NOT NULL,
        umbral             INTEGER DEFAULT 100,
        permite_paralelo   INTEGER DEFAULT 0,
        liberacion_forzada INTEGER DEFAULT 0,
        forzada_por        TEXT,
        forzada_fecha      TEXT,
        forzada_nota       TEXT,
        creado             TEXT
    );

    -- Usuarios del portal de proveedores (uno o varios por empresa)
    CREATE TABLE IF NOT EXISTS usuarios (
        id        INTEGER PRIMARY KEY AUTOINCREMENT,
        usuario   TEXT UNIQUE,     -- lo que teclean para entrar
        clave     TEXT,            -- contraseña simple (se guarda con hash)
        proveedor TEXT,            -- a qué proveedor/departamento pertenece (ve solo lo suyo)
        rol       TEXT DEFAULT 'proveedor',  -- 'proveedor' | 'admin'
        mundo     TEXT DEFAULT 'obra',       -- 'obra' (proveedor externo) | 'interno' (departamento HAP)
        clave_cambiada INTEGER DEFAULT 0,     -- 1 si el usuario ya cambió la contraseña que le diste
        creado    TEXT
    );

    -- Bitácora de avisos para el panel de validación del admin
    CREATE TABLE IF NOT EXISTS avisos (
        id        INTEGER PRIMARY KEY AUTOINCREMENT,
        fecha     TEXT,
        proveedor TEXT,
        usuario   TEXT,
        tipo      TEXT,            -- 'propuesta' | 'avance'
        detalle   TEXT,
        visto     INTEGER DEFAULT 0
    );

    -- Suscripciones a notificaciones push (celular/navegador) del admin,
    -- para avisar en el momento cuando un proveedor reporta algo.
    CREATE TABLE IF NOT EXISTS push_subscriptions (
        id        INTEGER PRIMARY KEY AUTOINCREMENT,
        usuario   TEXT,
        endpoint  TEXT UNIQUE,
        p256dh    TEXT,
        auth      TEXT,
        creado    TEXT
    );

    -- Catálogo de causas de retraso (7 base + las que agregue el usuario)
    CREATE TABLE IF NOT EXISTS causas (
        id     INTEGER PRIMARY KEY AUTOINCREMENT,
        nombre TEXT UNIQUE,
        base   INTEGER DEFAULT 0,
        mundo  TEXT DEFAULT 'obra'
    );

    -- Fotografías semanales de avance por partida (para comparar semana vs semana)
    CREATE TABLE IF NOT EXISTS snapshots (
        id           INTEGER PRIMARY KEY AUTOINCREMENT,
        semana       TEXT,      -- etiqueta 'AAAA-Www' (año-semana ISO)
        fecha        TEXT,      -- fecha en que se tomó la foto
        actividad_id INTEGER,
        avance       INTEGER
    );

    CREATE TABLE IF NOT EXISTS historial (
        id           INTEGER PRIMARY KEY AUTOINCREMENT,
        actividad_id INTEGER,
        campo        TEXT,
        valor_antes  TEXT,
        valor_despues TEXT,
        fecha        TEXT,
        quien        TEXT
    );

    -- Catálogo de áreas/departamentos que intervienen para dejar un espacio funcional
    CREATE TABLE IF NOT EXISTS tipos_internos (
        id     INTEGER PRIMARY KEY AUTOINCREMENT,
        nombre TEXT UNIQUE
    );

    CREATE TABLE IF NOT EXISTS involucrados (
        id     INTEGER PRIMARY KEY AUTOINCREMENT,
        nombre TEXT UNIQUE
    );

    -- Qué involucrados le tocan a cada área (para dejarla funcional)
    CREATE TABLE IF NOT EXISTS area_involucrados (
        id           INTEGER PRIMARY KEY AUTOINCREMENT,
        area         TEXT,
        involucrado  TEXT,
        UNIQUE(area, involucrado)
    );

    CREATE TABLE IF NOT EXISTS config (
        clave TEXT PRIMARY KEY,
        valor TEXT
    );

    -- Pines de actividades sobre el plano de obra (x,y en fracción 0..1 del plano)
    CREATE TABLE IF NOT EXISTS plano_pines (
        id           INTEGER PRIMARY KEY AUTOINCREMENT,
        actividad_id INTEGER NOT NULL,
        x            REAL NOT NULL,
        y            REAL NOT NULL,
        nota         TEXT,
        creado       TEXT,
        creado_por   TEXT
    );

    -- Trazos sobre el plano: una línea (muro, pintura, tablaroca) o una zona (plafón, piso).
    -- Cada trazo puede tener varias actividades de distintos gremios asignadas.
    CREATE TABLE IF NOT EXISTS plano_trazos (
        id         INTEGER PRIMARY KEY AUTOINCREMENT,
        tipo       TEXT NOT NULL,          -- 'linea' | 'zona'
        nombre     TEXT,
        puntos     TEXT NOT NULL,          -- JSON [[x,y],...] en fracción 0..1 del plano
        creado     TEXT,
        creado_por TEXT
    );
    CREATE TABLE IF NOT EXISTS plano_trazo_act (
        trazo_id     INTEGER NOT NULL,
        actividad_id INTEGER NOT NULL,
        PRIMARY KEY (trazo_id, actividad_id)
    );
    -- Bloque/área(s) de tu lista a las que está ligado el trazo (una zona puede abarcar varias áreas)
    CREATE TABLE IF NOT EXISTS plano_trazo_ub (
        trazo_id INTEGER NOT NULL,
        bloque   TEXT NOT NULL,
        area     TEXT NOT NULL DEFAULT '',
        PRIMARY KEY (trazo_id, bloque, area)
    );

    -- Catálogo de proveedores con tipo (interno/externo) y de qué se encarga
    CREATE TABLE IF NOT EXISTS proveedores (
        id       INTEGER PRIMARY KEY AUTOINCREMENT,
        nombre   TEXT UNIQUE,
        tipo     TEXT DEFAULT 'Externo',
        funcion  TEXT
    );

    -- Catálogo simple para bloques, áreas y giros (nombre + nota de para qué es)
    CREATE TABLE IF NOT EXISTS catalogo (
        id     INTEGER PRIMARY KEY AUTOINCREMENT,
        clase  TEXT,      -- 'bloque' | 'area' | 'giro'
        nombre TEXT,
        nota   TEXT,
        UNIQUE(clase, nombre)
    );
    """
    )
    # --- Migración: los trazos del plano guardan su mundo, tipo de elemento, especialidad principal y si traen solos las actividades de sus áreas ---
    tcols = [r[1] for r in con.execute("PRAGMA table_info(plano_trazos)").fetchall()]
    for c, tipo in (("giro", "TEXT"), ("tipo_elem", "TEXT"), ("mundo", "TEXT DEFAULT 'obra'"), ("auto", "INTEGER DEFAULT 0")):
        if c not in tcols:
            con.execute(f"ALTER TABLE plano_trazos ADD COLUMN {c} {tipo}")
    con.execute("UPDATE plano_trazos SET mundo='obra' WHERE mundo IS NULL")
    # --- Migración: agrega columnas nuevas si la BD ya existía sin ellas ---
    cols = [r[1] for r in con.execute("PRAGMA table_info(actividades)").fetchall()]
    if "tipo_partida" not in cols:
        con.execute("ALTER TABLE actividades ADD COLUMN tipo_partida TEXT DEFAULT 'Construcción'")
        con.execute("UPDATE actividades SET tipo_partida='Construcción' WHERE tipo_partida IS NULL")
    if "definido" not in cols:
        con.execute("ALTER TABLE actividades ADD COLUMN definido TEXT DEFAULT 'NO'")
        con.execute("UPDATE actividades SET definido='NO' WHERE definido IS NULL")
    for c in ("causa_retraso", "nota_proveedor", "causa_por", "causa_fecha"):
        if c not in cols:
            con.execute(f"ALTER TABLE actividades ADD COLUMN {c} TEXT")
    # Portal de proveedores y rechazos
    portal_cols = {
        "origen": "TEXT DEFAULT 'oficial'",
        "estado_val": "TEXT DEFAULT 'validado'",
        "avance_decl": "INTEGER",
        "avance_decl_por": "TEXT",
        "avance_decl_fecha": "TEXT",
        "definido_por": "TEXT",
        "creado_por": "TEXT",
        "reconocida": "TEXT DEFAULT 'NO'",
        "reconocida_por": "TEXT",
        "reconocida_fecha": "TEXT",
        "no_reconocida_nota": "TEXT",
        "mundo": "TEXT DEFAULT 'obra'",
        "departamento": "TEXT",
        "tipo_interno": "TEXT",
        "fuera_zona": "INTEGER DEFAULT 0",
        "rechazo_motivo": "TEXT",
        "rechazado_por": "TEXT",
        "rechazado_fecha": "TEXT",
        "avance_decl_rechazado": "INTEGER DEFAULT 0",
    }
    for c, tipo in portal_cols.items():
        if c not in cols:
            con.execute(f"ALTER TABLE actividades ADD COLUMN {c} {tipo}")
    # Pruebas de funcionamiento / revisión y entrega validada por interno
    # (ej. contacto con polaridad y tierra, apagador que prenda, llave que saque agua)
    pruebas_cols = {
        "requiere_pruebas": "TEXT DEFAULT 'NO'",   # 'SÍ' | 'NO'
        "valida_depto": "TEXT",                    # a quién (interno) le toca validar
        "origen_actividad_id": "INTEGER",           # solo en la actividad de validación auto-creada: liga a la original
    }
    for c, tipo in pruebas_cols.items():
        if c not in cols:
            con.execute(f"ALTER TABLE actividades ADD COLUMN {c} {tipo}")
    con.execute("UPDATE actividades SET requiere_pruebas='NO' WHERE requiere_pruebas IS NULL")
    con.execute("UPDATE actividades SET mundo='obra' WHERE mundo IS NULL")
    # UNA sola vez: todas las actividades de obra quedan con "requiere pruebas" = SÍ (el admin desmarca las excepciones)
    con.execute("CREATE TABLE IF NOT EXISTS config (clave TEXT PRIMARY KEY, valor TEXT)")
    if not con.execute("SELECT 1 FROM config WHERE clave='requiere_pruebas_si_v1'").fetchone():
        con.execute("UPDATE actividades SET requiere_pruebas='SÍ' WHERE COALESCE(mundo,'obra')='obra' "
                    "AND (eliminada IS NULL OR eliminada=0)") if "eliminada" in cols else \
            con.execute("UPDATE actividades SET requiere_pruebas='SÍ' WHERE COALESCE(mundo,'obra')='obra'")
        con.execute("INSERT INTO config (clave,valor) VALUES ('requiere_pruebas_si_v1', date('now'))")
    # las partidas que ya existían son oficiales y validadas
    con.execute("UPDATE actividades SET origen='oficial' WHERE origen IS NULL")
    con.execute("UPDATE actividades SET estado_val='validado' WHERE estado_val IS NULL")
    # todas arrancan SIN reconocer: el proveedor tiene que aceptarlas sí o sí
    con.execute("UPDATE actividades SET reconocida='NO' WHERE reconocida IS NULL")

    # borrado suave: columna eliminada
    if "eliminada" not in cols:
        con.execute("ALTER TABLE actividades ADD COLUMN eliminada INTEGER DEFAULT 0")
    if "eliminada_por" not in cols:
        con.execute("ALTER TABLE actividades ADD COLUMN eliminada_por TEXT")
    if "eliminada_fecha" not in cols:
        con.execute("ALTER TABLE actividades ADD COLUMN eliminada_fecha TEXT")

    # historial: columna 'quien' si falta
    hcols = [r[1] for r in con.execute("PRAGMA table_info(historial)").fetchall()]
    if "quien" not in hcols:
        con.execute("ALTER TABLE historial ADD COLUMN quien TEXT")

    # usuarios: columnas mundo y clave_cambiada si faltan
    ucols = [r[1] for r in con.execute("PRAGMA table_info(usuarios)").fetchall()]
    if "mundo" not in ucols:
        con.execute("ALTER TABLE usuarios ADD COLUMN mundo TEXT DEFAULT 'obra'")
        con.execute("UPDATE usuarios SET mundo='obra' WHERE mundo IS NULL")
    if "clave_cambiada" not in ucols:
        con.execute("ALTER TABLE usuarios ADD COLUMN clave_cambiada INTEGER DEFAULT 0")
    if "activo" not in ucols:
        con.execute("ALTER TABLE usuarios ADD COLUMN activo INTEGER DEFAULT 1")
        con.execute("UPDATE usuarios SET activo=1 WHERE activo IS NULL")
    if "num_logins" not in ucols:
        con.execute("ALTER TABLE usuarios ADD COLUMN num_logins INTEGER DEFAULT 0")
    if "ultimo_login" not in ucols:
        con.execute("ALTER TABLE usuarios ADD COLUMN ultimo_login TEXT")
    if "telefono" not in ucols:
        con.execute("ALTER TABLE usuarios ADD COLUMN telefono TEXT")

    # Expediente de proveedores: columnas de contacto
    pcols = [r[1] for r in con.execute("PRAGMA table_info(proveedores)").fetchall()]
    for c in ("empresa", "contacto", "telefono", "correo", "notas"):
        if c not in pcols:
            con.execute(f"ALTER TABLE proveedores ADD COLUMN {c} TEXT")
    if "activo" not in pcols:
        con.execute("ALTER TABLE proveedores ADD COLUMN activo INTEGER DEFAULT 1")
        con.execute("UPDATE proveedores SET activo=1 WHERE activo IS NULL")

    # Sembrar las 7 causas base (solo si la tabla está vacía)
    n_causas = con.execute("SELECT COUNT(*) FROM causas").fetchone()[0]
    if n_causas == 0:
        base = [
            "Definición pendiente de proyecto",
            "Disponibilidad de personal",
            "Suministro o entrega de material",
            "Dependencia de otro proveedor",
            "Cambio solicitado por el hospital",
            "Acceso o liberación del área",
            "Condición no prevista en sitio",
        ]
        for nombre in base:
            con.execute("INSERT OR IGNORE INTO causas (nombre,base,mundo) VALUES (?,1,'obra')", (nombre,))

    # migración: columna mundo en causas si falta
    ccols = [r[1] for r in con.execute("PRAGMA table_info(causas)").fetchall()]
    if "mundo" not in ccols:
        con.execute("ALTER TABLE causas ADD COLUMN mundo TEXT DEFAULT 'obra'")
        con.execute("UPDATE causas SET mundo='obra' WHERE mundo IS NULL")

    # Causas propias del mundo interno (departamentos HAP) — distintas a las de obra
    n_causas_int = con.execute("SELECT COUNT(*) FROM causas WHERE mundo='interno'").fetchone()[0]
    if n_causas_int == 0:
        base_int = [
            "Compras no ha surtido el material",
            "Compras dio fecha de entrega lejana",
            "Falta autorización de la dirección",
            "Presupuesto no liberado",
            "Depende de que termine la obra en el área",
            "Falta definición técnica del equipo",
            "Personal del departamento ocupado en otra prioridad",
            "Equipo en revisión o garantía",
        ]
        for nombre in base_int:
            con.execute("INSERT OR IGNORE INTO causas (nombre,base,mundo) VALUES (?,1,'interno')", (nombre,))

    # Catálogo de involucrados (áreas del hospital que intervienen para dejar funcional un espacio)
    n_inv = con.execute("SELECT COUNT(*) FROM involucrados").fetchone()[0]
    if n_inv == 0:
        base_inv = ["Sistemas", "Intendencia", "Biomédica", "Publicidad", "Compras",
                    "Finanzas", "Contabilidad", "Mantenimiento", "Enfermería", "Gerencia médica"]
        for nombre in base_inv:
            con.execute("INSERT OR IGNORE INTO involucrados (nombre) VALUES (?)", (nombre,))

    # Tipos de tarea del mundo interno (editables por el admin)
    n_ti = con.execute("SELECT COUNT(*) FROM tipos_internos").fetchone()[0]
    if n_ti == 0:
        for nombre in ["Instalación", "Adecuación del área", "Solicitar a compras", "Puesta en marcha / prueba"]:
            con.execute("INSERT OR IGNORE INTO tipos_internos (nombre) VALUES (?)", (nombre,))

    # Usuario admin inicial: solo se crea en una base SIN usuarios (servidor nuevo).
    # La contraseña sale de la variable ADMIN_CLAVE_INICIAL; si no existe, se
    # genera una al azar y se muestra una sola vez en los logs del servidor.
    n_users = con.execute("SELECT COUNT(*) FROM usuarios").fetchone()[0]
    if n_users == 0 and es_principal:
        clave_inicial = os.environ.get("ADMIN_CLAVE_INICIAL") or secrets.token_urlsafe(9)
        con.execute(
            "INSERT INTO usuarios (usuario,clave,proveedor,rol,creado) VALUES (?,?,?,?,?)",
            ("admin", crear_hash_clave(clave_inicial), None, "admin",
             datetime.datetime.now().isoformat(timespec="seconds")))
        if os.environ.get("ADMIN_CLAVE_INICIAL"):
            print("  Usuario 'admin' creado con la contraseña de ADMIN_CLAVE_INICIAL.")
        else:
            print("  Usuario 'admin' creado. Contraseña inicial:", clave_inicial,
                  "— cámbiala al entrar.")
    if ruta == DB_PATH_TORRE_A:
        crear_catalogo_torre_a(con)
    crear_tablas_actas(con)
    crear_tablas_dashboard(con)
    crear_tablas_sh(con)
    if es_principal:
        crear_catalogo_sh(con)
    pendientes = contar_claves_antiguas(con)
    if pendientes > 0:
        print("  Contraseñas en formato antiguo pendientes de migrar:", pendientes,
              "(se migran solas al iniciar sesión; no quites SAL_CLAVES hasta que sea 0)")
    con.commit()
    con.close()
    generar_respaldo_bd(ruta)


# ----------------------------------------------------------------------------
# Utilidades
# ----------------------------------------------------------------------------
def hoy():
    return datetime.date.today().isoformat()


def parse_date(s):
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.datetime.strptime(str(s)[:10], fmt).date()
        except ValueError:
            continue
    return None


def registrar_historial(db, aid, campo, antes, despues, quien=None):
    """Guarda un cambio en el historial, con quién lo hizo."""
    if str(antes) == str(despues):
        return
    db.execute(
        "INSERT INTO historial (actividad_id,campo,valor_antes,valor_despues,fecha,quien) "
        "VALUES (?,?,?,?,?,?)",
        (aid, campo, "" if antes is None else str(antes),
         "" if despues is None else str(despues),
         datetime.datetime.now().isoformat(timespec="seconds"), quien))


def estatus_por_avance(av):
    if av <= 0:
        return "Pendiente"
    if av >= 100:
        return "Listo"
    return "En proceso"


# --- Torre A en la pantalla principal (v1.21) -------------------------------------------------
# Las actividades de Torre A guardan bloque y area como etiqueta "01 · Quirófanos" y
# "01.01 · Quirófano ..." (con ceros para que ordenen bien como texto). El numero sale de la etiqueta.
_RE_BLOQUE_TA = re.compile(r"^\s*(\d+)\s*·")
_RE_AREA_TA = re.compile(r"^\s*(\d+)\.(\d+)\s*·")


def etiqueta_bloque_ta(num, nombre):
    return "%02d · %s" % (num, nombre)


def etiqueta_area_ta(bloque, num, nombre):
    return "%02d.%02d · %s" % (bloque, num, nombre)


def catalogo_ta_para_pantalla(db):
    """Listas de Torre A para los desplegables de la pantalla principal."""
    bloques, areas, mapa = [], [], {}
    for b in db.execute("SELECT num,nombre FROM ta_bloques ORDER BY num").fetchall():
        eb = etiqueta_bloque_ta(b["num"], b["nombre"])
        bloques.append(eb)
        mapa[eb] = []
    nombres_b = {r["num"]: r["nombre"] for r in db.execute("SELECT num,nombre FROM ta_bloques")}
    for a in db.execute("SELECT bloque,num,nombre FROM ta_areas ORDER BY bloque,num").fetchall():
        if a["bloque"] not in nombres_b:
            continue
        ea = etiqueta_area_ta(a["bloque"], a["num"], a["nombre"])
        areas.append(ea)
        mapa[etiqueta_bloque_ta(a["bloque"], nombres_b[a["bloque"]])].append(ea)
    return bloques, areas, mapa


def numero_proveedor_ta(db, nombre, interno=False):
    """Numero fijo del responsable en Torre A. Si todavia no existe, se da de alta con el siguiente numero.
    Compara sin importar mayusculas ni acentos (Oscar = Óscar)."""
    nombre = " ".join((nombre or "").split())
    if not nombre:
        return None
    clave = sin_acentos(nombre)
    for r in db.execute("SELECT num,nombre FROM ta_proveedores").fetchall():
        if sin_acentos(r["nombre"]) == clave:
            return r["num"]
    try:
        cur = db.execute("INSERT INTO ta_proveedores (nombre,interno) VALUES (?,?)", (nombre, 1 if interno else 0))
        return cur.lastrowid
    except sqlite3.IntegrityError:
        r = db.execute("SELECT num FROM ta_proveedores WHERE lower(nombre)=lower(?)", (nombre,)).fetchone()
        return r["num"] if r else None


def nombre_proveedor_ta(db, num):
    r = db.execute("SELECT nombre FROM ta_proveedores WHERE num=?", (num,)).fetchone()
    return r["nombre"] if r else None


def partes_codigo_ta(db, data):
    """Valida bloque/area/responsable de una actividad nueva de Torre A.
    Devuelve ((bloque, area, proveedor), None) o (None, mensaje)."""
    mb = _RE_BLOQUE_TA.match(data.get("bloque") or "")
    ma = _RE_AREA_TA.match(data.get("area") or "")
    if not mb:
        return None, "Elige el bloque de la lista (los 20 bloques de Torre A)."
    if not ma:
        return None, "Elige el área de la lista (la de ese bloque)."
    b, a = int(mb.group(1)), int(ma.group(2))
    if int(ma.group(1)) != b:
        return None, "El área elegida no pertenece a ese bloque."
    if not db.execute("SELECT 1 FROM ta_areas WHERE bloque=? AND num=?", (b, a)).fetchone():
        return None, "Esa área no existe en el catálogo de Torre A."
    interno = (data.get("mundo") or "obra") == "interno"
    resp = (data.get("departamento") if interno else data.get("proveedor")) or ""
    p = numero_proveedor_ta(db, resp, interno)
    if not p:
        return None, "Elige el responsable (proveedor o departamento)."
    # se guardan SIEMPRE las etiquetas del catalogo (asi "Ver zona" y el plano siempre las encuentran)
    rb = db.execute("SELECT nombre FROM ta_bloques WHERE num=?", (b,)).fetchone()
    ra = db.execute("SELECT nombre FROM ta_areas WHERE bloque=? AND num=?", (b, a)).fetchone()
    if rb:
        data["bloque"] = etiqueta_bloque_ta(b, rb["nombre"])
    if ra:
        data["area"] = etiqueta_area_ta(b, a, ra["nombre"])
    data["departamento" if interno else "proveedor"] = nombre_proveedor_ta(db, p) or resp
    return (b, a, p), None


def siguiente_codigo_ta(db, b, a, p):
    """b.a.p.N con N = siguiente actividad de ese bloque/area/responsable (los numeros no se reutilizan:
    cuentan tambien las eliminadas)."""
    prefijo = "%d.%d.%d." % (b, a, p)
    mayor = 0
    for f in db.execute("SELECT codigo FROM actividades WHERE codigo LIKE ?", (prefijo + "%",)):
        m = re.match(r"^%s(\d+)(?:-V)?$" % re.escape(prefijo), (f["codigo"] or "").strip())
        if m:
            mayor = max(mayor, int(m.group(1)))
    return "%s%d" % (prefijo, mayor + 1)


def fecha_entrega_actual():
    """Torre B: la fecha fija de siempre. Torre A: la que el admin defina (puede no haber)."""
    if proyecto_actual() == "torre_a":
        try:
            return (_config_get(get_db(), "fecha_entrega", "") or "").strip() or None
        except Exception:
            return None
    return FECHA_ENTREGA


def siguiente_codigo(db):
    """Calcula el próximo código ACT-#### disponible: busca el número más alto
    entre TODOS los códigos con forma ACT-#### o PROP-####, ignorando filas sin
    código o con sufijo (como las tareas de validación interna, que llevan
    "-V" al final: ACT-1051-V). Se usa cada vez que se crea una actividad
    (nueva o duplicada) sin un código ya definido, para que nunca quede en
    blanco.

    El código provisional que se le da a una propuesta de proveedor (ver
    api_proponer_actividad) es un timestamp de 10 dígitos (PROP-MMDDHHMMSS),
    no un consecutivo — se limita a 6 dígitos el número que se considera aquí
    para que esos códigos NUNCA se cuelen en este cálculo y arrastren el
    consecutivo de ACT-#### a números absurdos como ACT-921230507 (bug real
    detectado en producción: una vez que existió una propuesta con timestamp,
    todo lo creado después heredó ese número gigante)."""
    num = 1
    for fila in db.execute("SELECT codigo FROM actividades WHERE codigo IS NOT NULL"):
        m = re.match(r"^(?:ACT|PROP)-(\d{1,6})$", (fila["codigo"] or "").strip())
        if m:
            num = max(num, int(m.group(1)) + 1)
    return f"ACT-{num:04d}"


def crear_validacion_interna_si_aplica(db, actividad_id):
    """Si la actividad llegó a 100% y tiene marcado 'requiere pruebas de
    funcionamiento / revisión y entrega' con un departamento asignado para
    validarla (ej. Biomédica revisa que un contacto tenga polaridad y tierra,
    que un apagador prenda y apague, que una llave saque agua), se crea UNA
    actividad interna ligada a ella para que ese departamento la trabaje y
    la marque como terminada cuando de verdad ya funcione. No duplica si ya
    existe una ligada a esta misma actividad."""
    a = db.execute("SELECT * FROM actividades WHERE id=?", (actividad_id,)).fetchone()
    if not a or (a["avance"] or 0) < 100:
        return
    if (a["requiere_pruebas"] or "NO") != "SÍ":
        return
    depto = (a["valida_depto"] or "").strip()
    if not depto:
        return
    ya = db.execute(
        "SELECT id FROM actividades WHERE origen_actividad_id=? AND (eliminada IS NULL OR eliminada=0)",
        (actividad_id,)).fetchone()
    if ya:
        return
    ahora = datetime.datetime.now().isoformat(timespec="seconds")
    partida_val = "Validar: " + (a["partida"] or a["codigo"] or f"actividad #{actividad_id}")
    # Mismo consecutivo que la actividad original, con sufijo "-V" (ej. ACT-1051 -> ACT-1051-V),
    # para que en cualquier pantalla (portal, listas, reportes) se vea clarísimo que es
    # la validación de esa actividad y no una actividad nueva sin relación.
    codigo_val = (a["codigo"].strip() + "-V") if a["codigo"] else None
    if codigo_val and proyecto_actual() == "torre_a":
        # en Torre A el codigo es unico: si ya hubo una validacion eliminada con ese codigo, se numera (-V2, -V3...)
        n = 1
        while db.execute("SELECT 1 FROM actividades WHERE codigo=?", (codigo_val,)).fetchone():
            n += 1
            codigo_val = "%s-V%d" % (a["codigo"].strip(), n)
    db.execute(
        """INSERT INTO actividades
        (codigo,bloque,area,giro,departamento,mundo,tipo_interno,partida,tipo_partida,aplica,avance,
         estatus,definido,origen,estado_val,reconocida,origen_actividad_id,actualizado)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (codigo_val, a["bloque"], a["area"], a["giro"], depto, "interno", "Validación",
         partida_val, "Puesta en marcha", "SÍ", 0, "Pendiente", "NO",
         "oficial", "validado", "NO", actividad_id, ahora))
    db.commit()


def dias_restantes():
    fecha = fecha_entrega_actual()
    if not fecha:
        return None
    try:
        return (parse_date(fecha) - datetime.date.today()).days
    except Exception:
        return None


def evaluar_dependencias_area(db, area, actividad_tipo_o_giro=None):
    """
    Evalúa el estado de las dependencias para un área dada (o para un tipo específico dentro del área).
    Devuelve dict con:
      bloqueada: bool,
      liberada: bool,
      estado: str,
      detalle: str,
      predecesoras_avance: float,
      umbral: int,
      forzada: bool
    """
    if not area:
        return {"bloqueada": False, "liberada": True, "estado": "Sin dependencias", "detalle": "", "forzada": False}
    
    area_norm = sin_acentos(area)
    all_deps = db.execute("SELECT * FROM dependencias").fetchall()
    deps = [d for d in all_deps if sin_acentos(d["area"]) == area_norm or sin_acentos(d["area"]) in area_norm or area_norm in sin_acentos(d["area"])]
    if not deps:
        return {"bloqueada": False, "liberada": True, "estado": "Sin dependencias", "detalle": "", "forzada": False}
    
    # Si se pide para un tipo específico, filtrar dependencias que apliquen a ese tipo sucesor
    if actividad_tipo_o_giro:
        act_tipo_norm = sin_acentos(actividad_tipo_o_giro)
        deps_filtradas = [d for d in deps if sin_acentos(d["tipo_sucesor"]) in act_tipo_norm or act_tipo_norm in sin_acentos(d["tipo_sucesor"])]
        if not deps_filtradas:
            return {"bloqueada": False, "liberada": True, "estado": "Sin dependencias", "detalle": "", "forzada": False}
        deps = deps_filtradas

    # Evaluar las dependencias activas
    for d in deps:
        if d["liberacion_forzada"]:
            return {
                "bloqueada": False,
                "liberada": True,
                "estado": "Liberada (anticipada)",
                "detalle": f"Liberación anticipada por {d['forzada_por'] or 'admin'} ({d['forzada_fecha'] or ''}): {d['forzada_nota'] or ''}",
                "forzada": True,
                "umbral": d["umbral"] or 100,
                "dep_id": d["id"]
            }
        
        # Obtener lista de tipos predecesores
        pred_txt = d["tipos_predecesores"] or ""
        preds = [sin_acentos(p.strip()) for p in pred_txt.split(",") if p.strip()]
        
        # Buscar avance oficial de las actividades predecesoras en esa área
        filas_area = db.execute(
            "SELECT avance, giro, tipo, tipo_partida, partida FROM actividades WHERE area=? AND (aplica IS NULL OR aplica<>'NO') AND (eliminada IS NULL OR eliminada=0)",
            (area,)
        ).fetchall()
        
        avances_pred = []
        for fa in filas_area:
            tps = [sin_acentos(fa["giro"]), sin_acentos(fa["tipo"]), sin_acentos(fa["tipo_partida"]), sin_acentos(fa["partida"])]
            texto_comp = " ".join(tps)
            if any(p in texto_comp for p in preds):
                avances_pred.append(fa["avance"] or 0)
                
        if not avances_pred:
            # Si no hay predecesoras declaradas o encontradas en esa área, continúa evaluando
            continue
            
        prom_avance = sum(avances_pred) / len(avances_pred)
        umbral = d["umbral"] or 100
        
        if prom_avance < umbral:
            return {
                "bloqueada": True,
                "liberada": False,
                "estado": f"Bloqueada (Predecesoras al {prom_avance:.0f}%, requiere {umbral}%)",
                "detalle": f"Esperando que '{pred_txt}' alcance el {umbral}% en {area} (avance actual: {prom_avance:.0f}%)",
                "predecesoras_avance": round(prom_avance, 1),
                "umbral": umbral,
                "forzada": False,
                "dep_id": d["id"]
            }
            
    return {
        "bloqueada": False,
        "liberada": True,
        "estado": "Liberada (cumplió umbral)",
        "detalle": "Predecesoras completadas al umbral requerido",
        "forzada": False
    }


# ----------------------------------------------------------------------------
# Rutas de páginas
# ----------------------------------------------------------------------------
@app.route("/api/tipos_internos", methods=["GET"])
@requiere_login
def api_tipos_internos():
    db = get_db()
    rows = db.execute("SELECT nombre FROM tipos_internos ORDER BY id").fetchall()
    return jsonify([r["nombre"] for r in rows])


@app.route("/api/tipos_internos", methods=["POST"])
@requiere_gestor
def api_tipos_internos_add():
    db = get_db()
    nombre = (request.get_json().get("nombre") or "").strip()
    if not nombre:
        return jsonify({"error": "Escribe un nombre"}), 400
    db.execute("INSERT OR IGNORE INTO tipos_internos (nombre) VALUES (?)", (nombre,))
    db.commit()
    return jsonify({"ok": True})


@app.route("/involucrados")
def involucrados_page():
    return render_template("involucrados.html")


@app.route("/api/involucrados/catalogo", methods=["GET"])
@requiere_gestor
def api_inv_catalogo():
    db = get_db()
    rows = db.execute("SELECT nombre FROM involucrados ORDER BY nombre").fetchall()
    return jsonify([r["nombre"] for r in rows])


@app.route("/api/involucrados/catalogo", methods=["POST"])
@requiere_gestor
def api_inv_catalogo_add():
    db = get_db()
    nombre = (request.get_json().get("nombre") or "").strip()
    if not nombre:
        return jsonify({"error": "Escribe un nombre"}), 400
    db.execute("INSERT OR IGNORE INTO involucrados (nombre) VALUES (?)", (nombre,))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/involucrados/por_area", methods=["GET"])
@requiere_gestor
def api_inv_por_area():
    """Devuelve, por cada área, qué involucrados tiene encendidos (y el catálogo completo)."""
    db = get_db()
    catalogo = [r["nombre"] for r in db.execute("SELECT nombre FROM involucrados ORDER BY nombre").fetchall()]
    areas = db.execute(
        "SELECT DISTINCT bloque, area FROM actividades WHERE area IS NOT NULL AND area<>'' "
        "ORDER BY bloque, area").fetchall()
    rel = {}
    for r in db.execute("SELECT area, involucrado FROM area_involucrados").fetchall():
        rel.setdefault(r["area"], set()).add(r["involucrado"])
    salida = []
    for a in areas:
        encendidos = sorted(rel.get(a["area"], set()))
        salida.append({"bloque": a["bloque"] or "", "area": a["area"], "encendidos": encendidos})
    return jsonify({"catalogo": catalogo, "areas": salida})


@app.route("/api/involucrados/toggle", methods=["POST"])
@requiere_gestor
def api_inv_toggle():
    """Prende o apaga un involucrado en una área."""
    db = get_db()
    data = request.get_json()
    area = data.get("area"); who = data.get("involucrado")
    if not area or not who:
        return jsonify({"error": "faltan datos"}), 400
    existe = db.execute("SELECT 1 FROM area_involucrados WHERE area=? AND involucrado=?", (area, who)).fetchone()
    if existe:
        db.execute("DELETE FROM area_involucrados WHERE area=? AND involucrado=?", (area, who))
        estado = "off"
    else:
        db.execute("INSERT OR IGNORE INTO area_involucrados (area,involucrado) VALUES (?,?)", (area, who))
        estado = "on"
    db.commit()
    return jsonify({"ok": True, "estado": estado})


@app.route("/login")
def login_page():
    return render_template("login.html")


@app.after_request
def _sin_cache_html(resp):
    """Las pantallas HTML no se guardan en la memoria del navegador: evita ver una versión vieja o un 'atrás' que no responde."""
    if resp.mimetype == "text/html":
        resp.headers["Cache-Control"] = "no-store"
    return resp


@app.route("/sw.js")
def service_worker():
    """El 'service worker' que recibe las notificaciones push y las muestra
    como notificación del sistema aunque la pestaña esté cerrada. Se sirve
    desde la raíz (no desde /static/) para que su alcance cubra todo el
    sitio sin configuración extra."""
    resp = send_file(os.path.join(BASE_DIR, "static", "sw.js"), mimetype="application/javascript")
    resp.headers["Service-Worker-Allowed"] = "/"
    resp.headers["Cache-Control"] = "no-cache"
    return resp


@app.route("/")
def index():
    # si no hay sesión, al login; si es proveedor, a su portal
    if "usuario" not in session:
        return redirect("/login")
    if not es_gestor(session.get("rol")):
        return redirect("/portal")
    return render_template("index.html")


@app.route("/reporte")
def reporte():
    return render_template("reporte.html")


@app.route("/reportes")
def reportes():
    # esta pantalla es de uso interno (hojas por proveedor, captura masiva,
    # importar avances) — antes no pedía sesión en absoluto.
    if "usuario" not in session:
        return redirect("/login")
    return render_template("reportes.html")


@app.route("/expediente")
def expediente_page():
    # esta pantalla se unificó con Usuarios: la ficha del proveedor y su
    # acceso a la plataforma ahora viven juntos en /usuarios
    return redirect("/usuarios")


# ----------------------------------------------------------------------------
# API — Actividades
# ----------------------------------------------------------------------------
@app.route("/api/actividades")
@requiere_login
def api_actividades():
    db = get_db()
    # n_pines: cuántos pines tiene la actividad en el plano de obra (para el 📍 de la tabla)
    q = ("SELECT actividades.*, (SELECT COUNT(*) FROM plano_pines pp WHERE pp.actividad_id=actividades.id) AS n_pines "
         "FROM actividades")
    cond = ["(eliminada IS NULL OR eliminada=0)"]
    args = []
    bloque = request.args.get("bloque")
    area = request.args.get("area")
    giro = request.args.get("giro")
    proveedor = request.args.get("proveedor")
    tipo_partida = request.args.get("tipo_partida")
    estatus = request.args.get("estatus")
    buscar = request.args.get("buscar")
    # mundo: 'obra' (default), 'interno' o 'todos'
    mundo = request.args.get("mundo", "obra")
    # en 'interno' el responsable vive en la columna departamento, no proveedor
    col_resp = "departamento" if mundo == "interno" else "proveedor"
    if mundo != "todos":
        cond.append("(mundo = ? OR (mundo IS NULL AND ? = 'obra'))")
        args += [mundo, mundo]
    if bloque:
        cond.append("bloque = ?"); args.append(bloque)
    if area:
        cond.append("area = ?"); args.append(area)
    if giro:
        cond.append("giro = ?"); args.append(giro)
    if proveedor:
        cond.append(f"{col_resp} = ?"); args.append(proveedor)
    if tipo_partida:
        cond.append("tipo_partida = ?"); args.append(tipo_partida)
    if estatus:
        cond.append("estatus = ?"); args.append(estatus)
    avance_min = request.args.get("avance_min")
    if avance_min not in (None, ""):
        try:
            cond.append("avance >= ?"); args.append(max(0, min(100, int(avance_min))))
        except (TypeError, ValueError):
            pass
    avance_max = request.args.get("avance_max")
    if avance_max not in (None, ""):
        try:
            cond.append("avance <= ?"); args.append(max(0, min(100, int(avance_max))))
        except (TypeError, ValueError):
            pass
    if buscar:
        b = "%" + sin_acentos(buscar) + "%"
        cond.append(f"(sinac(partida) LIKE ? OR sinac(area) LIKE ? OR sinac({col_resp}) LIKE ? OR sinac(bloque) LIKE ? OR sinac(codigo) LIKE ?)")
        args += [b] * 5
    # filtro por pin en el plano: 'con' = ya ubicadas, 'sin' = faltan por ubicar
    pin = request.args.get("pin")
    if pin == "con":
        cond.append("EXISTS (SELECT 1 FROM plano_pines pf WHERE pf.actividad_id=actividades.id)")
    elif pin == "sin":
        cond.append("NOT EXISTS (SELECT 1 FROM plano_pines pf WHERE pf.actividad_id=actividades.id)")
    # por defecto no mostramos las rechazadas en la tabla principal
    if request.args.get("incluir_rechazadas") != "1":
        cond.append("(estado_val IS NULL OR estado_val<>'rechazada')")
    if cond:
        q += " WHERE " + " AND ".join(cond)
    q += " ORDER BY id"
    rows = db.execute(q, args).fetchall()
    salida = []
    for r in rows:
        d = dict(r)
        dep_st = evaluar_dependencias_area(db, r["area"], r["tipo_partida"] or r["giro"])
        d["dep_bloqueada"] = dep_st["bloqueada"]
        d["dep_estado"] = dep_st["estado"]
        d["dep_detalle"] = dep_st["detalle"]
        salida.append(d)
    return jsonify(salida)


@app.route("/api/actividad/<int:aid>", methods=["GET"])
@requiere_login
def api_actividad(aid):
    db = get_db()
    r = db.execute("SELECT * FROM actividades WHERE id=?", (aid,)).fetchone()
    if not r:
        return jsonify({"error": "no existe"}), 404
    return jsonify(dict(r))


@app.route("/api/actividad/<int:aid>/historial")
@requiere_login
def api_actividad_historial(aid):
    db = get_db()
    a = db.execute("SELECT codigo,partida,area,bloque FROM actividades WHERE id=?", (aid,)).fetchone()
    rows = db.execute(
        "SELECT campo,valor_antes,valor_despues,fecha,quien FROM historial "
        "WHERE actividad_id=? ORDER BY fecha DESC, id DESC", (aid,)).fetchall()
    return jsonify({
        "actividad": dict(a) if a else None,
        "historial": [dict(r) for r in rows],
    })


@app.route("/api/actividad/<int:aid>", methods=["PUT"])
@requiere_gestor
def api_actualizar(aid):
    db = get_db()
    data = request.get_json()
    actual = db.execute("SELECT * FROM actividades WHERE id=?", (aid,)).fetchone()
    if not actual:
        return jsonify({"error": "no existe"}), 404
    campos = ["proveedor", "partida", "giro", "aplica", "avance", "f_inicio",
              "f_fin", "duracion_dias", "estatus", "depende_de", "notas", "tipo",
              "area", "bloque", "tipo_partida", "definido",
              "causa_retraso", "nota_proveedor",
              "mundo", "departamento", "tipo_interno",
              "requiere_pruebas", "valida_depto"]
    # si mandan causa_retraso, sellamos quién y cuándo
    if "causa_retraso" in data and (data.get("causa_retraso") or "").strip():
        data["causa_por"] = data.get("causa_por") or "Registrado en plataforma"
        data["causa_fecha"] = datetime.date.today().isoformat()
        campos += ["causa_por", "causa_fecha"]
    # candado: el avance nunca debe guardarse fuera de 0-100 (un típo como
    # "150" en vez de "15" antes se guardaba tal cual y descuadraba promedios).
    if "avance" in data:
        try:
            data["avance"] = max(0, min(100, int(data.get("avance") or 0)))
        except (TypeError, ValueError):
            data["avance"] = 0
    updates = []
    args = []
    for c in campos:
        if c in data:
            antes = actual[c]
            despues = data[c]
            if str(antes) != str(despues):
                db.execute(
                    "INSERT INTO historial (actividad_id,campo,valor_antes,valor_despues,fecha) VALUES (?,?,?,?,?)",
                    (aid, c, str(antes), str(despues), datetime.datetime.now().isoformat(timespec="seconds")),
                )
            updates.append(f"{c}=?")
            args.append(despues)
    # auto estatus si mandan avance y no mandan estatus explícito
    if "avance" in data and "estatus" not in data:
        updates.append("estatus=?")
        args.append(estatus_por_avance(int(data["avance"] or 0)))
    # auto-reconocer si admin pone avance > 0 y no estaba reconocida
    if "avance" in data and int(data["avance"] or 0) > 0 and actual["reconocida"] != "SÍ":
        updates.append("reconocida=?")
        args.append("SÍ")
        updates.append("reconocida_por=?")
        args.append("admin (avance)")
        updates.append("reconocida_fecha=?")
        args.append(datetime.date.today().isoformat())
        updates.append("no_reconocida_nota=?")
        args.append(None)
    # Si el gestor corrige el avance y el proveedor tenia un avance DECLARADO distinto,
    # se alinea el declarado. Sin esto, Validar copiaba el declarado viejo sobre el
    # avance corregido y la actividad "se regresaba" (ACT-0430 / ACT-0432).
    if ("avance" in data and actual["avance_decl"] is not None
            and data["avance"] != (actual["avance"] or 0)
            and actual["avance_decl"] != data["avance"]):
        db.execute(
            "INSERT INTO historial (actividad_id,campo,valor_antes,valor_despues,fecha) VALUES (?,?,?,?,?)",
            (aid, "avance_decl", str(actual["avance_decl"]), str(data["avance"]),
             datetime.datetime.now().isoformat(timespec="seconds")))
        updates.append("avance_decl=?")
        args.append(data["avance"])
        updates.append("avance_decl_rechazado=?")
        args.append(0)
    updates.append("actualizado=?")
    args.append(datetime.datetime.now().isoformat(timespec="seconds"))
    args.append(aid)
    db.execute(f"UPDATE actividades SET {','.join(updates)} WHERE id=?", args)
    db.commit()
    crear_validacion_interna_si_aplica(db, aid)
    r = db.execute("SELECT * FROM actividades WHERE id=?", (aid,)).fetchone()
    return jsonify(dict(r))


@app.route("/api/actividad", methods=["POST"])
@requiere_gestor
def api_crear():
    db = get_db()
    data = request.get_json()
    # Si no llega un código ya definido (caso normal: "Nueva actividad" y
    # "Duplicar" nunca lo mandan), se le asigna aquí el siguiente ACT-####
    # disponible. Antes se guardaba lo que mandara el front (nunca nada),
    # así que toda actividad nueva o duplicada se quedaba sin código.
    codigo = (data.get("codigo") or "").strip()
    partes_ta = None
    if proyecto_actual() == "torre_a":
        # Torre A: codigo bloque.area.proveedor.actividad (1.1.1.1); bloque, area y responsable son obligatorios
        partes_ta, error_ta = partes_codigo_ta(db, data)
        if error_ta:
            db.rollback()
            return jsonify({"error": error_ta}), 400
        if not (data.get("partida") or "").strip():
            db.rollback()
            return jsonify({"error": "Escribe la actividad (partida)."}), 400
    elif not codigo:
        codigo = siguiente_codigo(db)
    sql_alta = """INSERT INTO actividades
        (codigo,bloque,area,giro,proveedor,partida,tipo,tipo_partida,aplica,avance,
         f_inicio,f_fin,duracion_dias,estatus,depende_de,definido,causa_retraso,notas,
         mundo,departamento,tipo_interno,requiere_pruebas,valida_depto,actualizado)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)"""
    for intento in range(8):
        if partes_ta:
            codigo = siguiente_codigo_ta(db, *partes_ta)
        try:
            cur = db.execute(sql_alta, (
                codigo, data.get("bloque"), data.get("area"),
                data.get("giro"), data.get("proveedor"), data.get("partida"),
                data.get("tipo"), data.get("tipo_partida", "Construcción"),
                data.get("aplica", "SÍ"),
                int(data.get("avance", 0) or 0), data.get("f_inicio"),
                data.get("f_fin"), data.get("duracion_dias"),
                data.get("estatus", "Pendiente"), data.get("depende_de"),
                data.get("definido", "NO"), data.get("causa_retraso"), data.get("notas"),
                data.get("mundo", "obra"), data.get("departamento"), data.get("tipo_interno"),
                data.get("requiere_pruebas", "SÍ"), data.get("valida_depto"),
                datetime.datetime.now().isoformat(timespec="seconds"),
            ))
            break
        except sqlite3.IntegrityError:
            if not partes_ta or intento == 7:
                raise   # solo en Torre A el codigo es unico; ahi se recalcula y se reintenta
    db.commit()
    crear_validacion_interna_si_aplica(db, cur.lastrowid)
    r = db.execute("SELECT * FROM actividades WHERE id=?", (cur.lastrowid,)).fetchone()
    return jsonify(dict(r))


@app.route("/api/actividad/<int:aid>", methods=["DELETE"])
@requiere_admin
def api_borrar(aid):
    db = get_db()
    db.execute("UPDATE actividades SET eliminada=1, eliminada_por=?, eliminada_fecha=? WHERE id=?",
               (session.get("usuario", "admin"), datetime.date.today().isoformat(), aid))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/actividades/borrar", methods=["POST"])
@requiere_admin
def api_borrar_multiple():
    db = get_db()
    data = request.get_json()
    ids = data.get("ids", [])
    if not ids:
        return jsonify({"error": "sin ids"}), 400
    marcas = ",".join("?" * len(ids))
    quien = session.get("usuario", "admin")
    fecha = datetime.date.today().isoformat()
    db.execute(f"UPDATE actividades SET eliminada=1, eliminada_por=?, eliminada_fecha=? WHERE id IN ({marcas})",
               [quien, fecha] + ids)
    db.commit()
    return jsonify({"ok": True, "borradas": len(ids)})


# ----------------------------------------------------------------------------
@app.route("/api/actividades/eliminadas")
@requiere_admin
def api_eliminadas():
    db = get_db()
    rows = db.execute(
        "SELECT * FROM actividades WHERE eliminada=1 ORDER BY eliminada_fecha DESC"
    ).fetchall()
    return jsonify([dict(r) for r in rows])


@app.route("/api/actividad/<int:aid>/restaurar", methods=["POST"])
@requiere_admin
def api_restaurar(aid):
    db = get_db()
    db.execute("UPDATE actividades SET eliminada=0, eliminada_por=NULL, eliminada_fecha=NULL WHERE id=?", (aid,))
    db.commit()
    return jsonify({"ok": True})


# API — Catálogos y resumen
# ----------------------------------------------------------------------------
@app.route("/api/catalogos")
@requiere_login
def api_catalogos():
    db = get_db()
    def distintos(col):
        rows = db.execute(
            f"SELECT DISTINCT {col} FROM actividades WHERE {col} IS NOT NULL AND {col}<>'' ORDER BY {col}"
        ).fetchall()
        return [r[0] for r in rows]
    # departamentos = los del catálogo de involucrados + los que ya tienen actividades internas
    invol = [r[0] for r in db.execute("SELECT nombre FROM involucrados ORDER BY nombre").fetchall()]
    dept_act = distintos("departamento")
    departamentos = sorted(set(invol) | set(dept_act))
    # mapa bloque -> áreas (para el candado estricto en el formulario)
    mapa = {}
    for r in db.execute(
        "SELECT DISTINCT bloque, area FROM actividades "
        "WHERE bloque IS NOT NULL AND bloque<>'' AND area IS NOT NULL AND area<>'' "
        "ORDER BY bloque, area").fetchall():
        mapa.setdefault(r["bloque"], []).append(r["area"])
    bloques_l, areas_l, giros_l = distintos("bloque"), distintos("area"), distintos("giro")
    proveedores_l = distintos("proveedor")
    if proyecto_actual() == "torre_a":
        # Torre A: los desplegables salen del catalogo fijo (20 bloques / 85 areas), de los
        # proveedores numerados y de las especialidades de Torre A.
        bloques_l, areas_l, mapa = catalogo_ta_para_pantalla(db)
        giros_l = sorted(set(giros_l) | {r[0] for r in db.execute("SELECT nombre FROM ta_especialidades")})
        proveedores_l = [r[0] for r in db.execute("SELECT nombre FROM ta_proveedores WHERE interno=0 ORDER BY num")]
    return jsonify({
        "bloques": bloques_l,
        "areas": areas_l,
        "giros": giros_l,
        "proveedores": proveedores_l,
        "departamentos": departamentos,
        "mapa_bloque_areas": mapa,
        "tipos_partida": ["Construcción", "Mobiliario y equipo", "Puesta en marcha", "Detalles finales"],
        "estatus": ["Pendiente", "En proceso", "Listo", "Post-apertura"],
    })


# ----------------------------------------------------------------------------
# API — Catálogo de proveedores (con tipo interno/externo y función)
# ----------------------------------------------------------------------------
@app.route("/api/proveedores", methods=["GET"])
@requiere_login
def api_proveedores():
    db = get_db()
    rows = db.execute("SELECT * FROM proveedores ORDER BY nombre").fetchall()
    return jsonify([dict(r) for r in rows])


@app.route("/api/proveedores", methods=["POST"])
@requiere_gestor
def api_crear_proveedor():
    db = get_db()
    data = request.get_json()
    nombre = (data.get("nombre") or "").strip()
    if not nombre:
        return jsonify({"error": "nombre requerido"}), 400
    if proyecto_actual() == "torre_a":
        numero_proveedor_ta(db, nombre, data.get("tipo") == "Interno")   # toma su numero fijo desde el alta
    try:
        db.execute(
            "INSERT INTO proveedores (nombre,tipo,funcion) VALUES (?,?,?)",
            (nombre, data.get("tipo", "Externo"), data.get("funcion", "")),
        )
        db.commit()
    except sqlite3.IntegrityError:
        # ya existe: actualiza tipo y función
        db.execute(
            "UPDATE proveedores SET tipo=?, funcion=? WHERE nombre=?",
            (data.get("tipo", "Externo"), data.get("funcion", ""), nombre),
        )
        db.commit()
    r = db.execute("SELECT * FROM proveedores WHERE nombre=?", (nombre,)).fetchone()
    return jsonify(dict(r))


# ----------------------------------------------------------------------------
# API — Catálogo genérico (bloque / area / giro) con nota
# ----------------------------------------------------------------------------
@app.route("/api/proveedores/<nombre>/activo", methods=["POST"])
@requiere_admin
def api_proveedor_activo(nombre):
    """Activa o desactiva un proveedor del catalogo (nunca se borra, para
    conservar el historial de sus actividades y reportes)."""
    db = get_db()
    data = request.get_json() or {}
    activo = 1 if data.get("activo") else 0
    existe = db.execute("SELECT id FROM proveedores WHERE nombre=?", (nombre,)).fetchone()
    if not existe:
        db.execute("INSERT INTO proveedores (nombre, tipo, activo) VALUES (?,?,?)",
                   (nombre, "Externo", activo))
    else:
        db.execute("UPDATE proveedores SET activo=? WHERE nombre=?", (activo, nombre))
    db.commit()
    return jsonify({"ok": True, "nombre": nombre, "activo": activo})


def alta_catalogo_ta(db, clase, nombre, data):
    """'+ Agregar bloque / area / especialidad' en Torre A: bloque y area toman el siguiente numero."""
    nombre = " ".join(nombre.split())[:120]
    if clase == "giro":
        db.execute("INSERT OR IGNORE INTO ta_especialidades (nombre) VALUES (?)", (nombre,))
        db.commit()
        return jsonify({"ok": True, "clase": clase, "nombre": nombre, "etiqueta": nombre})
    if clase == "bloque":
        for r in db.execute("SELECT num,nombre FROM ta_bloques").fetchall():
            if sin_acentos(r["nombre"]) == sin_acentos(nombre):    # ya existe: no se duplica
                return jsonify({"ok": True, "clase": clase, "nombre": r["nombre"], "etiqueta": etiqueta_bloque_ta(r["num"], r["nombre"]), "existia": True})
        num = db.execute("SELECT COALESCE(MAX(num),0)+1 FROM ta_bloques").fetchone()[0]
        db.execute("INSERT INTO ta_bloques (num,nombre) VALUES (?,?)", (num, nombre))
        db.commit()
        return jsonify({"ok": True, "clase": clase, "nombre": nombre, "etiqueta": etiqueta_bloque_ta(num, nombre)})
    mb = _RE_BLOQUE_TA.match(data.get("bloque") or "")
    if not mb or not db.execute("SELECT 1 FROM ta_bloques WHERE num=?", (int(mb.group(1)),)).fetchone():
        return jsonify({"error": "Primero elige el bloque en el que va el área nueva."}), 400
    b = int(mb.group(1))
    for r in db.execute("SELECT num,nombre FROM ta_areas WHERE bloque=?", (b,)).fetchall():
        if sin_acentos(r["nombre"]) == sin_acentos(nombre):         # ya existe en ese bloque: no se duplica
            return jsonify({"ok": True, "clase": clase, "nombre": r["nombre"], "etiqueta": etiqueta_area_ta(b, r["num"], r["nombre"]), "existia": True})
    num = db.execute("SELECT COALESCE(MAX(num),0)+1 FROM ta_areas WHERE bloque=?", (b,)).fetchone()[0]
    db.execute("INSERT INTO ta_areas (bloque,num,nombre) VALUES (?,?,?)", (b, num, nombre))
    db.commit()
    return jsonify({"ok": True, "clase": clase, "nombre": nombre, "etiqueta": etiqueta_area_ta(b, num, nombre)})


@app.route("/api/catalogo/<clase>", methods=["POST"])
@requiere_gestor
def api_crear_catalogo(clase):
    if clase not in ("bloque", "area", "giro"):
        return jsonify({"error": "clase inválida"}), 400
    db = get_db()
    data = request.get_json()
    nombre = (data.get("nombre") or "").strip()
    if not nombre:
        return jsonify({"error": "nombre requerido"}), 400
    if proyecto_actual() == "torre_a":
        return alta_catalogo_ta(db, clase, nombre, data)
    try:
        db.execute(
            "INSERT INTO catalogo (clase,nombre,nota) VALUES (?,?,?)",
            (clase, nombre, data.get("nota", "")),
        )
        db.commit()
    except sqlite3.IntegrityError:
        db.execute(
            "UPDATE catalogo SET nota=? WHERE clase=? AND nombre=?",
            (data.get("nota", ""), clase, nombre),
        )
        db.commit()
    return jsonify({"ok": True, "clase": clase, "nombre": nombre})


# ----------------------------------------------------------------------------
# API — Causas de retraso (catálogo: 7 base + agregadas por el usuario)
# ----------------------------------------------------------------------------
@app.route("/api/causas", methods=["GET"])
@requiere_login
def api_causas():
    db = get_db()
    mundo = request.args.get("mundo", "obra")
    rows = db.execute(
        "SELECT nombre FROM causas WHERE (mundo = ? OR (mundo IS NULL AND ? = 'obra')) "
        "ORDER BY base DESC, nombre", (mundo, mundo)).fetchall()
    return jsonify([r[0] for r in rows])


@app.route("/api/causas", methods=["POST"])
@requiere_gestor
def api_crear_causa():
    db = get_db()
    data = request.get_json()
    nombre = (data.get("nombre") or "").strip()
    mundo = data.get("mundo", "obra")
    if not nombre:
        return jsonify({"error": "nombre requerido"}), 400
    db.execute("INSERT OR IGNORE INTO causas (nombre,base,mundo) VALUES (?,0,?)", (nombre, mundo))
    db.commit()
    rows = db.execute(
        "SELECT nombre FROM causas WHERE (mundo = ? OR (mundo IS NULL AND ? = 'obra')) "
        "ORDER BY base DESC, nombre", (mundo, mundo)).fetchall()
    return jsonify([r[0] for r in rows])


# ----------------------------------------------------------------------------
# API — Resumen de causas AGREGADAS (para el reporte a dirección, sin señalar)
# ----------------------------------------------------------------------------
@app.route("/api/resumen_causas")
@requiere_login
def api_resumen_causas():
    db = get_db()
    rows = db.execute(
        """SELECT causa_retraso, COUNT(*) n
           FROM actividades
           WHERE causa_retraso IS NOT NULL AND causa_retraso<>''
             AND (aplica IS NULL OR aplica<>'NO') AND (eliminada IS NULL OR eliminada=0)
           GROUP BY causa_retraso ORDER BY n DESC"""
    ).fetchall()
    total = sum(r["n"] for r in rows)
    salida = [{"causa": r["causa_retraso"], "partidas": r["n"],
               "porcentaje": round(100 * r["n"] / total, 1) if total else 0}
              for r in rows]
    return jsonify({"total_con_causa": total, "causas": salida})


# ----------------------------------------------------------------------------
# API — Expediente de proveedores (datos de empresa y contacto)
# ----------------------------------------------------------------------------
@app.route("/api/proveedores_usuarios")
@requiere_admin
def api_proveedores_usuarios():
    """Vista unificada: cada proveedor/departamento con su ficha (expediente)
    y su(s) acceso(s) a la plataforma, si ya tiene. No incluye admin/supervisores
    (esos son 'accesos generales', no estan ligados a un proveedor concreto)."""
    db = get_db()
    nombres = set()
    for r in db.execute("SELECT DISTINCT proveedor FROM actividades WHERE proveedor IS NOT NULL AND proveedor<>''"):
        nombres.add(r[0])
    for r in db.execute("SELECT DISTINCT departamento FROM actividades WHERE departamento IS NOT NULL AND departamento<>''"):
        nombres.add(r[0])
    fichas = {r["nombre"]: dict(r) for r in db.execute("SELECT * FROM proveedores").fetchall()}
    nombres |= set(fichas.keys())
    usuarios_por_prov = {}
    for r in db.execute("SELECT * FROM usuarios WHERE rol='proveedor' AND proveedor IS NOT NULL AND proveedor<>''"):
        usuarios_por_prov.setdefault(r["proveedor"], []).append(dict(r))
    nombres |= set(usuarios_por_prov.keys())

    salida = []
    for nombre in sorted(nombres, key=lambda x: x.lower()):
        f = fichas.get(nombre, {})
        stats = db.execute(
            "SELECT COUNT(*) n, ROUND(AVG(avance),1) av FROM actividades WHERE (proveedor=? OR departamento=?) AND (eliminada IS NULL OR eliminada=0)",
            (nombre, nombre)).fetchone()
        usus = usuarios_por_prov.get(nombre, [])
        salida.append({
            "nombre": nombre,
            "activo": 1 if f.get("activo") is None else f.get("activo"),
            "empresa": f.get("empresa") or "", "tipo": f.get("tipo") or "Externo",
            "funcion": f.get("funcion") or "", "contacto": f.get("contacto") or "",
            "telefono": f.get("telefono") or "", "correo": f.get("correo") or "",
            "notas": f.get("notas") or "",
            "partidas": stats["n"] or 0, "avance": stats["av"] or 0,
            "usuarios": [{
                "id": u["id"], "usuario": u["usuario"],
                "activo": 1 if u.get("activo") is None else u.get("activo"),
                "num_logins": u.get("num_logins") or 0,
                "ultimo_login": u.get("ultimo_login"),
                "telefono": u.get("telefono"), "mundo": u.get("mundo") or "obra",
            } for u in usus],
        })
    return jsonify(salida)


@app.route("/api/expediente")
@requiere_admin
def api_expediente():
    """Lista proveedores usados en actividades + su ficha (si existe)."""
    db = get_db()
    # proveedores que aparecen en actividades
    usados = db.execute(
        "SELECT proveedor, COUNT(*) n, ROUND(AVG(avance),1) av FROM actividades "
        "WHERE proveedor IS NOT NULL AND proveedor<>'' AND (eliminada IS NULL OR eliminada=0) GROUP BY proveedor ORDER BY proveedor"
    ).fetchall()
    fichas = {r["nombre"]: dict(r) for r in db.execute("SELECT * FROM proveedores").fetchall()}
    salida = []
    for u in usados:
        f = fichas.get(u["proveedor"], {})
        salida.append({
            "nombre": u["proveedor"], "partidas": u["n"], "avance": u["av"],
            "empresa": f.get("empresa", ""), "tipo": f.get("tipo", ""),
            "funcion": f.get("funcion", ""), "contacto": f.get("contacto", ""),
            "telefono": f.get("telefono", ""), "correo": f.get("correo", ""),
            "notas": f.get("notas", ""),
        })
    return jsonify(salida)


@app.route("/api/expediente", methods=["POST"])
@requiere_admin
def api_guardar_expediente():
    db = get_db()
    data = request.get_json()
    nombre = (data.get("nombre") or "").strip()
    if not nombre:
        return jsonify({"error": "nombre requerido"}), 400
    existe = db.execute("SELECT id FROM proveedores WHERE nombre=?", (nombre,)).fetchone()
    campos = ("empresa", "tipo", "funcion", "contacto", "telefono", "correo", "notas")
    if existe:
        sets = ",".join(f"{c}=?" for c in campos)
        db.execute(f"UPDATE proveedores SET {sets} WHERE nombre=?",
                   [data.get(c, "") for c in campos] + [nombre])
    else:
        db.execute(
            "INSERT INTO proveedores (nombre,empresa,tipo,funcion,contacto,telefono,correo,notas) "
            "VALUES (?,?,?,?,?,?,?,?)",
            [nombre] + [data.get(c, "") for c in campos])
    db.commit()
    return jsonify({"ok": True})


# ----------------------------------------------------------------------------
# API — Snapshots semanales (comparar semana vs semana)
# ----------------------------------------------------------------------------
@app.route("/api/snapshot", methods=["POST"])
@requiere_gestor
def api_snapshot():
    """Guarda una foto del avance de todas las partidas hoy."""
    db = get_db()
    hoy_d = datetime.date.today()
    semana = hoy_d.strftime("%G-W%V")  # año-semana ISO
    fecha = hoy_d.isoformat()
    # borra foto previa de la misma semana para no duplicar
    db.execute("DELETE FROM snapshots WHERE semana=?", (semana,))
    rows = db.execute("SELECT id, avance FROM actividades WHERE (eliminada IS NULL OR eliminada=0)").fetchall()
    for r in rows:
        db.execute(
            "INSERT INTO snapshots (semana,fecha,actividad_id,avance) VALUES (?,?,?,?)",
            (semana, fecha, r["id"], r["avance"] or 0))
    db.commit()
    return jsonify({"ok": True, "semana": semana, "partidas": len(rows)})


@app.route("/api/comparar_semanas")
@requiere_login
def api_comparar_semanas():
    """Compara el avance actual contra la última foto guardada."""
    db = get_db()
    semanas = db.execute(
        "SELECT DISTINCT semana FROM snapshots ORDER BY semana DESC").fetchall()
    if not semanas:
        return jsonify({"hay_foto": False})
    ultima = semanas[0]["semana"]
    # avance guardado en la última foto, por proveedor
    prev = db.execute(
        """SELECT a.proveedor prov, AVG(s.avance) av
           FROM snapshots s JOIN actividades a ON a.id=s.actividad_id
           WHERE s.semana=? GROUP BY a.proveedor""", (ultima,)).fetchall()
    prev_map = {r["prov"]: r["av"] for r in prev}
    # avance actual por proveedor
    act = db.execute(
        "SELECT proveedor prov, AVG(avance) av FROM actividades WHERE (eliminada IS NULL OR eliminada=0) GROUP BY proveedor").fetchall()
    salida = []
    for r in act:
        p = r["prov"] or "— Sin proveedor —"
        antes = prev_map.get(r["prov"], 0) or 0
        ahora = r["av"] or 0
        salida.append({
            "proveedor": p, "antes": round(antes, 1), "ahora": round(ahora, 1),
            "cambio": round(ahora - antes, 1),
        })
    salida.sort(key=lambda x: x["cambio"])
    return jsonify({"hay_foto": True, "semana_previa": ultima, "proveedores": salida})



# ============================================================
#  HISTÓRICO — consultar cómo estaba la obra en un día pasado
#  v1.5 · Reconstruye el pasado desde las fotos semanales y,
#  si no hay foto, deshaciendo los cambios del historial.
# ============================================================

def _fotos_disponibles(db):
    """Fechas en las que hay foto guardada, de la más nueva a la más vieja."""
    return [dict(r) for r in db.execute(
        "SELECT semana, MIN(fecha) fecha, COUNT(*) partidas "
        "FROM snapshots GROUP BY semana ORDER BY fecha DESC")]


def _avance_a_fecha(db, fecha, mundo="obra"):
    """Devuelve {actividad_id: avance} tal como estaba ESE DÍA.
    1) Si hay foto de ese día o anterior, la usa.
    2) Si no, reconstruye: parte del avance de hoy y deshace
       los cambios que ocurrieron DESPUÉS de la fecha pedida."""
    fila = db.execute(
        "SELECT semana, fecha FROM snapshots WHERE fecha<=? "
        "ORDER BY fecha DESC LIMIT 1", (fecha,)).fetchone()
    origen = "reconstruido"
    mapa = {}
    if fila:
        origen = "foto"
        for r in db.execute("SELECT actividad_id, avance FROM snapshots WHERE semana=?",
                            (fila["semana"],)):
            mapa[r["actividad_id"]] = r["avance"] or 0
    # actividades vigentes del mundo pedido
    acts = db.execute(
        "SELECT id, avance FROM actividades "
        "WHERE (aplica IS NULL OR aplica<>'NO') "
        "AND (estado_val IS NULL OR estado_val='validado') "
        "AND (eliminada IS NULL OR eliminada=0) "
        "AND (mundo = ? OR (mundo IS NULL AND ? = 'obra'))",
        (mundo, mundo)).fetchall()
    for a in acts:
        if a["id"] in mapa:
            continue
        # reconstruir desde el historial: el primer cambio POSTERIOR a la
        # fecha nos dice cuánto valía ANTES de ese cambio
        h = db.execute(
            "SELECT valor_antes FROM historial "
            "WHERE actividad_id=? AND (campo LIKE '%avance%') AND fecha> ? "
            "ORDER BY fecha ASC LIMIT 1", (a["id"], fecha)).fetchone()
        if h and h["valor_antes"] not in (None, "", "None"):
            try:
                mapa[a["id"]] = int(float(h["valor_antes"]))
            except (ValueError, TypeError):
                mapa[a["id"]] = a["avance"] or 0
        else:
            # no hubo cambios después de esa fecha: sigue igual que hoy
            mapa[a["id"]] = a["avance"] or 0
    return mapa, origen, (fila["fecha"] if fila else None)


@app.route("/api/historico/fechas")
@requiere_gestor
def api_historico_fechas():
    db = get_db()
    fotos = _fotos_disponibles(db)
    rango = db.execute(
        "SELECT MIN(substr(fecha,1,10)) ini, MAX(substr(fecha,1,10)) fin FROM historial").fetchone()
    return jsonify({
        "fotos": fotos,
        "historial_desde": rango["ini"] if rango else None,
        "historial_hasta": rango["fin"] if rango else None,
        "hoy": datetime.date.today().isoformat(),
    })


@app.route("/api/historico")
@requiere_gestor
def api_historico():
    """Cómo estaba la obra en una fecha dada."""
    fecha = request.args.get("fecha") or datetime.date.today().isoformat()
    mundo = request.args.get("mundo", "obra")
    db = get_db()
    mapa, origen, fecha_foto = _avance_a_fecha(db, fecha, mundo)
    filas = db.execute(
        "SELECT id, codigo, bloque, area, giro, proveedor, departamento, mundo, partida, avance "
        "FROM actividades WHERE (aplica IS NULL OR aplica<>'NO') "
        "AND (estado_val IS NULL OR estado_val='validado') "
        "AND (eliminada IS NULL OR eliminada=0) "
        "AND (mundo = ? OR (mundo IS NULL AND ? = 'obra'))",
        (mundo, mundo)).fetchall()
    if not filas:
        return jsonify({"total": 0, "fecha": fecha, "origen": origen})
    por_bloque, por_resp, detalle = {}, {}, []
    suma = 0
    for f in filas:
        ent = mapa.get(f["id"], 0)
        hoy_av = f["avance"] or 0
        suma += ent
        b = f["bloque"] or "— Sin bloque —"
        quien = (f["departamento"] if f["mundo"] == "interno" else f["proveedor"]) or "— Sin asignar —"
        por_bloque.setdefault(b, []).append(ent)
        por_resp.setdefault(quien, []).append((ent, hoy_av))
        detalle.append({
            "codigo": f["codigo"], "bloque": b, "area": f["area"],
            "responsable": quien, "partida": f["partida"],
            "avance_fecha": ent, "avance_hoy": hoy_av, "cambio": hoy_av - ent,
        })
    n = len(filas)
    bloques = [{"bloque": k, "total": len(v), "avance": round(sum(v) / len(v), 1)}
               for k, v in por_bloque.items()]
    bloques.sort(key=lambda x: -x["avance"])
    resp = []
    for k, v in por_resp.items():
        ent = sum(x[0] for x in v) / len(v)
        hoy_ = sum(x[1] for x in v) / len(v)
        resp.append({"responsable": k, "total": len(v), "avance": round(ent, 1),
                     "avance_hoy": round(hoy_, 1), "cambio": round(hoy_ - ent, 1)})
    resp.sort(key=lambda x: -x["total"])
    detalle.sort(key=lambda x: -abs(x["cambio"]))
    return jsonify({
        "fecha": fecha, "origen": origen, "fecha_foto": fecha_foto,
        "total": n, "avance_global": round(suma / n, 1),
        "avance_hoy": round(sum((f["avance"] or 0) for f in filas) / n, 1),
        "bloques": bloques, "responsables": resp, "detalle": detalle[:400],
    })


@app.route("/api/historico/comparar")
@requiere_gestor
def api_historico_comparar():
    """Cuánto se movió cada responsable entre dos fechas."""
    desde = request.args.get("desde")
    hasta = request.args.get("hasta") or datetime.date.today().isoformat()
    mundo = request.args.get("mundo", "obra")
    if not desde:
        return jsonify({"error": "Falta la fecha inicial"}), 400
    db = get_db()
    m1, o1, _ = _avance_a_fecha(db, desde, mundo)
    m2, o2, _ = _avance_a_fecha(db, hasta, mundo)
    filas = db.execute(
        "SELECT id, bloque, proveedor, departamento, mundo FROM actividades "
        "WHERE (aplica IS NULL OR aplica<>'NO') "
        "AND (estado_val IS NULL OR estado_val='validado') "
        "AND (eliminada IS NULL OR eliminada=0) "
        "AND (mundo = ? OR (mundo IS NULL AND ? = 'obra'))",
        (mundo, mundo)).fetchall()
    if not filas:
        return jsonify({"total": 0})
    por_resp, por_bloque = {}, {}
    for f in filas:
        quien = (f["departamento"] if f["mundo"] == "interno" else f["proveedor"]) or "— Sin asignar —"
        a1, a2 = m1.get(f["id"], 0), m2.get(f["id"], 0)
        por_resp.setdefault(quien, []).append((a1, a2))
        por_bloque.setdefault(f["bloque"] or "— Sin bloque —", []).append((a1, a2))
    def arma(d, etiqueta):
        out = []
        for k, v in d.items():
            p1 = sum(x[0] for x in v) / len(v)
            p2 = sum(x[1] for x in v) / len(v)
            out.append({etiqueta: k, "total": len(v), "antes": round(p1, 1),
                        "despues": round(p2, 1), "cambio": round(p2 - p1, 1)})
        out.sort(key=lambda x: x["cambio"])
        return out
    g1 = sum(m1.get(f["id"], 0) for f in filas) / len(filas)
    g2 = sum(m2.get(f["id"], 0) for f in filas) / len(filas)
    return jsonify({
        "desde": desde, "hasta": hasta, "origen_desde": o1, "origen_hasta": o2,
        "global_antes": round(g1, 1), "global_despues": round(g2, 1),
        "global_cambio": round(g2 - g1, 1),
        "responsables": arma(por_resp, "responsable"),
        "bloques": arma(por_bloque, "bloque"),
    })


def _foto_automatica(ruta=None):
    """Guarda la foto de la semana si todavía no existe. Se llama al arrancar,
    para que el histórico no dependa de que alguien se acuerde de picarle."""
    try:
        con = sqlite3.connect(ruta or DB_PATH)
        con.row_factory = sqlite3.Row
        hoy_d = datetime.date.today()
        semana = hoy_d.strftime("%G-W%V")
        ya = con.execute("SELECT COUNT(*) FROM snapshots WHERE semana=?", (semana,)).fetchone()[0]
        if ya == 0:
            filas = con.execute("SELECT id, avance FROM actividades").fetchall()
            for r in filas:
                con.execute(
                    "INSERT INTO snapshots (semana,fecha,actividad_id,avance) VALUES (?,?,?,?)",
                    (semana, hoy_d.isoformat(), r["id"], r["avance"] or 0))
            con.commit()
            print(f"  Foto semanal automatica guardada ({semana}, {len(filas)} partidas)")
        con.close()
    except Exception as e:
        print("  (no se pudo guardar la foto automatica:", e, ")")


@app.route("/historico")
def historico_page():
    return render_template("historico.html")


def _resumen_data():
    db = get_db()
    mundo = request.args.get("mundo", "obra")
    rows = db.execute(
        "SELECT * FROM actividades WHERE (aplica IS NULL OR aplica<>'NO') "
        "AND (estado_val IS NULL OR estado_val='validado') "
        "AND (eliminada IS NULL OR eliminada=0) "
        "AND (mundo = ? OR (mundo IS NULL AND ? = 'obra'))",
        (mundo, mundo)
    ).fetchall()
    total = len(rows)
    if total == 0:
        # sin actividades (Torre A al inicio): todo en cero, pero con la fecha de entrega para el contador
        return {"total": 0, "avance_global": 0, "estatus": {}, "bloques": [], "giros": [], "tipos": [],
                "retrasadas": [], "en_riesgo": [], "dias_restantes": dias_restantes(),
                "fecha_entrega": fecha_entrega_actual(), "fecha_hoy": datetime.date.today().isoformat()}
    suma_av = sum((r["avance"] or 0) for r in rows)
    avance_global = round(suma_av / total, 1)

    # por estatus
    est = {}
    for r in rows:
        e = r["estatus"] or "Pendiente"
        est[e] = est.get(e, 0) + 1

    # por bloque
    bloques = {}
    for r in rows:
        b = r["bloque"] or "—"
        bloques.setdefault(b, {"total": 0, "suma": 0})
        bloques[b]["total"] += 1
        bloques[b]["suma"] += (r["avance"] or 0)
    bloques_out = []
    for b, d in bloques.items():
        bloques_out.append({
            "bloque": b, "total": d["total"],
            "avance": round(d["suma"] / d["total"], 1),
        })
    bloques_out.sort(key=lambda x: x["avance"])

    # por giro
    giros = {}
    for r in rows:
        gg = r["giro"] or "—"
        giros.setdefault(gg, {"total": 0, "suma": 0})
        giros[gg]["total"] += 1
        giros[gg]["suma"] += (r["avance"] or 0)
    giros_out = [{"giro": k, "total": v["total"], "avance": round(v["suma"]/v["total"],1)} for k, v in giros.items()]
    giros_out.sort(key=lambda x: x["avance"])

    # por tipo de partida (Construcción, Mobiliario y equipo, etc.)
    tipos = {}
    for r in rows:
        tp = (r["tipo_partida"] if "tipo_partida" in r.keys() else None) or "Construcción"
        tipos.setdefault(tp, {"total": 0, "suma": 0})
        tipos[tp]["total"] += 1
        tipos[tp]["suma"] += (r["avance"] or 0)
    tipos_out = [{"tipo": k, "total": v["total"], "avance": round(v["suma"]/v["total"],1)} for k, v in tipos.items()]
    tipos_out.sort(key=lambda x: x["avance"])

    # retrasos: fecha fin pasada y avance < 100
    hoy_d = datetime.date.today()
    retrasadas = []
    en_riesgo = []
    for r in rows:
        fin = parse_date(r["f_fin"])
        av = r["avance"] or 0
        if fin and av < 100:
            dif = (fin - hoy_d).days
            item = {
                "id": r["id"], "area": r["area"], "partida": r["partida"],
                "proveedor": r["proveedor"], "giro": r["giro"],
                "f_fin": r["f_fin"], "avance": av, "dias": dif,
            }
            if dif < 0:
                retrasadas.append(item)
            elif dif <= 7:
                en_riesgo.append(item)
    retrasadas.sort(key=lambda x: x["dias"])
    en_riesgo.sort(key=lambda x: x["dias"])

    return {
        "total": total,
        "avance_global": avance_global,
        "estatus": est,
        "bloques": bloques_out,
        "giros": giros_out,
        "tipos": tipos_out,
        "retrasadas": retrasadas,
        "en_riesgo": en_riesgo,
        "dias_restantes": dias_restantes(),
        "fecha_entrega": fecha_entrega_actual(),
        "fecha_hoy": hoy_d.isoformat(),
    }


@app.route("/api/resumen")
def api_resumen():
    return jsonify(_resumen_data())


def _diagnostico_conteos_data():
    """Explica por qué la pantalla principal (KPI 'Actividades') y la pantalla
    de Seguimiento pueden mostrar un total distinto para el mismo mundo.

    La pantalla principal (/api/resumen) solo cuenta actividades que:
      (a) SÍ aplican (aplica <> 'NO'), y
      (b) ya están validadas (estado_val = 'validado' o NULL) — es decir,
          EXCLUYE las 'propuestas' que un proveedor/departamento capturó y
          que todavía esperan que el admin las apruebe en /validacion.

    La pantalla de Seguimiento (/api/seguimiento) cuenta TODAS las
    actividades asignadas a cada proveedor/departamento (sin filtrar por
    aplica ni por estado_val), pero SOLO las que tienen un responsable
    (proveedor o departamento) asignado; una actividad sin responsable no
    aparece en ningún acordeón de Seguimiento aunque sí la cuente la
    pantalla principal.

    Por eso el número no siempre coincide: no es necesariamente que "el
    proveedor no ha aceptado el trabajo" (eso es el campo 'reconocida', que
    ninguna de las dos pantallas usa para el total), sino la mezcla de estas
    dos reglas de conteo distintas."""
    db = get_db()
    resultado = {}
    for mundo, etiqueta, col in (("obra", "Obra (externos)", "proveedor"),
                                  ("interno", "Interno (departamentos)", "departamento")):
        filas = db.execute(
            "SELECT aplica, estado_val, proveedor, departamento, reconocida FROM actividades "
            "WHERE (eliminada IS NULL OR eliminada=0) AND (mundo=? OR (mundo IS NULL AND ?='obra'))",
            (mundo, mundo)
        ).fetchall()
        total = len(filas)
        cuenta_principal = cuenta_seguimiento = 0
        no_aplica = propuesta_sin_validar = sin_responsable = sin_reconocer = 0
        solo_en_seguimiento = solo_en_principal = 0
        for f in filas:
            es_no_aplica = (f["aplica"] == "NO")
            es_propuesta = (f["estado_val"] == "propuesta")
            resp = f[col]
            tiene_resp = bool(resp and resp.strip())
            en_principal = (not es_no_aplica) and (not es_propuesta)
            en_segui = tiene_resp
            if es_no_aplica:
                no_aplica += 1
            if es_propuesta:
                propuesta_sin_validar += 1
            if not tiene_resp:
                sin_responsable += 1
            if f["reconocida"] != "SÍ":
                sin_reconocer += 1
            if en_principal:
                cuenta_principal += 1
            if en_segui:
                cuenta_seguimiento += 1
            if en_segui and not en_principal:
                solo_en_seguimiento += 1
            if en_principal and not en_segui:
                solo_en_principal += 1
        resultado[mundo] = {
            "etiqueta": etiqueta, "total": total,
            "cuenta_pantalla_principal": cuenta_principal,
            "cuenta_seguimiento": cuenta_seguimiento,
            "diferencia": cuenta_seguimiento - cuenta_principal,
            "no_aplica": no_aplica,
            "propuesta_sin_validar": propuesta_sin_validar,
            "sin_responsable_asignado": sin_responsable,
            "sin_reconocer": sin_reconocer,
            "solo_en_seguimiento": solo_en_seguimiento,
            "solo_en_principal": solo_en_principal,
        }
    return resultado


@app.route("/diagnostico")
def diagnostico_page():
    return render_template("diagnostico.html")


@app.route("/api/diagnostico_conteos")
@requiere_gestor
def api_diagnostico_conteos():
    return jsonify(_diagnostico_conteos_data())


# ----------------------------------------------------------------------------
# API — Resumen por proveedor (para el centro de reportes)
# ----------------------------------------------------------------------------
@app.route("/api/resumen_proveedores")
@requiere_login
def api_resumen_proveedores():
    db = get_db()
    rows = db.execute(
        "SELECT * FROM actividades WHERE (aplica IS NULL OR aplica<>'NO') AND (eliminada IS NULL OR eliminada=0)"
    ).fetchall()
    hoy_d = datetime.date.today()
    prov = {}
    for r in rows:
        p = r["proveedor"] or "— Sin proveedor —"
        prov.setdefault(p, {"total": 0, "suma": 0, "listas": 0, "retrasadas": 0})
        prov[p]["total"] += 1
        prov[p]["suma"] += (r["avance"] or 0)
        if (r["avance"] or 0) >= 100:
            prov[p]["listas"] += 1
        fin = parse_date(r["f_fin"])
        if fin and (r["avance"] or 0) < 100 and (fin - hoy_d).days < 0:
            prov[p]["retrasadas"] += 1
    salida = []
    for p, d in prov.items():
        salida.append({
            "proveedor": p, "total": d["total"],
            "avance": round(d["suma"] / d["total"], 1) if d["total"] else 0,
            "listas": d["listas"], "retrasadas": d["retrasadas"],
        })
    salida.sort(key=lambda x: x["avance"])
    return jsonify(salida)


def _filtrar_actividades(args):
    """Devuelve filas filtradas según los mismos parámetros de la tabla."""
    db = get_db()
    q = "SELECT * FROM actividades"
    cond, a = ["(eliminada IS NULL OR eliminada=0)"], []
    for campo in ("bloque", "area", "giro", "proveedor", "tipo_partida", "estatus"):
        v = args.get(campo)
        if v:
            cond.append(f"{campo} = ?"); a.append(v)
    if args.get("solo_retraso") == "1":
        pass  # se filtra abajo con fecha
    if cond:
        q += " WHERE " + " AND ".join(cond)
    q += " ORDER BY area, id"
    rows = db.execute(q, a).fetchall()
    if args.get("solo_retraso") == "1":
        hoy_d = datetime.date.today()
        rows = [r for r in rows if parse_date(r["f_fin"]) and (r["avance"] or 0) < 100
                and (parse_date(r["f_fin"]) - hoy_d).days < 0]
    av_min = args.get("avance_min")
    if av_min not in (None, ""):
        try:
            av_min = max(0, min(100, int(av_min)))
            rows = [r for r in rows if (r["avance"] or 0) >= av_min]
        except (TypeError, ValueError):
            pass
    av_max = args.get("avance_max")
    if av_max not in (None, ""):
        try:
            av_max = max(0, min(100, int(av_max)))
            rows = [r for r in rows if (r["avance"] or 0) <= av_max]
        except (TypeError, ValueError):
            pass
    return rows


# ----------------------------------------------------------------------------
# Hoja para proveedor — Excel (para llenar) y PDF (para anotar a mano)
# ----------------------------------------------------------------------------
@app.route("/api/hoja_proveedor.xlsx")
@requiere_login
def hoja_proveedor_xlsx():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    prov = request.args.get("proveedor", "")
    titulo_txt = request.args.get("titulo", "")
    rows = _filtrar_actividades(request.args)
    hoy_d = datetime.date.today()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Actualización"
    torre_txt = "Torre A" if proyecto_actual() == "torre_a" else "Torre B"
    encabezado = titulo_txt or (f"HOJA DE ACTUALIZACIÓN DE AVANCE — {prov}" if prov else "REPORTE DE ACTIVIDADES")
    ws["A1"] = encabezado
    ws["A1"].font = Font(bold=True, size=13, color="1F4E78")
    ws.merge_cells("A1:H1")
    ws["A2"] = f"Fecha: {hoy_d.strftime('%d/%m/%Y')}   ·   Favor de anotar avance, fecha compromiso y observaciones."
    ws["A2"].font = Font(size=9, italic=True, color="555555")
    ws.merge_cells("A2:H2")
    cols = ["Código", "Área", "Partida", "Avance actual",
            "NUEVO avance %", "Fecha compromiso", "¿Retrasada?", "Observaciones"]
    hdr = 4
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=hdr, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    rojo = PatternFill("solid", fgColor="F8D7DA")
    i = hdr + 1
    for r in rows:
        fin = parse_date(r["f_fin"])
        retrasada = fin and (r["avance"] or 0) < 100 and (fin - hoy_d).days < 0
        vals = [r["codigo"], r["area"], r["partida"], f"{r['avance'] or 0}%",
                "", r["f_fin"] or "", "SÍ" if retrasada else "", ""]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.border = bd
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if retrasada:
                cell.fill = rojo
        i += 1
    anchos = [10, 22, 40, 12, 13, 15, 11, 30]
    for j, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A5"
    safe = "".join(ch for ch in (prov or "proveedor") if ch.isalnum() or ch in " _-")[:30].strip() or "proveedor"
    out = ruta_salida_temporal("hoja_proveedor.xlsx")
    wb.save(out)
    return send_file(out, as_attachment=True, download_name=f"HOJA_{safe}_{hoy()}.xlsx")


@app.route("/api/hoja_proveedor.pdf")
@requiere_login
def hoja_proveedor_pdf():
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    prov = request.args.get("proveedor", "")
    titulo_txt = request.args.get("titulo", "")
    rows = _filtrar_actividades(request.args)
    hoy_d = datetime.date.today()
    out = ruta_salida_temporal("hoja_proveedor.pdf")
    doc = SimpleDocTemplate(out, pagesize=landscape(letter),
                            leftMargin=1*cm, rightMargin=1*cm, topMargin=1*cm, bottomMargin=1*cm)
    styles = getSampleStyleSheet()
    titulo = ParagraphStyle("t", parent=styles["Title"], fontSize=14, textColor=colors.HexColor("#1F4E78"))
    small = ParagraphStyle("s", parent=styles["Normal"], fontSize=7, leading=8)
    encabezado = titulo_txt or (f"Hoja de actualización de avance — {prov}" if prov else "Reporte de actividades")
    elems = [Paragraph(encabezado, titulo),
             Paragraph(f"Fecha: {hoy_d.strftime('%d/%m/%Y')} · Anote avance, fecha compromiso y observaciones.", styles["Normal"]),
             Spacer(1, 8)]
    data = [["Código", "Área", "Partida", "Avance\nactual", "NUEVO\navance", "Fecha\ncompromiso", "Retra-\nsada", "Observaciones"]]
    filas_retraso = []
    for idx, r in enumerate(rows):
        fin = parse_date(r["f_fin"])
        retrasada = fin and (r["avance"] or 0) < 100 and (fin - hoy_d).days < 0
        data.append([r["codigo"] or "", Paragraph(r["area"] or "", small),
                     Paragraph(r["partida"] or "", small), f"{r['avance'] or 0}%",
                     "______", r["f_fin"] or "__________", "SÍ" if retrasada else "",
                     "________________"])
        if retrasada:
            filas_retraso.append(idx + 1)
    t = Table(data, colWidths=[2*cm, 3.5*cm, 6*cm, 1.5*cm, 1.6*cm, 2.2*cm, 1.3*cm, 5*cm], repeatRows=1)
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7F8")]),
    ]
    for fr in filas_retraso:
        estilo.append(("BACKGROUND", (0, fr), (-1, fr), colors.HexColor("#F8D7DA")))
    t.setStyle(TableStyle(estilo))
    elems.append(t)
    doc.build(elems)
    safe = "".join(ch for ch in (prov or "proveedor") if ch.isalnum() or ch in " _-")[:30].strip() or "proveedor"
    return send_file(out, as_attachment=True, download_name=f"HOJA_{safe}_{hoy()}.pdf")


# ----------------------------------------------------------------------------
# Exportar a Excel
# ----------------------------------------------------------------------------
@app.route("/api/exportar")
@requiere_login
def api_exportar():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    db = get_db()
    rows = db.execute("SELECT * FROM actividades WHERE (eliminada IS NULL OR eliminada=0) ORDER BY id").fetchall()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AVANCE"
    cols = ["ID", "Código", "Bloque", "Área", "Giro", "Proveedor", "Partida",
            "Tipo", "¿Aplica?", "% Avance", "F. Inicio", "F. Fin",
            "Duración (días)", "Estatus", "Notas"]
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for i, r in enumerate(rows, 2):
        vals = [r["id"], r["codigo"], r["bloque"], r["area"], r["giro"],
                r["proveedor"], r["partida"], r["tipo"], r["aplica"],
                r["avance"], r["f_inicio"], r["f_fin"], r["duracion_dias"],
                r["estatus"], r["notas"]]
        for j, v in enumerate(vals, 1):
            ws.cell(row=i, column=j, value=v)
    widths = [6, 10, 20, 24, 14, 22, 38, 12, 9, 9, 12, 12, 14, 14, 30]
    from openpyxl.utils import get_column_letter
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    out = ruta_salida_temporal("export_avance.xlsx")
    wb.save(out)
    return send_file(out, as_attachment=True,
                     download_name=f"AVANCE_HAP_{hoy()}.xlsx")


# ----------------------------------------------------------------------------
# Plantilla de captura para llenar rápido (con desplegables) y reimportar
# ----------------------------------------------------------------------------
@app.route("/api/plantilla_captura.xlsx")
@requiere_login
def api_plantilla_captura():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    rows = _filtrar_actividades(request.args)
    hoy_d = datetime.date.today()
    # causas para el desplegable
    dbc = get_db()
    causas = [r[0] for r in dbc.execute("SELECT nombre FROM causas ORDER BY base DESC,nombre").fetchall()]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Captura"

    ws["A1"] = "PLANTILLA DE CAPTURA DE AVANCE — HAP"
    ws["A1"].font = Font(bold=True, size=13, color="1F4E78")
    ws.merge_cells("A1:H1")
    ws["A2"] = ("Llena las columnas AMARILLAS. NO cambies la columna Código. "
                "Al terminar, súbela en la plataforma con 'Importar avances'.")
    ws["A2"].font = Font(size=9, italic=True, color="B00020")
    ws.merge_cells("A2:H2")

    # Encabezados: las primeras 4 son de solo lectura (referencia); las 3 amarillas se llenan
    cols = ["Código", "Área", "Partida", "Avance actual",
            "NUEVO avance %", "Fecha compromiso (AAAA-MM-DD)", "Causa de retraso"]
    hdr = 4
    azul = PatternFill("solid", fgColor="1F4E78")
    amar = PatternFill("solid", fgColor="FFF2CC")
    gris = PatternFill("solid", fgColor="EDEDED")
    thin = Side(style="thin", color="CCCCCC")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=hdr, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = azul
        cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        cell.border = bd
    ws.row_dimensions[hdr].height = 30

    i = hdr + 1
    for r in rows:
        ws.cell(row=i, column=1, value=r["codigo"]).border = bd
        ws.cell(row=i, column=2, value=r["area"]).border = bd
        ws.cell(row=i, column=3, value=r["partida"]).border = bd
        c4 = ws.cell(row=i, column=4, value=f"{r['avance'] or 0}%"); c4.border = bd; c4.fill = gris
        # columnas amarillas para llenar
        for col in (5, 6, 7):
            cc = ws.cell(row=i, column=col); cc.border = bd; cc.fill = amar
        i += 1
    ultima = i - 1

    # Validaciones: avance por lista 0/10/20…100, causa por lista de causas
    if ultima >= hdr + 1:
        dv_av = DataValidation(type="list", formula1='"0,10,20,30,40,50,60,70,80,90,100"', allow_blank=True)
        ws.add_data_validation(dv_av)
        dv_av.add(f"E{hdr+1}:E{ultima}")
        # causas: si son pocas caben en formula directa; si no, usar hoja aparte
        causas_txt = ",".join(c.replace(",", " ") for c in causas)
        if len(causas_txt) < 240:
            dv_ca = DataValidation(type="list", formula1=f'"{causas_txt}"', allow_blank=True)
            ws.add_data_validation(dv_ca)
            dv_ca.add(f"G{hdr+1}:G{ultima}")
        else:
            ws2 = wb.create_sheet("causas")
            for k, c in enumerate(causas, 1):
                ws2.cell(row=k, column=1, value=c)
            dv_ca = DataValidation(type="list",
                                   formula1=f"causas!$A$1:$A${len(causas)}", allow_blank=True)
            ws.add_data_validation(dv_ca)
            dv_ca.add(f"G{hdr+1}:G{ultima}")

    anchos = [11, 22, 42, 12, 14, 20, 26]
    for j, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A5"
    out = ruta_salida_temporal("plantilla_captura.xlsx")
    wb.save(out)
    prov = request.args.get("proveedor", "")
    safe = "".join(ch for ch in (prov or "todos") if ch.isalnum() or ch in " _-")[:30].strip() or "todos"
    return send_file(out, as_attachment=True, download_name=f"CAPTURA_{safe}_{hoy()}.xlsx")


@app.route("/api/importar_avances", methods=["POST"])
@requiere_gestor
def api_importar_avances():
    """Lee un Excel con Código + NUEVO avance / fecha / causa y actualiza."""
    import openpyxl
    if "archivo" not in request.files:
        return jsonify({"error": "No llegó ningún archivo"}), 400
    f = request.files["archivo"]
    try:
        wb = openpyxl.load_workbook(f, data_only=True)
    except Exception as e:
        return jsonify({"error": f"No se pudo leer el Excel: {e}"}), 400
    ws = wb["Captura"] if "Captura" in wb.sheetnames else wb.active

    # localizar la fila de encabezados y las columnas por su nombre
    encabezados = {}
    fila_hdr = None
    for ri in range(1, min(12, ws.max_row + 1)):
        vals = {(ws.cell(row=ri, column=ci).value or "").__str__().strip().lower(): ci
                for ci in range(1, ws.max_column + 1)}
        # la fila de encabezados tiene una celda que es EXACTAMENTE "código"/"codigo"
        if "código" in vals or "codigo" in vals:
            encabezados = vals; fila_hdr = ri; break
    if not fila_hdr:
        return jsonify({"error": "No encontré la fila de encabezados con la columna 'Código'"}), 400

    def col_de(*claves):
        for k, ci in encabezados.items():
            for clave in claves:
                if clave in k:
                    return ci
        return None
    c_cod = encabezados.get("código") or encabezados.get("codigo")
    c_av = col_de("nuevo avance", "nuevo")
    c_fe = col_de("fecha compromiso", "compromiso")
    c_ca = col_de("causa")

    db = get_db()
    actualizadas, sin_match, errores = 0, [], 0
    for ri in range(fila_hdr + 1, ws.max_row + 1):
        codigo = ws.cell(row=ri, column=c_cod).value
        if not codigo:
            continue
        codigo = str(codigo).strip()
        act = db.execute("SELECT * FROM actividades WHERE codigo=?", (codigo,)).fetchone()
        if not act:
            sin_match.append(codigo); continue
        cambios, args = [], []
        # avance
        if c_av:
            v = ws.cell(row=ri, column=c_av).value
            if v is not None and str(v).strip() != "":
                try:
                    av = int(float(str(v).replace("%", "").strip()))
                    av = max(0, min(100, av))
                    if av != (act["avance"] or 0):
                        cambios.append("avance=?"); args.append(av)
                        cambios.append("estatus=?"); args.append(estatus_por_avance(av))
                        # mismo motivo que en api_actualizar: alinear el avance declarado
                        if act["avance_decl"] is not None and act["avance_decl"] != av:
                            cambios.append("avance_decl=?"); args.append(av)
                except ValueError:
                    errores += 1
        # fecha compromiso
        if c_fe:
            v = ws.cell(row=ri, column=c_fe).value
            if v is not None and str(v).strip() != "":
                d = parse_date(str(v)[:10]) if not isinstance(v, datetime.datetime) else v.date()
                if d:
                    iso = d.isoformat()
                    if iso != (act["f_fin"] or ""):
                        cambios.append("f_fin=?"); args.append(iso)
        # causa
        if c_ca:
            v = ws.cell(row=ri, column=c_ca).value
            if v is not None and str(v).strip() != "":
                causa = str(v).strip()
                if causa != (act["causa_retraso"] or ""):
                    cambios.append("causa_retraso=?"); args.append(causa)
                    cambios.append("causa_por=?"); args.append("Importado de Excel")
                    cambios.append("causa_fecha=?"); args.append(datetime.date.today().isoformat())
        if cambios:
            cambios.append("actualizado=?")
            args.append(datetime.datetime.now().isoformat(timespec="seconds"))
            args.append(act["id"])
            db.execute(f"UPDATE actividades SET {','.join(cambios)} WHERE id=?", args)
            actualizadas += 1
    db.commit()
    return jsonify({
        "ok": True, "actualizadas": actualizadas,
        "sin_coincidencia": sin_match, "errores": errores,
    })


# ----------------------------------------------------------------------------
# Autenticación y portal de proveedores
# ----------------------------------------------------------------------------
@app.route("/api/login", methods=["POST"])
def api_login():
    data = request.get_json(silent=True) or {}
    usuario = (data.get("usuario") or "").strip()
    clave = data.get("clave") or ""
    db = get_db_usuarios()
    u = db.execute("SELECT * FROM usuarios WHERE usuario=?", (usuario,)).fetchone()
    if not u:
        verificar_clave(_HASH_RELLENO, clave)  # mismo tiempo de respuesta que un usuario real
        return jsonify({"error": "Usuario o contraseña incorrectos"}), 401
    if not verificar_clave(u["clave"], clave):
        return jsonify({"error": "Usuario o contraseña incorrectos"}), 401
    # candado: usuario desactivado no puede entrar
    if "activo" in u.keys() and u["activo"] == 0:
        return jsonify({"error": "Tu acceso está desactivado. Contacta al administrador."}), 403
    # migración gradual: si la contraseña estaba en formato antiguo, se guarda en el nuevo
    actualizar_hash_si_antiguo(db, u["id"], u["clave"], clave)
    # contar el acceso
    try:
        db.execute("UPDATE usuarios SET num_logins=COALESCE(num_logins,0)+1, ultimo_login=? WHERE id=?",
                   (datetime.datetime.now().isoformat(timespec="seconds"), u["id"]))
        db.commit()
    except Exception:
        pass
    session["usuario"] = u["usuario"]
    session["rol"] = u["rol"]
    session["proveedor"] = u["proveedor"]
    session["mundo"] = u["mundo"] if "mundo" in u.keys() else "obra"
    session["proyecto"] = "torre_b"
    return jsonify({"ok": True, "usuario": u["usuario"], "rol": u["rol"],
                    "proveedor": u["proveedor"], "mundo": session["mundo"]})


@app.route("/api/logout", methods=["POST"])
def api_logout():
    session.clear()
    return jsonify({"ok": True})


@app.route("/api/quien_soy")
def api_quien_soy():
    if "usuario" not in session:
        return jsonify({"login": False})
    return jsonify({"login": True, "usuario": session["usuario"],
                    "rol": session["rol"], "proveedor": session.get("proveedor"),
                    "mundo": session.get("mundo", "obra")})


# --- Gestión de usuarios (solo admin) ---
@app.route("/api/usuarios", methods=["GET"])
@requiere_admin
def api_usuarios():
    db = get_db_usuarios()
    rows = db.execute("SELECT id,usuario,proveedor,rol,mundo,clave_cambiada,creado,activo,num_logins,ultimo_login,telefono FROM usuarios ORDER BY rol,mundo,usuario").fetchall()
    return jsonify([dict(r) for r in rows])


@app.route("/api/usuarios", methods=["POST"])
@requiere_admin
def api_crear_usuario():
    db = get_db_usuarios()
    data = request.get_json()
    usuario = (data.get("usuario") or "").strip()
    clave = data.get("clave") or ""
    proveedor = data.get("proveedor") or None
    mundo = data.get("mundo", "obra")
    telefono = (data.get("telefono") or "").strip() or None
    # rol: por defecto proveedor. El admin puede crear supervisores.
    rol = data.get("rol", "proveedor")
    if rol not in ("proveedor", "supervisor", "supervisor_obra", "supervisor_depto", "directivo"):
        rol = "proveedor"
    # el supervisor de obra vive en mundo 'obra'; el de depto en 'interno'; el directivo (solo dashboard) en 'obra'
    if rol in ("supervisor", "supervisor_obra", "directivo"):
        mundo = "obra"
    elif rol == "supervisor_depto":
        mundo = "interno"
    if not usuario or not clave:
        return jsonify({"error": "Usuario y contraseña son obligatorios"}), 400
    try:
        db.execute(
            "INSERT INTO usuarios (usuario,clave,proveedor,rol,mundo,telefono,activo,num_logins,creado) VALUES (?,?,?,?,?,?,1,0,?)",
            (usuario, crear_hash_clave(clave), proveedor, rol, mundo, telefono,
             datetime.datetime.now().isoformat(timespec="seconds")))
        db.commit()
    except sqlite3.IntegrityError:
        return jsonify({"error": "Ese usuario ya existe"}), 400
    return jsonify({"ok": True})


@app.route("/api/usuario/<int:uid>", methods=["DELETE"])
@requiere_admin
def api_borrar_usuario(uid):
    db = get_db_usuarios()
    u = db.execute("SELECT rol FROM usuarios WHERE id=?", (uid,)).fetchone()
    if u and u["rol"] == "admin":
        return jsonify({"error": "No se puede borrar al administrador"}), 400
    db.execute("DELETE FROM usuarios WHERE id=?", (uid,))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/usuario/<int:uid>/clave", methods=["POST"])
@requiere_admin
def api_cambiar_clave(uid):
    db = get_db_usuarios()
    data = request.get_json()
    nueva = data.get("clave") or ""
    if not nueva:
        return jsonify({"error": "Escribe la nueva contraseña"}), 400
    db.execute("UPDATE usuarios SET clave=? WHERE id=?", (crear_hash_clave(nueva), uid))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/usuario/<int:uid>/activo", methods=["POST"])
@requiere_admin
def api_usuario_activo(uid):
    """Activa o desactiva el acceso de un usuario (sin borrarlo)."""
    db = get_db_usuarios()
    data = request.get_json() or {}
    activo = 1 if data.get("activo") else 0
    u = db.execute("SELECT rol FROM usuarios WHERE id=?", (uid,)).fetchone()
    if u and u["rol"] == "admin":
        return jsonify({"error": "No se puede desactivar al administrador"}), 400
    db.execute("UPDATE usuarios SET activo=? WHERE id=?", (activo, uid))
    db.commit()
    return jsonify({"ok": True, "activo": activo})


@app.route("/api/usuario/<int:uid>/telefono", methods=["POST"])
@requiere_admin
def api_usuario_telefono(uid):
    """Guarda o actualiza el telefono (para el aviso por WhatsApp)."""
    db = get_db_usuarios()
    data = request.get_json() or {}
    tel = (data.get("telefono") or "").strip() or None
    db.execute("UPDATE usuarios SET telefono=? WHERE id=?", (tel, uid))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/usuario/<int:uid>/whatsapp", methods=["GET"])
@requiere_admin
def api_usuario_whatsapp(uid):
    """Arma el link de WhatsApp con un mensaje listo avisando de nuevas asignaciones.
    No envia nada: devuelve el link para que el admin le de click."""
    db = get_db_usuarios()
    u = db.execute("SELECT * FROM usuarios WHERE id=?", (uid,)).fetchone()
    if not u:
        return jsonify({"error": "Usuario no existe"}), 404
    tel = (u["telefono"] or "").strip() if "telefono" in u.keys() else ""
    if not tel:
        return jsonify({"error": "Este usuario no tiene telefono registrado"}), 400
    # contar nuevas por reconocer de ese proveedor
    proveedor = u["proveedor"]
    mundo = u["mundo"] if "mundo" in u.keys() else "obra"
    col = "departamento" if mundo == "interno" else "proveedor"
    n = db.execute(
        f"SELECT COUNT(*) FROM actividades WHERE {col}=? AND reconocida='NO' "
        "AND (estado_val IS NULL OR estado_val='validado') "
        "AND (mundo=? OR (mundo IS NULL AND ?='obra'))",
        (proveedor, mundo, mundo)).fetchone()[0]
    # normalizar telefono: solo digitos; si no trae lada pais, anteponer 52 (Mexico)
    solo_num = "".join(ch for ch in tel if ch.isdigit())
    if len(solo_num) == 10:
        solo_num = "52" + solo_num
    liga = request.host_url.rstrip("/") + "/portal"
    if n > 0:
        texto = (f"Hola {proveedor or u['usuario']}, tienes {n} "
                 f"{'nueva actividad asignada' if n == 1 else 'nuevas actividades asignadas'} "
                 f"en el Control de Obra HAP. Entra a reconocerlas y reportar tu avance: {liga}")
    else:
        texto = (f"Hola {proveedor or u['usuario']}, te recuerdo revisar tus actividades "
                 f"en el Control de Obra HAP: {liga}")
    import urllib.parse
    url = "https://wa.me/" + solo_num + "?text=" + urllib.parse.quote(texto)
    return jsonify({"ok": True, "url": url, "mensaje": texto, "nuevas": n, "telefono": tel})


# ----------------------------------------------------------------------------
# API — Avisos (bitácora) y notificaciones push para el admin
# ----------------------------------------------------------------------------
@app.route("/api/avisos")
@requiere_gestor
def api_avisos():
    db = get_db()
    try:
        limite = max(1, min(200, int(request.args.get("limite", 50))))
    except (TypeError, ValueError):
        limite = 50
    filas = db.execute("SELECT * FROM avisos ORDER BY id DESC LIMIT ?", (limite,)).fetchall()
    no_vistos = db.execute("SELECT COUNT(*) FROM avisos WHERE visto=0").fetchone()[0]
    return jsonify({"avisos": [dict(r) for r in filas], "no_vistos": no_vistos})


@app.route("/api/avisos/marcar_visto", methods=["POST"])
@requiere_gestor
def api_avisos_marcar_visto():
    db = get_db()
    db.execute("UPDATE avisos SET visto=1 WHERE visto=0")
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/push/vapid_public_key")
@requiere_login
def api_push_vapid_public_key():
    _, pub_b64 = _vapid_keys()
    if not pub_b64:
        return jsonify({"error": "Las notificaciones push no están disponibles en este servidor"}), 503
    return jsonify({"key": pub_b64})


@app.route("/api/push/subscribe", methods=["POST"])
@requiere_gestor
def api_push_subscribe():
    usuario, rol, proveedor = usuario_actual()
    data = request.get_json() or {}
    endpoint = data.get("endpoint")
    keys = data.get("keys") or {}
    if not endpoint or not keys.get("p256dh") or not keys.get("auth"):
        return jsonify({"error": "Suscripción inválida"}), 400
    db = get_db_usuarios()
    db.execute(
        "INSERT INTO push_subscriptions (usuario,endpoint,p256dh,auth,creado) VALUES (?,?,?,?,?) "
        "ON CONFLICT(endpoint) DO UPDATE SET usuario=excluded.usuario, p256dh=excluded.p256dh, auth=excluded.auth",
        (usuario, endpoint, keys["p256dh"], keys["auth"], datetime.datetime.now().isoformat(timespec="seconds")))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/push/unsubscribe", methods=["POST"])
@requiere_gestor
def api_push_unsubscribe():
    data = request.get_json() or {}
    endpoint = data.get("endpoint")
    if endpoint:
        db = get_db_usuarios()
        db.execute("DELETE FROM push_subscriptions WHERE endpoint=?", (endpoint,))
        db.commit()
    return jsonify({"ok": True})


@app.route("/api/push/estado")
@requiere_gestor
def api_push_estado():
    """Para que el botón sepa si YA hay algo suscrito desde este mismo
    usuario (informativo; el navegador manda su propio endpoint para saber
    si ESTE dispositivo específico ya está suscrito)."""
    usuario, rol, proveedor = usuario_actual()
    db = get_db_usuarios()
    n = db.execute("SELECT COUNT(*) FROM push_subscriptions WHERE usuario=?", (usuario,)).fetchone()[0]
    return jsonify({"dispositivos_activos": n})


@app.route("/api/push/test", methods=["POST"])
@requiere_gestor
def api_push_test():
    enviar_push_admins("🔔 Aviso de prueba", "Si ves esto, las notificaciones ya están funcionando en este celular/navegador.", "/")
    return jsonify({"ok": True})


# --- Portal del proveedor: ve y reporta SOLO lo suyo ---
@app.route("/portal")
def portal_page():
    # blindaje: si por cualquier motivo un gestor cae aqui, lo regresamos al panel
    if "usuario" in session and es_gestor(session.get("rol")):
        return redirect("/")
    return render_template("portal.html")


@app.route("/api/portal/mapa")
@requiere_login
def api_portal_mapa():
    """Mapa bloque->áreas para el formulario de nueva actividad.
    Externo: solo sus bloques. Interno: todo el mapa de la obra (su labor cruza el hospital)."""
    _, rol, proveedor = usuario_actual()
    db = get_db()
    col, mundo = col_duenio()
    mapa = {}
    if mundo == "interno":
        # todo el mapa de la obra
        filas = db.execute(
            "SELECT DISTINCT bloque, area FROM actividades "
            "WHERE bloque IS NOT NULL AND bloque<>'' AND area IS NOT NULL AND area<>'' "
            "ORDER BY bloque, area").fetchall()
        # bloques donde el depto YA tiene trabajo (para marcar lo que es 'fuera de zona')
        suyos = [r[0] for r in db.execute(
            f"SELECT DISTINCT bloque FROM actividades WHERE {col}=? AND bloque IS NOT NULL",
            (proveedor,)).fetchall()]
    else:
        filas = db.execute(
            f"SELECT DISTINCT bloque, area FROM actividades WHERE {col}=? "
            "AND bloque IS NOT NULL AND bloque<>'' AND area IS NOT NULL AND area<>'' "
            "ORDER BY bloque, area", (proveedor,)).fetchall()
        suyos = None  # externo: todos los que ve ya son suyos
    for r in filas:
        mapa.setdefault(r["bloque"], []).append(r["area"])
    return jsonify({"mapa": mapa, "mundo": mundo, "bloques_propios": suyos})


@app.route("/api/cambiar_mi_clave", methods=["POST"])
@requiere_login
def api_cambiar_mi_clave():
    usuario, _, _ = usuario_actual()
    db = get_db_usuarios()
    data = request.get_json() or {}
    clave_actual = data.get("clave_actual") or ""
    clave_nueva = data.get("clave_nueva") or ""
    if not clave_actual or not clave_nueva:
        return jsonify({"error": "Ingresa tu contraseña actual y la nueva contraseña"}), 400
    if len(clave_nueva) < 4:
        return jsonify({"error": "La nueva contraseña debe tener al menos 4 caracteres"}), 400
    
    u = db.execute("SELECT * FROM usuarios WHERE usuario=?", (usuario,)).fetchone()
    if not u or not verificar_clave(u["clave"], clave_actual):
        return jsonify({"error": "La contraseña actual es incorrecta"}), 400

    db.execute("UPDATE usuarios SET clave=?, clave_cambiada=1 WHERE usuario=?",
               (crear_hash_clave(clave_nueva), usuario))
    db.commit()
    return jsonify({"ok": True, "mensaje": "Contraseña actualizada exitosamente"})


@app.route("/api/portal/avance_zona", methods=["GET"])
@requiere_login
def api_portal_avance_zona():
    """Vista de solo lectura del avance de todos los gremios en la zona del proveedor."""
    usuario, rol, proveedor = usuario_actual()
    db = get_db()
    col, mundo = col_duenio()
    
    bloque = request.args.get("bloque")
    area = request.args.get("area")
    
    # Determinar qué bloques/áreas puede consultar
    if mundo == "interno":
        # Interno: ve todo el hospital
        cond = ["(aplica IS NULL OR aplica<>'NO')", "(estado_val IS NULL OR estado_val='validado')", "(eliminada IS NULL OR eliminada=0)"]
        args = []
        bloques_disp = [r[0] for r in db.execute("SELECT DISTINCT bloque FROM actividades WHERE bloque IS NOT NULL AND bloque<>'' ORDER BY bloque").fetchall()]
    else:
        # Externo: ve los bloques donde tiene partidas asignadas
        bloques_suyos = [r[0] for r in db.execute(
            f"SELECT DISTINCT bloque FROM actividades WHERE {col}=? AND bloque IS NOT NULL AND bloque<>''",
            (proveedor,)).fetchall()]
        if not bloques_suyos:
            return jsonify({"bloques_disponibles": [], "actividades": []})
        marcas = ",".join("?" * len(bloques_suyos))
        cond = [f"bloque IN ({marcas})", "(aplica IS NULL OR aplica<>'NO')", "(estado_val IS NULL OR estado_val='validado')", "(eliminada IS NULL OR eliminada=0)"]
        args = list(bloques_suyos)
        bloques_disp = sorted(bloques_suyos)
        
    if bloque:
        cond.append("bloque = ?"); args.append(bloque)
    if area:
        cond.append("area = ?"); args.append(area)
        
    q = f"SELECT id, codigo, bloque, area, giro, proveedor, departamento, mundo, partida, tipo, tipo_partida, avance, estatus, f_inicio, f_fin FROM actividades WHERE {' AND '.join(cond)} ORDER BY bloque, area, id"
    filas = db.execute(q, args).fetchall()
    
    return jsonify({
        "bloques_disponibles": bloques_disp,
        "actividades": [dict(r) for r in filas]
    })


@app.route("/api/portal/mis_actividades")
@requiere_login
def api_portal_mis():
    _, rol, proveedor = usuario_actual()
    if es_gestor(rol):
        return jsonify({"error": "Los gestores usan el panel principal"}), 400
    db = get_db()
    col, mundo = col_duenio()
    rows = db.execute(
        f"SELECT * FROM actividades WHERE {col}=? AND (mundo=? OR (mundo IS NULL AND ?='obra')) "
        "AND (reconocida IS NULL OR reconocida<>'RECHAZADA') "
        "AND (eliminada IS NULL OR eliminada=0) "
        "ORDER BY estado_val DESC, area, id",
        (proveedor, mundo, mundo)).fetchall()
    
    salida = []
    for r in rows:
        d = dict(r)
        dep_st = evaluar_dependencias_area(db, r["area"], r["tipo_partida"] or r["giro"])
        d["dep_bloqueada"] = dep_st["bloqueada"]
        d["dep_estado"] = dep_st["estado"]
        d["dep_detalle"] = dep_st["detalle"]
        salida.append(d)
    return jsonify(salida)


@app.route("/api/portal/reportar_avance/<int:aid>", methods=["POST"])
@requiere_login
def api_portal_reportar(aid):
    """El proveedor declara avance en una partida oficial suya. Entra directo como 'declarado'."""
    usuario, rol, proveedor = usuario_actual()
    db = get_db()
    a = db.execute("SELECT * FROM actividades WHERE id=?", (aid,)).fetchone()
    col, _ = col_duenio()
    if not a or a[col] != proveedor:
        return jsonify({"error": "Esa partida no es tuya"}), 403
    # candado: no puede reportar avance si no ha reconocido la actividad
    if a["reconocida"] != "SÍ":
        return jsonify({"error": "Primero tienes que reconocer esta actividad"}), 400
    data = request.get_json() or {}
    # candado: si ya está validada al 100% (y no hay un reporte pendiente de
    # revisión), no se deja seguir reportando — antes se podía reenviar avance
    # indefinidamente aunque la actividad ya estuviera cerrada.
    ya_completa = (a["avance"] or 0) >= 100 and (a["avance_decl"] is None or a["avance_decl"] == a["avance"])
    if ya_completa and "avance_decl" in data:
        return jsonify({"error": "Esta actividad ya está validada al 100% — si necesitas corregir algo, contacta al administrador."}), 400
    campos, args = [], []
    if "avance_decl" in data:
        av = max(0, min(100, int(data.get("avance_decl") or 0)))
        registrar_historial(db, aid, "avance reportado", a["avance_decl"], av, usuario)
        campos += ["avance_decl=?", "avance_decl_por=?", "avance_decl_fecha=?", "avance_decl_rechazado=0", "rechazo_motivo=NULL"]
        args += [av, usuario, datetime.date.today().isoformat()]
    for c in ("f_inicio", "f_fin", "definido", "definido_por", "nota_proveedor"):
        if c in data:
            registrar_historial(db, aid, c, a[c], data[c], usuario)
            campos.append(f"{c}=?"); args.append(data[c])
    if campos:
        campos.append("actualizado=?"); args.append(datetime.datetime.now().isoformat(timespec="seconds"))
        args.append(aid)
        db.execute(f"UPDATE actividades SET {','.join(campos)} WHERE id=?", args)
        # aviso para el admin
        db.execute("INSERT INTO avisos (fecha,proveedor,usuario,tipo,detalle,visto) VALUES (?,?,?,?,?,0)",
                   (datetime.datetime.now().isoformat(timespec="seconds"), proveedor, usuario,
                    "avance", f"Reportó avance en {a['codigo']} · {(a['partida'] or '')[:50]}"))
        db.commit()
        enviar_push_admins(f"📋 {proveedor}", f"Reportó avance en {a['codigo']}: {(a['partida'] or '')[:60]}", "/validacion")
    return jsonify({"ok": True})


@app.route("/api/portal/reconocer/<int:aid>", methods=["POST"])
@requiere_login
def api_portal_reconocer(aid):
    """El proveedor acepta que una partida es su trabajo. Sin esto no puede reportar avance."""
    usuario, rol, proveedor = usuario_actual()
    db = get_db()
    a = db.execute("SELECT * FROM actividades WHERE id=?", (aid,)).fetchone()
    col, _ = col_duenio()
    if not a or a[col] != proveedor:
        return jsonify({"error": "Esa partida no es tuya"}), 403
    registrar_historial(db, aid, "reconocimiento", "sin reconocer", "reconocida", usuario)
    db.execute("UPDATE actividades SET reconocida='SÍ', reconocida_por=?, reconocida_fecha=?, "
               "no_reconocida_nota=NULL WHERE id=?",
               (usuario, datetime.date.today().isoformat(), aid))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/portal/reconocer_bloque", methods=["POST"])
@requiere_login
def api_portal_reconocer_bloque():
    """Reconoce en lote todas las pendientes de un bloque (o todas si no se manda bloque)."""
    usuario, rol, proveedor = usuario_actual()
    db = get_db()
    data = request.get_json() or {}
    bloque = data.get("bloque")
    col, mundo = col_duenio()
    cond = f"{col}=? AND reconocida='NO' AND (estado_val IS NULL OR estado_val='validado') AND (mundo=? OR (mundo IS NULL AND ?='obra'))"
    args = [proveedor, mundo, mundo]
    if bloque:
        cond += " AND bloque=?"; args.append(bloque)
    filas = db.execute(f"SELECT id FROM actividades WHERE {cond}", args).fetchall()
    for f in filas:
        registrar_historial(db, f["id"], "reconocimiento", "sin reconocer", "reconocida (en lote)", usuario)
    db.execute(f"UPDATE actividades SET reconocida='SÍ', reconocida_por=?, reconocida_fecha=? WHERE {cond}",
               [usuario, datetime.date.today().isoformat()] + args)
    db.commit()
    return jsonify({"ok": True, "reconocidas": len(filas)})


@app.route("/api/portal/no_reconozco/<int:aid>", methods=["POST"])
@requiere_login
def api_portal_no_reconozco(aid):
    """El proveedor dice que esa partida NO es suya. Te avisa a ti para corregir."""
    usuario, rol, proveedor = usuario_actual()
    db = get_db()
    a = db.execute("SELECT * FROM actividades WHERE id=?", (aid,)).fetchone()
    col, _ = col_duenio()
    if not a or a[col] != proveedor:
        return jsonify({"error": "Esa partida no es tuya"}), 403
    data = request.get_json() or {}
    nota = (data.get("nota") or "").strip()
    registrar_historial(db, aid, "no reconocida", "", nota or "el proveedor no reconoce esta actividad", usuario)
    db.execute("UPDATE actividades SET no_reconocida_nota=?, reconocida='RECHAZADA' WHERE id=?", (nota or "No reconocida", aid))
    db.execute("INSERT INTO avisos (fecha,proveedor,usuario,tipo,detalle,visto) VALUES (?,?,?,?,?,0)",
               (datetime.datetime.now().isoformat(timespec="seconds"), proveedor, usuario,
                "no_reconocida", f"No reconoce {a['codigo']}: {(a['partida'] or '')[:50]}"))
    db.commit()
    enviar_push_admins(f"⚠️ {proveedor}", f"No reconoce {a['codigo']}: {(a['partida'] or '')[:60]}", "/validacion")
    return jsonify({"ok": True})


@app.route("/api/portal/pendientes_reconocer")
@requiere_login
def api_portal_pend_reconocer():
    """Conteo de lo que el proveedor tiene sin reconocer (para el botón de 'nuevas asignadas')."""
    _, rol, proveedor = usuario_actual()
    db = get_db()
    col, mundo = col_duenio()
    n = db.execute(f"SELECT COUNT(*) FROM actividades WHERE {col}=? AND reconocida='NO' "
                   "AND (estado_val IS NULL OR estado_val='validado') "
                   "AND (eliminada IS NULL OR eliminada=0) "
                   "AND (mundo=? OR (mundo IS NULL AND ?='obra'))",
                   (proveedor, mundo, mundo)).fetchone()[0]
    return jsonify({"pendientes": n})


@app.route("/api/portal/nueva", methods=["POST"])
@requiere_login
def api_portal_nueva():
    """El proveedor propone una tarea NUEVA. Queda en estado 'propuesta' hasta que el admin valide."""
    usuario, rol, proveedor = usuario_actual()
    db = get_db()
    data = request.get_json()
    partida = (data.get("partida") or "").strip()
    if not partida:
        return jsonify({"error": "Escribe qué actividad es"}), 400
    col, mundo = col_duenio()
    # detectar ANTES de insertar si el interno propone en un bloque donde NO tenía trabajo
    fuera_zona = 0
    if mundo == "interno" and data.get("bloque"):
        tiene = db.execute(
            f"SELECT COUNT(*) FROM actividades WHERE {col}=? AND bloque=?",
            (proveedor, data.get("bloque"))).fetchone()[0]
        fuera_zona = 1 if tiene == 0 else 0
    # el giro y tipo se ponen solos según el proveedor/departamento (los más comunes de sus partidas)
    ref = db.execute(
        f"SELECT giro, tipo_partida, tipo_interno FROM actividades WHERE {col}=? "
        "GROUP BY giro ORDER BY COUNT(*) DESC LIMIT 1", (proveedor,)).fetchone()
    giro = (ref["giro"] if ref else "") or data.get("giro", "")
    tipo_partida = (ref["tipo_partida"] if ref else "Construcción")
    tipo_interno = data.get("tipo_interno") or (ref["tipo_interno"] if ref else None)
    # columnas dueñas: en interno se llena departamento; en obra, proveedor
    prov_val = proveedor if mundo == "obra" else None
    depto_val = proveedor if mundo == "interno" else None
    # código provisional. Antes usaba resolución de 1 segundo, así que dos
    # propuestas que llegaran en el mismo segundo (doble clic, reintento de
    # red) se quedaban con el MISMO código — bug real detectado en
    # producción: 5 propuestas idénticas con el código PROP-0921230448.
    # Ahora usa microsegundos (colisión prácticamente imposible) y además
    # revisa que no exista ya ese código antes de usarlo, por si acaso.
    base_codigo = "PROP-" + datetime.datetime.now().strftime("%m%d%H%M%S%f")
    codigo = base_codigo
    sufijo = 1
    while db.execute("SELECT 1 FROM actividades WHERE codigo=?", (codigo,)).fetchone():
        sufijo += 1
        codigo = f"{base_codigo}-{sufijo}"
    cur = db.execute(
        """INSERT INTO actividades
        (codigo,bloque,area,giro,proveedor,departamento,mundo,tipo_interno,partida,tipo_partida,aplica,avance,
         avance_decl,f_inicio,f_fin,estatus,definido,definido_por,nota_proveedor,
         origen,estado_val,creado_por,reconocida,reconocida_por,reconocida_fecha,fuera_zona,requiere_pruebas,actualizado)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (codigo, data.get("bloque"), data.get("area"), giro, prov_val, depto_val, mundo, tipo_interno, partida,
         tipo_partida, "SÍ", 0, int(data.get("avance_decl") or 0),
         data.get("f_inicio"), data.get("f_fin"), "Pendiente",
         data.get("definido", "NO"), data.get("definido_por"), data.get("nota_proveedor"),
         "propuesta", "propuesta", usuario,
         "SÍ", usuario, datetime.date.today().isoformat(), fuera_zona,
         "SÍ" if mundo == "obra" else "NO",
         datetime.datetime.now().isoformat(timespec="seconds")))
    detalle = f"Propuso nueva actividad: {partida[:60]}"
    if fuera_zona:
        detalle = f"⚠️ FUERA DE SU ZONA ({data.get('bloque')}): {partida[:50]}"
    db.execute("INSERT INTO avisos (fecha,proveedor,usuario,tipo,detalle,visto) VALUES (?,?,?,?,?,0)",
               (datetime.datetime.now().isoformat(timespec="seconds"), proveedor, usuario,
                "propuesta", detalle))
    db.commit()
    enviar_push_admins(f"🆕 {proveedor}", detalle, "/validacion")
    return jsonify({"ok": True, "codigo": codigo, "fuera_zona": bool(fuera_zona)})


@app.route("/api/validacion/atencion")
@requiere_gestor
def api_val_atencion():
    """Actividades que necesitan tu atención: rechazadas y las que un proveedor no reconoce."""
    db = get_db()
    rechazadas = db.execute(
        "SELECT * FROM actividades WHERE (estado_val='rechazada' OR avance_decl_rechazado=1) AND (eliminada IS NULL OR eliminada=0) ORDER BY proveedor, actualizado DESC").fetchall()
    no_recon = db.execute(
        "SELECT * FROM actividades WHERE no_reconocida_nota IS NOT NULL AND no_reconocida_nota<>'' "
        "AND reconocida<>'SÍ' AND (eliminada IS NULL OR eliminada=0) ORDER BY proveedor, actualizado DESC").fetchall()
    return jsonify({
        "rechazadas": [dict(r) for r in rechazadas],
        "no_reconocidas": [dict(r) for r in no_recon],
    })


@app.route("/api/validacion/reactivar/<int:aid>", methods=["POST"])
@requiere_gestor
def api_val_reactivar(aid):
    """Reactiva una actividad rechazada: vuelve a propuesta para revisarla de nuevo."""
    db = get_db()
    quien_admin = session.get("usuario", "admin")
    registrar_historial(db, aid, "actividad propuesta", "rechazada", "reactivada", quien_admin)
    db.execute("UPDATE actividades SET estado_val='propuesta', avance_decl_rechazado=0, actualizado=? WHERE id=?",
               (datetime.datetime.now().isoformat(timespec="seconds"), aid))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/validacion/limpiar_no_reconocida/<int:aid>", methods=["POST"])
@requiere_gestor
def api_val_limpiar_noreco(aid):
    """Después de corregir una actividad que el proveedor no reconocía, limpias la marca."""
    db = get_db()
    db.execute("UPDATE actividades SET no_reconocida_nota=NULL, reconocida='NO', reconocida_por=NULL, reconocida_fecha=NULL WHERE id=?", (aid,))
    db.commit()
    return jsonify({"ok": True})


# --- Validación (admin) ---
@app.route("/api/validacion/pendientes")
@requiere_gestor
def api_val_pendientes():
    db = get_db()
    props = db.execute(
        "SELECT * FROM actividades WHERE estado_val='propuesta' AND (eliminada IS NULL OR eliminada=0) ORDER BY proveedor, actualizado DESC").fetchall()
    # partidas oficiales con avance declarado distinto al validado
    decl = db.execute(
        "SELECT * FROM actividades WHERE origen='oficial' AND avance_decl IS NOT NULL "
        "AND avance_decl <> avance AND (eliminada IS NULL OR eliminada=0) ORDER BY proveedor, actualizado DESC").fetchall()
    return jsonify({
        "propuestas": [dict(r) for r in props],
        "avances": [dict(r) for r in decl],
    })


@app.route("/api/validacion/aviso_conteo")
@requiere_gestor
def api_val_conteo():
    db = get_db()
    p = db.execute("SELECT COUNT(*) FROM actividades WHERE estado_val='propuesta' "
                   "AND (eliminada IS NULL OR eliminada=0)").fetchone()[0]
    a = db.execute("SELECT COUNT(*) FROM actividades WHERE origen='oficial' "
                   "AND avance_decl IS NOT NULL AND avance_decl <> avance "
                   "AND (eliminada IS NULL OR eliminada=0)").fetchone()[0]
    # "Necesitan tu atención": rechazadas + no reconocidas. Antes el numerito
    # del menú NO las contaba, así que nunca cuadraba con la suma de las 3
    # listas de la pantalla de Validación (por ejemplo 10+1+15 pero el badge
    # decía 24). Ahora sí suma las 3 categorías.
    aten = db.execute(
        "SELECT COUNT(*) FROM actividades WHERE "
        "((estado_val='rechazada' OR avance_decl_rechazado=1) "
        " OR (no_reconocida_nota IS NOT NULL AND no_reconocida_nota<>'' AND reconocida<>'SÍ')) "
        "AND (eliminada IS NULL OR eliminada=0)").fetchone()[0]
    return jsonify({"propuestas": p, "avances": a, "atencion": aten, "total": p + a + aten})


@app.route("/api/validacion/aprobar/<int:aid>", methods=["POST"])
@requiere_gestor
def api_val_aprobar(aid):
    db = get_db()
    a = db.execute("SELECT * FROM actividades WHERE id=?", (aid,)).fetchone()
    if not a:
        return jsonify({"error": "no existe"}), 404
    data = request.get_json() or {}
    quien_admin = session.get("usuario", "admin")
    if a["estado_val"] == "propuesta":
        # aprobar la tarea nueva: pasa a oficial y validada
        av = a["avance_decl"] if a["avance_decl"] is not None else 0
        registrar_historial(db, aid, "actividad propuesta", "propuesta", "aprobada y validada", quien_admin)
        db.execute("UPDATE actividades SET origen='oficial', estado_val='validado', "
                   "avance=?, estatus=?, avance_decl_rechazado=0, rechazo_motivo=NULL, actualizado=? WHERE id=?",
                   (av, estatus_por_avance(av),
                    datetime.datetime.now().isoformat(timespec="seconds"), aid))
    else:
        # firmar el avance declarado: el avance oficial toma el valor declarado
        av = a["avance_decl"] if a["avance_decl"] is not None else a["avance"]
        registrar_historial(db, aid, "avance validado", a["avance"], av, quien_admin)
        db.execute("UPDATE actividades SET avance=?, estatus=?, avance_decl_rechazado=0, rechazo_motivo=NULL, actualizado=? WHERE id=?",
                   (av, estatus_por_avance(av),
                    datetime.datetime.now().isoformat(timespec="seconds"), aid))
    db.commit()
    crear_validacion_interna_si_aplica(db, aid)
    return jsonify({"ok": True})


@app.route("/api/validacion/rechazar/<int:aid>", methods=["POST"])
@requiere_gestor
def api_val_rechazar(aid):
    db = get_db()
    a = db.execute("SELECT * FROM actividades WHERE id=?", (aid,)).fetchone()
    if not a:
        return jsonify({"error": "no existe"}), 404
    data = request.get_json() or {}
    motivo = (data.get("motivo") or "Rechazado por administración").strip()
    quien_admin = session.get("usuario", "admin")
    fecha_hoy = datetime.date.today().isoformat()
    
    if a["estado_val"] == "propuesta":
        # rechazar tarea nueva: se marca rechazada con motivo
        registrar_historial(db, aid, "actividad propuesta", "propuesta", f"rechazada: {motivo}", quien_admin)
        db.execute("UPDATE actividades SET estado_val='rechazada', rechazo_motivo=?, rechazado_por=?, "
                   "rechazado_fecha=?, actualizado=? WHERE id=?",
                   (motivo, quien_admin, fecha_hoy, datetime.datetime.now().isoformat(timespec="seconds"), aid))
    else:
        # rechazar avance declarado: se guarda el rechazo con su motivo para que el proveedor se entere en su portal
        registrar_historial(db, aid, "avance reportado", a["avance_decl"], f"rechazado por admin ({motivo})", quien_admin)
        db.execute("UPDATE actividades SET avance_decl_rechazado=1, rechazo_motivo=?, rechazado_por=?, "
                   "rechazado_fecha=?, avance_decl=NULL, avance_decl_por=NULL, avance_decl_fecha=NULL, "
                   "actualizado=? WHERE id=?",
                   (motivo, quien_admin, fecha_hoy, datetime.datetime.now().isoformat(timespec="seconds"), aid))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/validacion/aprobar_todas", methods=["POST"])
@requiere_gestor
def api_val_aprobar_todas():
    db = get_db()
    data = request.get_json() or {}
    ids = data.get("ids", [])
    if not ids:
        return jsonify({"error": "sin ids"}), 400
    n = 0
    quien_admin = session.get("usuario", "admin")
    for aid in ids:
        a = db.execute("SELECT * FROM actividades WHERE id=?", (aid,)).fetchone()
        if not a:
            continue
        av = a["avance_decl"] if a["avance_decl"] is not None else a["avance"]
        if a["estado_val"] == "propuesta":
            registrar_historial(db, aid, "actividad propuesta", "propuesta", "aprobada y validada", quien_admin)
            db.execute("UPDATE actividades SET origen='oficial', estado_val='validado', "
                       "avance=?, estatus=?, avance_decl_rechazado=0, rechazo_motivo=NULL, actualizado=? WHERE id=?",
                       (av, estatus_por_avance(av),
                        datetime.datetime.now().isoformat(timespec="seconds"), aid))
        else:
            registrar_historial(db, aid, "avance validado", a["avance"], av, quien_admin)
            db.execute("UPDATE actividades SET avance=?, estatus=?, avance_decl_rechazado=0, rechazo_motivo=NULL, actualizado=? WHERE id=?",
                       (av, estatus_por_avance(av),
                        datetime.datetime.now().isoformat(timespec="seconds"), aid))
        n += 1
    db.commit()
    for aid in ids:
        crear_validacion_interna_si_aplica(db, aid)
    return jsonify({"ok": True, "aprobadas": n})


# ----------------------------------------------------------------------------
# API — Dependencias (Modelo por área)
# ----------------------------------------------------------------------------
@app.route("/api/dependencias", methods=["GET"])
@requiere_login
def api_dependencias_list():
    db = get_db()
    area = request.args.get("area")
    q = "SELECT * FROM dependencias"
    args = []
    if area:
        q += " WHERE area=?"; args.append(area)
    q += " ORDER BY area, id"
    rows = db.execute(q, args).fetchall()
    
    salida = []
    for r in rows:
        d = dict(r)
        st = evaluar_dependencias_area(db, r["area"], r["tipo_sucesor"])
        d["evaluacion"] = st
        salida.append(d)
    return jsonify(salida)


@app.route("/api/dependencias", methods=["POST"])
@requiere_gestor
def api_dependencias_crear():
    db = get_db()
    data = request.get_json() or {}
    area = (data.get("area") or "").strip()
    sucesor = (data.get("tipo_sucesor") or "").strip()
    predecesores = (data.get("tipos_predecesores") or "").strip()
    umbral = max(1, min(100, int(data.get("umbral") or 100)))
    paralelo = 1 if data.get("permite_paralelo") else 0
    
    if not area or not sucesor or not predecesores:
        return jsonify({"error": "Área, tipo sucesor y tipos predecesores son obligatorios"}), 400
        
    db.execute(
        "INSERT INTO dependencias (area, tipo_sucesor, tipos_predecesores, umbral, permite_paralelo, creado) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        (area, sucesor, predecesores, umbral, paralelo, datetime.datetime.now().isoformat(timespec="seconds"))
    )
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/dependencias/<int:did>", methods=["DELETE"])
@requiere_gestor
def api_dependencias_borrar(did):
    db = get_db()
    db.execute("DELETE FROM dependencias WHERE id=?", (did,))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/dependencias/<int:did>/forzar_liberacion", methods=["POST"])
@requiere_gestor
def api_dependencias_forzar(did):
    db = get_db()
    data = request.get_json() or {}
    nota = (data.get("nota") or "").strip()
    if not nota:
        return jsonify({"error": "Se requiere una nota justificando la liberación anticipada"}), 400
        
    admin_user = session.get("usuario", "admin")
    fecha_hoy = datetime.date.today().isoformat()
    
    dep = db.execute("SELECT * FROM dependencias WHERE id=?", (did,)).fetchone()
    if not dep:
        return jsonify({"error": "Dependencia no encontrada"}), 404
        
    db.execute(
        "UPDATE dependencias SET liberacion_forzada=1, forzada_por=?, forzada_fecha=?, forzada_nota=? WHERE id=?",
        (admin_user, fecha_hoy, nota, did)
    )
    db.execute(
        "INSERT INTO avisos (fecha, proveedor, usuario, tipo, detalle, visto) VALUES (?, ?, ?, ?, ?, 0)",
        (datetime.datetime.now().isoformat(timespec="seconds"), "Admin", admin_user, "dependencia",
         f"Liberación anticipada forzada de {dep['tipo_sucesor']} en {dep['area']}: {nota}")
    )
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/dependencias/<int:did>/deshacer_forzar", methods=["POST"])
@requiere_gestor
def api_dependencias_deshacer_forzar(did):
    db = get_db()
    db.execute("UPDATE dependencias SET liberacion_forzada=0, forzada_por=NULL, forzada_fecha=NULL, forzada_nota=NULL WHERE id=?", (did,))
    db.commit()
    return jsonify({"ok": True})


# ----------------------------------------------------------------------------
# API — Reportes para Departamentos Internos (PDF / Excel)
# ----------------------------------------------------------------------------
@app.route("/api/hoja_departamento.pdf")
@requiere_login
def hoja_departamento_pdf():
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    
    depto = request.args.get("departamento", "")
    titulo_txt = request.args.get("titulo", "")
    
    db = get_db()
    cond = ["(mundo='interno' OR departamento IS NOT NULL)", "(eliminada IS NULL OR eliminada=0)"]
    args = []
    if depto:
        cond.append("departamento = ?"); args.append(depto)
    for campo in ("bloque", "area", "estatus"):
        v = request.args.get(campo)
        if v:
            cond.append(f"{campo} = ?"); args.append(v)
            
    q = f"SELECT * FROM actividades WHERE {' AND '.join(cond)} ORDER BY bloque, area, id"
    rows = db.execute(q, args).fetchall()
    
    hoy_d = datetime.date.today()
    out = ruta_salida_temporal("hoja_departamento.pdf")
    doc = SimpleDocTemplate(out, pagesize=landscape(letter),
                            leftMargin=1*cm, rightMargin=1*cm, topMargin=1*cm, bottomMargin=1*cm)
    styles = getSampleStyleSheet()
    titulo = ParagraphStyle("t", parent=styles["Title"], fontSize=13, textColor=colors.HexColor("#3A2C52"))
    small = ParagraphStyle("s", parent=styles["Normal"], fontSize=7, leading=8)
    
    torre_txt = "Torre A" if proyecto_actual() == "torre_a" else "Torre B"
    encabezado = titulo_txt or (f"Hoja de actualización de avance — Departamento: {depto}" if depto else "Reporte de actividades — Departamentos Internos HAP")
    elems = [Paragraph(encabezado, titulo),
             Paragraph(f"Fecha: {hoy_d.strftime('%d/%m/%Y')} · Unidad Quirúrgica {torre_txt} · Uso interno HAP", styles["Normal"]),
             Spacer(1, 8)]
             
    data = [["Código", "Bloque / Área", "Tipo tarea", "Actividad / Partida", "Avance\nactual", "NUEVO\navance", "Fecha\ncompromiso", "Retra-\nsada", "Observaciones"]]
    filas_retraso = []
    
    for idx, r in enumerate(rows):
        fin = parse_date(r["f_fin"])
        retrasada = fin and (r["avance"] or 0) < 100 and (fin - hoy_d).days < 0
        bloque_area = f"{r['bloque'] or ''}\n{r['area'] or ''}"
        data.append([
            r["codigo"] or "",
            Paragraph(bloque_area, small),
            Paragraph(r["tipo_interno"] or r["tipo"] or "Instalación", small),
            Paragraph(r["partida"] or "", small),
            f"{r['avance'] or 0}%",
            "______",
            r["f_fin"] or "__________",
            "SÍ" if retrasada else "",
            "________________"
        ])
        if retrasada:
            filas_retraso.append(idx + 1)
            
    t = Table(data, colWidths=[2*cm, 4*cm, 2.8*cm, 6.2*cm, 1.5*cm, 1.6*cm, 2.2*cm, 1.3*cm, 4.5*cm], repeatRows=1)
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3A2C52")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9F7FB")]),
    ]
    for fr in filas_retraso:
        estilo.append(("BACKGROUND", (0, fr), (-1, fr), colors.HexColor("#F8D7DA")))
    t.setStyle(TableStyle(estilo))
    elems.append(t)
    doc.build(elems)
    safe = "".join(ch for ch in (depto or "departamento") if ch.isalnum() or ch in " _-")[:30].strip() or "depto"
    return send_file(out, as_attachment=True, download_name=f"HOJA_DEPTO_{safe}_{hoy()}.pdf")


@app.route("/api/hoja_departamento.xlsx")
@requiere_login
def hoja_departamento_xlsx():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    
    depto = request.args.get("departamento", "")
    db = get_db()
    cond = ["(mundo='interno' OR departamento IS NOT NULL)", "(eliminada IS NULL OR eliminada=0)"]
    args = []
    if depto:
        cond.append("departamento = ?"); args.append(depto)
    for campo in ("bloque", "area", "estatus"):
        v = request.args.get(campo)
        if v:
            cond.append(f"{campo} = ?"); args.append(v)
            
    q = f"SELECT * FROM actividades WHERE {' AND '.join(cond)} ORDER BY bloque, area, id"
    rows = db.execute(q, args).fetchall()
    
    hoy_d = datetime.date.today()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Actualización Depto"
    encabezado = f"HOJA DE ACTUALIZACIÓN DE AVANCE — DEPARTAMENTO: {depto or 'INTERNO'}"
    ws["A1"] = encabezado
    ws["A1"].font = Font(bold=True, size=13, color="3A2C52")
    ws.merge_cells("A1:I1")
    ws["A2"] = f"Fecha: {hoy_d.strftime('%d/%m/%Y')}   ·   Uso interno HAP — Anotar nuevo avance y fecha compromiso."
    ws["A2"].font = Font(size=9, italic=True, color="555555")
    ws.merge_cells("A2:I2")
    
    cols = ["Código", "Bloque", "Área", "Tipo Tarea", "Actividad / Partida", "Avance actual",
            "NUEVO avance %", "Fecha compromiso", "Observaciones"]
    hdr = 4
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=hdr, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3A2C52")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    
    i = hdr + 1
    for r in rows:
        vals = [r["codigo"], r["bloque"], r["area"], r["tipo_interno"] or r["tipo"] or "Instalación",
                r["partida"], f"{r['avance'] or 0}%", "", r["f_fin"] or "", ""]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.border = bd
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        i += 1
    anchos = [10, 18, 22, 16, 38, 12, 13, 16, 28]
    for j, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A5"
    safe = "".join(ch for ch in (depto or "departamento") if ch.isalnum() or ch in " _-")[:30].strip() or "depto"
    out = ruta_salida_temporal("hoja_departamento.xlsx")
    wb.save(out)
    return send_file(out, as_attachment=True, download_name=f"HOJA_DEPTO_{safe}_{hoy()}.xlsx")


# ----------------------------------------------------------------------------
# API — Ajuste de avances a múltiplos de 10 (una sola vez, solo admin)
# Pedido de proveedores y arquitecto: capturar avances de 10 en 10.
# Regla: toda actividad vigente de OBRA con avance entre 1 y 99 que no sea
# múltiplo de 10 sube al siguiente múltiplo de 10 (25->30, 75->80, 1->10).
# 0 y 100 no se tocan. No toca eliminadas, "no aplica", propuestas ni Interno.
# ----------------------------------------------------------------------------
def _siguiente_decena(av):
    av = int(av or 0)
    if av <= 0 or av >= 100 or av % 10 == 0:
        return av
    return min(100, ((av // 10) + 1) * 10)


def _filas_vigentes_obra(db):
    return db.execute(
        "SELECT id, codigo, bloque, area, proveedor, departamento, avance, avance_decl "
        "FROM actividades WHERE (eliminada IS NULL OR eliminada=0) "
        "AND (aplica IS NULL OR aplica<>'NO') "
        "AND (estado_val IS NULL OR estado_val='validado') "
        "AND (mundo='obra' OR mundo IS NULL)").fetchall()


def _calcular_ajuste_decenas(db):
    filas = _filas_vigentes_obra(db)
    total = len(filas)
    suma_antes = sum((f["avance"] or 0) for f in filas)
    cambios, transiciones, por_resp = [], {}, {}
    suma_despues = 0
    for f in filas:
        antes = f["avance"] or 0
        despues = _siguiente_decena(antes)
        suma_despues += despues
        resp = f["proveedor"] or f["departamento"] or "Sin asignar"
        r = por_resp.setdefault(resp, {"nombre": resp, "n": 0, "suma_antes": 0, "suma_despues": 0, "cambian": 0})
        r["n"] += 1; r["suma_antes"] += antes; r["suma_despues"] += despues
        if despues != antes:
            r["cambian"] += 1
            cambios.append((f, antes, despues))
            transiciones[(antes, despues)] = transiciones.get((antes, despues), 0) + 1
    resumen = {
        "total_vigentes": total,
        "avance_actual": round(suma_antes / total, 1) if total else 0,
        "avance_proyectado": round(suma_despues / total, 1) if total else 0,
        "n_cambian": len(cambios),
        "transiciones": [{"de": k[0], "a": k[1], "n": v} for k, v in sorted(transiciones.items())],
        "por_responsable": sorted(
            [{"nombre": r["nombre"], "n": r["n"], "cambian": r["cambian"],
              "antes": round(r["suma_antes"] / r["n"], 1),
              "despues": round(r["suma_despues"] / r["n"], 1)} for r in por_resp.values()],
            key=lambda x: -x["n"]),
    }
    return resumen, cambios


@app.route("/api/ajuste_decenas/vista_previa")
@requiere_admin
def api_ajuste_decenas_vista_previa():
    resumen, _ = _calcular_ajuste_decenas(get_db())
    return jsonify(resumen)


@app.route("/api/ajuste_decenas/aplicar", methods=["POST"])
@requiere_admin
def api_ajuste_decenas_aplicar():
    data = request.get_json(silent=True) or {}
    if data.get("confirmar") is not True:
        return jsonify({"error": "Falta confirmar el ajuste."}), 400
    db = get_db()
    resumen_antes, cambios = _calcular_ajuste_decenas(db)
    if not cambios:
        return jsonify({"ok": True, "cambiadas": 0, "mensaje": "No había nada que ajustar: todo ya está en múltiplos de 10.",
                        "avance_actual": resumen_antes["avance_actual"]})
    # respaldo automático ANTES de tocar nada; si falla, no se aplica
    respaldo = generar_respaldo_bd()
    if not respaldo:
        return jsonify({"error": "No se pudo generar el respaldo previo; no se aplicó ningún cambio."}), 500
    usuario = session.get("usuario") or "admin"
    ahora = datetime.datetime.now().isoformat(timespec="seconds")
    try:
        for f, antes, despues in cambios:
            nuevo_decl = despues if (f["avance_decl"] is not None and f["avance_decl"] == antes) else f["avance_decl"]
            db.execute("UPDATE actividades SET avance=?, estatus=?, avance_decl=?, actualizado=? WHERE id=?",
                       (despues, estatus_por_avance(despues), nuevo_decl, ahora, f["id"]))
            registrar_historial(db, f["id"], "avance (ajuste a múltiplos de 10)", antes, despues, usuario)
        db.commit()
    except Exception as e:
        db.rollback()
        return jsonify({"error": f"No se pudo aplicar el ajuste: {e}"}), 500
    resumen_despues, _ = _calcular_ajuste_decenas(db)
    return jsonify({"ok": True, "cambiadas": len(cambios),
                    "avance_antes": resumen_antes["avance_actual"],
                    "avance_actual": resumen_despues["avance_actual"],
                    "respaldo": os.path.basename(respaldo)})


# ----------------------------------------------------------------------------
# Regresar a OBRA actividades que se movieron por error a INTERNO (solo admin)
# Se apoya en el historial: cada movimiento Obra->Interno dejó guardado el valor
# anterior de proveedor, departamento, tipo, etc., así que se deshace con precisión
# (sin tocar lo que ya era Interno desde antes ni lo de otros proveedores).
# ----------------------------------------------------------------------------
_CAMPOS_MOVIMIENTO = ("mundo", "proveedor", "departamento", "tipo_interno", "tipo_partida",
                      "requiere_pruebas", "valida_depto")


def _norm_nombre(t):
    return " ".join(str(t or "").split()).casefold()


def _valor_hist(v):
    """El historial guarda str(None) como 'None'; se regresa a NULL real."""
    return None if v in (None, "", "None") else v


def _candidatas_regreso_obra(db, proveedores):
    """Actividades hoy en Interno que antes eran de Obra de alguno de `proveedores`.
    Devuelve lista de dicts con la actividad y los valores a restaurar."""
    quer = {_norm_nombre(p) for p in proveedores if str(p or "").strip()}
    salida = []
    filas = db.execute(
        "SELECT * FROM actividades WHERE mundo='interno' AND (eliminada IS NULL OR eliminada=0)").fetchall()
    for a in filas:
        mov = db.execute(
            "SELECT fecha FROM historial WHERE actividad_id=? AND campo='mundo' "
            "AND valor_antes='obra' AND valor_despues='interno' ORDER BY fecha DESC, id DESC LIMIT 1",
            (a["id"],)).fetchone()
        if not mov:
            continue
        f0 = datetime.datetime.fromisoformat(mov["fecha"])
        f1 = (f0 + datetime.timedelta(seconds=3)).isoformat(timespec="seconds")
        rows = db.execute(
            "SELECT campo, valor_antes FROM historial WHERE actividad_id=? AND fecha>=? AND fecha<=? "
            "AND campo IN (%s) ORDER BY id" % ",".join("?" * len(_CAMPOS_MOVIMIENTO)),
            (a["id"], mov["fecha"], f1, *_CAMPOS_MOVIMIENTO)).fetchall()
        restaurar = {}
        for r in rows:
            restaurar.setdefault(r["campo"], _valor_hist(r["valor_antes"]))
        restaurar["mundo"] = "obra"
        prov_orig = restaurar.get("proveedor") if "proveedor" in restaurar else a["proveedor"]
        if _norm_nombre(prov_orig) not in quer:
            continue
        restaurar.setdefault("departamento", None)
        restaurar["departamento"] = None
        restaurar["tipo_interno"] = None
        restaurar["proveedor"] = prov_orig
        salida.append({"a": a, "restaurar": restaurar, "proveedor_original": prov_orig, "fecha_movimiento": mov["fecha"]})
    return salida


def _proveedores_regreso(request_args_or_data):
    crudo = request_args_or_data
    if isinstance(crudo, str):
        crudo = [x for x in crudo.split(",")]
    lista = [str(x).strip() for x in (crudo or []) if str(x).strip()]
    return lista or ["Marquinox", "Muñoz"]


@app.route("/api/regresar_a_obra/vista_previa")
@requiere_admin
def api_regresar_obra_vista_previa():
    db = get_db()
    provs = _proveedores_regreso(request.args.get("proveedores", ""))
    cand = _candidatas_regreso_obra(db, provs)
    por_prov = {}
    for c in cand:
        por_prov[c["proveedor_original"]] = por_prov.get(c["proveedor_original"], 0) + 1
    internas = db.execute(
        "SELECT COUNT(*) FROM actividades WHERE mundo='interno' AND (eliminada IS NULL OR eliminada=0)").fetchone()[0]
    return jsonify({
        "proveedores_buscados": provs,
        "total": len(cand),
        "por_proveedor": [{"proveedor": k, "n": v} for k, v in sorted(por_prov.items())],
        "internas_total": internas,
        "se_quedan_en_interno": internas - len(cand),
        "actividades": [{
            "id": c["a"]["id"], "codigo": c["a"]["codigo"], "partida": c["a"]["partida"],
            "area": c["a"]["area"], "avance": c["a"]["avance"] or 0,
            "departamento_actual": c["a"]["departamento"], "proveedor_original": c["proveedor_original"],
            "movida": c["fecha_movimiento"]} for c in cand],
    })


@app.route("/api/regresar_a_obra/aplicar", methods=["POST"])
@requiere_admin
def api_regresar_obra_aplicar():
    data = request.get_json(silent=True) or {}
    if data.get("confirmar") is not True:
        return jsonify({"error": "Falta confirmar."}), 400
    db = get_db()
    provs = _proveedores_regreso(data.get("proveedores") or [])
    cand = _candidatas_regreso_obra(db, provs)
    if not cand:
        return jsonify({"ok": True, "regresadas": 0, "mensaje": "No había actividades por regresar."})
    respaldo = generar_respaldo_bd()
    if not respaldo:
        return jsonify({"error": "No se pudo generar el respaldo previo; no se cambió nada."}), 500
    usuario = session.get("usuario") or "admin"
    ahora = datetime.datetime.now().isoformat(timespec="seconds")
    try:
        for c in cand:
            a, rest = c["a"], c["restaurar"]
            sets, args = [], []
            for campo in _CAMPOS_MOVIMIENTO:
                if campo in rest:
                    sets.append(f"{campo}=?"); args.append(rest[campo])
                    registrar_historial(db, a["id"], f"{campo} (regreso a Obra)", a[campo], rest[campo], usuario)
            sets.append("actualizado=?"); args.append(ahora); args.append(a["id"])
            db.execute(f"UPDATE actividades SET {','.join(sets)} WHERE id=?", args)
        db.commit()
    except Exception as e:
        db.rollback()
        return jsonify({"error": f"No se pudo regresar: {e}"}), 500
    por_prov = {}
    for c in cand:
        por_prov[c["proveedor_original"]] = por_prov.get(c["proveedor_original"], 0) + 1
    return jsonify({"ok": True, "regresadas": len(cand), "por_proveedor": por_prov,
                    "respaldo": os.path.basename(respaldo)})


# ----------------------------------------------------------------------------
# API — Proyecto activo (Torre B / Torre A). Torre A solo la ve el admin.
# ----------------------------------------------------------------------------
@app.route("/api/proyecto", methods=["GET"])
@requiere_login
def api_proyecto_estado():
    es_admin = session.get("rol") == "admin"
    return jsonify({
        "proyecto": proyecto_actual(),
        "puede_cambiar": es_admin,
        "proyectos": ([{"id": "torre_b", "nombre": "Torre B"},
                       {"id": "torre_a", "nombre": "Torre A"}] if es_admin else []),
    })


@app.route("/api/proyecto", methods=["POST"])
@requiere_admin
def api_proyecto_cambiar():
    data = request.get_json(silent=True) or {}
    p = data.get("proyecto")
    if p not in PROYECTOS:
        return jsonify({"error": "proyecto_invalido"}), 400
    session["proyecto"] = p
    return jsonify({"ok": True, "proyecto": p})


# ----------------------------------------------------------------------------
# API — Respaldos de base de datos
# ----------------------------------------------------------------------------
@app.route("/api/respaldo", methods=["POST"])
@requiere_admin
def api_respaldo():
    res = generar_respaldo_bd()
    if res:
        return jsonify({"ok": True, "archivo": os.path.basename(res)})
    return jsonify({"error": "No se pudo generar el respaldo"}), 500


@app.route("/api/respaldo/descargar", methods=["GET"])
@requiere_admin
def api_respaldo_descargar():
    # Genera una copia consistente (API de backup de SQLite) y la manda al
    # navegador; además queda guardada en el servidor (data/backups).
    dest = generar_respaldo_bd()
    if dest and os.path.exists(dest):
        return send_file(dest, as_attachment=True,
                         download_name=f"{'obra_torre_a_backup_' if proyecto_actual() == 'torre_a' else 'obra_backup_'}{datetime.datetime.now().strftime('%Y-%m-%d_%H%M')}.db")
    return jsonify({"error": "No se pudo generar el respaldo"}), 500


@app.route("/api/validacion/panel")
def validacion_page():
    return render_template("validacion.html")


@app.route("/api/torre-a/catalogo", methods=["GET"])
@requiere_admin
def api_torre_a_catalogo():
    if proyecto_actual() != "torre_a":
        return jsonify({"error": "no_es_torre_a"}), 409
    db = get_db()
    bloques = []
    for b in db.execute("SELECT num,nombre FROM ta_bloques ORDER BY num").fetchall():
        areas = [{"num": a["num"], "clave": "%d.%d" % (b["num"], a["num"]), "nombre": a["nombre"]}
                 for a in db.execute("SELECT num,nombre FROM ta_areas WHERE bloque=? ORDER BY num",
                                     (b["num"],)).fetchall()]
        bloques.append({"num": b["num"], "clave": str(b["num"]), "nombre": b["nombre"], "areas": areas})
    return jsonify({"bloques": bloques,
                    "total_bloques": len(bloques),
                    "total_areas": sum(len(b["areas"]) for b in bloques)})


def _solo_torre_a():
    """None si se esta en Torre A; si no, respuesta 409 (nunca escribir en la base de Torre B)."""
    if proyecto_actual() != "torre_a":
        return jsonify({"error": "no_es_torre_a"}), 409
    return None


def _codigo_ta(r):
    base = "%d.%d.%d.%d" % (r["bloque"], r["area"], r["proveedor"], r["actividad"])
    return base if r["detalle"] == 0 else "%s.%d" % (base, r["detalle"])


def _recalcular_actividad_ta(db, b, a, p, act):
    """El avance de la actividad general = promedio simple de sus detalles vigentes.
    Si no tiene detalles, conserva el avance que se le capturo directo."""
    fila = db.execute(
        "SELECT AVG(avance) AS m, COUNT(*) AS n FROM ta_actividades "
        "WHERE bloque=? AND area=? AND proveedor=? AND actividad=? AND detalle>0 AND eliminada=0",
        (b, a, p, act)).fetchone()
    if fila["n"] > 0:
        db.execute(
            "UPDATE ta_actividades SET avance=? WHERE bloque=? AND area=? AND proveedor=? "
            "AND actividad=? AND detalle=0",
            (int(fila["m"] + 0.5), b, a, p, act))


def _avance_valido(v):
    try:
        n = int(v)
    except (TypeError, ValueError):
        return None
    return n if 0 <= n <= 100 else None


@app.route("/api/torre-a/proveedores", methods=["GET"])
@requiere_admin
def api_ta_proveedores():
    err = _solo_torre_a()
    if err:
        return err
    filas = get_db().execute("SELECT num,nombre FROM ta_proveedores ORDER BY num").fetchall()
    return jsonify({"proveedores": [{"num": f["num"], "nombre": f["nombre"]} for f in filas]})


@app.route("/api/torre-a/proveedores", methods=["POST"])
@requiere_admin
def api_ta_proveedor_alta():
    err = _solo_torre_a()
    if err:
        return err
    nombre = " ".join(str((request.get_json(silent=True) or {}).get("nombre", "")).split())
    if not nombre or len(nombre) > 80:
        return jsonify({"error": "Escribe el nombre del proveedor (maximo 80 letras)"}), 400
    db = get_db()
    for f in db.execute("SELECT num,nombre FROM ta_proveedores").fetchall():
        if sin_acentos(f["nombre"]) == sin_acentos(nombre):
            return jsonify({"error": "Ese proveedor ya existe con el numero %d" % f["num"]}), 409
    cur = db.execute("INSERT INTO ta_proveedores (nombre) VALUES (?)", (nombre,))
    db.commit()
    return jsonify({"ok": True, "num": cur.lastrowid, "nombre": nombre})


@app.route("/api/torre-a/actividades", methods=["GET"])
@requiere_admin
def api_ta_actividades():
    err = _solo_torre_a()
    if err:
        return err
    db = get_db()
    filas = db.execute(
        "SELECT a.*, b.nombre AS bloque_nombre, ar.nombre AS area_nombre, p.nombre AS proveedor_nombre "
        "FROM ta_actividades a "
        "JOIN ta_bloques b ON b.num=a.bloque "
        "JOIN ta_areas ar ON ar.bloque=a.bloque AND ar.num=a.area "
        "JOIN ta_proveedores p ON p.num=a.proveedor "
        "WHERE a.eliminada=0 ORDER BY a.bloque,a.area,a.proveedor,a.actividad,a.detalle").fetchall()
    con_detalles = {(f["bloque"], f["area"], f["proveedor"], f["actividad"])
                    for f in filas if f["detalle"] > 0}
    salida, hojas = [], []
    for f in filas:
        llave = (f["bloque"], f["area"], f["proveedor"], f["actividad"])
        es_hoja = f["detalle"] > 0 or llave not in con_detalles
        if es_hoja:
            hojas.append(f["avance"])
        salida.append({
            "id": f["id"], "codigo": _codigo_ta(f), "nivel": 4 if f["detalle"] == 0 else 5,
            "bloque": f["bloque"], "bloque_nombre": f["bloque_nombre"],
            "area": f["area"], "area_nombre": f["area_nombre"],
            "proveedor": f["proveedor"], "proveedor_nombre": f["proveedor_nombre"],
            "actividad": f["actividad"], "detalle": f["detalle"],
            "nombre": f["nombre"], "avance": f["avance"], "editable": es_hoja,
        })
    return jsonify({
        "actividades": salida,
        "total": len(salida),
        "avance_global": int(sum(hojas) / len(hojas) + 0.5) if hojas else 0,
    })


@app.route("/api/torre-a/actividades", methods=["POST"])
@requiere_admin
def api_ta_actividad_alta():
    """Alta de actividad general (bloque+area+proveedor+nombre) o de un detalle (padre_id+nombre).
    Los numeros de actividad y de detalle los asigna el sistema."""
    err = _solo_torre_a()
    if err:
        return err
    d = request.get_json(silent=True) or {}
    nombre = " ".join(str(d.get("nombre", "")).split())
    if not nombre or len(nombre) > 200:
        return jsonify({"error": "Escribe el nombre (maximo 200 letras)"}), 400
    db = get_db()
    ahora = datetime.datetime.now().isoformat(timespec="seconds")
    padre_id = d.get("padre_id")
    if padre_id:
        padre = db.execute("SELECT * FROM ta_actividades WHERE id=? AND detalle=0 AND eliminada=0",
                           (padre_id,)).fetchone()
        if not padre:
            return jsonify({"error": "La actividad general no existe"}), 404
        b, a, p, act = padre["bloque"], padre["area"], padre["proveedor"], padre["actividad"]
    else:
        try:
            b, a, p = int(d.get("bloque")), int(d.get("area")), int(d.get("proveedor"))
        except (TypeError, ValueError):
            return jsonify({"error": "Elige bloque, area y proveedor"}), 400
        if not db.execute("SELECT 1 FROM ta_areas WHERE bloque=? AND num=?", (b, a)).fetchone():
            return jsonify({"error": "Ese bloque/area no existe en el catalogo"}), 400
        if not db.execute("SELECT 1 FROM ta_proveedores WHERE num=?", (p,)).fetchone():
            return jsonify({"error": "Ese proveedor no existe"}), 400
        act = None
    for _intento in range(3):
        try:
            if act is None:
                n_act = db.execute(
                    "SELECT COALESCE(MAX(actividad),0)+1 FROM ta_actividades "
                    "WHERE bloque=? AND area=? AND proveedor=?", (b, a, p)).fetchone()[0]
                n_det = 0
            else:
                n_act = act
                n_det = db.execute(
                    "SELECT COALESCE(MAX(detalle),0)+1 FROM ta_actividades "
                    "WHERE bloque=? AND area=? AND proveedor=? AND actividad=?",
                    (b, a, p, act)).fetchone()[0]
            cur = db.execute(
                "INSERT INTO ta_actividades (bloque,area,proveedor,actividad,detalle,nombre,avance,creada) "
                "VALUES (?,?,?,?,?,?,0,?)", (b, a, p, n_act, n_det, nombre, ahora))
            if n_det > 0:
                _recalcular_actividad_ta(db, b, a, p, n_act)
            db.commit()
            codigo = "%d.%d.%d.%d" % (b, a, p, n_act) + (".%d" % n_det if n_det else "")
            return jsonify({"ok": True, "id": cur.lastrowid, "codigo": codigo})
        except sqlite3.IntegrityError:
            db.rollback()
    return jsonify({"error": "No se pudo asignar el numero; intenta de nuevo"}), 409


@app.route("/api/torre-a/actividades/<int:aid>", methods=["PUT"])
@requiere_admin
def api_ta_actividad_editar(aid):
    err = _solo_torre_a()
    if err:
        return err
    d = request.get_json(silent=True) or {}
    db = get_db()
    f = db.execute("SELECT * FROM ta_actividades WHERE id=? AND eliminada=0", (aid,)).fetchone()
    if not f:
        return jsonify({"error": "No existe"}), 404
    if "nombre" in d:
        nombre = " ".join(str(d["nombre"]).split())
        if not nombre or len(nombre) > 200:
            return jsonify({"error": "Nombre invalido"}), 400
        db.execute("UPDATE ta_actividades SET nombre=? WHERE id=?", (nombre, aid))
    if "avance" in d:
        nuevo = _avance_valido(d["avance"])
        if nuevo is None:
            return jsonify({"error": "El avance debe ser un numero entero de 0 a 100"}), 400
        if f["detalle"] == 0 and db.execute(
                "SELECT 1 FROM ta_actividades WHERE bloque=? AND area=? AND proveedor=? AND actividad=? "
                "AND detalle>0 AND eliminada=0", (f["bloque"], f["area"], f["proveedor"], f["actividad"])).fetchone():
            return jsonify({"error": "Esta actividad tiene detalles: su avance es el promedio de ellos"}), 400
        if nuevo != f["avance"]:
            db.execute("UPDATE ta_actividades SET avance=? WHERE id=?", (nuevo, aid))
            db.execute("INSERT INTO ta_historial (act_id,antes,despues,fecha) VALUES (?,?,?,?)",
                       (aid, f["avance"], nuevo, datetime.datetime.now().isoformat(timespec="seconds")))
            if f["detalle"] > 0:
                _recalcular_actividad_ta(db, f["bloque"], f["area"], f["proveedor"], f["actividad"])
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/torre-a/actividades/<int:aid>", methods=["DELETE"])
@requiere_admin
def api_ta_actividad_baja(aid):
    """Baja logica: el numero queda ocupado para siempre (no se reutiliza)."""
    err = _solo_torre_a()
    if err:
        return err
    db = get_db()
    f = db.execute("SELECT * FROM ta_actividades WHERE id=? AND eliminada=0", (aid,)).fetchone()
    if not f:
        return jsonify({"error": "No existe"}), 404
    if f["detalle"] == 0:
        db.execute("UPDATE ta_actividades SET eliminada=1 WHERE bloque=? AND area=? AND proveedor=? AND actividad=?",
                   (f["bloque"], f["area"], f["proveedor"], f["actividad"]))
    else:
        db.execute("UPDATE ta_actividades SET eliminada=1 WHERE id=?", (aid,))
        _recalcular_actividad_ta(db, f["bloque"], f["area"], f["proveedor"], f["actividad"])
    db.commit()
    return jsonify({"ok": True})


def _nombre_ta(valor, maximo=200):
    n = " ".join(str(valor or "").split())
    return n if n and len(n) <= maximo else None


@app.route("/api/torre-a/bloques", methods=["POST"])
@requiere_admin
def api_ta_bloque_alta():
    """Bloque nuevo (por si el plano cambia). Toma el siguiente numero; los existentes no se mueven."""
    err = _solo_torre_a()
    if err:
        return err
    nombre = _nombre_ta((request.get_json(silent=True) or {}).get("nombre"))
    if not nombre:
        return jsonify({"error": "Escribe el nombre del bloque"}), 400
    db = get_db()
    num = db.execute("SELECT COALESCE(MAX(num),0)+1 FROM ta_bloques").fetchone()[0]
    db.execute("INSERT INTO ta_bloques (num,nombre) VALUES (?,?)", (num, nombre))
    db.commit()
    return jsonify({"ok": True, "num": num})


@app.route("/api/torre-a/bloques/<int:num>", methods=["PUT"])
@requiere_admin
def api_ta_bloque_renombrar(num):
    err = _solo_torre_a()
    if err:
        return err
    nombre = _nombre_ta((request.get_json(silent=True) or {}).get("nombre"))
    if not nombre:
        return jsonify({"error": "Nombre invalido"}), 400
    db = get_db()
    viejo = db.execute("SELECT nombre FROM ta_bloques WHERE num=?", (num,)).fetchone()
    cur = db.execute("UPDATE ta_bloques SET nombre=? WHERE num=?", (nombre, num))
    if viejo:   # las actividades y trazos guardan la etiqueta "01 · Nombre": se actualiza tambien
        _renombrar_etiqueta_ta(db, "bloque", etiqueta_bloque_ta(num, viejo["nombre"]), etiqueta_bloque_ta(num, nombre))
    db.commit()
    return jsonify({"ok": True}) if cur.rowcount else (jsonify({"error": "No existe"}), 404)


@app.route("/api/torre-a/areas", methods=["POST"])
@requiere_admin
def api_ta_area_alta():
    err = _solo_torre_a()
    if err:
        return err
    d = request.get_json(silent=True) or {}
    nombre = _nombre_ta(d.get("nombre"))
    if not nombre:
        return jsonify({"error": "Escribe el nombre del area"}), 400
    try:
        b = int(d.get("bloque"))
    except (TypeError, ValueError):
        return jsonify({"error": "Elige el bloque"}), 400
    db = get_db()
    if not db.execute("SELECT 1 FROM ta_bloques WHERE num=?", (b,)).fetchone():
        return jsonify({"error": "Ese bloque no existe"}), 404
    num = db.execute("SELECT COALESCE(MAX(num),0)+1 FROM ta_areas WHERE bloque=?", (b,)).fetchone()[0]
    db.execute("INSERT INTO ta_areas (bloque,num,nombre) VALUES (?,?,?)", (b, num, nombre))
    db.commit()
    return jsonify({"ok": True, "clave": "%d.%d" % (b, num)})


@app.route("/api/torre-a/areas/<int:bloque>/<int:num>", methods=["PUT"])
@requiere_admin
def api_ta_area_renombrar(bloque, num):
    err = _solo_torre_a()
    if err:
        return err
    nombre = _nombre_ta((request.get_json(silent=True) or {}).get("nombre"))
    if not nombre:
        return jsonify({"error": "Nombre invalido"}), 400
    db = get_db()
    viejo = db.execute("SELECT nombre FROM ta_areas WHERE bloque=? AND num=?", (bloque, num)).fetchone()
    cur = db.execute("UPDATE ta_areas SET nombre=? WHERE bloque=? AND num=?", (nombre, bloque, num))
    if viejo:
        _renombrar_etiqueta_ta(db, "area", etiqueta_area_ta(bloque, num, viejo["nombre"]), etiqueta_area_ta(bloque, num, nombre))
    db.commit()
    return jsonify({"ok": True}) if cur.rowcount else (jsonify({"error": "No existe"}), 404)


def _renombrar_etiqueta_ta(db, columna, viejo, nuevo):
    """Cambia la etiqueta de bloque/area en actividades y en los trazos del plano (columna = 'bloque' | 'area')."""
    if viejo == nuevo or columna not in ("bloque", "area"):
        return
    db.execute("UPDATE actividades SET %s=? WHERE %s=?" % (columna, columna), (nuevo, viejo))
    try:
        db.execute("UPDATE plano_trazo_ub SET %s=? WHERE %s=?" % (columna, columna), (nuevo, viejo))
        db.execute("UPDATE plano_trazos SET nombre=? WHERE nombre=?", (nuevo, viejo))          # contornos que se llamaban como el bloque/area
        if columna == "area":
            db.execute("UPDATE dependencias SET area=? WHERE area=?", (nuevo, viejo))
            db.execute("UPDATE OR IGNORE area_involucrados SET area=? WHERE area=?", (nuevo, viejo))
    except sqlite3.Error:
        pass


@app.route("/api/fecha-entrega", methods=["GET"])
@requiere_login
def api_fecha_entrega():
    return jsonify({"fecha_entrega": fecha_entrega_actual(), "dias_restantes": dias_restantes(),
                    "editable": proyecto_actual() == "torre_a" and session.get("rol") == "admin"})


@app.route("/api/fecha-entrega", methods=["POST"])
@requiere_admin
def api_fecha_entrega_guardar():
    """Solo Torre A: la fecha de entrega de Torre B es fija. Se manda {fecha: 'AAAA-MM-DD'} (vacia = quitar)."""
    err = _solo_torre_a()
    if err:
        return err
    fecha = str((request.get_json(silent=True) or {}).get("fecha") or "").strip()
    if fecha:
        try:
            fecha = datetime.date.fromisoformat(fecha).isoformat()
        except ValueError:
            return jsonify({"error": "Fecha no valida. Usa AAAA-MM-DD, por ejemplo 2026-12-15."}), 400
    db = get_db()
    _config_set(db, "fecha_entrega", fecha)
    db.commit()
    return jsonify({"ok": True, "fecha_entrega": fecha or None, "dias_restantes": dias_restantes()})


def _puntos_validos(pts, minimo=3, maximo=300):
    """Lista de puntos [x,y] con x,y entre 0 y 1 (entre minimo y maximo). Devuelve la lista limpia o None."""
    if not isinstance(pts, list) or not (minimo <= len(pts) <= maximo):
        return None
    limpio = []
    for p in pts:
        try:
            x, y = float(p[0]), float(p[1])
        except (TypeError, ValueError, IndexError, KeyError):
            return None
        if not (0.0 <= x <= 1.0 and 0.0 <= y <= 1.0):
            return None
        limpio.append([round(x, 5), round(y, 5)])
    return limpio


@app.route("/api/torre-a/zonas", methods=["GET"])
@requiere_admin
def api_ta_zonas():
    err = _solo_torre_a()
    if err:
        return err
    filas = get_db().execute(
        "SELECT z.id,z.bloque,z.area,z.puntos,a.nombre AS area_nombre,b.nombre AS bloque_nombre "
        "FROM ta_zonas z JOIN ta_areas a ON a.bloque=z.bloque AND a.num=z.area "
        "JOIN ta_bloques b ON b.num=z.bloque ORDER BY z.bloque,z.area,z.id").fetchall()
    return jsonify({"zonas": [{
        "id": f["id"], "bloque": f["bloque"], "area": f["area"], "clave": "%d.%d" % (f["bloque"], f["area"]),
        "area_nombre": f["area_nombre"], "bloque_nombre": f["bloque_nombre"],
        "puntos": json.loads(f["puntos"])} for f in filas]})


@app.route("/api/torre-a/zonas", methods=["POST"])
@requiere_admin
def api_ta_zona_guardar():
    """Agrega una pieza dibujada a un area (un area puede tener varias piezas).
    Con reemplazar=true primero quita las piezas que ya tenia."""
    err = _solo_torre_a()
    if err:
        return err
    d = request.get_json(silent=True) or {}
    try:
        b, a = int(d.get("bloque")), int(d.get("area"))
    except (TypeError, ValueError):
        return jsonify({"error": "Elige bloque y area"}), 400
    pts = _puntos_validos(d.get("puntos"))
    if pts is None:
        return jsonify({"error": "La zona necesita al menos 3 puntos dentro del plano"}), 400
    db = get_db()
    if not db.execute("SELECT 1 FROM ta_areas WHERE bloque=? AND num=?", (b, a)).fetchone():
        return jsonify({"error": "Esa area no existe en el catalogo"}), 400
    ahora = datetime.datetime.now().isoformat(timespec="seconds")
    if d.get("reemplazar"):
        db.execute("DELETE FROM ta_zonas WHERE bloque=? AND area=?", (b, a))
    cur = db.execute("INSERT INTO ta_zonas (bloque,area,puntos,creada) VALUES (?,?,?,?)",
                     (b, a, json.dumps(pts), ahora))
    db.commit()
    return jsonify({"ok": True, "id": cur.lastrowid})


@app.route("/api/torre-a/bloque-zonas", methods=["GET"])
@requiere_admin
def api_ta_bloque_zonas():
    err = _solo_torre_a()
    if err:
        return err
    filas = get_db().execute(
        "SELECT z.id,z.bloque,z.puntos,b.nombre FROM ta_bloque_zonas z JOIN ta_bloques b ON b.num=z.bloque "
        "ORDER BY z.bloque,z.id").fetchall()
    return jsonify({"bloque_zonas": [{
        "id": f["id"], "bloque": f["bloque"], "bloque_nombre": f["nombre"],
        "puntos": json.loads(f["puntos"])} for f in filas]})


@app.route("/api/torre-a/bloque-zonas", methods=["POST"])
@requiere_admin
def api_ta_bloque_zona_guardar():
    """Agrega una pieza del contorno de un bloque (un bloque puede tener varias piezas)."""
    err = _solo_torre_a()
    if err:
        return err
    d = request.get_json(silent=True) or {}
    try:
        b = int(d.get("bloque"))
    except (TypeError, ValueError):
        return jsonify({"error": "Elige el bloque"}), 400
    pts = _puntos_validos(d.get("puntos"))
    if pts is None:
        return jsonify({"error": "El bloque necesita al menos 3 puntos dentro del plano"}), 400
    db = get_db()
    if not db.execute("SELECT 1 FROM ta_bloques WHERE num=?", (b,)).fetchone():
        return jsonify({"error": "Ese bloque no existe en el catalogo"}), 400
    ahora = datetime.datetime.now().isoformat(timespec="seconds")
    if d.get("reemplazar"):
        db.execute("DELETE FROM ta_bloque_zonas WHERE bloque=?", (b,))
    cur = db.execute("INSERT INTO ta_bloque_zonas (bloque,puntos,creada) VALUES (?,?,?)",
                     (b, json.dumps(pts), ahora))
    db.commit()
    return jsonify({"ok": True, "id": cur.lastrowid})


@app.route("/api/torre-a/bloque-zonas/<int:zid>", methods=["DELETE"])
@requiere_admin
def api_ta_bloque_zona_borrar(zid):
    err = _solo_torre_a()
    if err:
        return err
    db = get_db()
    cur = db.execute("DELETE FROM ta_bloque_zonas WHERE id=?", (zid,))
    db.commit()
    return jsonify({"ok": True}) if cur.rowcount else (jsonify({"error": "No existe"}), 404)


@app.route("/api/torre-a/zonas/<int:zid>", methods=["DELETE"])
@requiere_admin
def api_ta_zona_borrar(zid):
    """Quita solo una pieza del dibujo; las actividades del area no se tocan."""
    err = _solo_torre_a()
    if err:
        return err
    db = get_db()
    cur = db.execute("DELETE FROM ta_zonas WHERE id=?", (zid,))
    db.commit()
    return jsonify({"ok": True}) if cur.rowcount else (jsonify({"error": "No existe"}), 404)


import re as _re_capas
TIPOS_CAPA_TA = {"punto": (1, 1), "linea": (2, 300), "zona": (3, 300)}
CAPAS_SUGERIDAS_TA = (
    ("Contactos", "punto", "#2563eb"), ("Apagadores", "punto", "#7c3aed"),
    ("Luminarias", "punto", "#d97706"), ("Muros con pintura", "linea", "#0d9488"),
    ("Muros con yeso", "linea", "#a16207"), ("Muros con tapiz", "linea", "#db2777"),
    ("Plafón", "zona", "#64748b"),
)


def _datos_capa_ta(db, d, parcial=False):
    """Valida y limpia los campos de una capa. Devuelve (dict, None) o (None, respuesta_error)."""
    out = {}
    if "nombre" in d or not parcial:
        n = _nombre_ta(d.get("nombre"), 60)
        if not n:
            return None, (jsonify({"error": "Escribe el nombre de la capa (maximo 60 letras)"}), 400)
        out["nombre"] = n
    if "color" in d or not parcial:
        c = str(d.get("color") or "#2563eb")
        if not _re_capas.fullmatch(r"#[0-9a-fA-F]{6}", c):
            return None, (jsonify({"error": "Color invalido"}), 400)
        out["color"] = c.lower()
    if "proveedor" in d or not parcial:
        pv = d.get("proveedor")
        if pv in (None, "", 0, "0"):
            out["proveedor"] = None
        else:
            try:
                pv = int(pv)
            except (TypeError, ValueError):
                return None, (jsonify({"error": "Proveedor invalido"}), 400)
            if not db.execute("SELECT 1 FROM ta_proveedores WHERE num=?", (pv,)).fetchone():
                return None, (jsonify({"error": "Ese proveedor no existe"}), 400)
            out["proveedor"] = pv
    if "actividad_id" in d or not parcial:
        ai = d.get("actividad_id")
        if ai in (None, "", 0, "0"):
            out["actividad_id"] = None
        else:
            try:
                ai = int(ai)
            except (TypeError, ValueError):
                return None, (jsonify({"error": "Actividad invalida"}), 400)
            if not db.execute("SELECT 1 FROM ta_actividades WHERE id=? AND eliminada=0", (ai,)).fetchone():
                return None, (jsonify({"error": "Esa actividad no existe"}), 400)
            out["actividad_id"] = ai
    if "descripcion" in d or not parcial:
        ds = " ".join(str(d.get("descripcion") or "").split())
        if len(ds) > 300:
            return None, (jsonify({"error": "La descripcion es muy larga (maximo 300 letras)"}), 400)
        out["descripcion"] = ds
    return out, None


def _nombre_capa_repetido(db, nombre, excepto=None):
    for f in db.execute("SELECT id,nombre FROM ta_capas").fetchall():
        if f["id"] != excepto and sin_acentos(f["nombre"]) == sin_acentos(nombre):
            return True
    return False


@app.route("/api/torre-a/capas", methods=["GET"])
@requiere_admin
def api_ta_capas():
    err = _solo_torre_a()
    if err:
        return err
    db = get_db()
    capas = []
    for c in db.execute(
            "SELECT c.*, p.nombre AS proveedor_nombre, "
            "(SELECT COUNT(*) FROM ta_elementos e WHERE e.capa_id=c.id) AS n "
            "FROM ta_capas c LEFT JOIN ta_proveedores p ON p.num=c.proveedor ORDER BY c.id").fetchall():
        capas.append({"id": c["id"], "nombre": c["nombre"], "tipo": c["tipo"], "color": c["color"],
                      "proveedor": c["proveedor"], "proveedor_nombre": c["proveedor_nombre"],
                      "actividad_id": c["actividad_id"], "descripcion": c["descripcion"] or "", "n": c["n"]})
    elementos = [{"id": e["id"], "capa_id": e["capa_id"], "puntos": json.loads(e["puntos"]), "nota": e["nota"] or ""}
                 for e in db.execute("SELECT * FROM ta_elementos ORDER BY id").fetchall()]
    return jsonify({"capas": capas, "elementos": elementos})


@app.route("/api/torre-a/capas", methods=["POST"])
@requiere_admin
def api_ta_capa_alta():
    err = _solo_torre_a()
    if err:
        return err
    d = request.get_json(silent=True) or {}
    db = get_db()
    if d.get("sugeridas"):
        creadas = 0
        for nombre, tipo, color in CAPAS_SUGERIDAS_TA:
            if _nombre_capa_repetido(db, nombre):
                continue
            db.execute("INSERT INTO ta_capas (nombre,tipo,color,creada) VALUES (?,?,?,?)",
                       (nombre, tipo, color, datetime.datetime.now().isoformat(timespec="seconds")))
            creadas += 1
        db.commit()
        return jsonify({"ok": True, "creadas": creadas})
    if d.get("tipo") not in TIPOS_CAPA_TA:
        return jsonify({"error": "El tipo debe ser punto, linea o zona"}), 400
    datos, e = _datos_capa_ta(db, d)
    if e:
        return e
    if _nombre_capa_repetido(db, datos["nombre"]):
        return jsonify({"error": "Ya existe una capa con ese nombre"}), 409
    cur = db.execute(
        "INSERT INTO ta_capas (nombre,tipo,color,proveedor,actividad_id,descripcion,creada) VALUES (?,?,?,?,?,?,?)",
        (datos["nombre"], d["tipo"], datos["color"], datos["proveedor"], datos["actividad_id"], datos["descripcion"],
         datetime.datetime.now().isoformat(timespec="seconds")))
    db.commit()
    return jsonify({"ok": True, "id": cur.lastrowid})


@app.route("/api/torre-a/capas/<int:cid>", methods=["PUT"])
@requiere_admin
def api_ta_capa_editar(cid):
    err = _solo_torre_a()
    if err:
        return err
    db = get_db()
    if not db.execute("SELECT 1 FROM ta_capas WHERE id=?", (cid,)).fetchone():
        return jsonify({"error": "No existe"}), 404
    datos, e = _datos_capa_ta(db, request.get_json(silent=True) or {}, parcial=True)
    if e:
        return e
    if "nombre" in datos and _nombre_capa_repetido(db, datos["nombre"], excepto=cid):
        return jsonify({"error": "Ya existe una capa con ese nombre"}), 409
    if datos:
        sets = ", ".join("%s=?" % k for k in datos)
        db.execute("UPDATE ta_capas SET %s WHERE id=?" % sets, list(datos.values()) + [cid])
        db.commit()
    return jsonify({"ok": True})


@app.route("/api/torre-a/capas/<int:cid>", methods=["DELETE"])
@requiere_admin
def api_ta_capa_baja(cid):
    """Quita la capa y sus elementos dibujados (no toca las actividades)."""
    err = _solo_torre_a()
    if err:
        return err
    db = get_db()
    if not db.execute("SELECT 1 FROM ta_capas WHERE id=?", (cid,)).fetchone():
        return jsonify({"error": "No existe"}), 404
    db.execute("DELETE FROM ta_elementos WHERE capa_id=?", (cid,))
    db.execute("DELETE FROM ta_capas WHERE id=?", (cid,))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/torre-a/elementos", methods=["POST"])
@requiere_admin
def api_ta_elemento_alta():
    err = _solo_torre_a()
    if err:
        return err
    d = request.get_json(silent=True) or {}
    db = get_db()
    try:
        cid = int(d.get("capa_id"))
    except (TypeError, ValueError):
        return jsonify({"error": "Elige la capa"}), 400
    capa = db.execute("SELECT tipo FROM ta_capas WHERE id=?", (cid,)).fetchone()
    if not capa:
        return jsonify({"error": "La capa no existe"}), 404
    minimo, maximo = TIPOS_CAPA_TA[capa["tipo"]]
    pts = _puntos_validos(d.get("puntos"), minimo, maximo)
    if pts is None:
        return jsonify({"error": "Puntos invalidos para una capa de tipo %s" % capa["tipo"]}), 400
    nota = _nombre_ta(d.get("nota"), 120) or ""
    cur = db.execute("INSERT INTO ta_elementos (capa_id,puntos,nota,creada) VALUES (?,?,?,?)",
                     (cid, json.dumps(pts), nota, datetime.datetime.now().isoformat(timespec="seconds")))
    db.commit()
    return jsonify({"ok": True, "id": cur.lastrowid})


@app.route("/api/torre-a/elementos/<int:eid>", methods=["PUT"])
@requiere_admin
def api_ta_elemento_editar(eid):
    err = _solo_torre_a()
    if err:
        return err
    d = request.get_json(silent=True) or {}
    db = get_db()
    fila = db.execute("SELECT e.id, c.tipo FROM ta_elementos e JOIN ta_capas c ON c.id=e.capa_id WHERE e.id=?",
                      (eid,)).fetchone()
    if not fila:
        return jsonify({"error": "No existe"}), 404
    if "nota" in d:
        db.execute("UPDATE ta_elementos SET nota=? WHERE id=?", (_nombre_ta(d.get("nota"), 120) or "", eid))
    if "puntos" in d:
        minimo, maximo = TIPOS_CAPA_TA[fila["tipo"]]
        pts = _puntos_validos(d["puntos"], minimo, maximo)
        if pts is None:
            return jsonify({"error": "Puntos invalidos"}), 400
        db.execute("UPDATE ta_elementos SET puntos=? WHERE id=?", (json.dumps(pts), eid))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/torre-a/elementos/<int:eid>", methods=["DELETE"])
@requiere_admin
def api_ta_elemento_baja(eid):
    err = _solo_torre_a()
    if err:
        return err
    db = get_db()
    cur = db.execute("DELETE FROM ta_elementos WHERE id=?", (eid,))
    db.commit()
    return jsonify({"ok": True}) if cur.rowcount else (jsonify({"error": "No existe"}), 404)


@app.route("/api/torre-a/escala", methods=["GET"])
@requiere_admin
def api_ta_escala():
    err = _solo_torre_a()
    if err:
        return err
    crudo = _config_get(get_db(), "ta_escala", "")
    try:
        return jsonify({"escala": json.loads(crudo) if crudo else None})
    except ValueError:
        return jsonify({"escala": None})


@app.route("/api/torre-a/escala", methods=["POST"])
@requiere_admin
def api_ta_escala_guardar():
    """Calibracion: el admin marca dos puntos del plano y dice cuantos metros hay entre ellos.
    Se guarda cuantos metros mide el ANCHO completo de la imagen; con eso se calculan m y m2."""
    err = _solo_torre_a()
    if err:
        return err
    d = request.get_json(silent=True) or {}
    try:
        m_ancho = float(d.get("m_por_ancho"))
        ref_m = float(d.get("ref_m"))
    except (TypeError, ValueError):
        return jsonify({"error": "Escala invalida"}), 400
    if not (0.5 <= m_ancho <= 5000) or not (0.01 <= ref_m <= 5000):
        return jsonify({"error": "La escala sale fuera de lo razonable; marca de nuevo los dos puntos y los metros"}), 400
    ref_pts = _puntos_validos(d.get("ref_pts"), 2, 2)
    if ref_pts is None:
        return jsonify({"error": "Faltan los dos puntos de referencia"}), 400
    db = get_db()
    _config_set(db, "ta_escala", json.dumps({"m_por_ancho": round(m_ancho, 4), "ref_m": round(ref_m, 3), "ref_pts": ref_pts}))
    db.commit()
    return jsonify({"ok": True})


ALTURA_MURO_POR_DEFECTO = 2.6   # metros; promedio editable (config ta_altura_muro)


def _altura_muro_global(db):
    try:
        v = float(_config_get(db, "ta_altura_muro", "") or ALTURA_MURO_POR_DEFECTO)
    except ValueError:
        v = ALTURA_MURO_POR_DEFECTO
    return v if 0.5 <= v <= 30 else ALTURA_MURO_POR_DEFECTO


@app.route("/api/torre-a/lados", methods=["GET"])
@requiere_gestor
def api_ta_lados():
    """Especialidades por lado de cada area dibujada + alturas de muro + colores de las especialidades."""
    err = _solo_torre_a()
    if err:
        return err
    db = get_db()
    lados = {}
    for r in db.execute("SELECT trazo_id, lado, giro FROM ta_lados ORDER BY trazo_id, lado, giro"):
        lados.setdefault(str(r["trazo_id"]), {}).setdefault(str(r["lado"]), []).append(r["giro"])
    alturas = {str(r["trazo_id"]): r["altura"] for r in db.execute("SELECT trazo_id, altura FROM ta_alturas")}
    colores = _mapa_colores_giro(db)
    usados = {g for m in lados.values() for gs in m.values() for g in gs}
    for g in usados:
        if g not in colores:
            _asegurar_color_giro(db, g)
    if usados:
        db.commit()
        colores = _mapa_colores_giro(db)
    return jsonify({"lados": lados, "alturas": alturas, "altura_global": _altura_muro_global(db),
                    "colores": {g: colores.get(g, GIRO_COLOR_OTRAS) for g in usados}})


@app.route("/api/torre-a/lados/<int:tid>", methods=["PUT"])
@requiere_admin
def api_ta_lados_guardar(tid):
    """Cuerpo: {lados: {"0": ["Pintura","Yeso"], "1": []...}, altura: 2.8 | null}. Reemplaza lo de ese trazo."""
    err = _solo_torre_a()
    if err:
        return err
    db = get_db()
    t = db.execute("SELECT tipo, tipo_elem, puntos FROM plano_trazos WHERE id=?", (tid,)).fetchone()
    if not t:
        return jsonify({"error": "El trazo no existe."}), 404
    if t["tipo"] != "zona" or t["tipo_elem"] != "Área":
        return jsonify({"error": "Los lados se marcan solo en las áreas dibujadas (tipo Área)."}), 400
    try:
        n = len(json.loads(t["puntos"]))
    except ValueError:
        n = 0
    data = request.get_json(silent=True) or {}
    filas = []
    if "lados" in data:
        crudo = data.get("lados")
        if not isinstance(crudo, dict):
            return jsonify({"error": "Formato de lados no válido."}), 400
        validas = {r[0] for r in db.execute("SELECT nombre FROM ta_especialidades")}
        for k, gs in crudo.items():
            try:
                i = int(k)
            except (TypeError, ValueError):
                return jsonify({"error": "Lado no válido."}), 400
            if not (0 <= i < n) or not isinstance(gs, list):
                return jsonify({"error": "Ese lado no existe en este trazo."}), 400
            for g in gs:
                g = " ".join(str(g or "").split())[:80]
                if not g:
                    continue
                if g not in validas:
                    return jsonify({"error": "La especialidad «%s» no está en la lista de Torre A." % g}), 400
                if (i, g) not in filas:
                    filas.append((i, g))
        db.execute("DELETE FROM ta_lados WHERE trazo_id=?", (tid,))
        for i, g in filas:
            db.execute("INSERT OR IGNORE INTO ta_lados (trazo_id,lado,giro) VALUES (?,?,?)", (tid, i, g))
            _asegurar_color_giro(db, g)
    if "altura" in data:
        a = data.get("altura")
        if a in (None, ""):
            db.execute("DELETE FROM ta_alturas WHERE trazo_id=?", (tid,))
        else:
            try:
                a = float(a)
            except (TypeError, ValueError):
                return jsonify({"error": "La altura debe ser un número en metros."}), 400
            if not (0.5 <= a <= 30):
                return jsonify({"error": "La altura debe estar entre 0.5 y 30 metros."}), 400
            db.execute("INSERT OR REPLACE INTO ta_alturas (trazo_id,altura) VALUES (?,?)", (tid, round(a, 3)))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/torre-a/altura-muro", methods=["PUT"])
@requiere_admin
def api_ta_altura_muro():
    """Altura promedio de muro de Torre A (para m2 de pintura, yeso, etc.)."""
    err = _solo_torre_a()
    if err:
        return err
    try:
        a = float((request.get_json(silent=True) or {}).get("altura"))
    except (TypeError, ValueError):
        return jsonify({"error": "La altura debe ser un número en metros."}), 400
    if not (0.5 <= a <= 30):
        return jsonify({"error": "La altura debe estar entre 0.5 y 30 metros."}), 400
    db = get_db()
    _config_set(db, "ta_altura_muro", str(round(a, 3)))
    db.commit()
    return jsonify({"ok": True, "altura_global": a})


# ----------------------------------------------------------------------------
# Torre A - Actas de entrega-recepcion y pruebas de funcionamiento (v1.24)
# Un acta junta varias actividades (de distintos proveedores) de un bloque o area, lleva
# cuantos elementos (contactos, apagadores, tomas de gases, alarmas...) se probaron de los
# que se deben probar, quien los probo y el resultado, y sale en un solo PDF para firmar.
# Es INDEPENDIENTE de la actividad: solo va ligada a ella, no cambia su avance.
# ----------------------------------------------------------------------------
SUPERVISION_ACTA_DEFECTO = "Héctor Longoria"      # supervisión de proyecto: siempre el mismo en todas las actas nuevas
RESULTADOS_ACTA = ("pendiente", "conforme", "observacion")
ETIQUETA_RESULTADO_ACTA = {"pendiente": "PENDIENTE", "conforme": "CONFORME", "observacion": "CON OBSERVACIÓN"}
CAMPOS_TEXTO_ACTA = {"titulo": 160, "recibe_obra": 120, "recibe_area": 120, "proveedor_rep": 120,
                     "supervision": 120, "observaciones": 3000}


ROLES_CONSULTA_ACTAS = ("admin", "supervisor", "supervisor_obra", "supervisor_depto")     # ver actas sin poder cambiarlas
FUNCIONES_PERSONA = ("interno", "supervisor", "directivo", "proveedor")


def requiere_consulta_actas(fn):
    """Admin y supervisores pueden VER las actas (solo lectura). Cambiar algo sigue siendo solo del admin."""
    @wraps(fn)
    def envoltura(*a, **k):
        if "usuario" not in session:
            return jsonify({"error": "no_login"}), 401
        if session.get("rol") not in ROLES_CONSULTA_ACTAS:
            return jsonify({"error": "sin_permiso"}), 403
        return fn(*a, **k)
    return envoltura


def _acta_texto(v, largo):
    return (str(v).strip() if v is not None else "")[:largo]


def _acta_fecha_valida(v):
    try:
        datetime.date.fromisoformat(str(v))
        return True
    except (TypeError, ValueError):
        return False


def _acta_catalogo(db):
    """(bloques, pares) validos para un acta. Torre A: catalogo fijo (20 bloques / 85 areas).
    Torre B: los bloques y areas que ya existen en sus actividades de obra."""
    if proyecto_actual() == "torre_a":
        bloques = {etiqueta_bloque_ta(b["num"], b["nombre"]) for b in db.execute("SELECT num,nombre FROM ta_bloques")}
        pares = {(etiqueta_bloque_ta(b["num"], b["nombre"]), etiqueta_area_ta(a["bloque"], a["num"], a["nombre"]))
                 for b in db.execute("SELECT num,nombre FROM ta_bloques")
                 for a in db.execute("SELECT bloque,num,nombre FROM ta_areas WHERE bloque=?", (b["num"],))}
        return bloques, pares
    bloques, pares = set(), set()
    for f in db.execute("SELECT DISTINCT bloque, area FROM actividades WHERE (eliminada IS NULL OR eliminada=0) "
                        "AND COALESCE(mundo,'obra')='obra' AND bloque IS NOT NULL AND bloque<>''"):
        bloques.add(f["bloque"])
        if f["area"]:
            pares.add((f["bloque"], f["area"]))
    return bloques, pares


def _acta_items(db, acta_id):
    filas = db.execute(
        "SELECT i.*, a.codigo, a.partida, a.proveedor, a.departamento AS depto_act, a.valida_depto, "
        "a.avance, a.area AS area_act, a.requiere_pruebas, a.giro, a.bloque AS bloque_act "
        "FROM ta_acta_items i LEFT JOIN actividades a ON a.id=i.actividad_id "
        "WHERE i.acta_id=? ORDER BY i.orden, i.id", (acta_id,)).fetchall()
    salida = []
    for f in filas:
        salida.append({
            "id": f["id"], "actividad_id": f["actividad_id"], "codigo": f["codigo"] or "",
            "partida": f["partida"] or "(actividad eliminada)", "proveedor": f["proveedor"] or "",
            "area": f["area_act"] or "", "valida_depto": f["valida_depto"] or "",
            "avance": f["avance"] or 0, "unidad": f["unidad"], "cant_total": f["cant_total"],
            "cant_ok": f["cant_ok"], "resultado": f["resultado"], "nota": f["nota"] or "",
            "probado_por": f["probado_por"] or "", "fecha_prueba": f["fecha_prueba"] or "",
            "activo": f["activo"] != 0, "giro": f["giro"] or "", "bloque": f["bloque_act"] or "",
        })
    return salida


def _acta_areas(a):
    """Lista de áreas del acta: varias (areas_json), una sola (area) o ninguna (todo el bloque)."""
    try:
        lista = json.loads(a["areas_json"]) if a["areas_json"] else []
    except (TypeError, ValueError):
        lista = []
    if not isinstance(lista, list):
        lista = []
    if not lista and a["area"]:
        lista = [a["area"]]
    return [x for x in lista if isinstance(x, str)]


def _acta_resumen(items):
    items = [i for i in items if i.get("activo", True)]      # los desactivados no cuentan
    return {
        "actividades": len(items),
        "conformes": sum(1 for i in items if i["resultado"] == "conforme"),
        "observaciones": sum(1 for i in items if i["resultado"] == "observacion"),
        "pendientes": sum(1 for i in items if i["resultado"] == "pendiente"),
        "probados": sum(i["cant_ok"] for i in items),
        "por_probar": sum(i["cant_total"] for i in items),
    }


def _acta_a_dict(db, a, con_items=True):
    items = _acta_items(db, a["id"])
    d = {k: a[k] for k in ("id", "folio", "bloque", "area", "titulo", "estado", "fecha", "recibe_obra",
                           "recibe_area", "proveedor_rep", "supervision", "observaciones", "creada",
                           "creada_por", "cerrada", "cerrada_por")}
    d["parcial_n"] = a["parcial_n"]
    cont = db.execute("SELECT id,folio FROM ta_actas WHERE id=? AND eliminada=0", (a["continua_id"],)).fetchone() if a["continua_id"] else None
    prev = db.execute("SELECT id,folio FROM ta_actas WHERE continua_id=? AND eliminada=0", (a["id"],)).fetchone()
    d["continua"] = {"id": cont["id"], "folio": cont["folio"]} if cont else None
    d["viene_de"] = {"id": prev["id"], "folio": prev["folio"]} if prev else None
    d["areas"] = _acta_areas(a)
    d["area_txt"] = ", ".join(d["areas"]) if len(d["areas"]) > 1 else (d["area"] or "")
    d["resumen"] = _acta_resumen(items)
    d["desactivadas"] = sum(1 for i in items if not i["activo"])
    if con_items:
        d["items"] = items
    return d


def _acta_o_error(db, aid):
    a = db.execute("SELECT * FROM ta_actas WHERE id=? AND eliminada=0", (aid,)).fetchone()
    if not a:
        return None, (jsonify({"error": "No existe esa acta."}), 404)
    return a, None


def _acta_editable_o_error(a):
    if a["estado"] != "borrador":
        return jsonify({"error": "El acta ya está cerrada. Reábrela para cambiarla."}), 409
    return None


def _folio_acta(db, bloque, area):
    """Torre A: ACTA-BB[.AA]-NNN (por numero de bloque/area). Torre B: ACTA-B-NNN consecutivo."""
    if proyecto_actual() == "torre_a":
        bnum = int(bloque.split(" · ")[0])
        prefijo = "ACTA-%02d" % bnum
        if area:
            prefijo = "ACTA-%02d.%02d" % (bnum, int(area.split(" · ")[0].split(".")[1]))
    else:
        prefijo = "ACTA-B"
    mayor = 0
    patron = re.compile("^" + re.escape(prefijo) + r"-(\d+)(?:-\d+)?$")      # ACTA-B-010 y también ACTA-B-010-2
    for f in db.execute("SELECT folio FROM ta_actas WHERE folio LIKE ?", (prefijo + "-%",)).fetchall():
        m = patron.match(f["folio"])
        if m:
            mayor = max(mayor, int(m.group(1)))
    return "%s-%03d" % (prefijo, mayor + 1)


def _titulo_acta_por_defecto(bloque, area):
    texto = area or bloque
    if proyecto_actual() == "torre_a" and " · " in texto:
        return texto.split(" · ", 1)[1]
    return texto


def _ocupacion_en_acta(db, act_id, acta_propia=None):
    """Otra acta que ya tiene esta actividad: {'folio','item_id','movible'} o None.
    movible = esta en un BORRADOR y todavia no esta conforme: se puede pasar a otra acta sin perder nada."""
    filas = db.execute(
        "SELECT t.folio, t.estado, i.resultado, t.id AS acta_id, i.id AS item_id FROM ta_acta_items i "
        "JOIN ta_actas t ON t.id=i.acta_id WHERE i.actividad_id=? AND t.eliminada=0 AND i.activo=1", (act_id,)).fetchall()
    for f in filas:
        if acta_propia is not None and f["acta_id"] == acta_propia:
            continue
        if f["estado"] == "borrador" or f["resultado"] == "conforme":
            return {"folio": f["folio"], "item_id": f["item_id"], "acta_id": f["acta_id"],
                    "movible": f["estado"] == "borrador" and f["resultado"] != "conforme"}
    return None


def _actividad_ocupada_en_acta(db, act_id, acta_propia=None):
    """Folio de otra acta que ya tiene esta actividad (conforme en acta vigente o en un borrador), o None."""
    o = _ocupacion_en_acta(db, act_id, acta_propia)
    return o["folio"] if o else None


@app.route("/api/actas", methods=["GET"])
@requiere_consulta_actas
def api_ta_actas_lista():
    db = get_db()
    filas = db.execute("SELECT * FROM ta_actas WHERE eliminada=0 ORDER BY id DESC").fetchall()
    return jsonify({"actas": [_acta_a_dict(db, a, con_items=False) for a in filas]})


def _persona_a_dict(f):
    return {"id": f["id"], "nombre": f["nombre"], "funcion": f["funcion"], "empresa": f["empresa"] or "", "activo": f["activo"] != 0}


@app.route("/api/actas/personas", methods=["GET"])
@requiere_consulta_actas
def api_ta_personas_lista():
    db = get_db()
    filas = db.execute("SELECT * FROM ta_personas ORDER BY nombre COLLATE NOCASE").fetchall()
    return jsonify({"personas": [_persona_a_dict(f) for f in filas]})


@app.route("/api/actas/personas", methods=["POST"])
@requiere_admin
def api_ta_personas_crear():
    """Alta rápida desde el acta (sin salir de la pantalla). Si ya existe (aunque cambie mayúsculas o acentos) se reutiliza."""
    db = get_db()
    data = request.get_json(silent=True) or {}
    nombre = " ".join(str(data.get("nombre") or "").split())[:120]
    if len(nombre) < 2 or nombre.startswith("__"):
        return jsonify({"error": "Escribe el nombre de la persona."}), 400
    funcion = data.get("funcion") if data.get("funcion") in FUNCIONES_PERSONA else "interno"
    empresa = _acta_texto(data.get("empresa"), 120)
    for f in db.execute("SELECT * FROM ta_personas").fetchall():
        if sin_acentos(f["nombre"]) == sin_acentos(nombre):
            if f["activo"] == 0:
                db.execute("UPDATE ta_personas SET activo=1 WHERE id=?", (f["id"],))
                db.commit()
                f = db.execute("SELECT * FROM ta_personas WHERE id=?", (f["id"],)).fetchone()
            return jsonify({"ok": True, "existente": True, "persona": _persona_a_dict(f)})
    cur = db.execute("INSERT INTO ta_personas (nombre,funcion,empresa,creada,creada_por) VALUES (?,?,?,?,?)",
                     (nombre, funcion, empresa, datetime.datetime.now().isoformat(timespec="seconds"), session.get("usuario")))
    db.commit()
    f = db.execute("SELECT * FROM ta_personas WHERE id=?", (cur.lastrowid,)).fetchone()
    return jsonify({"ok": True, "existente": False, "persona": _persona_a_dict(f)})


@app.route("/api/actas/personas/<int:pid>", methods=["PUT"])
@requiere_admin
def api_ta_personas_editar(pid):
    """Cambiar la función (interno / supervisor / directivo / proveedor), la empresa o desactivarla. El nombre no cambia
    (ya está escrito en las actas)."""
    db = get_db()
    f = db.execute("SELECT * FROM ta_personas WHERE id=?", (pid,)).fetchone()
    if not f:
        return jsonify({"error": "No existe esa persona."}), 404
    data = request.get_json(silent=True) or {}
    cambios = {}
    if "funcion" in data:
        if data["funcion"] not in FUNCIONES_PERSONA:
            return jsonify({"error": "Función no válida."}), 400
        cambios["funcion"] = data["funcion"]
    if "empresa" in data:
        cambios["empresa"] = _acta_texto(data["empresa"], 120)
    if "activo" in data:
        cambios["activo"] = 1 if data["activo"] else 0
    if cambios:
        db.execute("UPDATE ta_personas SET %s WHERE id=?" % ", ".join("%s=?" % c for c in cambios), list(cambios.values()) + [pid])
        db.commit()
    return jsonify({"ok": True, "persona": _persona_a_dict(db.execute("SELECT * FROM ta_personas WHERE id=?", (pid,)).fetchone())})


@app.route("/api/actas/pendientes", methods=["GET"])
@requiere_consulta_actas
def api_ta_actas_pendientes():
    """Inventario de lo que falta por entregar: (1) actividades que están en un acta como pendiente o con observación y
    todavía no tienen una entrega CONFORME en un acta cerrada, (2) cuántas actividades de obra no están en ningún acta."""
    db = get_db()
    entregadas = {r[0] for r in db.execute(
        "SELECT i.actividad_id FROM ta_acta_items i JOIN ta_actas t ON t.id=i.acta_id "
        "WHERE t.eliminada=0 AND t.estado='cerrada' AND i.activo=1 AND i.resultado='conforme'")}
    ultimo = {}
    for f in db.execute(
            "SELECT i.id, i.actividad_id, i.resultado, i.nota, i.cant_ok, i.cant_total, i.unidad, t.id AS acta_id, t.folio, t.estado "
            "FROM ta_acta_items i JOIN ta_actas t ON t.id=i.acta_id WHERE t.eliminada=0 AND i.activo=1 ORDER BY i.id"):
        ultimo[f["actividad_id"]] = f                    # el renglón más reciente de cada actividad
    pendientes = []
    for act_id, f in ultimo.items():
        if act_id in entregadas or f["resultado"] == "conforme":
            continue
        a = db.execute("SELECT codigo,partida,proveedor,bloque,area FROM actividades WHERE id=?", (act_id,)).fetchone()
        if not a:
            continue
        pendientes.append({"actividad_id": act_id, "codigo": a["codigo"] or "", "partida": a["partida"] or "",
                           "proveedor": a["proveedor"] or "", "bloque": a["bloque"] or "", "area": a["area"] or "",
                           "resultado": f["resultado"], "nota": f["nota"] or "", "cant_ok": f["cant_ok"],
                           "cant_total": f["cant_total"], "unidad": f["unidad"], "acta_id": f["acta_id"],
                           "folio": f["folio"], "estado_acta": f["estado"]})
    pendientes.sort(key=lambda x: (x["proveedor"], x["bloque"], x["area"], x["codigo"]))
    en_acta = {r[0] for r in db.execute("SELECT i.actividad_id FROM ta_acta_items i JOIN ta_actas t ON t.id=i.acta_id "
                                        "WHERE t.eliminada=0 AND i.activo=1")}
    sin_acta = {}
    for f in db.execute("SELECT id,bloque FROM actividades WHERE (eliminada IS NULL OR eliminada=0) AND COALESCE(mundo,'obra')='obra' "
                        "AND COALESCE(estado_val,'validado')='validado' AND requiere_pruebas='SÍ' AND bloque IS NOT NULL AND bloque<>''"):
        if f["id"] not in en_acta:
            sin_acta[f["bloque"]] = sin_acta.get(f["bloque"], 0) + 1
    return jsonify({"pendientes": pendientes, "entregadas": len(entregadas),
                    "sin_acta": [{"bloque": b, "n": n} for b, n in sorted(sin_acta.items())],
                    "sin_acta_total": sum(sin_acta.values())})


@app.route("/api/actas/candidatas", methods=["GET"])
@requiere_admin
def api_ta_actas_candidatas():
    """Actividades de obra de un bloque (y de una o varias áreas) que se pueden meter al acta."""
    db = get_db()
    bloques, pares = _acta_catalogo(db)
    bloque = request.args.get("bloque", "")
    areas = [x for x in request.args.getlist("area") if x]
    if bloque and bloque not in bloques:
        return jsonify({"error": "Elige un bloque de la lista."}), 400
    for ar in areas:
        if (bloque, ar) not in pares:
            return jsonify({"error": "Elige un área de la lista."}), 400
    todas = request.args.get("todas") == "1"
    sql = ("SELECT id,codigo,partida,proveedor,bloque,area,giro,requiere_pruebas,valida_depto,avance FROM actividades "
           "WHERE (eliminada IS NULL OR eliminada=0) AND COALESCE(mundo,'obra')='obra' "
           "AND COALESCE(estado_val,'validado')='validado' AND bloque IS NOT NULL AND bloque<>''")
    args = []
    if bloque:
        sql += " AND bloque=?"
        args.append(bloque)
    if areas:
        sql += " AND area IN (%s)" % ",".join("?" * len(areas))
        args += areas
    if not todas:
        sql += " AND requiere_pruebas='SÍ'"
    sql += " ORDER BY area, codigo, id"
    salida = []
    for f in db.execute(sql, args).fetchall():
        salida.append({"id": f["id"], "codigo": f["codigo"] or "", "partida": f["partida"] or "",
                       "proveedor": f["proveedor"] or "", "bloque": f["bloque"] or "", "area": f["area"] or "", "giro": f["giro"] or "",
                       "requiere_pruebas": f["requiere_pruebas"] == "SÍ", "valida_depto": f["valida_depto"] or "",
                       "avance": f["avance"] or 0, "ocupada_en": None, "movible": False})
        o = _ocupacion_en_acta(db, f["id"])
        if o:
            salida[-1]["ocupada_en"] = o["folio"]
            salida[-1]["movible"] = o["movible"]
    return jsonify({"actividades": salida})


def _borrar_actas_vacias(db, ids_acta):
    """Un borrador que se quedó sin actividades activas (porque se pasaron a otra acta) se borra solo."""
    borradas = []
    for aid in ids_acta:
        a = db.execute("SELECT folio,estado FROM ta_actas WHERE id=? AND eliminada=0", (aid,)).fetchone()
        if not a or a["estado"] != "borrador":
            continue
        if db.execute("SELECT 1 FROM ta_acta_items WHERE acta_id=? AND activo=1", (aid,)).fetchone():
            continue
        db.execute("UPDATE ta_actas SET eliminada=1 WHERE id=?", (aid,))
        db.execute("UPDATE ta_actas SET continua_id=NULL WHERE continua_id=?", (aid,))
        borradas.append(a["folio"])
    return borradas


@app.route("/api/actas/<int:aid>/actividad_nueva", methods=["POST"])
@requiere_admin
def api_ta_acta_actividad_nueva(aid):
    """Torre B: da de alta una actividad que faltaba (ej. pintura que sí se hizo) y la mete al acta.
    Nace POR VALIDAR (propuesta): no suma en ningún total hasta que se apruebe en Validación."""
    db = get_db()
    a, e = _acta_o_error(db, aid)
    if e:
        return e
    e = _acta_editable_o_error(a)
    if e:
        return e
    if proyecto_actual() == "torre_a":
        return jsonify({"error": "En Torre A solo se entregan actividades ya documentadas."}), 400
    data = request.get_json(silent=True) or {}
    bloques, pares = _acta_catalogo(db)
    bloque = _acta_texto(data.get("bloque"), 200)
    area = _acta_texto(data.get("area"), 300)
    giro = _acta_texto(data.get("giro"), 120)
    prov = _acta_texto(data.get("proveedor"), 120)
    partida = _acta_texto(data.get("partida"), 300)
    if bloque not in bloques or (area and (bloque, area) not in pares):
        return jsonify({"error": "Elige el bloque y el área de la lista."}), 400
    if not prov or not partida:
        return jsonify({"error": "Falta el responsable o la actividad."}), 400
    existe = db.execute("SELECT 1 FROM actividades WHERE proveedor=? AND bloque=? AND COALESCE(area,'')=? AND partida=? "
                        "AND (eliminada IS NULL OR eliminada=0)", (prov, bloque, area, partida)).fetchone()
    if existe:
        return jsonify({"error": "Esa actividad ya existe: búscala en 'Agregar actividades'."}), 409
    codigo = siguiente_codigo(db)
    est = datos_estado_carga(True)
    ahora = datetime.datetime.now().isoformat(timespec="seconds")
    cur = db.execute(
        """INSERT INTO actividades (codigo,bloque,area,giro,proveedor,partida,tipo,tipo_partida,aplica,avance,
           estatus,definido,mundo,requiere_pruebas,origen,estado_val,creado_por,actualizado)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (codigo, bloque, area or None, giro or None, prov, partida, "Construcción", "Construcción", "SÍ", 0,
         "Pendiente", "NO", "obra", "SÍ", est[0], est[1], est[2], ahora))
    nid = cur.lastrowid
    orden = (db.execute("SELECT COALESCE(MAX(orden),0) FROM ta_acta_items WHERE acta_id=?", (aid,)).fetchone()[0] or 0) + 1
    db.execute("INSERT INTO ta_acta_items (acta_id,actividad_id,orden) VALUES (?,?,?)", (aid, nid, orden))
    db.commit()
    return jsonify({"ok": True, "id": nid, "codigo": codigo}), 201


@app.route("/api/actas/auto", methods=["POST"])
@requiere_admin
def api_ta_actas_auto():
    """Genera UN acta (borrador) por cada área con actividades libres. Una actividad nunca queda en dos actas:
    se saltan las que ya están en un borrador o conformes en otra acta. Sin fecha (se pone al abrirla),
    sin nombres de quien recibe y con la supervisión de siempre. 'simular': solo cuenta, no crea nada."""
    data = request.get_json(silent=True) or {}
    db = get_db()
    bloques, pares = _acta_catalogo(db)
    solo_bloque = _acta_texto(data.get("bloque"), 200)
    if solo_bloque and solo_bloque not in bloques:
        return jsonify({"error": "Elige un bloque de la lista."}), 400
    solo_term = bool(data.get("solo_terminadas"))
    simular = bool(data.get("simular"))
    sql = ("SELECT id,bloque,area,avance FROM actividades WHERE (eliminada IS NULL OR eliminada=0) "
           "AND COALESCE(mundo,'obra')='obra' AND COALESCE(estado_val,'validado')='validado' "
           "AND requiere_pruebas='SÍ' AND bloque IS NOT NULL AND bloque<>''")
    args = []
    if solo_bloque:
        sql += " AND bloque=?"
        args.append(solo_bloque)
    if solo_term:
        sql += " AND avance>=100"
    sql += " ORDER BY bloque, area, codigo, id"
    grupos, ocupadas = {}, 0
    for f in db.execute(sql, args).fetchall():
        if f["bloque"] not in bloques:
            continue
        area = f["area"] if f["area"] else None
        if area and (f["bloque"], area) not in pares:
            continue
        if _ocupacion_en_acta(db, f["id"]):
            ocupadas += 1
            continue
        grupos.setdefault((f["bloque"], area), []).append(f["id"])
    resumen = [{"bloque": b, "area": a, "actividades": len(ids)} for (b, a), ids in grupos.items()]
    if simular:
        return jsonify({"ok": True, "simulacion": True, "actas": resumen, "ya_en_otra_acta": ocupadas})
    ahora = datetime.datetime.now().isoformat(timespec="seconds")
    creadas = []
    for (bloque, area), ids in grupos.items():
        folio = _folio_acta(db, bloque, area)
        titulo = ("%s — %s" % (area, bloque)) if area else bloque
        cur = db.execute(
            "INSERT INTO ta_actas (folio,bloque,area,titulo,estado,fecha,supervision,creada,creada_por) "
            "VALUES (?,?,?,?,?,?,?,?,?)",
            (folio, bloque, area, titulo[:160], "borrador", None, SUPERVISION_ACTA_DEFECTO, ahora, session.get("usuario")))
        for n, aid in enumerate(ids, 1):
            db.execute("INSERT INTO ta_acta_items (acta_id,actividad_id,orden) VALUES (?,?,?)", (cur.lastrowid, aid, n))
        creadas.append({"id": cur.lastrowid, "folio": folio, "bloque": bloque, "area": area, "actividades": len(ids)})
    db.commit()
    return jsonify({"ok": True, "creadas": creadas, "ya_en_otra_acta": ocupadas}), 201


@app.route("/api/actas", methods=["POST"])
@requiere_admin
def api_ta_actas_crear():
    data = request.get_json(silent=True) or {}
    db = get_db()
    bloques, pares = _acta_catalogo(db)
    bloque = _acta_texto(data.get("bloque"), 200)
    areas = data.get("areas")
    if not isinstance(areas, list):
        areas = [data.get("area")] if data.get("area") else []
    areas = list(dict.fromkeys(_acta_texto(x, 300) for x in areas if x))
    if bloque not in bloques:
        return jsonify({"error": "Elige un bloque de la lista."}), 400
    for ar in areas:
        if (bloque, ar) not in pares:
            return jsonify({"error": "Esa área no es del bloque elegido."}), 400
    area = areas[0] if len(areas) == 1 else None       # un acta de una sola área conserva su folio por área
    ids = data.get("actividades")
    if not isinstance(ids, list) or not ids:
        return jsonify({"error": "Marca al menos una actividad para el acta."}), 400
    validas, rechazadas, por_mover, origen_actas = [], [], [], set()
    mover = bool(data.get("mover"))          # pasar a esta acta lo que está en otro BORRADOR (sin conformes)
    for raw in ids[:1500]:
        try:
            aid = int(raw)
        except (TypeError, ValueError):
            continue
        act = db.execute("SELECT id,bloque,area FROM actividades WHERE id=? AND (eliminada IS NULL OR eliminada=0) "
                         "AND COALESCE(mundo,'obra')='obra'", (aid,)).fetchone()
        if not act or act["bloque"] != bloque or (areas and act["area"] not in areas):
            rechazadas.append({"id": aid, "motivo": "no es de ese bloque/área"})
            continue
        oc = _ocupacion_en_acta(db, aid)
        if oc and not (mover and oc["movible"]):
            rechazadas.append({"id": aid, "motivo": "ya está en " + oc["folio"]})
            continue
        if oc:
            por_mover.append(oc["item_id"])
            origen_actas.add(oc["acta_id"])
        if aid not in validas:
            validas.append(aid)
    if not validas:
        return jsonify({"error": "Ninguna actividad se pudo agregar: " +
                        "; ".join("#%d %s" % (r["id"], r["motivo"]) for r in rechazadas[:5])}), 400
    ahora = datetime.datetime.now().isoformat(timespec="seconds")
    folio = _folio_acta(db, bloque, area)
    titulo = _acta_texto(data.get("titulo"), 160) or _titulo_acta_por_defecto(bloque, area)
    cur = db.execute(
        "INSERT INTO ta_actas (folio,bloque,area,areas_json,titulo,estado,fecha,supervision,creada,creada_por) "
        "VALUES (?,?,?,?,?,?,?,?,?,?)",
        (folio, bloque, area, json.dumps(areas, ensure_ascii=False) if len(areas) > 1 else None,
         titulo, "borrador", hoy(), SUPERVISION_ACTA_DEFECTO, ahora, session.get("usuario")))
    acta_id = cur.lastrowid
    for iid in por_mover:
        db.execute("UPDATE ta_acta_items SET activo=0 WHERE id=?", (iid,))
    for n, aid in enumerate(validas, 1):
        db.execute("INSERT INTO ta_acta_items (acta_id,actividad_id,orden) VALUES (?,?,?)", (acta_id, aid, n))
    vaciadas = _borrar_actas_vacias(db, origen_actas)
    db.commit()
    return jsonify({"ok": True, "id": acta_id, "folio": folio, "rechazadas": rechazadas, "movidas": len(por_mover),
                    "actas_vaciadas": vaciadas}), 201


@app.route("/api/actas/<int:aid>", methods=["GET"])
@requiere_consulta_actas
def api_ta_acta_detalle(aid):
    db = get_db()
    a, e = _acta_o_error(db, aid)
    if e:
        return e
    if a["estado"] == "borrador" and not (a["fecha"] or "").strip() and session.get("rol") == "admin":
        # actas automáticas nacen sin fecha: la fecha es la del día en que se abre por primera vez
        db.execute("UPDATE ta_actas SET fecha=? WHERE id=?", (hoy(), aid))
        db.commit()
        a = db.execute("SELECT * FROM ta_actas WHERE id=?", (aid,)).fetchone()
    return jsonify(_acta_a_dict(db, a))


@app.route("/api/actas/<int:aid>", methods=["PUT"])
@requiere_admin
def api_ta_acta_editar(aid):
    db = get_db()
    a, e = _acta_o_error(db, aid)
    if e:
        return e
    e = _acta_editable_o_error(a)
    if e:
        return e
    data = request.get_json(silent=True) or {}
    cambios = {}
    for campo, largo in CAMPOS_TEXTO_ACTA.items():
        if campo in data:
            cambios[campo] = _acta_texto(data[campo], largo)
    if "fecha" in data:
        f = _acta_texto(data["fecha"], 10)
        if f and not _acta_fecha_valida(f):
            return jsonify({"error": "La fecha no es válida."}), 400
        cambios["fecha"] = f
    if cambios:
        sets = ", ".join("%s=?" % c for c in cambios)
        db.execute("UPDATE ta_actas SET %s WHERE id=?" % sets, list(cambios.values()) + [aid])
        db.commit()
    return jsonify({"ok": True})


@app.route("/api/actas/<int:aid>/actividades", methods=["POST"])
@requiere_admin
def api_ta_acta_agregar(aid):
    db = get_db()
    a, e = _acta_o_error(db, aid)
    if e:
        return e
    e = _acta_editable_o_error(a)
    if e:
        return e
    cuerpo = request.get_json(silent=True) or {}
    ids = cuerpo.get("actividades")
    mover = bool(cuerpo.get("mover"))
    if not isinstance(ids, list) or not ids:
        return jsonify({"error": "Marca al menos una actividad."}), 400
    agregadas, rechazadas, movidas, actas_origen = 0, [], 0, set()
    orden = (db.execute("SELECT COALESCE(MAX(orden),0) FROM ta_acta_items WHERE acta_id=?", (aid,)).fetchone()[0] or 0)
    areas_acta = _acta_areas(a)
    for raw in ids[:1500]:
        try:
            act_id = int(raw)
        except (TypeError, ValueError):
            continue
        act = db.execute("SELECT id,bloque,area FROM actividades WHERE id=? AND (eliminada IS NULL OR eliminada=0) "
                         "AND COALESCE(mundo,'obra')='obra'", (act_id,)).fetchone()
        if not act:
            rechazadas.append({"id": act_id, "motivo": "la actividad no existe"})
            continue
        previo = db.execute("SELECT id,activo FROM ta_acta_items WHERE acta_id=? AND actividad_id=?", (aid, act_id)).fetchone()
        if previo and previo["activo"] != 0:
            continue
        oc = _ocupacion_en_acta(db, act_id, acta_propia=aid)
        if oc and not (mover and oc["movible"]):
            rechazadas.append({"id": act_id, "motivo": "ya está en " + oc["folio"]})
            continue
        if oc:
            db.execute("UPDATE ta_acta_items SET activo=0 WHERE id=?", (oc["item_id"],))
            movidas += 1
            actas_origen.add(oc["acta_id"])
        if previo:                       # estaba desactivada en esta misma acta: se reactiva con lo que ya tenía
            db.execute("UPDATE ta_acta_items SET activo=1 WHERE id=?", (previo["id"],))
            agregadas += 1
            continue
        orden += 1
        db.execute("INSERT INTO ta_acta_items (acta_id,actividad_id,orden) VALUES (?,?,?)", (aid, act_id, orden))
        agregadas += 1
    vaciadas = _borrar_actas_vacias(db, actas_origen)
    db.commit()
    return jsonify({"ok": True, "agregadas": agregadas, "rechazadas": rechazadas, "movidas": movidas, "actas_vaciadas": vaciadas})


def _entero_acta(v, minimo, maximo):
    try:
        n = int(v)
    except (TypeError, ValueError):
        return None
    return n if minimo <= n <= maximo else None


@app.route("/api/actas/<int:aid>/items/<int:iid>", methods=["PUT"])
@requiere_admin
def api_ta_acta_item_editar(aid, iid):
    db = get_db()
    a, e = _acta_o_error(db, aid)
    if e:
        return e
    e = _acta_editable_o_error(a)
    if e:
        return e
    it = db.execute("SELECT * FROM ta_acta_items WHERE id=? AND acta_id=?", (iid, aid)).fetchone()
    if not it:
        return jsonify({"error": "No existe ese renglón."}), 404
    if it["activo"] == 0:
        return jsonify({"error": "Esa actividad está desactivada en el acta. Reactívala para editarla."}), 409
    data = request.get_json(silent=True) or {}
    unidad = _acta_texto(data["unidad"], 30) or "piezas" if "unidad" in data else it["unidad"]
    total, ok_n = it["cant_total"], it["cant_ok"]
    if "cant_total" in data:
        total = _entero_acta(data["cant_total"], 1, 100000)
        if total is None:
            return jsonify({"error": "El total por probar debe ser un número entre 1 y 100000."}), 400
    if "cant_ok" in data:
        ok_n = _entero_acta(data["cant_ok"], 0, 100000)
        if ok_n is None:
            return jsonify({"error": "Los elementos probados deben ser un número de 0 en adelante."}), 400
    if ok_n > total and "cant_total" not in data:
        total = ok_n          # actividad generalizada (ej. "tomas de gases" = 4 tomas): el total sube con lo que se prueba
    if ok_n > total:
        return jsonify({"error": "No puedes probar más elementos (%d) de los que hay por probar (%d)." % (ok_n, total)}), 400
    nota = _acta_texto(data["nota"], 600) if "nota" in data else (it["nota"] or "")
    probado = _acta_texto(data["probado_por"], 120) if "probado_por" in data else (it["probado_por"] or "")
    fecha = _acta_texto(data["fecha_prueba"], 10) if "fecha_prueba" in data else (it["fecha_prueba"] or "")
    if fecha and not _acta_fecha_valida(fecha):
        return jsonify({"error": "La fecha de la prueba no es válida."}), 400
    resultado = it["resultado"]
    if "resultado" in data:
        resultado = data["resultado"]
        if resultado not in RESULTADOS_ACTA:
            return jsonify({"error": "Resultado no válido."}), 400
        if resultado == "conforme" and ok_n != total:
            return jsonify({"error": "Para marcar CONFORME deben estar probados todos los elementos (%d de %d)." % (ok_n, total)}), 400
        if resultado == "observacion" and not nota:
            return jsonify({"error": "Escribe en la nota qué observación quedó."}), 400
        if resultado != "pendiente":
            fecha = fecha or hoy()
    elif resultado == "conforme" and ok_n != total:
        resultado = "pendiente"          # cambió la cantidad: hay que volver a confirmar
    db.execute("UPDATE ta_acta_items SET unidad=?,cant_total=?,cant_ok=?,resultado=?,nota=?,probado_por=?,"
               "fecha_prueba=? WHERE id=?", (unidad, total, ok_n, resultado, nota, probado, fecha or None, iid))
    db.commit()
    fila = [i for i in _acta_items(db, aid) if i["id"] == iid][0]
    return jsonify({"ok": True, "item": fila, "resumen": _acta_resumen(_acta_items(db, aid))})


@app.route("/api/actas/<int:aid>/items/<int:iid>", methods=["DELETE"])
@requiere_admin
def api_ta_acta_item_quitar(aid, iid):
    """Desactiva la actividad en el acta (no se borra: se puede reactivar con todo lo que ya tenía capturado)."""
    db = get_db()
    a, e = _acta_o_error(db, aid)
    if e:
        return e
    e = _acta_editable_o_error(a)
    if e:
        return e
    db.execute("UPDATE ta_acta_items SET activo=0 WHERE id=? AND acta_id=?", (iid, aid))
    db.commit()
    return jsonify({"ok": True, "resumen": _acta_resumen(_acta_items(db, aid))})


@app.route("/api/actas/<int:aid>/items/<int:iid>/restaurar", methods=["POST"])
@requiere_admin
def api_ta_acta_item_restaurar(aid, iid):
    db = get_db()
    a, e = _acta_o_error(db, aid)
    if e:
        return e
    e = _acta_editable_o_error(a)
    if e:
        return e
    it = db.execute("SELECT actividad_id FROM ta_acta_items WHERE id=? AND acta_id=?", (iid, aid)).fetchone()
    if not it:
        return jsonify({"error": "No existe ese renglón."}), 404
    if not db.execute("SELECT 1 FROM actividades WHERE id=? AND (eliminada IS NULL OR eliminada=0)", (it["actividad_id"],)).fetchone():
        return jsonify({"error": "Esa actividad ya fue eliminada de la plataforma; no se puede reactivar."}), 409
    ocupada = _actividad_ocupada_en_acta(db, it["actividad_id"], acta_propia=aid)
    if ocupada:
        return jsonify({"error": "No se puede reactivar: la actividad ya está en " + ocupada + "."}), 409
    db.execute("UPDATE ta_acta_items SET activo=1 WHERE id=?", (iid,))
    db.commit()
    return jsonify({"ok": True, "resumen": _acta_resumen(_acta_items(db, aid))})


@app.route("/api/actas/<int:aid>/marcar_conformes", methods=["POST"])
@requiere_admin
def api_ta_acta_marcar_conformes(aid):
    """Marca CONFORME (con todos los elementos probados) los renglones pendientes indicados.
    Sin 'ids' lo hace con todos los pendientes activos. Los que ya tienen observación no se tocan."""
    db = get_db()
    a, e = _acta_o_error(db, aid)
    if e:
        return e
    e = _acta_editable_o_error(a)
    if e:
        return e
    data = request.get_json(silent=True) or {}
    ids = data.get("ids")
    if ids is not None:
        if not isinstance(ids, list):
            return jsonify({"error": "Lista de renglones no válida."}), 400
        ids = {int(x) for x in ids if str(x).lstrip("-").isdigit()}
    quien = _acta_texto(data.get("probado_por"), 120)
    fecha = _acta_texto(data.get("fecha_prueba"), 10) or hoy()
    if not _acta_fecha_valida(fecha):
        return jsonify({"error": "La fecha de la prueba no es válida."}), 400
    n = 0
    for f in db.execute("SELECT id,cant_total,probado_por FROM ta_acta_items WHERE acta_id=? AND activo=1 "
                        "AND resultado='pendiente'", (aid,)).fetchall():
        if ids is not None and f["id"] not in ids:
            continue
        db.execute("UPDATE ta_acta_items SET cant_ok=cant_total, resultado='conforme', fecha_prueba=COALESCE(NULLIF(fecha_prueba,''),?), "
                   "probado_por=CASE WHEN COALESCE(probado_por,'')='' THEN ? ELSE probado_por END WHERE id=?",
                   (fecha, quien, f["id"]))
        n += 1
    db.commit()
    return jsonify({"ok": True, "marcadas": n})


@app.route("/api/actas/<int:aid>/cerrar", methods=["POST"])
@requiere_admin
def api_ta_acta_cerrar(aid):
    """Cierra el acta. Si quedan pendientes se puede cerrar como ENTREGA PARCIAL (parcial=true): el acta queda
    ACTA-X-1 cerrada y los pendientes pasan a una acta nueva ACTA-X-2 (borrador) para que sigan los demás proveedores."""
    db = get_db()
    a, e = _acta_o_error(db, aid)
    if e:
        return e
    e = _acta_editable_o_error(a)
    if e:
        return e
    data = request.get_json(silent=True) or {}
    r = _acta_resumen(_acta_items(db, aid))
    if r["actividades"] == 0:
        return jsonify({"error": "El acta no tiene actividades."}), 400
    ahora = datetime.datetime.now().isoformat(timespec="seconds")
    if not r["pendientes"]:
        db.execute("UPDATE ta_actas SET estado='cerrada', cerrada=?, cerrada_por=?, fecha=COALESCE(NULLIF(fecha,''),?) WHERE id=?",
                   (ahora, session.get("usuario"), hoy(), aid))
        db.commit()
        return jsonify({"ok": True, "resumen": r, "continua": None})
    if not data.get("parcial"):
        return jsonify({"error": "Faltan %d actividad(es) por probar." % r["pendientes"], "permite_parcial": True}), 400
    if r["conformes"] + r["observaciones"] == 0:
        return jsonify({"error": "Para una entrega parcial marca al menos una actividad como Conforme o Con observación."}), 400
    raiz = a["origen_id"] or a["id"]
    base = re.sub(r"-\d+$", "", a["folio"]) if a["parcial_n"] else a["folio"]
    n_ant = a["parcial_n"] or 1
    n = n_ant + 1
    folio_nuevo = "%s-%d" % (base, n)
    while db.execute("SELECT 1 FROM ta_actas WHERE folio=?", (folio_nuevo,)).fetchone():
        n += 1
        folio_nuevo = "%s-%d" % (base, n)
    folio_actual = a["folio"] if a["parcial_n"] else "%s-1" % base
    if folio_actual != a["folio"] and db.execute("SELECT 1 FROM ta_actas WHERE folio=?", (folio_actual,)).fetchone():
        return jsonify({"error": "Ya existe el folio %s." % folio_actual}), 409
    cur = db.execute(
        "INSERT INTO ta_actas (folio,bloque,area,areas_json,titulo,estado,fecha,recibe_obra,recibe_area,proveedor_rep,supervision,"
        "observaciones,creada,creada_por,origen_id,parcial_n) VALUES (?,?,?,?,?,'borrador',NULL,?,?,?,?,NULL,?,?,?,?)",
        (folio_nuevo, a["bloque"], a["area"], a["areas_json"], a["titulo"], a["recibe_obra"], a["recibe_area"],
         a["proveedor_rep"], a["supervision"] or SUPERVISION_ACTA_DEFECTO, ahora, session.get("usuario"), raiz, n))
    nuevo_id = cur.lastrowid
    db.execute("UPDATE ta_acta_items SET acta_id=? WHERE acta_id=? AND activo=1 AND resultado='pendiente'", (nuevo_id, aid))
    db.execute("UPDATE ta_actas SET estado='cerrada', cerrada=?, cerrada_por=?, fecha=COALESCE(NULLIF(fecha,''),?), folio=?, "
               "origen_id=?, parcial_n=?, continua_id=? WHERE id=?",
               (ahora, session.get("usuario"), hoy(), folio_actual, raiz, n_ant, nuevo_id, aid))
    db.commit()
    return jsonify({"ok": True, "resumen": _acta_resumen(_acta_items(db, aid)), "movidas": r["pendientes"],
                    "continua": {"id": nuevo_id, "folio": folio_nuevo}, "folio": folio_actual})


@app.route("/api/actas/<int:aid>/reabrir", methods=["POST"])
@requiere_admin
def api_ta_acta_reabrir(aid):
    db = get_db()
    a, e = _acta_o_error(db, aid)
    if e:
        return e
    if a["continua_id"]:
        sig = db.execute("SELECT id,folio FROM ta_actas WHERE id=? AND eliminada=0", (a["continua_id"],)).fetchone()
        if sig:
            return jsonify({"error": "Esta entrega parcial ya continúa en %s. Haz los cambios allá." % sig["folio"]}), 409
    db.execute("UPDATE ta_actas SET estado='borrador', cerrada=NULL, cerrada_por=NULL WHERE id=?", (aid,))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/actas/<int:aid>", methods=["DELETE"])
@requiere_admin
def api_ta_acta_borrar(aid):
    db = get_db()
    a, e = _acta_o_error(db, aid)
    if e:
        return e
    if a["estado"] != "borrador":
        return jsonify({"error": "Solo se puede borrar un acta que sigue en borrador. Reábrela primero."}), 409
    db.execute("UPDATE ta_actas SET eliminada=1 WHERE id=?", (aid,))
    db.execute("UPDATE ta_actas SET continua_id=NULL WHERE continua_id=?", (aid,))
    db.execute("DELETE FROM ta_acta_items WHERE acta_id=?", (aid,))
    db.commit()
    return jsonify({"ok": True})


def _construir_pdf_acta(acta, ruta_salida):
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, KeepTogether
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from xml.sax.saxutils import escape

    OSCURO, AMARILLO, GRIS, ROJO, VERDE = (colors.HexColor("#2F3237"), colors.HexColor("#F5B400"),
                                           colors.HexColor("#E8E6E1"), colors.HexColor("#D0342C"),
                                           colors.HexColor("#1E7B3A"))
    est = getSampleStyleSheet()
    base = ParagraphStyle("b", parent=est["Normal"], fontName="Helvetica", fontSize=8.5, leading=10.5, textColor=OSCURO)
    peq = ParagraphStyle("p", parent=base, fontSize=7, leading=8.5, textColor=colors.HexColor("#6b6f75"))
    neg = ParagraphStyle("n", parent=base, fontName="Helvetica-Bold")
    rojo = ParagraphStyle("r", parent=base, fontSize=7.5, leading=9, textColor=ROJO, fontName="Helvetica-Bold")
    blanco = ParagraphStyle("w", parent=base, fontName="Helvetica-Bold", fontSize=13, leading=15, textColor=colors.white)
    esc = lambda t: escape(str(t or ""))

    ancho = landscape(letter)[0] - 2 * cm
    raya = Table([[""] * 48], colWidths=[ancho / 48.0] * 48, rowHeights=[0.22 * cm])
    raya.setStyle(TableStyle([("BACKGROUND", (i, 0), (i, 0), AMARILLO if i % 2 == 0 else OSCURO) for i in range(48)] +
                             [("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                              ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
    logo = Paragraph("<b>HAP</b>", ParagraphStyle("l", parent=blanco, textColor=OSCURO, fontSize=15))
    titulo = [Paragraph("ACTA DE ENTREGA-RECEPCIÓN Y PRUEBAS DE FUNCIONAMIENTO", blanco),
              Paragraph("Hospital Aranda de la Parra · Proyecto Remodelación de Quirófanos · Torre %s" % ("A" if proyecto_actual() == "torre_a" else "B"),
                        ParagraphStyle("s", parent=base, textColor=colors.HexColor("#D9D6CF"), fontSize=8))]
    folio = Paragraph("FOLIO<br/><font size=14 color='#F5B400'><b>%s</b></font>" % esc(acta["folio"]),
                      ParagraphStyle("f", parent=base, textColor=colors.white, alignment=2, fontSize=7, leading=17))
    cab = Table([[logo, titulo, folio]], colWidths=[2.0 * cm, ancho - 2.0 * cm - 5.6 * cm, 5.6 * cm])
    cab.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), OSCURO), ("BACKGROUND", (0, 0), (0, 0), AMARILLO),
                             ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (0, 0), "CENTER"),
                             ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))

    def celda(rot, valor):
        return Paragraph("<font size=6.5 color='#6b6f75'>%s</font><br/>%s" % (esc(rot).upper(), esc(valor) or "________________"), base)
    r = acta["resumen"]
    fecha_txt = acta["fecha"] and "/".join(reversed(acta["fecha"].split("-"))) or "__ / __ / ____"
    datos = Table([[celda("Bloque", acta["bloque"]), celda("Área", acta.get("area_txt") or "Todo el bloque"),
                    celda("Fecha de recepción", fecha_txt), celda("Estado", ("CERRADA · ENTREGA PARCIAL %d" % acta["parcial_n"] if acta.get("parcial_n") and acta.get("continua") else "CERRADA")
                          if acta["estado"] == "cerrada" else "BORRADOR")],
                   [celda("Recibe obra (Mantenimiento e Infraestructura)", acta["recibe_obra"]),
                    celda("Recibe operación del área", acta["recibe_area"]),
                    celda("Supervisión de proyecto", acta["supervision"]),
                    celda("Representante del proveedor", acta["proveedor_rep"])]],
                  colWidths=[ancho / 4.0] * 4)
    datos.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, OSCURO), ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cfcbc3")),
                               ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))

    estilo_kpi = ParagraphStyle("k", parent=base, leading=21)

    def kpi(num, txt, fondo, color=OSCURO):
        return Paragraph("<font size=17 color='%s'><b>%s</b></font><br/><font size=7.5 color='#444444'>%s</font>" % (color.hexval().replace("0x", "#"), num, txt), estilo_kpi)
    resumen = Table([[kpi(r["actividades"], "actividades en esta acta", AMARILLO),
                      kpi("%d / %d" % (r["conformes"], r["actividades"]), "conformes", GRIS),
                      kpi("%d / %d" % (r["probados"], r["por_probar"]), "elementos probados / por probar", GRIS),
                      kpi(r["observaciones"] + r["pendientes"], "con observación o pendientes", colors.HexColor("#FBE7E5"), ROJO)]],
                    colWidths=[ancho / 4.0] * 4)
    resumen.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, 0), AMARILLO), ("BACKGROUND", (1, 0), (2, 0), GRIS),
                                 ("BACKGROUND", (3, 0), (3, 0), colors.HexColor("#FBE7E5")), ("BOX", (0, 0), (-1, -1), 0.8, OSCURO),
                                 ("INNERGRID", (0, 0), (-1, -1), 0.6, colors.white),
                                 ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))

    parcial = None
    if acta.get("filtro_txt"):
        parcial = Paragraph("<b>ENTREGA PARCIAL DEL ACTA %s</b> · Solo se incluyen las actividades que cumplen: %s. "
                            "El acta completa no cambia." % (esc(acta["folio"]), esc(acta["filtro_txt"])),
                            ParagraphStyle("pa", parent=base, textColor=ROJO, fontSize=8.5))
    if acta.get("parcial_n") and (acta.get("continua") or acta.get("viene_de")):
        txt = "<b>ENTREGA PARCIAL %d.</b> " % acta["parcial_n"]
        if acta.get("viene_de"):
            txt += "Continúa el acta %s. " % esc(acta["viene_de"]["folio"])
        txt += ("Las actividades pendientes pasaron al acta %s." % esc(acta["continua"]["folio"])) if acta.get("continua") else "Es la entrega final de esta serie."
        cadena = Paragraph(txt, ParagraphStyle("ca", parent=base, textColor=ROJO, fontSize=8.5))
        parcial = Table([[parcial], [cadena]], colWidths=[ancho]) if parcial else cadena
    intro = Paragraph("Se hace constar que las actividades listadas fueron revisadas físicamente y sometidas a pruebas de "
                      "funcionamiento por el personal indicado. La columna <b>Probados / Total</b> indica cuántos elementos "
                      "(contactos, apagadores, tomas de gases, alarmas, etc.) se verificaron contra los que el proyecto requiere. "
                      "Cada renglón lleva fecha y nombre/firma de quien realizó la prueba.", base)

    filas = [[Paragraph("<b>NO.</b>", ParagraphStyle("h", parent=base, textColor=colors.white, fontSize=7)),
              Paragraph("<b>ACTIVIDAD (CÓDIGO · PROVEEDOR)</b>", ParagraphStyle("h", parent=base, textColor=colors.white, fontSize=7)),
              Paragraph("<b>PRUEBA LA REALIZA</b>", ParagraphStyle("h", parent=base, textColor=colors.white, fontSize=7)),
              Paragraph("<b>PROBADOS / TOTAL</b>", ParagraphStyle("h", parent=base, textColor=colors.white, fontSize=7, alignment=1)),
              Paragraph("<b>RESULTADO</b>", ParagraphStyle("h", parent=base, textColor=colors.white, fontSize=7, alignment=1)),
              Paragraph("<b>FECHA</b>", ParagraphStyle("h", parent=base, textColor=colors.white, fontSize=7)),
              Paragraph("<b>NOMBRE Y FIRMA DE QUIEN PRUEBA</b>", ParagraphStyle("h", parent=base, textColor=colors.white, fontSize=7))]]
    estilos_fila = []
    centro = ParagraphStyle("c", parent=base, alignment=1)
    for n, it in enumerate(acta["items"], 1):
        col = {"conforme": VERDE, "observacion": ROJO, "pendiente": colors.HexColor("#8a5a00")}[it["resultado"]]
        act = "<b>%s</b><br/><font size=7 color='#6b6f75'>%s%s%s</font>" % (
            esc(it["partida"]), esc(it["codigo"]), (" · " + esc(it["proveedor"])) if it["proveedor"] else "",
            (" · " + esc(it["area"])) if it.get("area") else "")
        celda_act = [Paragraph(act, base)]
        if it["nota"]:
            celda_act.append(Paragraph(esc(it["nota"]), rojo if it["resultado"] != "conforme" else peq))
        fch = it["fecha_prueba"] and "/".join(reversed(it["fecha_prueba"].split("-"))) or ""
        filas.append([Paragraph("<b>%d</b>" % n, centro), celda_act, Paragraph(esc(it["valida_depto"] or "—"), base),
                      Paragraph("<font size=11><b>%d</b></font><font color='#777777'>/%d</font><br/><font size=6.5 color='#6b6f75'>%s</font>"
                                % (it["cant_ok"], it["cant_total"], esc(it["unidad"])), centro),
                      Paragraph("<font size=7 color='%s'><b>%s</b></font>" % (col.hexval().replace("0x", "#"), ETIQUETA_RESULTADO_ACTA[it["resultado"]]), centro),
                      Paragraph(fch, base),
                      Paragraph(esc(it["probado_por"]) + "<br/>&nbsp;<br/>&nbsp;", peq)])
        estilos_fila.append(("LINEBEFORE", (0, n), (0, n), 3, col if it["resultado"] != "conforme" else AMARILLO))
    tabla = Table(filas, colWidths=[1.0 * cm, 8.6 * cm, 3.2 * cm, 2.4 * cm, 2.4 * cm, 2.3 * cm, ancho - 19.9 * cm], repeatRows=1)
    tabla.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), OSCURO), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("LINEBELOW", (0, 1), (-1, -1), 0.4, colors.HexColor("#cfcbc3")),
                               ("LINEAFTER", (4, 1), (5, -1), 0.4, colors.HexColor("#cfcbc3")),
                               ("BOX", (0, 0), (-1, -1), 0.6, OSCURO),
                               ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)] + estilos_fila))

    pendientes_txt = [it for it in acta["items"] if it["resultado"] != "conforme"]
    obs_gral = esc(acta["observaciones"]).replace("\n", "<br/>")
    # una fila por pendiente: la tabla se parte entre páginas (una sola celda gigante reventaba el PDF)
    filas_obs = [[Paragraph("<font size=7 color='#6b6f75'>OBSERVACIONES GENERALES Y PENDIENTES (FECHA COMPROMISO DEL PROVEEDOR)</font>", base)]]
    for n, it in enumerate(pendientes_txt, 1):
        filas_obs.append([Paragraph("<b>%d.</b> %s %s: %s Fecha compromiso: ____ / ____" % (
            n, esc(it["codigo"]), esc(it["partida"]), esc(it["nota"] or "pendiente de probar.")), base)])
    if not pendientes_txt:
        filas_obs.append([Paragraph("Sin pendientes.", base)])
    filas_obs.append([Paragraph(obs_gral or "&nbsp;", base)])
    caja_obs = Table(filas_obs, colWidths=[ancho])
    caja_obs.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, OSCURO), ("TOPPADDING", (0, 0), (-1, -1), 2),
                                  ("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("TOPPADDING", (0, 0), (-1, 0), 6),
                                  ("BOTTOMPADDING", (0, -1), (-1, -1), 12)]))
    decl = Table([[Paragraph("<b>Declaratoria.</b> Las partes firmantes reciben las actividades marcadas como <b>CONFORME</b> a su entera "
                             "satisfacción, con las pruebas de funcionamiento descritas. Las actividades <b>CON OBSERVACIÓN</b> quedan "
                             "pendientes y deberán cerrarse en una nueva acta complementaria. La firma de esta acta no libera al proveedor "
                             "de la garantía ni de las correcciones pendientes aquí anotadas.", base)]], colWidths=[ancho])
    decl.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FFF4CC")), ("LINEBEFORE", (0, 0), (0, 0), 4, AMARILLO),
                              ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))

    def firma(rol, nombre, cargo):
        return [Paragraph("<b>%s</b>" % rol, ParagraphStyle("fr", parent=base, alignment=1)),
                Paragraph(esc(nombre) or "Nombre: ______________", ParagraphStyle("fn", parent=base, alignment=1, fontSize=8)),
                Paragraph(cargo, ParagraphStyle("fc", parent=peq, alignment=1))]
    ancho_f, hueco = (ancho - 3 * 0.8 * cm) / 4.0, 0.8 * cm
    f1 = [firma("Recibe obra", acta["recibe_obra"], "Jefe de Mantenimiento e Infraestructura"),
          firma("Recibe operación del área", acta["recibe_area"], "Responsable del área"),
          firma("Entrega proveedor(es)", acta["proveedor_rep"], "Representante del contratista" + (" · " + acta["filtro_proveedor"] if acta.get("filtro_proveedor") else "")),
          firma("Supervisión de proyecto", acta["supervision"], "Supervisión")]
    fila_blanca = [""] * 7
    fila_firmas = [f1[0], "", f1[1], "", f1[2], "", f1[3]]
    firmas = Table([fila_blanca, fila_firmas], colWidths=[ancho_f, hueco, ancho_f, hueco, ancho_f, hueco, ancho_f],
                   rowHeights=[1.7 * cm, None])
    firmas.setStyle(TableStyle([("LINEABOVE", (c, 1), (c, 1), 1.2, OSCURO) for c in (0, 2, 4, 6)] +
                               [("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 1), (-1, 1), 3)]))

    cierre = [Spacer(1, 8), caja_obs, Spacer(1, 6), KeepTogether([decl, Spacer(1, 4), firmas])]

    def decorar(canv, doc):
        canv.saveState()
        if acta["estado"] != "cerrada":
            canv.setFont("Helvetica-Bold", 70)
            canv.setFillColor(colors.Color(0.85, 0.85, 0.85, alpha=0.35))
            canv.translate(landscape(letter)[0] / 2.0, landscape(letter)[1] / 2.0)
            canv.rotate(28)
            canv.drawCentredString(0, 0, "BORRADOR")
        canv.restoreState()
        canv.saveState()
        canv.setFont("Helvetica", 7)
        canv.setFillColor(colors.HexColor("#6b6f75"))
        canv.drawString(1 * cm, 0.55 * cm, "%s · Generada por la plataforma Control de Obra HAP · documento de evidencia" % acta["folio"])
        canv.drawRightString(landscape(letter)[0] - 1 * cm, 0.55 * cm, "Hoja %d" % doc.page)
        canv.restoreState()

    doc = SimpleDocTemplate(ruta_salida, pagesize=landscape(letter), leftMargin=1 * cm, rightMargin=1 * cm,
                            topMargin=1 * cm, bottomMargin=1.2 * cm, title="Acta " + acta["folio"])
    doc.build([raya, cab, datos, Spacer(1, 6), resumen, Spacer(1, 6)] + ([parcial, Spacer(1, 4)] if parcial else []) +
              [intro, Spacer(1, 6), tabla] + cierre, onFirstPage=decorar, onLaterPages=decorar)


FILTROS_ACTA = (("proveedor", "proveedor", "Proveedor"), ("area", "area", "Área"), ("giro", "giro", "Especialidad"),
                ("valida_depto", "valida_depto", "Prueba la hace"), ("resultado", "resultado", "Resultado"))


def _filtrar_items_acta(items, args):
    """Aplica a los renglones activos los filtros del PDF/pantalla. Devuelve (items, texto de filtros)."""
    items = [i for i in items if i["activo"]]
    partes = []
    for param, campo, rotulo in FILTROS_ACTA:
        val = (args.get(param) or "").strip()
        if val:
            items = [i for i in items if i[campo] == val]
            partes.append("%s = %s" % (rotulo, ETIQUETA_RESULTADO_ACTA.get(val, val) if param == "resultado" else val))
    q = sin_acentos((args.get("q") or "").strip())
    if q:
        items = [i for i in items if q in sin_acentos(" ".join([i["partida"], i["codigo"], i["proveedor"], i["nota"]]))]
        partes.append("Texto contiene “%s”" % args.get("q").strip())
    return items, "; ".join(partes)


@app.route("/api/actas/<int:aid>/pdf", methods=["GET"])
@requiere_consulta_actas
def api_ta_acta_pdf(aid):
    """PDF del acta. Con filtros (?proveedor=&area=&giro=&valida_depto=&resultado=&q=) sale una entrega parcial:
    solo esos renglones (el acta completa no se toca). Los desactivados nunca salen."""
    db = get_db()
    a, e = _acta_o_error(db, aid)
    if e:
        return e
    acta = _acta_a_dict(db, a)
    items, filtro = _filtrar_items_acta(acta["items"], request.args)
    if not items:
        return jsonify({"error": "No hay actividades con esos filtros." if filtro else "El acta no tiene actividades."}), 400
    acta["items"] = items
    acta["resumen"] = _acta_resumen(items)
    nombre = acta["folio"]
    if filtro:
        acta["filtro_txt"] = filtro
        prov = (request.args.get("proveedor") or "").strip()
        if prov:
            acta["filtro_proveedor"] = prov
        extra = re.sub(r"[^A-Za-z0-9]+", "_", sin_acentos(prov or request.args.get("area") or "parcial")).strip("_")[:30]
        nombre = "%s_%s" % (acta["folio"], extra or "parcial")
    salida = ruta_salida_temporal("acta.pdf")
    _construir_pdf_acta(acta, salida)
    return send_file(salida, as_attachment=True, download_name="%s.pdf" % nombre, mimetype="application/pdf")


# ============================================================
#  DASHBOARD PARA DIRECTIVOS (solo lectura + notas)
#  El rol "directivo" solo puede entrar aqui. Ve porcentajes por bloque/area, nunca
#  proveedores ni actividades sueltas. Cada torre sale de su propia base de datos.
# ============================================================
from contextlib import closing as _closing

ROL_DIRECTIVO = "directivo"
ROLES_DASHBOARD = ("admin", ROL_DIRECTIVO)
DASH_UMBRAL_AVANCE = 2.0     # puntos ganados en 7 dias para decir "avanzando"
DASH_DIAS_TENDENCIA = 7
_RUTAS_DIRECTIVO = ("/api/dashboard/", "/static/", "/api/login", "/api/logout", "/api/quien_soy", "/api/cambiar_mi_clave", "/login", "/sw.js")


@app.before_request
def _candado_directivo():
    """Un directivo solo entra al dashboard: cualquier otra ruta se bloquea (aunque conozca la direccion)."""
    if session.get("rol") != ROL_DIRECTIVO:
        return None
    ruta = request.path
    if ruta == "/dashboard" or any(ruta.startswith(r) for r in _RUTAS_DIRECTIVO):
        return None
    if ruta.startswith("/api/"):
        return jsonify({"error": "solo_dashboard"}), 403
    return redirect("/dashboard")


def requiere_dashboard(fn):
    @wraps(fn)
    def envoltura(*a, **k):
        if session.get("rol") not in ROLES_DASHBOARD:
            return jsonify({"error": "sin_permiso"}), 403
        return fn(*a, **k)
    return envoltura


def _proyecto_dash(data=None):
    """('torre_a'|'torre_b', None) o (None, respuesta de error)."""
    p = request.args.get("proyecto") or ((data or {}).get("proyecto") if isinstance(data, dict) else None) or "torre_b"
    if p not in PROYECTOS:
        return None, (jsonify({"error": "Proyecto no válido."}), 400)
    return p, None


def _abrir_db_proyecto(p):
    con = sqlite3.connect(ruta_db(p), timeout=10.0)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA busy_timeout = 5000")
    return con


def _dash_actividades(db):
    """Mismas actividades que cuenta la pantalla principal (obra, aplican, validadas, vigentes)."""
    return db.execute(
        "SELECT id, bloque, area, avance FROM actividades WHERE (aplica IS NULL OR aplica<>'NO') "
        "AND (estado_val IS NULL OR estado_val='validado') AND (eliminada IS NULL OR eliminada=0) "
        "AND COALESCE(mundo,'obra')='obra'").fetchall()


def _dash_promedio(vals):
    return round(sum(vals) / len(vals), 1) if vals else 0.0


def _dash_semaforo(avance, delta):
    if avance >= 100:
        return "terminado"
    if delta >= DASH_UMBRAL_AVANCE:
        return "avanzando"
    if delta > 0:
        return "poco"
    return "sin_cambios"


_DASH_ORDEN_SEM = {"sin_cambios": 0, "poco": 1, "avanzando": 2, "terminado": 3}


def _dash_fecha_entrega(db, p):
    f = (_config_get(db, "fecha_entrega", "") or "").strip()
    if not f and p == "torre_b":
        f = FECHA_ENTREGA
    d = parse_date(f)
    if not d:
        return None, None
    return d.isoformat(), (d - datetime.date.today()).days


def _nombre_sesion():
    """Nombre de quien entro (el que escribio el admin al crear el usuario); si no hay, el usuario."""
    usuario = session.get("usuario")
    try:
        fila = get_db_usuarios().execute("SELECT proveedor FROM usuarios WHERE usuario=?", (usuario,)).fetchone()
    except sqlite3.Error:
        fila = None
    return ((fila["proveedor"] if fila else None) or usuario or "")[:80]


def _dash_notas(db, rol, usuario):
    try:
        filas = db.execute(
            "SELECT id,texto,bloque,area,aviso,autor,usuario,creada FROM dir_notas WHERE eliminada=0 "
            "ORDER BY aviso DESC, id DESC LIMIT 40").fetchall()
    except sqlite3.Error:
        return []
    return [{"id": f["id"], "texto": f["texto"], "bloque": f["bloque"] or "", "area": f["area"] or "",
             "aviso": bool(f["aviso"]), "autor": f["autor"] or f["usuario"] or "", "fecha": f["creada"] or "",
             "puedo_borrar": rol == "admin" or f["usuario"] == usuario} for f in filas]


def _dash_actas(db):
    sal = {"cerradas": 0, "abiertas": 0, "conformes": 0, "actividades": 0}
    try:
        for f in db.execute("SELECT estado, COUNT(*) n FROM ta_actas WHERE eliminada=0 GROUP BY estado"):
            if f["estado"] == "cerrada":
                sal["cerradas"] = f["n"]
            else:
                sal["abiertas"] += f["n"]
        r = db.execute(
            "SELECT COUNT(*) n, SUM(CASE WHEN i.resultado='conforme' THEN 1 ELSE 0 END) ok "
            "FROM ta_acta_items i JOIN ta_actas t ON t.id=i.acta_id WHERE t.eliminada=0 AND i.activo=1").fetchone()
        sal["actividades"] = r["n"] or 0
        sal["conformes"] = r["ok"] or 0
    except sqlite3.Error:
        pass
    return sal


@app.route("/dashboard")
def dashboard_pagina():
    if "usuario" not in session:
        return redirect("/login")
    if session.get("rol") not in ROLES_DASHBOARD:
        return redirect("/")
    return render_template("dashboard.html", es_admin=(session.get("rol") == "admin"))


@app.route("/api/dashboard/datos")
@requiere_dashboard
def api_dash_datos():
    p, err = _proyecto_dash()
    if err:
        return err
    with _closing(_abrir_db_proyecto(p)) as db:
        filas = _dash_actividades(db)
        fecha_ref = (datetime.date.today() - datetime.timedelta(days=DASH_DIAS_TENDENCIA)).isoformat()
        try:
            antes, _origen, _foto = _avance_a_fecha(db, fecha_ref, "obra")
        except Exception:
            antes = {}
        bloques = {}
        todos_ahora, todos_antes = [], []
        for f in filas:
            av = f["avance"] or 0
            av0 = antes.get(f["id"], av)
            todos_ahora.append(av)
            todos_antes.append(av0)
            b = bloques.setdefault(f["bloque"] or "—", {"ahora": [], "antes": [], "areas": {}})
            b["ahora"].append(av); b["antes"].append(av0)
            a = b["areas"].setdefault(f["area"] or "(sin área)", {"ahora": [], "antes": []})
            a["ahora"].append(av); a["antes"].append(av0)
        salida, conteo = [], {"avanzando": 0, "poco": 0, "sin_cambios": 0, "terminado": 0}
        for nombre, b in bloques.items():
            av, delta = _dash_promedio(b["ahora"]), round(_dash_promedio(b["ahora"]) - _dash_promedio(b["antes"]), 1)
            sem = _dash_semaforo(av, delta)
            conteo[sem] += 1
            areas = []
            for an, a in b["areas"].items():
                aav, ad = _dash_promedio(a["ahora"]), round(_dash_promedio(a["ahora"]) - _dash_promedio(a["antes"]), 1)
                areas.append({"area": an, "avance": aav, "delta": ad, "semaforo": _dash_semaforo(aav, ad)})
            areas.sort(key=lambda x: x["area"])
            salida.append({"bloque": nombre, "avance": av, "delta": delta, "semaforo": sem, "areas": areas})
        salida.sort(key=lambda x: (_DASH_ORDEN_SEM[x["semaforo"]], x["avance"], x["bloque"]))
        glob = _dash_promedio(todos_ahora)
        glob_delta = round(glob - _dash_promedio(todos_antes), 1) if filas else 0.0
        fecha_e, dias = _dash_fecha_entrega(db, p)
        rol, usuario = session.get("rol"), session.get("usuario")
        return jsonify({
            "proyecto": p, "torre": "A" if p == "torre_a" else "B",
            "actualizado": datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
            "total_actividades": len(filas), "avance_global": glob, "delta_global": glob_delta,
            "dias_tendencia": DASH_DIAS_TENDENCIA, "umbral": DASH_UMBRAL_AVANCE,
            "fecha_entrega": fecha_e, "dias_restantes": dias,
            "bloques": salida, "conteo": conteo, "actas": _dash_actas(db),
            "notas": _dash_notas(db, rol, usuario),
            "rol": rol, "usuario": usuario, "nombre": _nombre_sesion(), "puede_avisos": rol == "admin",
            "hay_plano": os.path.exists(PLANO_ARCHIVO_TORRE_A if p == "torre_a" else PLANO_ARCHIVO),
        })


def _area_poligono(pts):
    return abs(sum(pts[i][0] * pts[(i + 1) % len(pts)][1] - pts[(i + 1) % len(pts)][0] * pts[i][1]
                   for i in range(len(pts)))) / 2.0


@app.route("/api/dashboard/plano")
@requiere_dashboard
def api_dash_plano():
    """Zonas del plano pintadas por avance (solo nombre del area, porcentaje y cuantas actividades)."""
    p, err = _proyecto_dash()
    if err:
        return err
    archivo = PLANO_ARCHIVO_TORRE_A if p == "torre_a" else PLANO_ARCHIVO
    with _closing(_abrir_db_proyecto(p)) as db:
        por_area, por_bloque = {}, {}
        for f in _dash_actividades(db):
            av = f["avance"] or 0
            por_bloque.setdefault(f["bloque"] or "", []).append(av)
            if f["area"]:
                por_area.setdefault((f["bloque"] or "", f["area"]), []).append(av)
        ubs = {}
        for u in db.execute("SELECT trazo_id, bloque, area FROM plano_trazo_ub"):
            ubs.setdefault(u["trazo_id"], []).append((u["bloque"] or "", u["area"] or ""))
        zonas = []
        for t in db.execute(
                "SELECT id, nombre, puntos, tipo_elem FROM plano_trazos WHERE tipo='zona' "
                "AND (COALESCE(mundo,'obra')='obra' OR tipo_elem IN ('Bloque','Área')) ORDER BY id"):
            lista = ubs.get(t["id"])
            if not lista:
                continue
            try:
                pts = json.loads(t["puntos"])
            except ValueError:
                continue
            if not isinstance(pts, list) or len(pts) < 3:
                continue
            vals = []
            for b, a in lista:
                vals += por_area.get((b, a), []) if a else por_bloque.get(b, [])
            etiqueta = " + ".join((a or b) for b, a in lista)[:120]
            zonas.append({"id": t["id"], "nombre": etiqueta, "tipo_elem": t["tipo_elem"] or "",
                          "puntos": pts, "n": len(vals),
                          "avance": round(sum(vals) / len(vals), 1) if vals else None})
        # los contornos de bloque van al fondo y las zonas chicas arriba
        zonas.sort(key=lambda z: (0 if z["tipo_elem"] == "Bloque" else 1, -_area_poligono(z["puntos"])))
        version = _config_get(db, "plano_version", "0")
    return jsonify({"existe": os.path.exists(archivo), "zonas": zonas, "version": version})


@app.route("/api/dashboard/plano/imagen")
@requiere_dashboard
def api_dash_plano_imagen():
    p, err = _proyecto_dash()
    if err:
        return err
    archivo = PLANO_ARCHIVO_TORRE_A if p == "torre_a" else PLANO_ARCHIVO
    if not os.path.exists(archivo):
        return jsonify({"error": "sin_plano"}), 404
    with _closing(_abrir_db_proyecto(p)) as db:
        mime = _config_get(db, "plano_mime", "image/jpeg")
    return send_file(archivo, mimetype=mime, max_age=0, conditional=True)


@app.route("/api/dashboard/notas", methods=["POST"])
@requiere_dashboard
def api_dash_nota_crear():
    data = request.get_json(silent=True) or {}
    p, err = _proyecto_dash(data)
    if err:
        return err
    texto = _acta_texto(data.get("texto"), 600)
    if len(texto) < 2:
        return jsonify({"error": "Escribe la nota."}), 400
    aviso = 1 if data.get("aviso") else 0
    if aviso and session.get("rol") != "admin":
        return jsonify({"error": "Solo el administrador publica avisos."}), 403
    bloque = _acta_texto(data.get("bloque"), 200) or None
    area = _acta_texto(data.get("area"), 300) or None
    with _closing(_abrir_db_proyecto(p)) as db:
        if bloque:
            filas = _dash_actividades(db)
            if bloque not in {f["bloque"] for f in filas}:
                return jsonify({"error": "Ese bloque no existe."}), 400
            if area and (bloque, area) not in {(f["bloque"], f["area"]) for f in filas}:
                return jsonify({"error": "Esa área no es del bloque elegido."}), 400
        else:
            area = None
        usuario = session.get("usuario")
        autor = _nombre_sesion()
        cur = db.execute(
            "INSERT INTO dir_notas (texto,bloque,area,aviso,autor,usuario,rol,creada) VALUES (?,?,?,?,?,?,?,?)",
            (texto, bloque, area, aviso, autor, usuario, session.get("rol"),
             datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds")))
        db.commit()
    return jsonify({"ok": True, "id": cur.lastrowid}), 201


@app.route("/api/dashboard/notas/<int:nid>", methods=["DELETE"])
@requiere_dashboard
def api_dash_nota_borrar(nid):
    p, err = _proyecto_dash()
    if err:
        return err
    with _closing(_abrir_db_proyecto(p)) as db:
        n = db.execute("SELECT usuario FROM dir_notas WHERE id=? AND eliminada=0", (nid,)).fetchone()
        if not n:
            return jsonify({"error": "No existe esa nota."}), 404
        if session.get("rol") != "admin" and n["usuario"] != session.get("usuario"):
            return jsonify({"error": "Solo puedes borrar tus propias notas."}), 403
        db.execute("UPDATE dir_notas SET eliminada=1 WHERE id=?", (nid,))
        db.commit()
    return jsonify({"ok": True})


@app.route("/api/dashboard/fecha-entrega", methods=["POST"])
@requiere_admin
def api_dash_fecha_entrega():
    data = request.get_json(silent=True) or {}
    p, err = _proyecto_dash(data)
    if err:
        return err
    fecha = _acta_texto(data.get("fecha"), 10)
    if fecha and not _acta_fecha_valida(fecha):
        return jsonify({"error": "Fecha no válida."}), 400
    with _closing(_abrir_db_proyecto(p)) as db:
        _config_set(db, "fecha_entrega", fecha)
        db.commit()
        f, dias = _dash_fecha_entrega(db, p)
    return jsonify({"ok": True, "fecha_entrega": f, "dias_restantes": dias})


@app.route("/actas")
def actas_pagina():
    if "usuario" not in session:
        return redirect("/login")
    if session.get("rol") not in ROLES_CONSULTA_ACTAS:
        return redirect("/")
    return render_template("actas.html", torre=("A" if proyecto_actual() == "torre_a" else "B"),
                           solo_lectura=(session.get("rol") != "admin"))


@app.route("/torre-a/plano")
def torre_a_plano_pagina():
    """Ya no hay pagina aparte: el plano de Torre A es el Plano de obra de siempre."""
    return redirect("/plano")


@app.route("/torre-a")
def torre_a_pagina():
    """Ya no hay pagina aparte: Torre A usa la pantalla principal."""
    return redirect("/")


@app.route("/usuarios")
def usuarios_page():
    return render_template("usuarios.html")


@app.route("/validacion")
def validacion_page_corta():
    return render_template("validacion.html")



# ============================================================
#  SEGURIDAD E HIGIENE (visitas, hallazgos, doble validacion, respuesta del proveedor)
#  Un inspector captura; validan Mantenimiento (admin) y Arquitectos; despues el proveedor
#  responde con texto y fotos. Hallazgos/fotos viven en la base de cada torre; el catalogo,
#  las normas y los roles viven en la base principal (uno solo para las dos torres).
# ============================================================
import base64 as _b64
import re as _re_sh

SH_TIPOS = {"acto": "Acto inseguro", "condicion": "Condición insegura", "observacion": "Observación"}
SH_META_SEMANA = 2
SH_MAX_FOTOS = 4
SH_MAX_FOTO_BYTES = 3 * 1024 * 1024
_RUTAS_SEGURIDAD = ("/api/sh/", "/seguridad", "/static/", "/api/login", "/api/logout", "/api/quien_soy",
                    "/api/cambiar_mi_clave", "/login", "/sw.js")


@app.before_request
def _candado_seguridad():
    """El rol 'seguridad' (inspector) solo entra a su pantalla; lo demas se le bloquea."""
    if session.get("rol") != "seguridad":
        return None
    ruta = request.path
    if any(ruta == r or ruta.startswith(r) for r in _RUTAS_SEGURIDAD):
        return None
    if ruta.startswith("/api/"):
        return jsonify({"error": "solo_seguridad"}), 403
    return redirect("/seguridad")


def _sh_ahora():
    return datetime.datetime.utcnow().isoformat(timespec="seconds") + "Z"


def _sh_a_local(iso):
    """ISO UTC ('...Z') -> datetime hora de Mexico (centro)."""
    if not iso:
        return None
    try:
        d = datetime.datetime.fromisoformat(str(iso).replace("Z", "+00:00"))
    except ValueError:
        return None
    if d.tzinfo is None:
        d = d.replace(tzinfo=datetime.timezone.utc)
    try:
        from zoneinfo import ZoneInfo
        return d.astimezone(ZoneInfo("America/Mexico_City"))
    except Exception:
        return d.astimezone(datetime.timezone(datetime.timedelta(hours=-6)))


def _sh_hoy():
    return _sh_a_local(_sh_ahora()).date()


def _sh_fmt(iso, con_hora=True):
    d = _sh_a_local(iso)
    if not d:
        return ""
    return d.strftime("%d/%m/%Y %H:%M") if con_hora else d.strftime("%d/%m/%Y")


def _sh_fmt_fecha(f):
    d = parse_date(f) if f else None
    return d.strftime("%d/%m/%Y") if d else ""


def _sh_flags():
    """Permisos de quien pregunta: admin, inspector, arquitecto, proveedor."""
    usuario, rol, proveedor = usuario_actual()
    f = {"admin": rol == "admin", "inspector": rol == "seguridad", "arquitecto": False,
         "mantenimiento": rol == "admin", "proveedor": False, "usuario": usuario, "rol": rol, "nombre_prov": proveedor or ""}
    if not usuario:
        return f
    try:
        r = get_db_usuarios().execute("SELECT rol FROM sh_roles WHERE usuario=?", (usuario,)).fetchone()
        if r:
            if r["rol"] == "inspector":
                f["inspector"] = True
            elif r["rol"] == "arquitecto":
                f["arquitecto"] = True
            elif r["rol"] == "mantenimiento":
                f["mantenimiento"] = True
    except Exception:
        pass
    f["proveedor"] = rol == "proveedor" and bool(proveedor)
    f["acceso"] = f["admin"] or f["inspector"] or f["arquitecto"] or f["mantenimiento"] or f["proveedor"]
    f["interno"] = f["admin"] or f["inspector"] or f["arquitecto"] or f["mantenimiento"]
    return f


def requiere_sh(fn):
    @wraps(fn)
    def envoltura(*a, **k):
        if "usuario" not in session:
            return jsonify({"error": "no_login"}), 401
        if not _sh_flags()["acceso"]:
            return jsonify({"error": "sin_permiso"}), 403
        return fn(*a, **k)
    return envoltura


def _sh_proyecto(data=None):
    p = request.args.get("proyecto") or ((data or {}).get("proyecto") if isinstance(data, dict) else None) or "torre_b"
    if p not in PROYECTOS:
        return None, (jsonify({"error": "Proyecto no válido."}), 400)
    return p, None


def _sh_letra(p):
    return "A" if p == "torre_a" else "B"


def _sh_folio_h(p, hid):
    return "SH-%s-%04d" % (_sh_letra(p), hid)


def _sh_limpio(v, largo=600):
    return (str(v).strip() if v is not None else "")[:largo]


def _sh_normas_lista(v):
    if isinstance(v, str):
        v = v.split(",")
    claves = []
    for x in (v or []):
        x = _sh_limpio(x, 60)
        if x and x not in claves:
            claves.append(x)
    return claves[:8]


# ---------- fotos ----------
def _sh_guardar_fotos(db, hid, fase, fotos, usuario):
    """fotos: lista de 'data:image/jpeg;base64,...'. Solo JPEG (el navegador ya las reduce)."""
    ya = db.execute("SELECT COUNT(*) FROM sh_fotos WHERE hallazgo_id=? AND fase=?", (hid, fase)).fetchone()[0]
    guardadas = 0
    for d in (fotos or []):
        if ya + guardadas >= SH_MAX_FOTOS:
            break
        m = _re_sh.match(r"^data:image/(jpeg|jpg);base64,(.+)$", str(d or ""), _re_sh.S)
        if not m:
            continue
        try:
            raw = _b64.b64decode(m.group(2), validate=False)
        except Exception:
            continue
        if len(raw) < 200 or len(raw) > SH_MAX_FOTO_BYTES or raw[:3] != b"\xff\xd8\xff":
            continue
        db.execute("INSERT INTO sh_fotos (hallazgo_id,fase,mime,datos,subida_por,creada) VALUES (?,?,?,?,?,?)",
                   (hid, fase, "image/jpeg", raw, usuario, _sh_ahora()))
        guardadas += 1
    return guardadas


def _sh_fotos_ids(db, hid):
    out = {"hallazgo": [], "respuesta": []}
    for f in db.execute("SELECT id, fase FROM sh_fotos WHERE hallazgo_id=? ORDER BY id", (hid,)):
        out.setdefault(f["fase"], []).append(f["id"])
    return out


# ---------- visibilidad ----------
def _sh_puede_ver(fl, h):
    if fl["interno"]:
        return True
    return fl["proveedor"] and (h["proveedor"] or "") == fl["nombre_prov"] and h["estado"] in ("enviado", "respondido", "cerrado")


def _sh_dict(h, p, db=None, detalle=False):
    d = {k: h[k] for k in h.keys() if k != "eliminada"}
    d["folio"] = _sh_folio_h(p, h["id"])
    d["tipo_texto"] = SH_TIPOS.get(h["tipo"], h["tipo"])
    d["normas_lista"] = [x for x in (h["normas"] or "").split(",") if x]
    d["retroactivo"] = bool(h["retroactivo"])
    hoy = _sh_hoy().isoformat()
    comp = h["resp_compromiso"] or h["fecha_compromiso"] or ""
    d["vencido"] = bool(comp and comp < hoy and h["estado"] in ("registrado", "enviado", "devuelto"))
    if db is not None and detalle:
        d["fotos"] = _sh_fotos_ids(db, h["id"])
    return d


def _sh_obtener(db, hid):
    return db.execute("SELECT * FROM sh_hallazgos WHERE id=? AND eliminada=0", (hid,)).fetchone()


# ---------- avisos (push al admin; links de WhatsApp / correo listos) ----------
def _sh_texto_aviso(h, p, fase):
    cab = {"registro": "Nuevo hallazgo por validar", "enviado": "Hallazgo de seguridad e higiene para atender",
           "respuesta": "El proveedor respondió un hallazgo", "devuelto": "Hallazgo devuelto"}.get(fase, "Hallazgo")
    ub = " · ".join(x for x in (h["bloque"], h["area"]) if x)
    t = "HAP Seguridad e Higiene · %s\n%s (%s)\nTorre %s · %s\n%s" % (
        cab, _sh_folio_h(p, h["id"]), SH_TIPOS.get(h["tipo"], h["tipo"]), _sh_letra(p), ub or "sin ubicación",
        h["texto"])
    if h["proveedor"]:
        t += "\nProveedor: %s" % h["proveedor"]
    if h["accion_propuesta"]:
        t += "\nAcción propuesta: %s" % h["accion_propuesta"]
    if h["fecha_compromiso"]:
        t += "\nFecha compromiso: %s" % _sh_fmt_fecha(h["fecha_compromiso"])
    t += "\nEntra a: %s/seguridad" % request.host_url.rstrip("/")
    return t


def _sh_destinatarios(h, fase, excluir=None):
    dbu = get_db_usuarios()
    usuarios = {}
    if fase in ("registro", "respuesta"):
        for r in dbu.execute("SELECT usuario, proveedor, telefono FROM usuarios WHERE rol='admin' AND COALESCE(activo,1)=1"):
            usuarios[r["usuario"]] = r
        for r in dbu.execute("SELECT u.usuario, u.proveedor, u.telefono FROM usuarios u JOIN sh_roles s ON s.usuario=u.usuario "
                             "WHERE s.rol IN ('arquitecto','mantenimiento') AND COALESCE(u.activo,1)=1"):
            usuarios[r["usuario"]] = r
        if fase == "respuesta":
            for r in dbu.execute("SELECT u.usuario, u.proveedor, u.telefono FROM usuarios u LEFT JOIN sh_roles s ON s.usuario=u.usuario "
                                 "WHERE (u.rol='seguridad' OR s.rol='inspector') AND COALESCE(u.activo,1)=1"):
                usuarios[r["usuario"]] = r
    elif fase == "enviado" and h["proveedor"]:
        for r in dbu.execute("SELECT usuario, proveedor, telefono FROM usuarios WHERE rol='proveedor' AND proveedor=? "
                             "AND COALESCE(activo,1)=1", (h["proveedor"],)):
            usuarios[r["usuario"]] = r
    elif fase == "devuelto" and h["capturado_por"]:
        r = dbu.execute("SELECT usuario, proveedor, telefono FROM usuarios WHERE usuario=?", (h["capturado_por"],)).fetchone()
        if r:
            usuarios[r["usuario"]] = r
    return [u for k, u in usuarios.items() if k != excluir]


def _sh_wa(tel, texto):
    solo = "".join(ch for ch in (tel or "") if ch.isdigit())
    if not solo:
        return None
    if len(solo) == 10:
        solo = "52" + solo
    from urllib.parse import quote
    return "https://wa.me/%s?text=%s" % (solo, quote(texto))


def _sh_avisos(h, p, fase):
    texto = _sh_texto_aviso(h, p, fase)
    from urllib.parse import quote
    out = []
    for u in _sh_destinatarios(h, fase, excluir=session.get("usuario")):
        out.append({"usuario": u["usuario"], "nombre": u["proveedor"] or u["usuario"],
                    "whatsapp": _sh_wa(u["telefono"], texto), "tiene_telefono": bool((u["telefono"] or "").strip())})
    asunto = "HAP S. e H. · %s" % _sh_folio_h(p, h["id"])
    return {"texto": texto, "destinatarios": out,
            "correo": "mailto:?subject=%s&body=%s" % (quote(asunto), quote(texto))}


def _sh_push(h, p, fase):
    try:
        titulo = {"registro": "Hallazgo por validar", "respuesta": "Respuesta del proveedor"}.get(fase, "Seguridad e higiene")
        enviar_push_admins(titulo, "%s · %s" % (_sh_folio_h(p, h["id"]), h["texto"][:90]), "/seguridad")
    except Exception:
        pass


# ---------- ubicaciones, proveedores, plano ----------
def _sh_ubicaciones(db, p):
    if p == "torre_a":
        salida = []
        for b in db.execute("SELECT num,nombre FROM ta_bloques ORDER BY num").fetchall():
            areas = [etiqueta_area_ta(a["bloque"], a["num"], a["nombre"]) for a in
                     db.execute("SELECT bloque,num,nombre FROM ta_areas WHERE bloque=? ORDER BY num", (b["num"],))]
            salida.append({"nombre": etiqueta_bloque_ta(b["num"], b["nombre"]), "areas": areas})
        return salida
    mapa = {}
    for f in db.execute("SELECT DISTINCT bloque, area FROM actividades WHERE (eliminada IS NULL OR eliminada=0) "
                        "AND COALESCE(mundo,'obra')='obra' AND bloque IS NOT NULL AND bloque<>''"):
        mapa.setdefault(f["bloque"], set())
        if f["area"]:
            mapa[f["bloque"]].add(f["area"])
    return [{"nombre": b, "areas": sorted(a)} for b, a in sorted(mapa.items())]


def _sh_proveedores(db, p):
    nombres = set()
    if p == "torre_a":
        nombres |= {r["nombre"] for r in db.execute("SELECT nombre FROM ta_proveedores")}
    else:
        nombres |= {r["nombre"] for r in db.execute("SELECT nombre FROM proveedores WHERE nombre IS NOT NULL")}
        nombres |= {r["proveedor"] for r in db.execute(
            "SELECT DISTINCT proveedor FROM actividades WHERE COALESCE(mundo,'obra')='obra' AND proveedor IS NOT NULL AND proveedor<>''")}
    try:
        nombres |= {r["proveedor"] for r in get_db_usuarios().execute(
            "SELECT DISTINCT proveedor FROM usuarios WHERE rol='proveedor' AND proveedor IS NOT NULL AND proveedor<>''")}
    except Exception:
        pass
    return sorted(n for n in nombres if n)


@app.route("/seguridad")
def seguridad_pagina():
    if "usuario" not in session:
        return redirect("/login")
    if not _sh_flags()["acceso"]:
        return redirect("/")
    return render_template("seguridad.html")


@app.route("/api/sh/yo")
@requiere_sh
def api_sh_yo():
    fl = _sh_flags()
    volver = "/" if (fl["admin"] or session.get("rol") in ROLES_GESTORES) else ("/portal" if fl["proveedor"] else "")
    return jsonify({"usuario": fl["usuario"], "rol": fl["rol"], "admin": fl["admin"], "inspector": fl["inspector"],
                    "arquitecto": fl["arquitecto"], "mantenimiento": fl["mantenimiento"], "proveedor": fl["proveedor"], "nombre_prov": fl["nombre_prov"],
                    "volver": volver})


@app.route("/api/sh/base")
@requiere_sh
def api_sh_base():
    """Todo lo que la pantalla necesita para arrancar en una torre."""
    p, err = _sh_proyecto()
    if err:
        return err
    fl = _sh_flags()
    with _closing(_abrir_db_proyecto(p)) as db:
        ub = _sh_ubicaciones(db, p) if fl["interno"] else []
        provs = _sh_proveedores(db, p) if fl["interno"] else []
        version = _config_get(db, "plano_version", "0")
    archivo = PLANO_ARCHIVO_TORRE_A if p == "torre_a" else PLANO_ARCHIVO
    return jsonify({"ubicaciones": ub, "proveedores": provs, "plano": os.path.exists(archivo), "version": version,
                    "tipos": SH_TIPOS, "meta_semana": SH_META_SEMANA})


@app.route("/api/sh/plano")
@requiere_sh
def api_sh_plano():
    p, err = _sh_proyecto()
    if err:
        return err
    if not _sh_flags()["interno"]:
        return jsonify({"error": "sin_permiso"}), 403
    with _closing(_abrir_db_proyecto(p)) as db:
        ubs = {}
        for u in db.execute("SELECT trazo_id, bloque, area FROM plano_trazo_ub"):
            ubs.setdefault(u["trazo_id"], []).append((u["bloque"] or "", u["area"] or ""))
        zonas = []
        for t in db.execute("SELECT id, nombre, puntos, tipo_elem FROM plano_trazos WHERE tipo='zona' "
                            "AND (COALESCE(mundo,'obra')='obra' OR tipo_elem IN ('Bloque','Área')) ORDER BY id"):
            lista = ubs.get(t["id"])
            if not lista:
                continue
            try:
                pts = json.loads(t["puntos"])
            except ValueError:
                continue
            if not isinstance(pts, list) or len(pts) < 3:
                continue
            b, a = lista[0]
            zonas.append({"id": t["id"], "nombre": " + ".join((x or y) for y, x in lista)[:120], "bloque": b, "area": a,
                          "tipo_elem": t["tipo_elem"] or "", "puntos": pts})
        zonas.sort(key=lambda z: (0 if z["tipo_elem"] == "Bloque" else 1, -_area_poligono(z["puntos"])))
        version = _config_get(db, "plano_version", "0")
    archivo = PLANO_ARCHIVO_TORRE_A if p == "torre_a" else PLANO_ARCHIVO
    return jsonify({"existe": os.path.exists(archivo), "zonas": zonas, "version": version})


@app.route("/api/sh/plano/imagen")
@requiere_sh
def api_sh_plano_imagen():
    p, err = _sh_proyecto()
    if err:
        return err
    archivo = PLANO_ARCHIVO_TORRE_A if p == "torre_a" else PLANO_ARCHIVO
    if not os.path.exists(archivo):
        return jsonify({"error": "sin_plano"}), 404
    with _closing(_abrir_db_proyecto(p)) as db:
        mime = _config_get(db, "plano_mime", "image/jpeg")
    return send_file(archivo, mimetype=mime, max_age=0, conditional=True)


# ---------- catalogo (compartido por las dos torres) ----------
def _sh_puede_catalogo():
    fl = _sh_flags()
    return fl["admin"] or fl["inspector"]


@app.route("/api/sh/catalogo")
@requiere_sh
def api_sh_catalogo():
    fl = _sh_flags()
    if not fl["interno"]:
        return jsonify({"error": "sin_permiso"}), 403
    todo = request.args.get("todo") == "1" and _sh_puede_catalogo()
    db = get_db_usuarios()
    normas = [dict(r) for r in db.execute("SELECT id, clave, descripcion, activa FROM sh_normas "
                                          + ("" if todo else "WHERE activa=1 ") + "ORDER BY clave")]
    puntos = []
    for pt in db.execute("SELECT * FROM sh_puntos " + ("" if todo else "WHERE activo=1 ") + "ORDER BY orden, id"):
        ops = []
        for o in db.execute("SELECT * FROM sh_opciones WHERE punto_id=? " + ("" if todo else "AND activo=1 ") + "ORDER BY orden, id", (pt["id"],)):
            ops.append({"id": o["id"], "texto": o["texto"], "tipo": o["tipo"], "activo": o["activo"],
                        "normas": [x for x in (o["normas"] or "").split(",") if x]})
        puntos.append({"id": pt["id"], "categoria": pt["categoria"], "texto": pt["texto"], "activo": pt["activo"], "opciones": ops})
    return jsonify({"puntos": puntos, "normas": normas, "tipos": SH_TIPOS})


def _sh_exige_catalogo():
    if not _sh_puede_catalogo():
        return jsonify({"error": "Solo el administrador o Seguridad e Higiene editan el catálogo."}), 403
    return None


@app.route("/api/sh/catalogo/punto", methods=["POST"])
@requiere_sh
def api_sh_punto_crear():
    e = _sh_exige_catalogo()
    if e:
        return e
    d = request.get_json(silent=True) or {}
    cat, txt = _sh_limpio(d.get("categoria"), 80), _sh_limpio(d.get("texto"), 160)
    if not cat or not txt:
        return jsonify({"error": "Escribe la categoría y el punto a revisar."}), 400
    db = get_db_usuarios()
    n = db.execute("SELECT COALESCE(MAX(orden),0)+1 FROM sh_puntos").fetchone()[0]
    cur = db.execute("INSERT INTO sh_puntos (categoria,texto,orden,activo) VALUES (?,?,?,1)", (cat, txt, n))
    db.commit()
    return jsonify({"ok": True, "id": cur.lastrowid})


@app.route("/api/sh/catalogo/punto/<int:pid>", methods=["PUT"])
@requiere_sh
def api_sh_punto_editar(pid):
    e = _sh_exige_catalogo()
    if e:
        return e
    d = request.get_json(silent=True) or {}
    db = get_db_usuarios()
    r = db.execute("SELECT * FROM sh_puntos WHERE id=?", (pid,)).fetchone()
    if not r:
        return jsonify({"error": "No existe."}), 404
    txt = _sh_limpio(d.get("texto"), 160) or r["texto"]
    cat = _sh_limpio(d.get("categoria"), 80) or r["categoria"]
    act = r["activo"] if d.get("activo") is None else (1 if d.get("activo") else 0)
    db.execute("UPDATE sh_puntos SET texto=?, categoria=?, activo=? WHERE id=?", (txt, cat, act, pid))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/sh/catalogo/opcion", methods=["POST"])
@requiere_sh
def api_sh_opcion_crear():
    e = _sh_exige_catalogo()
    if e:
        return e
    d = request.get_json(silent=True) or {}
    db = get_db_usuarios()
    if not db.execute("SELECT 1 FROM sh_puntos WHERE id=?", (d.get("punto_id"),)).fetchone():
        return jsonify({"error": "El punto no existe."}), 400
    txt = _sh_limpio(d.get("texto"), 200)
    tipo = d.get("tipo") if d.get("tipo") in SH_TIPOS else "condicion"
    if not txt:
        return jsonify({"error": "Escribe el hallazgo."}), 400
    n = db.execute("SELECT COALESCE(MAX(orden),0)+1 FROM sh_opciones WHERE punto_id=?", (d["punto_id"],)).fetchone()[0]
    cur = db.execute("INSERT INTO sh_opciones (punto_id,texto,tipo,normas,orden,activo) VALUES (?,?,?,?,?,1)",
                     (d["punto_id"], txt, tipo, ",".join(_sh_normas_lista(d.get("normas"))), n))
    db.commit()
    return jsonify({"ok": True, "id": cur.lastrowid})


@app.route("/api/sh/catalogo/opcion/<int:oid>", methods=["PUT"])
@requiere_sh
def api_sh_opcion_editar(oid):
    e = _sh_exige_catalogo()
    if e:
        return e
    d = request.get_json(silent=True) or {}
    db = get_db_usuarios()
    r = db.execute("SELECT * FROM sh_opciones WHERE id=?", (oid,)).fetchone()
    if not r:
        return jsonify({"error": "No existe."}), 404
    txt = _sh_limpio(d.get("texto"), 200) or r["texto"]
    tipo = d.get("tipo") if d.get("tipo") in SH_TIPOS else r["tipo"]
    normas = r["normas"] if d.get("normas") is None else ",".join(_sh_normas_lista(d.get("normas")))
    act = r["activo"] if d.get("activo") is None else (1 if d.get("activo") else 0)
    db.execute("UPDATE sh_opciones SET texto=?, tipo=?, normas=?, activo=? WHERE id=?", (txt, tipo, normas, act, oid))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/sh/catalogo/norma", methods=["POST"])
@requiere_sh
def api_sh_norma_crear():
    e = _sh_exige_catalogo()
    if e:
        return e
    d = request.get_json(silent=True) or {}
    clave, desc = _sh_limpio(d.get("clave"), 60).upper(), _sh_limpio(d.get("descripcion"), 200)
    if not clave:
        return jsonify({"error": "Escribe la clave de la norma, por ejemplo NOM-017-STPS-2008."}), 400
    db = get_db_usuarios()
    try:
        db.execute("INSERT INTO sh_normas (clave,descripcion,activa) VALUES (?,?,1)", (clave, desc))
        db.commit()
    except sqlite3.IntegrityError:
        return jsonify({"error": "Esa norma ya está en la lista."}), 400
    return jsonify({"ok": True})


@app.route("/api/sh/catalogo/norma/<int:nid>", methods=["PUT"])
@requiere_sh
def api_sh_norma_editar(nid):
    e = _sh_exige_catalogo()
    if e:
        return e
    d = request.get_json(silent=True) or {}
    db = get_db_usuarios()
    r = db.execute("SELECT * FROM sh_normas WHERE id=?", (nid,)).fetchone()
    if not r:
        return jsonify({"error": "No existe."}), 404
    desc = r["descripcion"] if d.get("descripcion") is None else _sh_limpio(d.get("descripcion"), 200)
    act = r["activa"] if d.get("activa") is None else (1 if d.get("activa") else 0)
    db.execute("UPDATE sh_normas SET descripcion=?, activa=? WHERE id=?", (desc, act, nid))
    db.commit()
    return jsonify({"ok": True})


# ---------- visitas ----------
def _sh_semana(db):
    hoy = _sh_hoy()
    lunes = hoy - datetime.timedelta(days=hoy.weekday())
    n = 0
    for v in db.execute("SELECT inicio FROM sh_visitas"):
        d = _sh_a_local(v["inicio"])
        if d and d.date() >= lunes and d.date() <= lunes + datetime.timedelta(days=6):
            n += 1
    return {"n": n, "meta": SH_META_SEMANA, "lunes": lunes.isoformat()}


def _sh_lugar(bloque, area=""):
    """Clave del lugar de una revision: 'bloque' o 'bloque|area'."""
    b, a = (bloque or "").strip(), (area or "").strip()
    return (b + "|" + a) if (b and a) else b


_SQL_LUGAR = "(COALESCE(bloque,'') || CASE WHEN COALESCE(area,'')<>'' AND COALESCE(bloque,'')<>'' THEN '|' || area ELSE '' END)"


def _sh_visita_dict(v):
    return {"id": v["id"], "folio": v["folio"], "estado": v["estado"], "inicio": v["inicio"], "fin": v["fin"],
            "iniciada_por": v["iniciada_por"], "cerrada_por": v["cerrada_por"], "notas": v["notas"] or ""}


@app.route("/api/sh/visitas")
@requiere_sh
def api_sh_visitas():
    p, err = _sh_proyecto()
    if err:
        return err
    if not _sh_flags()["interno"]:
        return jsonify({"error": "sin_permiso"}), 403
    with _closing(_abrir_db_proyecto(p)) as db:
        act = db.execute("SELECT * FROM sh_visitas WHERE estado='en_curso' ORDER BY id DESC LIMIT 1").fetchone()
        recientes = [_sh_visita_dict(v) for v in db.execute("SELECT * FROM sh_visitas ORDER BY id DESC LIMIT 12")]
        respuestas = {}
        hall = {}
        lugares = []
        lugar = _sh_limpio(request.args.get("lugar"), 330)
        if act and lugar:
            for r in db.execute("SELECT punto_id, resultado FROM sh_checklist WHERE visita_id=? AND lugar=?", (act["id"], lugar)):
                respuestas[str(r["punto_id"])] = r["resultado"]
            for r in db.execute("SELECT punto_id, COUNT(*) AS n FROM sh_hallazgos WHERE visita_id=? AND eliminada=0 "
                                "AND punto_id IS NOT NULL AND " + _SQL_LUGAR + "=? GROUP BY punto_id", (act["id"], lugar)):
                hall[str(r["punto_id"])] = r["n"]
        if act:
            for r in db.execute("SELECT lugar, COUNT(*) AS n FROM sh_checklist WHERE visita_id=? AND lugar<>'' GROUP BY lugar ORDER BY MIN(cuando)", (act["id"],)):
                lugares.append({"lugar": r["lugar"], "revisados": r["n"]})
        de_visita = []
        if act:
            for r in db.execute("SELECT * FROM sh_hallazgos WHERE visita_id=? AND eliminada=0 ORDER BY id DESC", (act["id"],)):
                de_visita.append({"id": r["id"], "folio": _sh_folio_h(p, r["id"]), "tipo": r["tipo"], "tipo_texto": SH_TIPOS.get(r["tipo"], r["tipo"]),
                                  "texto": r["texto"], "bloque": r["bloque"] or "", "area": r["area"] or "", "proveedor": r["proveedor"] or "",
                                  "estado": r["estado"]})
        semana = _sh_semana(db)
    return jsonify({"actual": _sh_visita_dict(act) if act else None, "respuestas": respuestas, "lugares": lugares, "de_visita": de_visita,
                    "hallazgos_por_punto": hall, "semana": semana, "recientes": recientes})


@app.route("/api/sh/visitas/iniciar", methods=["POST"])
@requiere_sh
def api_sh_visita_iniciar():
    fl = _sh_flags()
    if not (fl["admin"] or fl["inspector"]):
        return jsonify({"error": "Solo el administrador o el inspector marcan el inicio y el fin de la visita."}), 403
    d = request.get_json(silent=True) or {}
    p, err = _sh_proyecto(d)
    if err:
        return err
    with _closing(_abrir_db_proyecto(p)) as db:
        if db.execute("SELECT 1 FROM sh_visitas WHERE estado='en_curso'").fetchone():
            return jsonify({"error": "Ya hay una visita en curso en esta torre. Ciérrala antes de iniciar otra."}), 409
        anio = _sh_hoy().year
        n = db.execute("SELECT COUNT(*) FROM sh_visitas WHERE folio LIKE ?", ("SV-%s-%d-%%" % (_sh_letra(p), anio),)).fetchone()[0] + 1
        folio = "SV-%s-%d-%03d" % (_sh_letra(p), anio, n)
        cur = db.execute("INSERT INTO sh_visitas (folio,estado,inicio,iniciada_por,notas) VALUES (?,?,?,?,?)",
                         (folio, "en_curso", _sh_ahora(), session.get("usuario"), _sh_limpio(d.get("notas"), 300)))
        db.commit()
        v = db.execute("SELECT * FROM sh_visitas WHERE id=?", (cur.lastrowid,)).fetchone()
        return jsonify({"ok": True, "visita": _sh_visita_dict(v), "semana": _sh_semana(db)})


@app.route("/api/sh/visitas/<int:vid>/cerrar", methods=["POST"])
@requiere_sh
def api_sh_visita_cerrar(vid):
    fl = _sh_flags()
    if not (fl["admin"] or fl["inspector"]):
        return jsonify({"error": "Solo el administrador o el inspector marcan el inicio y el fin de la visita."}), 403
    d = request.get_json(silent=True) or {}
    p, err = _sh_proyecto(d)
    if err:
        return err
    with _closing(_abrir_db_proyecto(p)) as db:
        v = db.execute("SELECT * FROM sh_visitas WHERE id=?", (vid,)).fetchone()
        if not v or v["estado"] != "en_curso":
            return jsonify({"error": "Esa visita no está en curso."}), 400
        db.execute("UPDATE sh_visitas SET estado='cerrada', fin=?, cerrada_por=?, notas=COALESCE(NULLIF(?,''),notas) WHERE id=?",
                   (_sh_ahora(), session.get("usuario"), _sh_limpio(d.get("notas"), 300), vid))
        db.commit()
        v = db.execute("SELECT * FROM sh_visitas WHERE id=?", (vid,)).fetchone()
        n_h = db.execute("SELECT COUNT(*) FROM sh_hallazgos WHERE visita_id=? AND eliminada=0", (vid,)).fetchone()[0]
        return jsonify({"ok": True, "visita": _sh_visita_dict(v), "hallazgos": n_h, "semana": _sh_semana(db)})


@app.route("/api/sh/visitas/<int:vid>/respuesta", methods=["POST"])
@requiere_sh
def api_sh_visita_respuesta(vid):
    d = request.get_json(silent=True) or {}
    p, err = _sh_proyecto(d)
    if err:
        return err
    fl = _sh_flags()
    if not (fl["admin"] or fl["inspector"]):
        return jsonify({"error": "Solo el inspector o el administrador marcan la revisión."}), 403
    res = d.get("resultado")
    if res not in ("cumple", "no", "na", ""):
        return jsonify({"error": "Resultado no válido."}), 400
    try:
        punto = int(d.get("punto_id"))
    except (TypeError, ValueError):
        return jsonify({"error": "Falta el punto a revisar."}), 400
    lugar = _sh_limpio(d.get("lugar"), 330)
    if not lugar:
        return jsonify({"error": "Elige primero el lugar (bloque o área) que estás revisando."}), 400
    with _closing(_abrir_db_proyecto(p)) as db:
        v = db.execute("SELECT estado FROM sh_visitas WHERE id=?", (vid,)).fetchone()
        if not v or v["estado"] != "en_curso":
            return jsonify({"error": "La visita no está en curso."}), 400
        if res == "":
            db.execute("DELETE FROM sh_checklist WHERE visita_id=? AND punto_id=? AND lugar=?", (vid, punto, lugar))
        else:
            db.execute("INSERT OR REPLACE INTO sh_checklist (visita_id,punto_id,lugar,resultado,por,cuando) VALUES (?,?,?,?,?,?)",
                       (vid, punto, lugar, res, session.get("usuario"), _sh_ahora()))
        db.commit()
    return jsonify({"ok": True})


# ---------- hallazgos ----------
def _sh_validar_ubicacion(d):
    bloque, area = _sh_limpio(d.get("bloque"), 160), _sh_limpio(d.get("area"), 160)
    return bloque, area


@app.route("/api/sh/hallazgos", methods=["GET"])
@requiere_sh
def api_sh_hallazgos():
    p, err = _sh_proyecto()
    if err:
        return err
    fl = _sh_flags()
    cond, par = ["eliminada=0"], []
    prov = _sh_limpio(request.args.get("proveedor"), 120)
    if fl["proveedor"] and not fl["interno"]:
        cond.append("proveedor=?")
        par.append(fl["nombre_prov"])
        cond.append("estado IN ('enviado','respondido','cerrado')")
    elif prov:
        cond.append("proveedor=?")
        par.append(prov)
    est = request.args.get("estado")
    if est == "pend_mtto":
        cond.append("estado='registrado' AND val_mtto_por IS NULL")
    elif est == "pend_arq":
        cond.append("estado='registrado' AND val_arq_por IS NULL")
    elif est == "abiertos":
        cond.append("estado<>'cerrado'")
    elif est in ("registrado", "devuelto", "enviado", "respondido", "cerrado"):
        cond.append("estado=?")
        par.append(est)
    if request.args.get("tipo") in SH_TIPOS:
        cond.append("tipo=?")
        par.append(request.args.get("tipo"))
    with _closing(_abrir_db_proyecto(p)) as db:
        filas = db.execute("SELECT * FROM sh_hallazgos WHERE " + " AND ".join(cond) + " ORDER BY id DESC LIMIT 400", par).fetchall()
        lista = []
        for h in filas:
            d = _sh_dict(h, p)
            d["n_fotos"] = db.execute("SELECT COUNT(*) FROM sh_fotos WHERE hallazgo_id=?", (h["id"],)).fetchone()[0]
            lista.append(d)
        cuenta = {"total": db.execute("SELECT COUNT(*) FROM sh_hallazgos WHERE eliminada=0").fetchone()[0]}
        for k, sql in (("pend_mtto", "estado='registrado' AND val_mtto_por IS NULL"),
                       ("pend_arq", "estado='registrado' AND val_arq_por IS NULL"),
                       ("abiertos", "estado<>'cerrado'"), ("respondidos", "estado='respondido'")):
            cuenta[k] = db.execute("SELECT COUNT(*) FROM sh_hallazgos WHERE eliminada=0 AND " + sql).fetchone()[0]
        if fl["proveedor"] and not fl["interno"]:
            cuenta = {"total": len(lista)}
    return jsonify({"hallazgos": lista, "cuenta": cuenta})


@app.route("/api/sh/pendientes")
@requiere_sh
def api_sh_pendientes():
    """Avisos dentro de la plataforma: lo que le toca atender a quien pregunta, por torre.
    - Mantenimiento / Arquitectos: hallazgos por validar de su lado.
    - Admin e inspector: respuestas del proveedor por revisar y hallazgos devueltos por corregir.
    - Proveedor: hallazgos enviados a el (nuevos o que Seguridad e Higiene no acepto) por atender."""
    fl = _sh_flags()
    torres, total = {}, 0
    for p in PROYECTOS:
        motivos = []
        try:
            with _closing(_abrir_db_proyecto(p)) as db:
                def cuenta(sql, par=()):
                    return db.execute("SELECT COUNT(*) FROM sh_hallazgos WHERE eliminada=0 AND " + sql, par).fetchone()[0]
                if fl["proveedor"] and not fl["interno"]:
                    motivos.append(("por_atender", "Hallazgos por atender", cuenta("estado='enviado' AND proveedor=?", (fl["nombre_prov"],))))
                else:
                    if fl["mantenimiento"]:
                        motivos.append(("pend_mtto", "Por validar (Mantenimiento)", cuenta("estado='registrado' AND val_mtto_por IS NULL")))
                    if fl["arquitecto"]:
                        motivos.append(("pend_arq", "Por validar (Arquitectos)", cuenta("estado='registrado' AND val_arq_por IS NULL")))
                    if fl["admin"] or fl["inspector"]:
                        motivos.append(("respondidos", "Respuestas del proveedor por revisar", cuenta("estado='respondido'")))
                        motivos.append(("devueltos", "Devueltos por corregir", cuenta("estado='devuelto'")))
        except Exception:
            motivos = []
        motivos = [{"clave": k, "texto": t, "n": n} for k, t, n in motivos if n]
        n_t = sum(m["n"] for m in motivos)
        torres[p] = {"n": n_t, "motivos": motivos}
        total += n_t
    return jsonify({"torres": torres, "total": total})


@app.route("/api/sh/hallazgos/<int:hid>", methods=["GET"])
@requiere_sh
def api_sh_hallazgo(hid):
    p, err = _sh_proyecto()
    if err:
        return err
    fl = _sh_flags()
    with _closing(_abrir_db_proyecto(p)) as db:
        h = _sh_obtener(db, hid)
        if not h or not _sh_puede_ver(fl, h):
            return jsonify({"error": "No existe o no tienes permiso."}), 404
        return jsonify(_sh_dict(h, p, db, detalle=True))


@app.route("/api/sh/hallazgos", methods=["POST"])
@requiere_sh
def api_sh_hallazgo_crear():
    d = request.get_json(silent=True) or {}
    p, err = _sh_proyecto(d)
    if err:
        return err
    fl = _sh_flags()
    if not (fl["admin"] or fl["inspector"]):
        return jsonify({"error": "Solo el inspector o el administrador registran hallazgos."}), 403
    tipo = d.get("tipo")
    if tipo not in SH_TIPOS:
        return jsonify({"error": "Elige si es acto inseguro, condición insegura u observación."}), 400
    texto = _sh_limpio(d.get("texto"), 800)
    if not texto:
        return jsonify({"error": "Escribe qué se vio."}), 400
    bloque, area = _sh_validar_ubicacion(d)
    if not bloque:
        return jsonify({"error": "Elige el bloque o el área donde se vio."}), 400
    retro = bool(d.get("retroactivo"))
    hoy = _sh_hoy()
    f_hall = hoy.isoformat()
    evidencia_ref = ""
    if retro:
        fd = parse_date(_sh_limpio(d.get("fecha_hallazgo"), 10))
        if not fd or fd > hoy:
            return jsonify({"error": "Para un registro retroactivo escribe la fecha real en que ocurrió (no puede ser futura)."}), 400
        f_hall = fd.isoformat()
        evidencia_ref = _sh_limpio(d.get("evidencia_ref"), 300)
        if not evidencia_ref and not d.get("fotos"):
            return jsonify({"error": "Un registro retroactivo necesita evidencia real: agrega una foto o escribe de dónde sale (correo, acta, mensaje)."}), 400
    comp = _sh_limpio(d.get("fecha_compromiso"), 10)
    if comp and not parse_date(comp):
        return jsonify({"error": "La fecha compromiso no es válida."}), 400
    with _closing(_abrir_db_proyecto(p)) as db:
        visita = db.execute("SELECT id FROM sh_visitas WHERE estado='en_curso' ORDER BY id DESC LIMIT 1").fetchone()
        if not retro and not visita:
            return jsonify({"error": "No hay una visita en curso. El administrador debe marcar INICIO de la visita (o usa registro retroactivo con evidencia)."}), 409
        punto_id = d.get("punto_id") if isinstance(d.get("punto_id"), int) else None
        cur = db.execute(
            "INSERT INTO sh_hallazgos (visita_id,punto_id,categoria,punto_texto,tipo,texto,normas,bloque,area,proveedor,"
            "fecha_hallazgo,fecha_captura,retroactivo,evidencia_ref,fecha_compromiso,estado,capturado_por,accion_propuesta) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (None if retro else visita["id"], punto_id, _sh_limpio(d.get("categoria"), 80), _sh_limpio(d.get("punto_texto"), 160),
             tipo, texto, ",".join(_sh_normas_lista(d.get("normas"))), bloque, area, _sh_limpio(d.get("proveedor"), 120),
             f_hall, _sh_ahora(), 1 if retro else 0, evidencia_ref, comp or None, "registrado", session.get("usuario"),
             _sh_limpio(d.get("accion_propuesta"), 600)))
        hid = cur.lastrowid
        _sh_guardar_fotos(db, hid, "hallazgo", d.get("fotos"), session.get("usuario"))
        if punto_id and visita and not retro:
            db.execute("INSERT OR REPLACE INTO sh_checklist (visita_id,punto_id,lugar,resultado,por,cuando) VALUES (?,?,?,?,?,?)",
                       (visita["id"], punto_id, _sh_lugar(bloque, area), "no", session.get("usuario"), _sh_ahora()))
        db.commit()
        h = _sh_obtener(db, hid)
        _sh_push(h, p, "registro")
        return jsonify({"ok": True, "id": hid, "folio": _sh_folio_h(p, hid), "avisos": _sh_avisos(h, p, "registro")})


@app.route("/api/sh/hallazgos/<int:hid>", methods=["PUT"])
@requiere_sh
def api_sh_hallazgo_editar(hid):
    """El inspector corrige un hallazgo mientras esta 'registrado' o 'devuelto'. Al editar, las validaciones se reinician."""
    d = request.get_json(silent=True) or {}
    p, err = _sh_proyecto(d)
    if err:
        return err
    fl = _sh_flags()
    if not (fl["admin"] or fl["inspector"]):
        return jsonify({"error": "Solo el inspector o el administrador editan hallazgos."}), 403
    with _closing(_abrir_db_proyecto(p)) as db:
        h = _sh_obtener(db, hid)
        if not h:
            return jsonify({"error": "No existe."}), 404
        if h["estado"] not in ("registrado", "devuelto"):
            return jsonify({"error": "Ya fue validado: no se puede editar. Pide que lo regresen o abre uno nuevo."}), 400
        tipo = d.get("tipo") if d.get("tipo") in SH_TIPOS else h["tipo"]
        texto = _sh_limpio(d.get("texto"), 800) or h["texto"]
        bloque = _sh_limpio(d.get("bloque"), 160) or h["bloque"]
        area = h["area"] if d.get("area") is None else _sh_limpio(d.get("area"), 160)
        normas = h["normas"] if d.get("normas") is None else ",".join(_sh_normas_lista(d.get("normas")))
        prov = h["proveedor"] if d.get("proveedor") is None else _sh_limpio(d.get("proveedor"), 120)
        comp = h["fecha_compromiso"] if d.get("fecha_compromiso") is None else (_sh_limpio(d.get("fecha_compromiso"), 10) or None)
        if comp and not parse_date(comp):
            return jsonify({"error": "La fecha compromiso no es válida."}), 400
        accion = h["accion_propuesta"] if d.get("accion_propuesta") is None else _sh_limpio(d.get("accion_propuesta"), 600)
        db.execute("UPDATE sh_hallazgos SET tipo=?,texto=?,bloque=?,area=?,normas=?,proveedor=?,fecha_compromiso=?,accion_propuesta=?,estado='registrado',"
                   "val_mtto_por=NULL,val_mtto_at=NULL,val_mtto_nota=NULL,val_arq_por=NULL,val_arq_at=NULL,val_arq_nota=NULL "
                   "WHERE id=?", (tipo, texto, bloque, area, normas, prov, comp, accion, hid))
        _sh_guardar_fotos(db, hid, "hallazgo", d.get("fotos"), session.get("usuario"))
        for fid in (d.get("quitar_fotos") or []):
            if isinstance(fid, int):
                db.execute("DELETE FROM sh_fotos WHERE id=? AND hallazgo_id=? AND fase='hallazgo'", (fid, hid))
        db.commit()
        h = _sh_obtener(db, hid)
        _sh_push(h, p, "registro")
        return jsonify({"ok": True, "avisos": _sh_avisos(h, p, "registro")})


@app.route("/api/sh/hallazgos/<int:hid>/validar", methods=["POST"])
@requiere_sh
def api_sh_hallazgo_validar(hid):
    d = request.get_json(silent=True) or {}
    p, err = _sh_proyecto(d)
    if err:
        return err
    fl = _sh_flags()
    por = d.get("por")
    if por == "mtto" and not fl["mantenimiento"]:
        return jsonify({"error": "La validación de Mantenimiento la hacen el administrador y los usuarios asignados como Mantenimiento."}), 403
    if por == "arq" and not fl["arquitecto"]:
        return jsonify({"error": "La validación de Arquitectos la hacen los usuarios asignados como Arquitectos."}), 403
    if por not in ("mtto", "arq"):
        return jsonify({"error": "Falta indicar quién valida."}), 400
    ok = bool(d.get("ok"))
    nota = _sh_limpio(d.get("nota"), 400)
    with _closing(_abrir_db_proyecto(p)) as db:
        h = _sh_obtener(db, hid)
        if not h:
            return jsonify({"error": "No existe."}), 404
        if h["estado"] != "registrado":
            return jsonify({"error": "Este hallazgo ya no está por validar."}), 400
        ahora, quien = _sh_ahora(), session.get("usuario")
        if not ok:
            if not nota:
                return jsonify({"error": "Escribe el motivo para regresarlo al inspector."}), 400
            db.execute("UPDATE sh_hallazgos SET estado='devuelto', devuelto_por=?, devuelto_at=?, devuelto_nota=? WHERE id=?",
                       (quien, ahora, nota, hid))
            db.commit()
            h = _sh_obtener(db, hid)
            return jsonify({"ok": True, "estado": "devuelto", "avisos": _sh_avisos(h, p, "devuelto")})
        col = "val_mtto" if por == "mtto" else "val_arq"
        db.execute("UPDATE sh_hallazgos SET %s_por=?, %s_at=?, %s_nota=? WHERE id=?" % (col, col, col), (quien, ahora, nota, hid))
        h = _sh_obtener(db, hid)
        enviado = bool(h["val_mtto_por"] and h["val_arq_por"])
        if enviado:
            db.execute("UPDATE sh_hallazgos SET estado='enviado', enviado_at=? WHERE id=?", (ahora, hid))
        db.commit()
        h = _sh_obtener(db, hid)
        resp = {"ok": True, "estado": h["estado"]}
        if enviado:
            resp["avisos"] = _sh_avisos(h, p, "enviado")
        return jsonify(resp)


@app.route("/api/sh/hallazgos/<int:hid>/responder", methods=["POST"])
@requiere_sh
def api_sh_hallazgo_responder(hid):
    d = request.get_json(silent=True) or {}
    p, err = _sh_proyecto(d)
    if err:
        return err
    fl = _sh_flags()
    if not fl["proveedor"]:
        return jsonify({"error": "Solo el proveedor responsable responde."}), 403
    texto = _sh_limpio(d.get("texto"), 1000)
    comp = _sh_limpio(d.get("compromiso"), 10)
    if not texto:
        return jsonify({"error": "Escribe qué se hizo o qué vas a hacer para atenderlo."}), 400
    fecha_comp = parse_date(comp) if comp else None
    if not fecha_comp:
        return jsonify({"error": "Indica la nueva fecha compromiso de atención."}), 400
    if fecha_comp < _sh_hoy():
        return jsonify({"error": "La fecha compromiso no puede ser anterior a hoy."}), 400
    with _closing(_abrir_db_proyecto(p)) as db:
        h = _sh_obtener(db, hid)
        if not h or (h["proveedor"] or "") != fl["nombre_prov"] or h["estado"] not in ("enviado", "respondido"):
            return jsonify({"error": "Este hallazgo no está disponible para responder."}), 404
        nuevas = _sh_guardar_fotos(db, hid, "respuesta", d.get("fotos"), session.get("usuario"))
        total = db.execute("SELECT COUNT(*) FROM sh_fotos WHERE hallazgo_id=? AND fase='respuesta'", (hid,)).fetchone()[0]
        if total == 0:
            db.rollback()
            return jsonify({"error": "Agrega al menos una fotografía como evidencia."}), 400
        db.execute("UPDATE sh_hallazgos SET estado='respondido', resp_texto=?, resp_compromiso=?, resp_por=?, resp_at=? WHERE id=?",
                   (texto, comp or None, session.get("usuario"), _sh_ahora(), hid))
        db.commit()
        h = _sh_obtener(db, hid)
        _sh_push(h, p, "respuesta")
        return jsonify({"ok": True, "avisos": _sh_avisos(h, p, "respuesta")})


@app.route("/api/sh/hallazgos/<int:hid>/cerrar", methods=["POST"])
@requiere_sh
def api_sh_hallazgo_cerrar(hid):
    """Verifica la respuesta del proveedor: cierra el hallazgo o lo regresa al proveedor."""
    d = request.get_json(silent=True) or {}
    p, err = _sh_proyecto(d)
    if err:
        return err
    fl = _sh_flags()
    if not (fl["admin"] or fl["inspector"]):
        return jsonify({"error": "Solo el inspector o el administrador cierran un hallazgo."}), 403
    nota = _sh_limpio(d.get("nota"), 400)
    with _closing(_abrir_db_proyecto(p)) as db:
        h = _sh_obtener(db, hid)
        if not h or h["estado"] != "respondido":
            return jsonify({"error": "Solo se cierra un hallazgo que el proveedor ya respondió."}), 400
        if d.get("aceptar") is False:
            if not nota:
                return jsonify({"error": "Escribe por qué no se acepta la respuesta."}), 400
            db.execute("UPDATE sh_hallazgos SET estado='enviado', cierre_nota=? WHERE id=?", (nota, hid))
            db.commit()
            h = _sh_obtener(db, hid)
            return jsonify({"ok": True, "estado": "enviado", "avisos": _sh_avisos(h, p, "enviado")})
        db.execute("UPDATE sh_hallazgos SET estado='cerrado', cerrado_por=?, cerrado_at=?, cierre_nota=? WHERE id=?",
                   (session.get("usuario"), _sh_ahora(), nota, hid))
        db.commit()
        return jsonify({"ok": True, "estado": "cerrado"})


@app.route("/api/sh/hallazgos/<int:hid>", methods=["DELETE"])
@requiere_sh
@requiere_admin
def api_sh_hallazgo_borrar(hid):
    p, err = _sh_proyecto()
    if err:
        return err
    with _closing(_abrir_db_proyecto(p)) as db:
        db.execute("UPDATE sh_hallazgos SET eliminada=1 WHERE id=?", (hid,))
        db.commit()
    return jsonify({"ok": True})


@app.route("/api/sh/hallazgos/<int:hid>/avisos")
@requiere_sh
def api_sh_hallazgo_avisos(hid):
    p, err = _sh_proyecto()
    if err:
        return err
    fl = _sh_flags()
    if not fl["interno"]:
        return jsonify({"error": "sin_permiso"}), 403
    fase = request.args.get("fase") if request.args.get("fase") in ("registro", "enviado", "respuesta", "devuelto") else "registro"
    with _closing(_abrir_db_proyecto(p)) as db:
        h = _sh_obtener(db, hid)
        if not h:
            return jsonify({"error": "No existe."}), 404
        return jsonify(_sh_avisos(h, p, fase))


@app.route("/api/sh/foto/<int:fid>")
@requiere_sh
def api_sh_foto(fid):
    p, err = _sh_proyecto()
    if err:
        return err
    fl = _sh_flags()
    with _closing(_abrir_db_proyecto(p)) as db:
        f = db.execute("SELECT * FROM sh_fotos WHERE id=?", (fid,)).fetchone()
        if not f:
            return jsonify({"error": "No existe."}), 404
        h = _sh_obtener(db, f["hallazgo_id"])
        if not h or not _sh_puede_ver(fl, h):
            return jsonify({"error": "No existe."}), 404
        datos, mime = bytes(f["datos"]), f["mime"]
    resp = Response(datos, mimetype=mime)
    resp.headers["Cache-Control"] = "private, max-age=86400"
    return resp


# ---------- equipo (admin) ----------
@app.route("/api/sh/equipo")
@requiere_sh
@requiere_admin
def api_sh_equipo():
    db = get_db_usuarios()
    roles = {r["usuario"]: r["rol"] for r in db.execute("SELECT usuario, rol FROM sh_roles")}
    filas = []
    for u in db.execute("SELECT usuario, proveedor, rol, telefono, COALESCE(activo,1) AS activo FROM usuarios ORDER BY rol, usuario"):
        if u["rol"] == "directivo":
            continue
        filas.append({"usuario": u["usuario"], "nombre": u["proveedor"] or "", "rol": u["rol"], "activo": u["activo"],
                      "telefono": bool((u["telefono"] or "").strip()), "sh": roles.get(u["usuario"], "")})
    return jsonify({"usuarios": filas})


@app.route("/api/sh/equipo/asignar", methods=["POST"])
@requiere_sh
@requiere_admin
def api_sh_equipo_asignar():
    d = request.get_json(silent=True) or {}
    usuario, rol = _sh_limpio(d.get("usuario"), 80), d.get("rol") or ""
    if rol not in ("", "inspector", "arquitecto", "mantenimiento"):
        return jsonify({"error": "Rol no válido."}), 400
    db = get_db_usuarios()
    u = db.execute("SELECT rol FROM usuarios WHERE usuario=?", (usuario,)).fetchone()
    if not u or u["rol"] in ("admin", "directivo"):
        return jsonify({"error": "Ese usuario no se puede asignar."}), 400
    if u["rol"] == "seguridad" and rol != "inspector":
        return jsonify({"error": "Un usuario de Seguridad e Higiene siempre es inspector."}), 400
    if rol == "":
        db.execute("DELETE FROM sh_roles WHERE usuario=?", (usuario,))
    else:
        db.execute("INSERT OR REPLACE INTO sh_roles (usuario, rol) VALUES (?,?)", (usuario, rol))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/sh/equipo/nuevo", methods=["POST"])
@requiere_sh
@requiere_admin
def api_sh_equipo_nuevo():
    """Crea el usuario del inspector de Seguridad e Higiene (solo entra a esta pantalla)."""
    d = request.get_json(silent=True) or {}
    usuario, clave = _sh_limpio(d.get("usuario"), 60), d.get("clave") or ""
    nombre, tel = _sh_limpio(d.get("nombre"), 80), _sh_limpio(d.get("telefono"), 30) or None
    if not usuario or not clave or not nombre:
        return jsonify({"error": "Escribe nombre, usuario y contraseña."}), 400
    if len(clave) < 4:
        return jsonify({"error": "La contraseña debe tener al menos 4 caracteres."}), 400
    db = get_db_usuarios()
    try:
        db.execute("INSERT INTO usuarios (usuario,clave,proveedor,rol,mundo,telefono,activo,num_logins,creado) VALUES (?,?,?,?,?,?,1,0,?)",
                   (usuario, crear_hash_clave(clave), nombre, "seguridad", "obra", tel,
                    datetime.datetime.now().isoformat(timespec="seconds")))
        db.execute("INSERT OR REPLACE INTO sh_roles (usuario, rol) VALUES (?,?)", (usuario, "inspector"))
        db.commit()
    except sqlite3.IntegrityError:
        return jsonify({"error": "Ese usuario ya existe."}), 400
    return jsonify({"ok": True})


# ---------- reporte PDF por proveedor ----------
def _sh_construir_pdf(p, hallazgos, fotos_por, proveedor, filtro_txt, ruta_salida):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, KeepTogether, Image)
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from xml.sax.saxutils import escape
    import io as _io

    OSCURO, AMARILLO, GRIS = colors.HexColor("#2F3237"), colors.HexColor("#F5B400"), colors.HexColor("#F1EFEA")
    ROJO, NARANJA, AZUL, VERDE = (colors.HexColor("#D0342C"), colors.HexColor("#C26A00"),
                                  colors.HexColor("#4F7396"), colors.HexColor("#1F7A52"))
    est = getSampleStyleSheet()
    base = ParagraphStyle("b", parent=est["Normal"], fontName="Helvetica", fontSize=9, leading=11.5, textColor=OSCURO)
    peq = ParagraphStyle("p", parent=base, fontSize=7.5, leading=9, textColor=colors.HexColor("#6b6f75"))
    blanco = ParagraphStyle("w", parent=base, fontName="Helvetica-Bold", fontSize=13, leading=15, textColor=colors.white)
    esc = lambda t: escape(str(t or ""))
    ancho = letter[0] - 2.4 * cm
    color_tipo = {"acto": NARANJA, "condicion": ROJO, "observacion": AZUL}

    hoy_txt = _sh_hoy().strftime("%d/%m/%Y")
    folio = "SH-R-%s-%s" % (_sh_letra(p), _sh_hoy().strftime("%Y%m%d"))
    cab = Table([[Paragraph("<b>HAP</b>", ParagraphStyle("l", parent=blanco, textColor=OSCURO, fontSize=15)),
                  [Paragraph("REPORTE DE HALLAZGOS · SEGURIDAD E HIGIENE", blanco),
                   Paragraph("Hospital Aranda de la Parra · Remodelación de Quirófanos · Torre %s" % _sh_letra(p),
                             ParagraphStyle("s", parent=base, textColor=colors.HexColor("#D9D6CF"), fontSize=8))],
                  Paragraph("FOLIO<br/><font size=11 color='#F5B400'><b>%s</b></font><br/>%s" % (folio, hoy_txt),
                            ParagraphStyle("f", parent=base, textColor=colors.white, alignment=2, fontSize=7, leading=12))]],
                colWidths=[1.8 * cm, ancho - 1.8 * cm - 4.6 * cm, 4.6 * cm])
    cab.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), OSCURO), ("BACKGROUND", (0, 0), (0, 0), AMARILLO),
                             ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (0, 0), "CENTER"),
                             ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))

    n = len(hallazgos)
    por_tipo = {k: sum(1 for h in hallazgos if h["tipo"] == k) for k in SH_TIPOS}
    doble = sum(1 for h in hallazgos if h["val_mtto_por"] and h["val_arq_por"])
    cerr = sum(1 for h in hallazgos if h["estado"] == "cerrado")
    retro = sum(1 for h in hallazgos if h["retroactivo"])

    def kpi(num, txt):
        return Paragraph("<font size=16><b>%s</b></font><br/><font size=7.5 color='#555555'>%s</font>" % (num, txt),
                         ParagraphStyle("k", parent=base, leading=19))
    tarjetas = [kpi(n, "hallazgos"), kpi(por_tipo["condicion"], "condiciones inseguras"),
                kpi(por_tipo["acto"], "actos inseguros"), kpi(por_tipo["observacion"], "observaciones"),
                kpi(doble, "con doble validación"), kpi(cerr, "cerrados")]
    kp = Table([tarjetas], colWidths=[ancho / 6.0] * 6)
    kp.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, OSCURO), ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d8d5ce")),
                            ("BACKGROUND", (0, 0), (-1, -1), GRIS), ("TOPPADDING", (0, 0), (-1, -1), 5),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    filtro = Table([[Paragraph("Proveedor: <b>%s</b> &nbsp;·&nbsp; %s" % (esc(proveedor) or "Todos", esc(filtro_txt)),
                               ParagraphStyle("fl", parent=base, textColor=colors.white))]], colWidths=[ancho])
    filtro.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), OSCURO), ("TOPPADDING", (0, 0), (-1, -1), 6),
                                ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))

    def foto_img(datos, ancho_max, alto_max):
        try:
            img = Image(_io.BytesIO(datos))
            w, h = img.imageWidth, img.imageHeight
            k = min(ancho_max / float(w), alto_max / float(h), 1.0 if False else 99)
            img.drawWidth, img.drawHeight = w * k, h * k
            return img
        except Exception:
            return Paragraph("(foto no disponible)", peq)

    def estado_val(por, at, rot):
        if por:
            return Paragraph("<font size=7 color='#555555'>%s</font><br/><font color='#1F7A52'><b>VALIDADO</b></font> · %s<br/>"
                             "<font size=7 color='#555555'>%s</font>" % (rot, esc(_sh_fmt(at, False)), esc(por)), base)
        return Paragraph("<font size=7 color='#555555'>%s</font><br/><font color='#C26A00'><b>Pendiente</b></font>" % rot, base)

    historia = [cab, Spacer(1, 6), filtro, Spacer(1, 6), kp, Spacer(1, 10)]
    if not hallazgos:
        historia.append(Paragraph("No hay hallazgos con este filtro.", base))
    for h in hallazgos:
        fh = fotos_por.get(h["id"], {})
        color = color_tipo.get(h["tipo"], OSCURO)
        ub = " · ".join(x for x in (h["bloque"], h["area"]) if x)
        etiqueta = Paragraph("<font color='white'><b>%s</b></font>" % esc(SH_TIPOS.get(h["tipo"], h["tipo"]).upper()),
                             ParagraphStyle("e", parent=base, fontSize=8))
        tit = Table([[etiqueta, Paragraph("<b>Hallazgo %s</b>%s" % (_sh_folio_h(p, h["id"]),
                                         " &nbsp;<font color='#C26A00' size=7.5>[REGISTRO RETROACTIVO]</font>" if h["retroactivo"] else ""), base),
                      Paragraph(esc(ub), ParagraphStyle("u", parent=base, alignment=2, fontSize=8))]],
                    colWidths=[3.3 * cm, ancho - 3.3 * cm - 6.2 * cm, 6.2 * cm])
        tit.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, 0), color), ("BACKGROUND", (1, 0), (-1, -1), GRIS),
                                 ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (0, 0), "CENTER"),
                                 ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
        primera = fh.get("hallazgo", [])
        col_foto = [foto_img(primera[0], 6.0 * cm, 4.6 * cm) if primera else Paragraph("Sin fotografía", peq)]
        if len(primera) > 1:
            col_foto.append(Spacer(1, 3))
            col_foto.append(Table([[foto_img(x, 1.9 * cm, 1.5 * cm) for x in primera[1:4]]]))
        normas_txt = ", ".join(x for x in (h["normas"] or "").split(",") if x)
        texto = [Paragraph(esc(h["texto"]), ParagraphStyle("t", parent=base, fontSize=10, leading=13)), Spacer(1, 5),
                 Paragraph("<b>Respaldo normativo:</b> %s" % (esc(normas_txt) or "Sin norma indicada"),
                           ParagraphStyle("n", parent=base, backColor=colors.HexColor("#FFF6D6"), borderPadding=4, leading=12)),
                 Spacer(1, 4)]
        if h["accion_propuesta"]:
            texto.append(Paragraph("<b>Acción propuesta:</b> %s" % esc(h["accion_propuesta"]),
                                   ParagraphStyle("ap", parent=base, backColor=colors.HexColor("#E7EEF4"), borderPadding=4, leading=12)))
            texto.append(Spacer(1, 4))
        fechas = "Fecha en que se detectó: <b>%s</b>" % (esc(_sh_fmt_fecha(h["fecha_hallazgo"])) or "—")
        if h["fecha_compromiso"]:
            fechas += " &nbsp;·&nbsp; Compromiso de atención: <b>%s</b>" % esc(_sh_fmt_fecha(h["fecha_compromiso"]))
        texto.append(Paragraph(fechas, peq))
        if h["retroactivo"]:
            texto.append(Paragraph("Capturado el %s · Evidencia: %s" % (esc(_sh_fmt(h["fecha_captura"], False)),
                                   esc(h["evidencia_ref"]) or "fotografía adjunta"), peq))
        cuerpo = Table([[col_foto, texto]], colWidths=[6.4 * cm, ancho - 6.4 * cm])
        cuerpo.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 6),
                                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
        vals = Table([[estado_val(h["val_mtto_por"], h["val_mtto_at"], "MANTENIMIENTO E INFRAESTRUCTURA"),
                       estado_val(h["val_arq_por"], h["val_arq_at"], "ARQUITECTOS · SUPERVISIÓN")]],
                     colWidths=[ancho / 2.0] * 2)
        vals.setStyle(TableStyle([("LINEABOVE", (0, 0), (-1, 0), 0.5, colors.HexColor("#cfcbc3")),
                                  ("LINEAFTER", (0, 0), (0, 0), 0.5, colors.HexColor("#cfcbc3")),
                                  ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        bloque = [tit, cuerpo, vals]
        if h["resp_texto"]:
            resp_f = fh.get("respuesta", [])
            fila = [[Paragraph("<font size=7 color='#555555'>RESPUESTA DEL PROVEEDOR · %s · %s</font><br/>%s%s" % (
                esc(h["resp_por"]), esc(_sh_fmt(h["resp_at"], False)), esc(h["resp_texto"]),
                ("<br/><font size=7.5>Compromiso: <b>%s</b></font>" % esc(_sh_fmt_fecha(h["resp_compromiso"]))) if h["resp_compromiso"] else ""), base),
                     Table([[foto_img(x, 2.6 * cm, 2.0 * cm) for x in resp_f[:3]]]) if resp_f else ""]]
            rt = Table(fila, colWidths=[ancho - 8.4 * cm, 8.4 * cm])
            rt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EAF5EF")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                    ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
            bloque.append(rt)
        if h["estado"] == "cerrado":
            ct = Table([[Paragraph("<font color='#1F7A52'><b>CERRADO Y VERIFICADO</b></font> · %s · %s%s" % (
                esc(h["cerrado_por"]), esc(_sh_fmt(h["cerrado_at"], False)),
                (" · " + esc(h["cierre_nota"])) if h["cierre_nota"] else ""), base)]], colWidths=[ancho])
            ct.setStyle(TableStyle([("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
            bloque.append(ct)
        caja = Table([[bloque]], colWidths=[ancho])
        caja.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#b9b5ac")),
                                  ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                                  ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
        historia += [KeepTogether([caja]), Spacer(1, 9)]

    firmas = Table([[Paragraph("Inspector de Seguridad e Higiene", base), Paragraph("Mantenimiento e Infraestructura", base),
                     Paragraph("Arquitectos · Supervisión", base), Paragraph("Proveedor: %s" % (esc(proveedor) or "________"), base)]],
                   colWidths=[ancho / 4.0] * 4, rowHeights=[1.5 * cm])
    firmas.setStyle(TableStyle([("LINEABOVE", (0, 0), (-1, 0), 0.7, OSCURO), ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
                                ("ALIGN", (0, 0), (-1, -1), "CENTER")]))
    historia += [Spacer(1, 12), KeepTogether([Paragraph("Normas citadas como referencia técnica; las confirma Seguridad e Higiene.", peq),
                                              Spacer(1, 4), firmas])]
    if retro:
        historia.insert(5, Paragraph("Incluye %d registro(s) retroactivo(s): llevan la fecha real en que ocurrieron y la evidencia que los respalda."
                                     % retro, peq))

    def pie(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawString(1.2 * cm, 0.8 * cm, "HAP · Seguridad e Higiene · %s · Torre %s" % (folio, _sh_letra(p)))
        canvas.drawRightString(letter[0] - 1.2 * cm, 0.8 * cm, "Página %d" % doc.page)
        canvas.restoreState()
    doc = SimpleDocTemplate(ruta_salida, pagesize=letter, leftMargin=1.2 * cm, rightMargin=1.2 * cm,
                            topMargin=1.1 * cm, bottomMargin=1.5 * cm, title="Reporte de hallazgos S. e H.")
    doc.build(historia, onFirstPage=pie, onLaterPages=pie)


def _sh_generar_reporte():
    """Arma el PDF segun los filtros de la peticion. Devuelve (ruta, nombre) o (None, respuesta_de_error)."""
    p, err = _sh_proyecto()
    if err:
        return None, err
    fl = _sh_flags()
    cond, par, filtro = ["eliminada=0"], [], []
    prov = _sh_limpio(request.args.get("proveedor"), 120)
    if fl["proveedor"] and not fl["interno"]:
        prov = fl["nombre_prov"]
        cond.append("estado IN ('enviado','respondido','cerrado')")
    if prov:
        cond.append("proveedor=?")
        par.append(prov)
    est = request.args.get("estado")
    nombres = {"abiertos": "abiertos", "registrado": "por validar", "devuelto": "devueltos", "enviado": "enviados al proveedor",
               "respondido": "respondidos", "cerrado": "cerrados"}
    if est == "abiertos":
        cond.append("estado<>'cerrado'")
    elif est in ("registrado", "devuelto", "enviado", "respondido", "cerrado"):
        cond.append("estado=?")
        par.append(est)
    filtro.append("Estado: " + nombres.get(est, "todos"))
    mes = _sh_limpio(request.args.get("mes"), 7)
    if mes and re.match(r"^\d{4}-\d{2}$", mes):
        cond.append("substr(COALESCE(cerrado_at, fecha_captura),1,7)=?")
        par.append(mes)
        filtro.append("Mes: " + mes)
    if request.args.get("solo_validados") == "1":
        cond.append("val_mtto_por IS NOT NULL AND val_arq_por IS NOT NULL")
        filtro.append("solo con doble validación")
    with _closing(_abrir_db_proyecto(p)) as db:
        filas = db.execute("SELECT * FROM sh_hallazgos WHERE " + " AND ".join(cond) + " ORDER BY id", par).fetchall()
        fotos = {}
        for h in filas:
            for f in db.execute("SELECT fase, datos FROM sh_fotos WHERE hallazgo_id=? ORDER BY id", (h["id"],)):
                fotos.setdefault(h["id"], {}).setdefault(f["fase"], []).append(bytes(f["datos"]))
    salida = ruta_salida_temporal("sh.pdf")
    _sh_construir_pdf(p, filas, fotos, prov, " · ".join(filtro), salida)
    nombre = "Reporte_SH_Torre%s_%s_%s.pdf" % (_sh_letra(p), (prov or "todos").replace(" ", "_")[:30], _sh_hoy().strftime("%Y%m%d"))
    return (salida, nombre), None


@app.route("/api/sh/reporte/pdf")
@requiere_sh
def api_sh_reporte_pdf():
    """El PDF. Sin 'descargar=1' se abre en el visor del navegador (para imprimir); con 'descargar=1' se guarda."""
    res, err = _sh_generar_reporte()
    if err:
        return err
    salida, nombre = res
    return send_file(salida, as_attachment=request.args.get("descargar") == "1", download_name=nombre, mimetype="application/pdf")


@app.route("/api/sh/reporte/vista")
@requiere_sh
def api_sh_reporte_vista():
    """Vista previa: las paginas del PDF como imagenes, para verlas antes de imprimir o guardar (sirve tambien en celular)."""
    res, err = _sh_generar_reporte()
    if err:
        return err
    salida, nombre = res
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf
    paginas, total = [], 0
    try:
        doc = pymupdf.open(salida)
        total = doc.page_count
        for i in range(min(total, 40)):
            pix = doc[i].get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
            paginas.append("data:image/jpeg;base64," + base64.b64encode(pix.tobytes("jpeg", jpg_quality=80)).decode("ascii"))
        doc.close()
    except Exception as e:
        return jsonify({"error": "No se pudo preparar la vista previa: %s" % e}), 500
    return jsonify({"paginas": paginas, "total": total, "nombre": nombre})


def _sembrar_si_vacia():
    """Primera vez en un servidor nuevo: si el disco esta vacio, copia la
    base semilla con las 955 actividades. Asi no hay que subirla a mano."""
    semilla = os.path.join(BASE_DIR, "semilla", "obra_inicial.db")
    if not os.path.exists(semilla):
        return
    try:
        necesita = not os.path.exists(DB_PATH)
        if not necesita:
            con = sqlite3.connect(DB_PATH)
            try:
                n = con.execute("SELECT COUNT(*) FROM actividades").fetchone()[0]
                necesita = (n == 0)
            except sqlite3.Error:
                necesita = True
            con.close()
        if necesita:
            import shutil
            shutil.copy(semilla, DB_PATH)
            print("  Base inicial sembrada desde semilla/obra_inicial.db")
    except Exception as e:
        print("  (no se pudo sembrar la base inicial:", e, ")")


# --- Arranque cuando corre en internet (gunicorn no ejecuta el bloque de abajo) ---
_ARRANQUE_MODULO = True
try:
    _sembrar_si_vacia()
    init_db()
    _foto_automatica()
    # Torre A: su propia base (vacia al inicio). Si falla, Torre B no se afecta.
    try:
        init_db(DB_PATH_TORRE_A)
        _foto_automatica(DB_PATH_TORRE_A)
    except Exception as _ea:
        print("Aviso al iniciar la base de Torre A:", _ea)
except Exception as _e:
    print("Aviso al iniciar la base:", _e)


def resolver_bloque_area_ta(db, bloque_txt, area_txt):
    """Excel de Torre A: acepta la etiqueta ('01 · Quirófanos'), el nombre ('Quirófanos') o el numero ('1').
    Devuelve (bloque, area, etiqueta_bloque, etiqueta_area) o (None, mensaje)."""
    bt, at = str(bloque_txt or "").strip(), str(area_txt or "").strip()
    if not bt or not at:
        return None, "faltan bloque o área"
    bloques = db.execute("SELECT num,nombre FROM ta_bloques ORDER BY num").fetchall()
    mb = _RE_BLOQUE_TA.match(bt)
    num_b = None
    if mb:
        num_b = int(mb.group(1))
    elif bt.isdigit():
        num_b = int(bt)
    else:
        for b in bloques:
            if sin_acentos(b["nombre"]) == sin_acentos(bt):
                num_b = b["num"]
                break
    bloque = next((b for b in bloques if b["num"] == num_b), None)
    if not bloque:
        return None, "el bloque «%s» no está en el catálogo" % bt
    areas = db.execute("SELECT num,nombre FROM ta_areas WHERE bloque=? ORDER BY num", (bloque["num"],)).fetchall()
    ma = _RE_AREA_TA.match(at)
    num_a = None
    if ma:
        num_a = int(ma.group(2)) if int(ma.group(1)) == bloque["num"] else -1
    elif re.match(r"^\d+\.\d+$", at):
        pb, pa = at.split(".")
        num_a = int(pa) if int(pb) == bloque["num"] else -1
    else:
        coinciden = [a["num"] for a in areas if sin_acentos(a["nombre"]) == sin_acentos(at)]
        if len(coinciden) == 1:
            num_a = coinciden[0]
        elif len(coinciden) > 1:
            return None, "el área «%s» se repite en ese bloque; escribe su número (ej. %d.%d)" % (at, bloque["num"], coinciden[0])
    area = next((a for a in areas if a["num"] == num_a), None)
    if not area:
        return None, "el área «%s» no está en el bloque %d" % (at, bloque["num"])
    return (bloque["num"], area["num"], etiqueta_bloque_ta(bloque["num"], bloque["nombre"]),
            etiqueta_area_ta(bloque["num"], area["num"], area["nombre"])), None


def datos_estado_carga(como_propuesta):
    """(origen, estado_val, creado_por) de una actividad creada por carga masiva.
    Por validar: nace como propuesta y no suma en ningun total hasta que el admin la apruebe."""
    if como_propuesta:
        return ("propuesta", "propuesta", session.get("usuario", "admin"))
    return ("oficial", "validado", None)


def carga_masiva_ta(db, acts, como_propuesta=False):
    """Carga masiva en Torre A: todo o nada. Valida cada fila y arma el codigo bloque.area.proveedor.actividad."""
    errores, filas = [], []
    for i, a in enumerate(acts, start=1):
        res, msj = resolver_bloque_area_ta(db, a.get("bloque"), a.get("area"))
        interno = (a.get("mundo") or "obra") == "interno"
        resp = str(a.get("proveedor") or "").strip()
        if not res:
            errores.append("Fila %d: %s" % (i, msj))
        elif not resp:
            errores.append("Fila %d: falta el responsable" % i)
        elif not str(a.get("partida") or "").strip():
            errores.append("Fila %d: falta la actividad (partida)" % i)
        else:
            filas.append((a, res, resp, interno))
    if errores:
        db.rollback()
        resto = " (y %d más)" % (len(errores) - 5) if len(errores) > 5 else ""
        return jsonify({"error": "No se cargó nada. " + "; ".join(errores[:5]) + resto, "errores": errores}), 400
    ahora = datetime.datetime.now().isoformat(timespec="seconds")
    codigos = []
    for a, (b, ar, eb, ea), resp, interno in filas:
        p = numero_proveedor_ta(db, resp, interno)
        resp = nombre_proveedor_ta(db, p) or resp
        codigo = siguiente_codigo_ta(db, b, ar, p)
        giro = str(a.get("giro") or "").strip()
        if giro:
            db.execute("INSERT OR IGNORE INTO ta_especialidades (nombre) VALUES (?)", (giro,))
        db.execute(
            """INSERT INTO actividades
            (codigo,bloque,area,giro,proveedor,departamento,partida,tipo,tipo_partida,aplica,avance,
             estatus,definido,mundo,requiere_pruebas,f_inicio,f_fin,origen,estado_val,creado_por,actualizado)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (codigo, eb, ea, giro, None if interno else resp, resp if interno else None,
             str(a.get("partida")).strip(), a.get("tipo_partida", "Construcción"), a.get("tipo_partida", "Construcción"),
             "SÍ", 0, "Pendiente", "NO", "interno" if interno else "obra", "NO" if interno else "SÍ",
             a.get("f_inicio") or None, a.get("f_fin") or None, *datos_estado_carga(como_propuesta), ahora))
        codigos.append(codigo)
    db.commit()
    return jsonify({"ok": True, "creadas": len(codigos), "codigos": codigos})


# ──── Carga masiva de actividades por Excel ────
@app.route("/api/carga_masiva", methods=["POST"])
@requiere_gestor
def api_carga_masiva():
    db = get_db()
    data = request.get_json()
    acts = data.get("actividades", [])
    if not acts:
        return jsonify({"error": "No hay actividades para subir"}), 400
    como_propuesta = bool(data.get("como_propuesta"))
    if proyecto_actual() == "torre_a":
        return carga_masiva_ta(db, acts, como_propuesta)

    # Siguiente código disponible (misma lógica que usa la creación normal
    # y la duplicación de actividades, ver siguiente_codigo()).
    num = int(siguiente_codigo(db).split("-")[1])

    ahora = datetime.datetime.now().isoformat(timespec="seconds")
    codigos = []
    for a in acts:
        cod = f"ACT-{num:04d}"
        db.execute(
            """INSERT INTO actividades
            (codigo,bloque,area,giro,proveedor,partida,tipo,tipo_partida,aplica,avance,
             estatus,definido,mundo,requiere_pruebas,f_inicio,f_fin,origen,estado_val,creado_por,actualizado)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (cod, a.get("bloque",""), a.get("area",""), a.get("giro",""),
             a.get("proveedor",""), a.get("partida",""),
             a.get("tipo_partida","Construcción"), a.get("tipo_partida","Construcción"),
             "SÍ", 0, "Pendiente", "NO", a.get("mundo","obra"),
             "SÍ" if a.get("mundo","obra") == "obra" else "NO",
             a.get("f_inicio") or None, a.get("f_fin") or None, *datos_estado_carga(como_propuesta), ahora))
        codigos.append(cod)
        num += 1
    db.commit()
    return jsonify({"ok": True, "creadas": len(codigos), "codigos": codigos})


@app.route("/api/plantilla_carga_masiva")
@requiere_gestor
def api_plantilla_carga_masiva():
    import io
    try:
        from openpyxl import Workbook
    except ImportError:
        return "openpyxl no disponible", 500
    wb = Workbook()
    ws = wb.active
    ws.title = "Actividades"
    ws.append(["Bloque", "Área", "Especialidad", "Proveedor", "Partida", "Tipo de partida", "Fecha compromiso", "Fecha inicio (opcional)"])
    if proyecto_actual() == "torre_a":
        ws.append(["Quirófanos", "Quirófano de procedimientos gineco obstétricos 1", "Eléctrico", "CEBSA", "Contactos a 110V", "Construcción", "2026-10-15"])
    else:
        ws.append(["Ejemplo: Cuneros", "Cuneros 1", "Eléctrico", "Daniel Contreras", "Contactos a 110V", "Construcción", "2026-10-15"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="plantilla_carga_masiva.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/seguimiento")
def seguimiento_page():
    return render_template("seguimiento.html")


@app.route("/api/seguimiento")
@requiere_gestor
def api_seguimiento():
    """Módulo de Seguimiento y Cumplimiento — solo lectura, no modifica nada."""
    db = get_db()
    fecha_hoy = datetime.date.today().isoformat()
    fecha_entrega = fecha_entrega_actual()

    # Todos los responsables (proveedores + departamentos)
    nombres = set()
    for r in db.execute("SELECT DISTINCT proveedor FROM actividades WHERE proveedor IS NOT NULL AND proveedor<>'' AND (eliminada IS NULL OR eliminada=0)"):
        nombres.add((r[0], "obra"))
    for r in db.execute("SELECT DISTINCT departamento FROM actividades WHERE departamento IS NOT NULL AND departamento<>'' AND (eliminada IS NULL OR eliminada=0)"):
        nombres.add((r[0], "interno"))

    prov_info = {}
    for r in db.execute("SELECT nombre, tipo, funcion FROM proveedores"):
        prov_info[r["nombre"]] = {"tipo": r["tipo"] or "Externo", "funcion": r["funcion"] or ""}

    usuarios_por = {}
    for r in db.execute("SELECT proveedor, usuario, num_logins, ultimo_login, activo FROM usuarios WHERE rol='proveedor' AND proveedor IS NOT NULL"):
        usuarios_por.setdefault(r["proveedor"], []).append(dict(r))

    responsables = []
    totales = {}
    for m in ("obra", "interno"):
        totales[m] = {"total":0,"terminadas":0,"proceso":0,"pendientes":0,"atrasadas":0,
                      "riesgo":0,"pend_validar":0,"propuestas":0,"sin_reconocer":0}

    fecha_riesgo = (datetime.date.today() + datetime.timedelta(days=14)).isoformat()

    for nombre, mundo in sorted(nombres, key=lambda x: x[0].lower()):
        col = "departamento" if mundo == "interno" else "proveedor"
        acts = db.execute(f"SELECT * FROM actividades WHERE {col}=? AND (eliminada IS NULL OR eliminada=0)", (nombre,)).fetchall()
        if not acts:
            continue

        n_total = len(acts)
        n_terminadas = sum(1 for a in acts if a["estatus"] == "Terminada")
        n_proceso = sum(1 for a in acts if a["estatus"] == "En proceso")
        n_pendientes = sum(1 for a in acts if a["estatus"] == "Pendiente")
        n_reconocidas = sum(1 for a in acts if a["reconocida"] == "SÍ")
        n_sin_reconocer = sum(1 for a in acts if a["reconocida"] != "SÍ")
        n_atrasadas = sum(1 for a in acts if a["f_fin"] and a["f_fin"] < fecha_hoy and a["estatus"] != "Terminada")
        n_riesgo = sum(1 for a in acts if a["f_fin"] and fecha_hoy <= a["f_fin"] <= fecha_riesgo
                       and (a["avance"] or 0) < 75 and a["estatus"] != "Terminada")
        n_pend_validar = sum(1 for a in acts if a["avance_decl"] is not None and a["avance_decl"] > (a["avance"] or 0))
        n_propuestas = sum(1 for a in acts if a["origen"] == "propuesta" and a["estado_val"] == "propuesta")

        av_oficial = round(sum(a["avance"] or 0 for a in acts) / n_total, 1)
        fechas_act = [a["actualizado"] for a in acts if a["actualizado"]]
        ultimo_reporte = max(fechas_act) if fechas_act else None
        dias_sin = None
        if ultimo_reporte:
            try:
                dias_sin = (datetime.date.today() - datetime.date.fromisoformat(ultimo_reporte[:10])).days
            except Exception:
                pass

        usus = usuarios_por.get(nombre, [])
        tiene_usuario = len(usus) > 0
        ultimo_login = None
        num_logins_total = 0
        for u in usus:
            num_logins_total += u["num_logins"] or 0
            if u["ultimo_login"] and (ultimo_login is None or u["ultimo_login"] > ultimo_login):
                ultimo_login = u["ultimo_login"]

        dias_sin_login = None
        if ultimo_login:
            try:
                dias_sin_login = (datetime.date.today() - datetime.date.fromisoformat(ultimo_login[:10])).days
            except Exception:
                pass

        # --- Semáforo ---
        motivos = []
        pend_admin = []
        if n_pend_validar > 0:
            pend_admin.append(f"{n_pend_validar} avances esperan validación del administrador")
        if n_propuestas > 0:
            pend_admin.append(f"{n_propuestas} propuestas esperan aprobación")

        if not tiene_usuario and n_total > 0:
            semaforo = "gris"
            motivos.append("No cuenta con usuario para la plataforma")
            if n_sin_reconocer > 0:
                motivos.append(f"{n_sin_reconocer} actividades sin reconocer")
        elif n_atrasadas >= 5 or (n_sin_reconocer > n_total * 0.5 and n_sin_reconocer > 10) or \
             (dias_sin_login is not None and dias_sin_login > 14) or \
             (tiene_usuario and num_logins_total == 0):
            semaforo = "rojo"
            if n_atrasadas > 0:
                motivos.append(f"{n_atrasadas} actividades vencidas")
            if n_sin_reconocer > 0:
                motivos.append(f"{n_sin_reconocer} actividades sin reconocer")
            if tiene_usuario and num_logins_total == 0:
                motivos.append("Tiene usuario pero nunca ha ingresado")
            elif dias_sin_login and dias_sin_login > 14:
                motivos.append(f"Último acceso hace {dias_sin_login} días")
            if n_riesgo > 0:
                motivos.append(f"{n_riesgo} actividades en riesgo de vencer")
        elif n_atrasadas > 0 or n_sin_reconocer > 5 or \
             (dias_sin is not None and dias_sin > 7) or n_riesgo > 3:
            semaforo = "amarillo"
            if n_atrasadas > 0:
                motivos.append(f"{n_atrasadas} actividades vencidas")
            if n_sin_reconocer > 0:
                motivos.append(f"{n_sin_reconocer} sin reconocer")
            if dias_sin and dias_sin > 7:
                motivos.append(f"{dias_sin} días sin actualización")
            if n_riesgo > 0:
                motivos.append(f"{n_riesgo} próximas a vencer")
        else:
            semaforo = "verde"
            motivos.append("Al corriente")

        accion = ""
        if semaforo == "gris":
            accion = "Crear usuario y solicitar reconocimiento de actividades."
        elif semaforo == "rojo":
            partes = []
            if n_sin_reconocer > 10:
                partes.append("reconocimiento")
            if n_atrasadas > 0:
                partes.append("actualización inmediata y fecha compromiso")
            if tiene_usuario and num_logins_total == 0:
                partes.append("que ingrese a la plataforma")
            accion = "Solicitar " + " y ".join(partes) + "." if partes else "Dar seguimiento urgente."
        elif semaforo == "amarillo":
            accion = "Dar seguimiento y solicitar actualización."

        info = prov_info.get(nombre, {"tipo": "Interno" if mundo == "interno" else "Externo", "funcion": ""})

        responsables.append({
            "nombre": nombre, "mundo": mundo, "tipo": info["tipo"],
            "funcion": info.get("funcion", ""),
            "total": n_total, "terminadas": n_terminadas, "proceso": n_proceso,
            "pendientes": n_pendientes, "reconocidas": n_reconocidas,
            "sin_reconocer": n_sin_reconocer, "atrasadas": n_atrasadas,
            "riesgo": n_riesgo, "pend_validar": n_pend_validar,
            "propuestas": n_propuestas, "av_oficial": av_oficial,
            "ultimo_reporte": ultimo_reporte, "ultimo_login": ultimo_login,
            "dias_sin_actualizar": dias_sin, "dias_sin_login": dias_sin_login,
            "tiene_usuario": tiene_usuario, "num_logins": num_logins_total,
            "semaforo": semaforo, "motivos": motivos, "pend_admin": pend_admin,
            "accion": accion,
        })

        t = totales[mundo]
        t["total"] += n_total
        t["terminadas"] += n_terminadas
        t["proceso"] += n_proceso
        t["pendientes"] += n_pendientes
        t["atrasadas"] += n_atrasadas
        t["riesgo"] += n_riesgo
        t["pend_validar"] += n_pend_validar
        t["propuestas"] += n_propuestas
        t["sin_reconocer"] += n_sin_reconocer

    try:
        dias_entrega = (datetime.date.fromisoformat(fecha_entrega) - datetime.date.today()).days
    except Exception:
        dias_entrega = None   # Torre A sin fecha de entrega definida
    totales["dias_para_entrega"] = dias_entrega
    totales["fecha_entrega"] = fecha_entrega

    return jsonify({"responsables": responsables, "totales": totales, "fecha": fecha_hoy})


@app.route("/api/seguimiento/detalle/<nombre>")
@requiere_gestor
def api_seguimiento_detalle(nombre):
    """Detalle de actividades de un responsable para el módulo de Seguimiento."""
    db = get_db()
    fecha_hoy = datetime.date.today().isoformat()
    acts = db.execute("SELECT * FROM actividades WHERE (proveedor=? OR departamento=?) AND (eliminada IS NULL OR eliminada=0)", (nombre, nombre)).fetchall()
    resultado = []
    for a in acts:
        pend_val = a["avance_decl"] is not None and a["avance_decl"] > (a["avance"] or 0)
        vencida = bool(a["f_fin"] and a["f_fin"] < fecha_hoy and a["estatus"] != "Terminada")
        resultado.append({
            "id": a["id"], "codigo": a["codigo"], "bloque": a["bloque"], "area": a["area"],
            "giro": a["giro"], "partida": a["partida"], "avance": a["avance"] or 0,
            "avance_decl": a["avance_decl"], "estatus": a["estatus"],
            "reconocida": a["reconocida"], "f_fin": a["f_fin"],
            "vencida": vencida, "pend_validar": pend_val,
            "origen": a["origen"], "estado_val": a["estado_val"],
            "actualizado": a["actualizado"],
        })
    return jsonify(resultado)


@app.route("/api/seguimiento/whatsapp/<nombre>")
@requiere_gestor
def api_seguimiento_whatsapp(nombre):
    """Genera mensaje de seguimiento por WhatsApp para un responsable."""
    db = get_db()
    fecha_hoy = datetime.date.today().isoformat()
    acts = db.execute("SELECT * FROM actividades WHERE (proveedor=? OR departamento=?) AND (eliminada IS NULL OR eliminada=0)", (nombre, nombre)).fetchall()
    n_total = len(acts)
    n_terminadas = sum(1 for a in acts if a["estatus"] == "Terminada")
    n_pend = sum(1 for a in acts if a["estatus"] != "Terminada" and (a["avance"] or 0) < 100)
    n_atrasadas = sum(1 for a in acts if a["f_fin"] and a["f_fin"] < fecha_hoy and a["estatus"] != "Terminada")
    n_sin_reco = sum(1 for a in acts if a["reconocida"] != "SÍ")
    av_global = round(sum(a["avance"] or 0 for a in acts) / n_total, 1) if n_total else 0

    # Lista de actividades pendientes en sí (no solo el conteo), para que el
    # proveedor vea de un vistazo QUÉ le falta sin tener que entrar al portal.
    # Orden de urgencia: primero las vencidas (más vieja primero), luego las
    # que tienen fecha compromiso (más próxima primero), y al final las que
    # no tienen fecha todavía.
    pendientes = [a for a in acts if a["estatus"] != "Terminada" and (a["avance"] or 0) < 100]

    def orden_urgencia(a):
        vencida = bool(a["f_fin"] and a["f_fin"] < fecha_hoy)
        return (0 if vencida else (1 if a["f_fin"] else 2), a["f_fin"] or "9999-99-99")
    pendientes.sort(key=orden_urgencia)

    TOPE_LISTA = 15  # para no mandar un WhatsApp interminable si son muchas
    lineas_pend = []
    for a in pendientes[:TOPE_LISTA]:
        vencida = bool(a["f_fin"] and a["f_fin"] < fecha_hoy)
        marca = "🔴" if vencida else "•"
        area = a["area"] or "—"
        partida = a["partida"] or "(sin descripción)"
        av = a["avance"] or 0
        fecha_txt = f" — vence {a['f_fin']}" if a["f_fin"] else ""
        lineas_pend.append(f"{marca} {area}: {partida} ({av}%){fecha_txt}")
    restantes = len(pendientes) - len(lineas_pend)

    lineas = []
    lineas.append(f"📋 *Control de Obra HAP*")
    lineas.append(f"Seguimiento y Cumplimiento")
    lineas.append(f"")
    lineas.append(f"Buen día. Al corte del *{fecha_hoy}*, el estatus de tus actividades es el siguiente:")
    lineas.append(f"")
    lineas.append(f"📊 *Resumen de {nombre}:*")
    lineas.append(f"• Total de actividades: *{n_total}*")
    lineas.append(f"• Terminadas: ✅ *{n_terminadas}*")
    lineas.append(f"• Pendientes de actualización: ⚠️ *{n_pend}*")
    if n_atrasadas:
        lineas.append(f"• Con atraso: 🔴 *{n_atrasadas}*")
    if n_sin_reco:
        lineas.append(f"• Sin reconocer: ❌ *{n_sin_reco}*")
    lineas.append(f"• Avance global: *{av_global}%*")
    if lineas_pend:
        lineas.append(f"")
        lineas.append(f"📝 *Tus actividades pendientes:*")
        lineas.extend(lineas_pend)
        if restantes > 0:
            lineas.append(f"…y {restantes} más — ingresa al portal para ver todas.")
    lineas.append(f"")
    lineas.append(f"Favor de ingresar a la plataforma, actualizar los avances y confirmar fechas compromiso.")
    lineas.append(f"")
    lineas.append(f"🔗 https://control-obra-hap.onrender.com/portal")

    msg = "\n".join(lineas)
    return jsonify({"mensaje": msg, "nombre": nombre})


# ──── RUTA TEMPORAL: alta masiva mármol (borrar después de usar) ────
@app.route("/api/migrar_marmol")
@requiere_admin
def migrar_marmol():
    db = get_db()
    # verificar que no se haya corrido ya
    ya = db.execute("SELECT COUNT(*) c FROM actividades WHERE codigo='ACT-1101'").fetchone()["c"]
    if ya:
        return jsonify({"error": "Ya se ejecutó esta migración (ACT-1101 ya existe)"}), 400
    nuevas = [
        ("ACT-1101","Club médico y descanso","Vestidor doctoras","Mármol","Luis Carlos López","Barra para lavabo vestidores Dras. 3.00×0.50 mts — Mármol Crema Marfil"),
        ("ACT-1102","Club médico y descanso","Vestidores doctores","Mármol","Luis Carlos López","Barra para lavabo vestidores Dres. 1.80×0.50 mts — Mármol Crema Marfil"),
        ("ACT-1103","Club médico y descanso","Baño discapacitados mujeres","Mármol","Luis Carlos López","Barra para lavabo 75×50 cm — Granito Blanco Confeti"),
        ("ACT-1104","Club médico y descanso","Baño discapacitados hombres","Mármol","Luis Carlos López","Barra para lavabo 75×50 cm — Granito Blanco Confeti"),
        ("ACT-1105","Club médico y descanso","Vestidores doctores","Mármol","Luis Carlos López","Barra para lavabo baño hombres club médico 1.20×0.60 mts — Mármol Crema Marfil"),
        ("ACT-1106","Club médico y descanso","Vestidor doctoras","Mármol","Luis Carlos López","Barra para lavabo baño mujeres club médico en escuadra 2.92×0.60 mts — Mármol Crema Marfil"),
        ("ACT-1107","Recuperación pre","Cubículo anestesiólogo","Mármol","Luis Carlos López","Barra para cubículo anestesiólogo en escuadra 4.19×0.60 mts — Granito Gris Castello"),
        ("ACT-1108","Club médico y descanso","Descanso médico","Mármol","Luis Carlos López","Barra mueble cafetera tipo grapa 1.45×0.65 + faldones + splash back — Granito Gris Castello"),
        ("ACT-1109","Club médico y descanso","Descanso médico","Mármol","Luis Carlos López","Barra mueble club médico 2.80×0.65 + splash back 2.80×0.68 — Granito Gris Castello"),
        ("ACT-1110","Club médico y descanso","Vestidores doctores","Mármol","Luis Carlos López","8 cubiertas para cabeceras mueble ropa sucia vestidores Dres. y Dras. — Granito Gris Castello"),
        ("ACT-1111","Administrativas y técnicas","Jefatura","Mármol","Luis Carlos López","Cubierta para mueble jefatura 2.81×0.50 mts — Granito Gris Castello"),
        ("ACT-1112","Administrativas y técnicas","Jefatura","Mármol","Luis Carlos López","Cubierta para mueble jefatura 1.20×0.50 mts — Granito Gris Castello"),
    ]
    ahora = datetime.datetime.now().isoformat(timespec="seconds")
    creadas = []
    for cod, bloque, area, giro, prov, partida in nuevas:
        db.execute(
            """INSERT INTO actividades
            (codigo,bloque,area,giro,proveedor,partida,tipo,tipo_partida,aplica,avance,
             estatus,definido,mundo,actualizado)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (cod, bloque, area, giro, prov, partida,
             "Construcción", "Construcción", "SÍ", 0,
             "Pendiente", "NO", "obra", ahora))
        creadas.append(cod)
    db.commit()
    return jsonify({"ok": True, "creadas": creadas, "total": len(creadas)})


# ──── RUTA TEMPORAL: reestructura de quirófanos sept-2026 (borrar después de usar) ────
@app.route("/api/migrar_quirofanos_reestructura")
@requiere_admin
def migrar_quirofanos_reestructura():
    """La Sala 5 se dividió en dos (nueva Sala 5 y Sala 6, ambas con la
    especialidad de la Sala 4), la Sala 6 original pasa a ser Sala 7 con la
    especialidad que tenía la Sala 5 original, la Sala 7 original pasa a
    Sala 8 con su misma especialidad, y la Sala 8 original pasa a Sala 9 con
    su misma especialidad. Las salas 1-4 no se tocan.

    Todo se hace RENOMBRANDO el campo 'area' de las actividades que ya
    existen — nunca se borra ni se vuelve a crear una fila existente — así
    que avance, estatus, proveedor, fechas y notas quedan exactamente
    igual que estaban. Solo se insertan actividades nuevas para la Sala 6,
    que es la única realmente nueva, copiando el catálogo (giro/partida/
    proveedor) de la Sala 4, todas en 0% y Pendiente.

    El orden de los renombres importa: se hace de la sala más alta a la más
    baja (8→9, luego 7→8, luego 6→7, luego 5→5) para que una sala recién
    renombrada no se vuelva a mezclar con la siguiente antes de su turno.
    Idempotente: si ya se corrió, no hace nada de nuevo."""
    db = get_db()
    ya = db.execute(
        "SELECT COUNT(*) c FROM actividades WHERE bloque='Quirófanos' AND area='Q9 (Neuro/Trauma)'"
    ).fetchone()["c"]
    if ya:
        return jsonify({"error": "Ya se ejecutó esta migración (ya existe la Sala 9)"}), 400

    renombres = [
        ("Q8 (Neuro/Trauma)", "Q9 (Neuro/Trauma)"),
        ("Q7 (Neuro/Trauma/Cardiología)", "Q8 (Neuro/Trauma/Cardiología)"),
        ("Q6 (Urología/Laparoscopia/Gastro/Tórax/Ortopedia)", "Q7 (Oncología/Gastro/ORL/Procto/Hernia)"),
        ("Q5 (Oncología/Gastro/ORL/Procto/Hernia)", "Q5 (Gastro/ORL/Procto/Hernia)"),
    ]
    resumen = {}
    for area_vieja, area_nueva in renombres:
        cur = db.execute(
            "UPDATE actividades SET area=? WHERE bloque='Quirófanos' AND area=?",
            (area_nueva, area_vieja),
        )
        resumen[f"{area_vieja} -> {area_nueva}"] = cur.rowcount
    db.commit()

    # Sala 6 nueva: copia del catálogo de la Sala 4 (misma especialidad), desde cero
    plantilla = db.execute(
        "SELECT giro,proveedor,partida,tipo,tipo_partida FROM actividades "
        "WHERE bloque='Quirófanos' AND area='Q4 (Gastro/ORL/Procto/Hernia)' "
        "AND (eliminada IS NULL OR eliminada=0)"
    ).fetchall()
    ahora = datetime.datetime.now().isoformat(timespec="seconds")
    num = int(siguiente_codigo(db).split("-")[1])
    nuevos_codigos = []
    for p in plantilla:
        cod = f"ACT-{num:04d}"
        db.execute(
            """INSERT INTO actividades
            (codigo,bloque,area,giro,proveedor,partida,tipo,tipo_partida,aplica,avance,
             estatus,definido,mundo,actualizado)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (cod, "Quirófanos", "Q6 (Gastro/ORL/Procto/Hernia)", p["giro"], p["proveedor"],
             p["partida"], p["tipo"], p["tipo_partida"], "SÍ", 0, "Pendiente", "NO", "obra", ahora),
        )
        nuevos_codigos.append(cod)
        num += 1
    db.commit()
    resumen["Sala 6 (nueva, copiada de Sala 4)"] = len(nuevos_codigos)

    return jsonify({"ok": True, "resumen": resumen, "codigos_nuevos_sala6": nuevos_codigos})


# =====================================================================
#  PLANO DE OBRA: el admin ubica cada actividad en el plano (pin) y el
#  proveedor ve en el plano SOLO sus actividades, con su avance, y desde
#  ahi abre la ventana de reportar avance.
# =====================================================================
PLANO_ARCHIVO = os.path.join(DATA_DIR, "plano_obra.img")
PLANO_ARCHIVO_TORRE_A = os.path.join(DATA_DIR, "plano_torre_a.img")


def ruta_plano():
    """Archivo del plano del proyecto de la sesion."""
    return PLANO_ARCHIVO_TORRE_A if proyecto_actual() == "torre_a" else PLANO_ARCHIVO
PLANO_MAX_BYTES = 30 * 1024 * 1024
PLANO_ANCHO_PX = 4600   # ancho al convertir el PDF a imagen


def _config_get(db, clave, defecto=None):
    r = db.execute("SELECT valor FROM config WHERE clave=?", (clave,)).fetchone()
    return r["valor"] if r and r["valor"] is not None else defecto


def _config_set(db, clave, valor):
    db.execute("INSERT INTO config (clave,valor) VALUES (?,?) "
               "ON CONFLICT(clave) DO UPDATE SET valor=excluded.valor", (clave, str(valor)))


@app.route("/plano")
def plano_page():
    return render_template("plano.html")


@app.route("/api/plano/info")
@requiere_login
def api_plano_info():
    db = get_db()
    existe = os.path.exists(ruta_plano())
    escala = None
    if proyecto_actual() == "torre_a":
        try:
            escala = json.loads(_config_get(db, "ta_escala", "") or "null")
        except ValueError:
            escala = None
    return jsonify({
        "proyecto": proyecto_actual(),
        "escala": escala,
        "existe": existe,
        "version": _config_get(db, "plano_version", "0"),
        "nombre": _config_get(db, "plano_nombre", ""),
        "rol": session.get("rol"),
        "puede_editar": session.get("rol") == "admin",
        "puede_trazar": es_gestor(session.get("rol")),
        "ver_otros": _config_get(db, "plano_ver_otros", "1") == "1",
    })


@app.route("/api/plano/config", methods=["POST"])
@requiere_admin
def api_plano_config():
    """Opciones del plano que decide el admin (hoy: si los proveedores ven a otros gremios)."""
    data = request.get_json(silent=True) or {}
    db = get_db()
    if "ver_otros" in data:
        _config_set(db, "plano_ver_otros", "1" if data.get("ver_otros") else "0")
    db.commit()
    return jsonify({"ok": True, "ver_otros": _config_get(db, "plano_ver_otros", "1") == "1"})


@app.route("/api/plano/conteo")
@requiere_login
def api_plano_conteo():
    """Cuántas actividades del mundo (obra/interno) ya tienen pin y cuántas faltan: para la ventanita de filtro."""
    db = get_db()
    mundo = request.args.get("mundo", "obra")
    cond = ["(a.eliminada IS NULL OR a.eliminada=0)", "(a.estado_val IS NULL OR a.estado_val<>'rechazada')"]
    args = []
    if mundo != "todos":
        cond.append("(a.mundo = ? OR (a.mundo IS NULL AND ? = 'obra'))"); args += [mundo, mundo]
    r = db.execute(
        "SELECT COUNT(*) AS total, SUM(CASE WHEN EXISTS (SELECT 1 FROM plano_pines p WHERE p.actividad_id=a.id) THEN 1 ELSE 0 END) AS con "
        f"FROM actividades a WHERE {' AND '.join(cond)}", args).fetchone()
    total, con = r["total"] or 0, r["con"] or 0
    return jsonify({"total": total, "con": con, "sin": total - con})


@app.route("/api/plano/imagen")
@requiere_login
def api_plano_imagen():
    archivo = ruta_plano()
    if not os.path.exists(archivo):
        return jsonify({"error": "sin_plano"}), 404
    mime = _config_get(get_db(), "plano_mime", "image/jpeg")
    # Torre A sin cache del navegador: asi nunca se mezcla con el plano de Torre B.
    edad = 0 if proyecto_actual() == "torre_a" else 86400
    resp = send_file(archivo, mimetype=mime, max_age=edad, conditional=True)
    return resp


@app.route("/api/plano/subir", methods=["POST"])
@requiere_admin
def api_plano_subir():
    """Sube el plano (PDF de una hoja o imagen JPG/PNG). El PDF se convierte a imagen."""
    f = request.files.get("archivo")
    if not f or not f.filename:
        return jsonify({"error": "Elige el archivo del plano (PDF, JPG o PNG)."}), 400
    datos = f.read()
    if len(datos) > PLANO_MAX_BYTES:
        return jsonify({"error": "El archivo pesa más de 30 MB."}), 400
    nombre = f.filename
    if datos[:5] == b"%PDF-":
        try:
            import pymupdf
        except ImportError:
            try:
                import fitz as pymupdf
            except ImportError:
                return jsonify({"error": "El servidor aún no puede convertir PDF. "
                                         "Exporta el plano como imagen JPG o PNG y súbelo así."}), 500
        try:
            doc = pymupdf.open(stream=datos, filetype="pdf")
            if doc.page_count < 1:
                return jsonify({"error": "El PDF no tiene páginas."}), 400
            pag = doc[0]
            zoom = PLANO_ANCHO_PX / float(pag.rect.width)
            pix = pag.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
            salida = pix.tobytes("jpg", jpg_quality=88)
            mime = "image/jpeg"
        except Exception as e:
            return jsonify({"error": f"No se pudo leer el PDF: {e}"}), 400
    elif datos[:3] == b"\xff\xd8\xff":
        salida, mime = datos, "image/jpeg"
    elif datos[:8] == b"\x89PNG\r\n\x1a\n":
        salida, mime = datos, "image/png"
    else:
        return jsonify({"error": "Formato no válido. Sube un PDF, JPG o PNG."}), 400
    destino = ruta_plano()
    tmp = destino + ".tmp"
    with open(tmp, "wb") as fh:
        fh.write(salida)
    os.replace(tmp, destino)
    db = get_db()
    _config_set(db, "plano_mime", mime)
    _config_set(db, "plano_nombre", nombre)
    _config_set(db, "plano_version", datetime.datetime.now().strftime("%Y%m%d%H%M%S"))
    if proyecto_actual() == "torre_a":
        _config_set(db, "ta_escala", "")   # plano nuevo: hay que volver a calibrar la escala
    db.commit()
    return jsonify({"ok": True, "version": _config_get(db, "plano_version"), "kb": len(salida) // 1024})


# Paleta de capas por responsable (azul, naranja, aqua, violeta, magenta; el resto va a "Otros").
PLANO_PALETA = ["#2a78d6", "#eb6834", "#1baf7a", "#4a3aa7", "#e87ba4"]
PLANO_COLOR_OTROS = "#8b949e"
_SQL_RESP = "COALESCE(NULLIF(TRIM(a.proveedor),''), NULLIF(TRIM(a.departamento),''), 'Sin responsable')"


def _filtro_mundo_sql(mundo, alias="a"):
    """Condición SQL para un mundo ('obra' | 'interno'); 'todos' o vacío = sin filtro."""
    if mundo in ("obra", "interno"):
        return f" AND ({alias}.mundo = '{mundo}' OR ({alias}.mundo IS NULL AND '{mundo}' = 'obra'))"
    return ""


def _capas_plano(db, mundo=None):
    """Asigna color a cada responsable según cuántos pines tiene en TODO el plano (los 5 con más
    pines tienen color propio; el resto se agrupa en 'Otros'). Es la misma asignación para el
    admin y para los proveedores, así un gremio se ve del mismo color en todas partes."""
    filas = db.execute(
        f"SELECT {_SQL_RESP} AS resp, COUNT(*) AS n FROM plano_pines p "
        "JOIN actividades a ON a.id=p.actividad_id "
        "WHERE (a.eliminada IS NULL OR a.eliminada=0)" + _filtro_mundo_sql(mundo) + " GROUP BY resp "
        "ORDER BY n DESC, resp COLLATE NOCASE").fetchall()
    mapa = {}
    for i, f in enumerate(filas):
        if i < len(PLANO_PALETA):
            mapa[f["resp"]] = (f["resp"], PLANO_PALETA[i])
        else:
            mapa[f["resp"]] = ("Otros", PLANO_COLOR_OTROS)
    return mapa


def _fila_pin(r, capas=None):
    resp = (r["proveedor"] or "").strip() or (r["departamento"] or "").strip() or "Sin responsable"
    capa, color = (capas or {}).get(resp, (resp, PLANO_COLOR_OTROS))
    return {
        "resp": resp, "capa": capa, "color": color,
        "id": r["pin_id"], "actividad_id": r["actividad_id"], "x": r["x"], "y": r["y"],
        "nota": r["nota"] or "", "codigo": r["codigo"], "partida": r["partida"],
        "bloque": r["bloque"], "area": r["area"], "giro": r["giro"],
        "proveedor": r["proveedor"], "departamento": r["departamento"],
        "avance": r["avance"] or 0, "avance_decl": r["avance_decl"],
        "estatus": r["estatus"], "reconocida": r["reconocida"], "estado_val": r["estado_val"],
    }


@app.route("/api/plano/pines")
@requiere_login
def api_plano_pines():
    """Pines visibles: gestores ven todos; proveedor/departamento solo los de sus actividades."""
    db = get_db()
    _, rol, proveedor = usuario_actual()
    cond = ["(a.eliminada IS NULL OR a.eliminada=0)"]
    args = []
    mundo_f = None
    if not es_gestor(rol):
        col, mundo = col_duenio()
        cond.append(f"a.{col}=?"); args.append(proveedor)
        cond.append("(a.reconocida IS NULL OR a.reconocida<>'RECHAZADA')")
    else:
        mundo_f = request.args.get("mundo")          # 'obra' | 'interno' | 'todos' (sin parámetro = todos)
        if mundo_f in ("obra", "interno"):
            cond.append("(a.mundo = ? OR (a.mundo IS NULL AND ? = 'obra'))"); args += [mundo_f, mundo_f]
    filas = db.execute(
        "SELECT p.id AS pin_id, p.actividad_id, p.x, p.y, p.nota, a.codigo, a.partida, a.bloque, "
        "a.area, a.giro, a.proveedor, a.departamento, a.avance, a.avance_decl, a.estatus, "
        "a.reconocida, a.estado_val "
        "FROM plano_pines p JOIN actividades a ON a.id=p.actividad_id "
        f"WHERE {' AND '.join(cond)} ORDER BY p.id", args).fetchall()
    capas = _capas_plano(db, mundo_f)
    return jsonify([_fila_pin(r, capas) for r in filas])


@app.route("/api/plano/otros")
@requiere_login
def api_plano_otros():
    """Solo consulta para proveedores/departamentos: dónde trabajan los OTROS gremios de su mismo
    mundo y cuánto llevan. No incluye códigos ni notas, y el admin puede apagarlo."""
    db = get_db()
    _, rol, proveedor = usuario_actual()
    if es_gestor(rol):
        return jsonify({"habilitado": True, "pines": []})
    if _config_get(db, "plano_ver_otros", "1") != "1":
        return jsonify({"habilitado": False, "pines": []})
    col, _m = col_duenio()
    filas = db.execute(
        "SELECT p.id AS pin_id, p.actividad_id, p.x, p.y, a.partida, a.bloque, a.area, a.giro, "
        "a.proveedor, a.departamento, a.avance "
        "FROM plano_pines p JOIN actividades a ON a.id=p.actividad_id "
        f"WHERE (a.eliminada IS NULL OR a.eliminada=0) AND a.{col} IS NOT NULL AND TRIM(a.{col})<>'' "
        f"AND a.{col}<>? AND (a.reconocida IS NULL OR a.reconocida<>'RECHAZADA') ORDER BY p.id",
        (proveedor,)).fetchall()
    capas = _capas_plano(db)
    sal = []
    for r in filas:
        resp = (r[col] or "").strip()
        capa, color = capas.get(resp, (resp, PLANO_COLOR_OTROS))
        sal.append({"id": r["pin_id"], "x": r["x"], "y": r["y"], "partida": r["partida"],
                    "bloque": r["bloque"], "area": r["area"], "avance": r["avance"] or 0,
                    "resp": resp, "capa": capa, "color": color})
    return jsonify({"habilitado": True, "pines": sal})


# ----------------------------------------------------------------------------
# Panel del plano del proveedor: prioridades y "quién trabaja en mis áreas"
# Fechas clave y empresas con prueba de funcionamiento: se cambian aquí.
# ----------------------------------------------------------------------------
PRIORIDAD_PLANO = {
    "fecha_limite_pruebas": "2026-10-24",   # quienes llevan prueba de funcionamiento terminan antes
    "hito_aires": "2026-10-25",             # prueba de aires (ACSA)
    "bloque_aires": "Quirófanos",           # bloque que debe estar cerrado para la prueba de aires
    "con_pruebas": ["ACSA", "CEBSA", "Óscar Crespo", "Daniel Contreras", "Ana Rada"],
    "cierre_aires": ["Aluoferta", "Óscar Crespo"],
}


def _prioridad_de(responsable):
    """Qué avisos de prioridad le tocan a un proveedor (solo los suyos)."""
    resp = (responsable or "").strip()
    return {
        "entrega": FECHA_ENTREGA,
        "fecha_limite_pruebas": PRIORIDAD_PLANO["fecha_limite_pruebas"],
        "hito_aires": PRIORIDAD_PLANO["hito_aires"],
        "bloque_aires": PRIORIDAD_PLANO["bloque_aires"],
        "con_pruebas": resp in PRIORIDAD_PLANO["con_pruebas"],
        "cierre_aires": resp in PRIORIDAD_PLANO["cierre_aires"],
    }


@app.route("/api/portal/mis_areas")
@requiere_login
def api_portal_mis_areas():
    """Para el panel del plano del proveedor: en cada área (bloque + área) donde tiene
    actividades, qué otras empresas trabajan ahí y su avance PROMEDIO. Solo totales por
    empresa: sin códigos, partidas ni notas de los demás. Respeta el permiso del admin
    para ver a otros gremios (plano_ver_otros)."""
    try:
        db = get_db()
        _, rol, proveedor = usuario_actual()
        col, mundo = col_duenio()
        prioridad = _prioridad_de(proveedor)
        if es_gestor(rol) or not proveedor:
            return jsonify({"habilitado": False, "areas": [], "prioridad": prioridad})
        filtro = ("(a.eliminada IS NULL OR a.eliminada=0) "
                  "AND COALESCE(a.aplica,'SÍ')<>'NO' "
                  "AND (a.reconocida IS NULL OR a.reconocida<>'RECHAZADA') "
                  "AND (a.estado_val IS NULL OR a.estado_val<>'rechazada') "
                  "AND COALESCE(a.mundo,'obra')=?")   # solo su mundo: lo post-obra (interno) no se mezcla
        mias = db.execute(
            f"SELECT a.bloque, a.area, COUNT(*) AS n, AVG(COALESCE(a.avance,0)) AS av, "
            f"SUM(CASE WHEN COALESCE(a.avance,0)<100 THEN 1 ELSE 0 END) AS pend "
            f"FROM actividades a WHERE {filtro} AND a.{col}=? "
            f"GROUP BY a.bloque, a.area", (mundo, proveedor)).fetchall()
        if _config_get(db, "plano_ver_otros", "1") != "1":
            areas = [{"bloque": m["bloque"] or "", "area": m["area"] or "", "mias_n": m["n"],
                      "mias_av": round(m["av"] or 0), "mias_pend": m["pend"], "otros": []} for m in mias]
            return jsonify({"habilitado": False, "areas": areas, "prioridad": prioridad})
        otros = db.execute(
            f"SELECT a.bloque, a.area, TRIM(a.{col}) AS resp, COUNT(*) AS n, AVG(COALESCE(a.avance,0)) AS av "
            f"FROM actividades a WHERE {filtro} AND a.{col} IS NOT NULL AND TRIM(a.{col})<>'' AND a.{col}<>? "
            f"AND LOWER(TRIM(a.{col})) NOT IN ('—','-','por definir') "
            f"GROUP BY a.bloque, a.area, TRIM(a.{col})", (mundo, proveedor)).fetchall()
        por_area = {}
        for o in otros:
            por_area.setdefault((o["bloque"] or "", o["area"] or ""), []).append(
                {"resp": o["resp"], "n": o["n"], "av": round(o["av"] or 0)})
        areas = []
        for m in mias:
            clave = (m["bloque"] or "", m["area"] or "")
            lista = sorted(por_area.get(clave, []), key=lambda x: (x["av"], -x["n"]))
            areas.append({"bloque": clave[0], "area": clave[1], "mias_n": m["n"],
                          "mias_av": round(m["av"] or 0), "mias_pend": m["pend"], "otros": lista})
        # primero las áreas donde más le falta, para que vea a sus vecinos de trabajo inmediato
        areas.sort(key=lambda x: (-x["mias_pend"], x["mias_av"]))
        return jsonify({"habilitado": True, "areas": areas, "prioridad": prioridad})
    except sqlite3.Error as e:
        print("  Error en /api/portal/mis_areas:", e)
        return jsonify({"error": "No se pudo leer el avance de tus áreas."}), 500


@app.route("/api/plano/actividades")
@requiere_gestor
def api_plano_actividades():
    """Buscador de actividades para ubicarlas en el plano. Busca igual que la pantalla
    principal (sin importar acentos ni mayúsculas, en código, partida, área, bloque y
    responsable) y trae los mismos filtros: bloque, área, especialidad y proveedor."""
    db = get_db()
    q = (request.args.get("q") or "").strip()
    bloque = (request.args.get("bloque") or "").strip()
    area = (request.args.get("area") or "").strip()
    giro = (request.args.get("giro") or "").strip()
    prov = (request.args.get("proveedor") or "").strip()
    solo_sin = request.args.get("solo_sin_pin") == "1"
    mundo_b = request.args.get("mundo")
    base = ["(a.eliminada IS NULL OR a.eliminada=0)", "(a.aplica IS NULL OR a.aplica<>'NO')"]
    if mundo_b in ("obra", "interno"):
        base.append(f"(a.mundo = '{mundo_b}' OR (a.mundo IS NULL AND '{mundo_b}' = 'obra'))")
    cond = list(base)
    args = []
    for palabra in q.split():
        b = "%" + sin_acentos(palabra) + "%"
        cond.append("(sinac(a.codigo) LIKE ? OR sinac(a.partida) LIKE ? OR sinac(a.area) LIKE ? OR sinac(a.bloque) LIKE ? "
                    "OR sinac(a.giro) LIKE ? OR sinac(a.proveedor) LIKE ? OR sinac(a.departamento) LIKE ?)")
        args += [b] * 7
    if bloque:
        cond.append("a.bloque=?"); args.append(bloque)
    if area:
        cond.append("a.area=?"); args.append(area)
    if giro:
        giros_sel = [g for g in giro.split("|") if g]      # varias especialidades a la vez: "Pintura|Tablaroca"
        cond.append("a.giro IN (" + ",".join("?" * len(giros_sel)) + ")"); args += giros_sel
    if prov:
        cond.append("(a.proveedor=? OR a.departamento=?)"); args += [prov, prov]
    if solo_sin:
        cond.append("NOT EXISTS (SELECT 1 FROM plano_pines p2 WHERE p2.actividad_id=a.id)")
    where = " AND ".join(cond)
    total = db.execute(f"SELECT COUNT(*) FROM actividades a WHERE {where}", args).fetchone()[0]
    filas = db.execute(
        "SELECT a.id, a.codigo, a.partida, a.bloque, a.area, a.giro, a.proveedor, a.departamento, a.mundo, "
        "a.avance, (SELECT COUNT(*) FROM plano_pines p WHERE p.actividad_id=a.id) AS n_pines "
        f"FROM actividades a WHERE {where} "
        "ORDER BY a.bloque, a.area, a.id LIMIT 300", args).fetchall()
    bw = " AND ".join(base)
    bloques = [r[0] for r in db.execute(f"SELECT DISTINCT a.bloque FROM actividades a WHERE {bw} AND a.bloque IS NOT NULL AND a.bloque<>'' ORDER BY 1")]
    areas = [{"bloque": r[0], "area": r[1]} for r in db.execute(
        f"SELECT DISTINCT a.bloque, a.area FROM actividades a WHERE {bw} AND a.area IS NOT NULL AND a.area<>'' ORDER BY 1,2")]
    giros = [r[0] for r in db.execute(f"SELECT DISTINCT a.giro FROM actividades a WHERE {bw} AND a.giro IS NOT NULL AND a.giro<>'' ORDER BY 1")]
    provs = [r[0] for r in db.execute(
        f"SELECT DISTINCT COALESCE(NULLIF(a.proveedor,''), a.departamento) FROM actividades a WHERE {bw} "
        "AND COALESCE(NULLIF(a.proveedor,''), a.departamento) IS NOT NULL AND COALESCE(NULLIF(a.proveedor,''), a.departamento)<>'' ORDER BY 1")]
    if proyecto_actual() == "torre_a":
        # Torre A: los desplegables salen del catalogo fijo (20 bloques / 85 areas) y de sus proveedores numerados
        bloques, _todas, mapa_ta = catalogo_ta_para_pantalla(db)
        areas = [{"bloque": b, "area": a} for b, lst in mapa_ta.items() for a in lst]
        giros = sorted(set(giros) | {r[0] for r in db.execute("SELECT nombre FROM ta_especialidades")})
        if mundo_b != "interno":
            provs = sorted(set(provs) | {r[0] for r in db.execute("SELECT nombre FROM ta_proveedores WHERE interno=0")})
    return jsonify({"actividades": [dict(r) for r in filas], "total": total, "limite": 300,
                    "bloques": bloques, "areas": areas, "giros": giros, "proveedores": provs})


def _leer_xy(data):
    try:
        x = float(data.get("x")); y = float(data.get("y"))
    except (TypeError, ValueError):
        return None, None
    if not (0 <= x <= 1 and 0 <= y <= 1):
        return None, None
    return round(x, 5), round(y, 5)


@app.route("/api/plano/pines", methods=["POST"])
@requiere_admin
def api_plano_pin_crear():
    db = get_db()
    data = request.get_json() or {}
    x, y = _leer_xy(data)
    if x is None:
        return jsonify({"error": "Posición inválida en el plano."}), 400
    a = db.execute("SELECT id FROM actividades WHERE id=?", (data.get("actividad_id"),)).fetchone()
    if not a:
        return jsonify({"error": "La actividad no existe."}), 404
    cur = db.execute(
        "INSERT INTO plano_pines (actividad_id,x,y,nota,creado,creado_por) VALUES (?,?,?,?,?,?)",
        (a["id"], x, y, (data.get("nota") or "").strip()[:300],
         datetime.datetime.now().isoformat(timespec="seconds"), session.get("usuario")))
    db.commit()
    return jsonify({"ok": True, "id": cur.lastrowid})


@app.route("/api/plano/pines/<int:pid>", methods=["PUT"])
@requiere_admin
def api_plano_pin_editar(pid):
    db = get_db()
    data = request.get_json() or {}
    if not db.execute("SELECT 1 FROM plano_pines WHERE id=?", (pid,)).fetchone():
        return jsonify({"error": "El pin no existe."}), 404
    if "x" in data or "y" in data:
        x, y = _leer_xy(data)
        if x is None:
            return jsonify({"error": "Posición inválida en el plano."}), 400
        db.execute("UPDATE plano_pines SET x=?, y=? WHERE id=?", (x, y, pid))
    if "nota" in data:
        db.execute("UPDATE plano_pines SET nota=? WHERE id=?", ((data.get("nota") or "").strip()[:300], pid))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/plano/pines/<int:pid>", methods=["DELETE"])
@requiere_admin
def api_plano_pin_borrar(pid):
    db = get_db()
    db.execute("DELETE FROM plano_pines WHERE id=?", (pid,))
    db.commit()
    return jsonify({"ok": True})


# ---------------------------------------------------------------------
#  TRAZOS: líneas (muros, pintura, tablaroca) y zonas (plafones, pisos)
#  dibujadas sobre el plano, con las actividades de varios gremios
#  asignadas. El color/avance del trazo sale del promedio de sus actividades.
#  Dibujan el admin y los supervisores; el proveedor solo ve los trazos donde
#  tiene al menos una actividad.
# ---------------------------------------------------------------------
TRAZO_TIPOS = ("linea", "zona")
TRAZO_MAX_PUNTOS = 400

# Color de cada trazo = color de su ESPECIALIDAD principal (nunca el del avance).
# Hay 8 colores validados; cada especialidad recibe el suyo la primera vez que un trazo la usa y
# ya no cambia. Las especialidades que llegan después comparten el gris de "Otras".
GIRO_PALETA = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7", "#e34948"]
GIRO_COLOR_OTRAS = "#8b949e"


def _mapa_colores_giro(db):
    try:
        m = json.loads(_config_get(db, "plano_giro_colores", "{}") or "{}")
    except ValueError:
        m = {}
    return m if isinstance(m, dict) else {}


def _asegurar_color_giro(db, giro):
    """Le da color a una especialidad la primera vez que se usa como principal de un trazo."""
    giro = (giro or "").strip()
    if not giro:
        return
    m = _mapa_colores_giro(db)
    if giro in m:
        return
    usados = {c for c in m.values() if c in GIRO_PALETA}
    libre = next((c for c in GIRO_PALETA if c not in usados), None)
    m[giro] = libre or GIRO_COLOR_OTRAS
    _config_set(db, "plano_giro_colores", json.dumps(m, ensure_ascii=False))


def _leer_puntos(data, tipo):
    """Valida la lista de puntos [[x,y],...]. Devuelve (lista, error)."""
    crudo = data.get("puntos")
    if not isinstance(crudo, list):
        return None, "Faltan los puntos del trazo."
    pts = []
    for p in crudo:
        try:
            x, y = float(p[0]), float(p[1])
        except (TypeError, ValueError, IndexError, KeyError):
            return None, "Hay un punto inválido en el trazo."
        if not (0 <= x <= 1 and 0 <= y <= 1):
            return None, "Hay un punto fuera del plano."
        pts.append([round(x, 5), round(y, 5)])
    minimo = 2 if tipo == "linea" else 3
    if len(pts) < minimo:
        return None, ("Una línea necesita al menos 2 puntos." if tipo == "linea"
                      else "Una zona necesita al menos 3 puntos.")
    if len(pts) > TRAZO_MAX_PUNTOS:
        return None, "El trazo tiene demasiados puntos."
    return pts, None


def _ids_actividades(data):
    ids = []
    for v in (data.get("actividades") or []):
        try:
            i = int(v)
        except (TypeError, ValueError):
            continue
        if i not in ids:
            ids.append(i)
    return ids[:300]


def _texto_corto(v, n=80):
    return (str(v).strip() if v is not None else "")[:n]


def _leer_ubicaciones(data):
    """[{bloque, area}, ...] del cuerpo de la petición (sin repetidos, máx. 20)."""
    sal, vistos = [], set()
    for u in (data.get("ubicaciones") or []):
        if not isinstance(u, dict):
            continue
        b, a = _texto_corto(u.get("bloque")), _texto_corto(u.get("area"))
        if not b or (b, a) in vistos:
            continue
        vistos.add((b, a)); sal.append({"bloque": b, "area": a})
    return sal[:20]


def _guardar_ubicaciones(db, trazo_id, ubs):
    db.execute("DELETE FROM plano_trazo_ub WHERE trazo_id=?", (trazo_id,))
    for u in ubs:
        db.execute("INSERT OR IGNORE INTO plano_trazo_ub (trazo_id,bloque,area) VALUES (?,?,?)", (trazo_id, u["bloque"], u["area"]))


def _mas_comun(valores):
    cuenta = {}
    for v in valores:
        v = (v or "").strip()
        if v:
            cuenta[v] = cuenta.get(v, 0) + 1
    return max(sorted(cuenta), key=lambda k: cuenta[k]) if cuenta else None


def _guardar_vinculos(db, trazo_id, ids):
    """Cambia las actividades que se asignaron a mano SIN tocar la forma del trazo. Si el trazo aún no
    tenía especialidad principal ni ubicación, las toma de lo que le asignaron (la más repetida);
    si ya las tenía, no cambian."""
    db.execute("DELETE FROM plano_trazo_act WHERE trazo_id=?", (trazo_id,))
    for i in ids:
        if db.execute("SELECT 1 FROM actividades WHERE id=?", (i,)).fetchone():
            db.execute("INSERT OR IGNORE INTO plano_trazo_act (trazo_id,actividad_id) VALUES (?,?)", (trazo_id, i))
    _completar_un_trazo(db, trazo_id)


def _completar_un_trazo(db, trazo_id):
    t = db.execute("SELECT giro FROM plano_trazos WHERE id=?", (trazo_id,)).fetchone()
    if not t:
        return
    filas = db.execute(
        "SELECT a.giro, a.bloque, a.area FROM plano_trazo_act ta JOIN actividades a ON a.id=ta.actividad_id "
        "WHERE ta.trazo_id=? AND (a.eliminada IS NULL OR a.eliminada=0)", (trazo_id,)).fetchall()
    if not filas:
        return
    if not (t["giro"] or "").strip():
        g = _mas_comun(f["giro"] for f in filas)
        if g:
            db.execute("UPDATE plano_trazos SET giro=? WHERE id=?", (g, trazo_id))
            _asegurar_color_giro(db, g)
    if not db.execute("SELECT 1 FROM plano_trazo_ub WHERE trazo_id=?", (trazo_id,)).fetchone():
        b = _mas_comun(f["bloque"] for f in filas)
        if b:
            a = _mas_comun(f["area"] for f in filas if (f["bloque"] or "").strip() == b) or ""
            db.execute("INSERT OR IGNORE INTO plano_trazo_ub (trazo_id,bloque,area) VALUES (?,?,?)", (trazo_id, b, a))


def _completar_trazos(db):
    """Trazos hechos antes de existir 'especialidad principal' y 'ubicación': se completan una vez
    con lo que ya tienen asignado (así toman su color y su área sin tocar su forma)."""
    pend = db.execute(
        "SELECT t.id FROM plano_trazos t WHERE ((t.giro IS NULL OR TRIM(t.giro)='') "
        "OR NOT EXISTS (SELECT 1 FROM plano_trazo_ub u WHERE u.trazo_id=t.id)) "
        "AND EXISTS (SELECT 1 FROM plano_trazo_act ta WHERE ta.trazo_id=t.id)").fetchall()
    for r in pend:
        _completar_un_trazo(db, r["id"])
    if pend:
        db.commit()


# ---- Tipos de elemento (Muro, Piso, Plafón…): catálogo que el admin da de alta y edita ----
TIPOS_ELEM_BASE = ["Muro", "Piso", "Plafón", "Tablaroca", "Pintura", "Azulejo"]


TIPOS_FIJOS_TA = ("Bloque", "Área")


def _tipos_elem(db):
    try:
        t = json.loads(_config_get(db, "plano_tipos_elem", "") or "null")
    except ValueError:
        t = None
    lista = t if isinstance(t, list) else list(TIPOS_ELEM_BASE)
    if proyecto_actual() == "torre_a":   # Torre A: "Bloque" y "Área" son tipos fijos (capas de contornos del catalogo)
        lista = [x for x in TIPOS_FIJOS_TA if x not in lista] + lista
    return lista


def _guardar_tipos_elem(db, lista):
    _config_set(db, "plano_tipos_elem", json.dumps(lista, ensure_ascii=False))


@app.route("/api/plano/tipos")
@requiere_login
def api_plano_tipos():
    db = get_db()
    uso = {r[0]: r[1] for r in db.execute("SELECT tipo_elem, COUNT(*) FROM plano_trazos WHERE tipo_elem IS NOT NULL AND tipo_elem<>'' GROUP BY tipo_elem")}
    return jsonify([{"nombre": n, "n": uso.get(n, 0)} for n in _tipos_elem(db)])


@app.route("/api/plano/tipos", methods=["POST"])
@requiere_gestor
def api_plano_tipo_alta():
    db = get_db()
    nombre = _texto_corto((request.get_json(silent=True) or {}).get("nombre"), 40)
    if not nombre:
        return jsonify({"error": "Escribe el nombre del tipo."}), 400
    lista = _tipos_elem(db)
    if any(x.lower() == nombre.lower() for x in lista):
        return jsonify({"error": "Ese tipo ya existe."}), 400
    lista.append(nombre)
    _guardar_tipos_elem(db, lista); db.commit()
    return jsonify({"ok": True, "nombre": nombre})


@app.route("/api/plano/tipos", methods=["PUT"])
@requiere_gestor
def api_plano_tipo_renombrar():
    db = get_db()
    d = request.get_json(silent=True) or {}
    viejo, nuevo = _texto_corto(d.get("viejo"), 40), _texto_corto(d.get("nuevo"), 40)
    lista = _tipos_elem(db)
    if proyecto_actual() == "torre_a" and viejo in TIPOS_FIJOS_TA:
        return jsonify({"error": "«%s» es un tipo fijo de Torre A; no se renombra." % viejo}), 400
    if viejo not in lista or not nuevo:
        return jsonify({"error": "No encontré ese tipo."}), 404
    if nuevo.lower() != viejo.lower() and any(x.lower() == nuevo.lower() for x in lista):
        return jsonify({"error": "Ya existe un tipo con ese nombre."}), 400
    lista[lista.index(viejo)] = nuevo
    _guardar_tipos_elem(db, lista)
    db.execute("UPDATE plano_trazos SET tipo_elem=? WHERE tipo_elem=?", (nuevo, viejo))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/plano/tipos", methods=["DELETE"])
@requiere_gestor
def api_plano_tipo_borrar():
    db = get_db()
    nombre = _texto_corto(request.args.get("nombre"), 40)
    lista = _tipos_elem(db)
    if proyecto_actual() == "torre_a" and nombre in TIPOS_FIJOS_TA:
        return jsonify({"error": "«%s» es un tipo fijo de Torre A; no se quita." % nombre}), 400
    if nombre not in lista:
        return jsonify({"error": "No encontré ese tipo."}), 404
    n = db.execute("SELECT COUNT(*) FROM plano_trazos WHERE tipo_elem=?", (nombre,)).fetchone()[0]
    if n and request.args.get("forzar") != "1":
        return jsonify({"error": f"Hay {n} trazos con este tipo.", "en_uso": n}), 409
    lista.remove(nombre)
    _guardar_tipos_elem(db, lista)
    db.execute("UPDATE plano_trazos SET tipo_elem='' WHERE tipo_elem=?", (nombre,))
    db.commit()
    return jsonify({"ok": True})


_SQL_ACT_TRAZO = ("a.id AS actividad_id, a.codigo, a.partida, a.giro, a.bloque, a.area, a.proveedor, a.departamento, "
                  "a.mundo, a.avance, a.avance_decl, a.reconocida, a.estado_val")


def _trazos_para(db, rol, proveedor, mundo=None):
    """Trazos con sus actividades (las asignadas a mano + las de sus áreas si el trazo es automático).
    Gestores: ven el mundo pedido ('obra' | 'interno' | 'todos'). Proveedor/departamento: solo su mundo
    y solo los trazos donde tiene al menos una actividad; de las demás ve gremio y avance (sin código)
    si el admin lo permite."""
    gestor = es_gestor(rol)
    col, mundo_ses = col_duenio()
    ver_otros = _config_get(db, "plano_ver_otros", "1") == "1"
    _completar_trazos(db)
    colores = _mapa_colores_giro(db)
    capas = _capas_plano(db, None if gestor and mundo not in ("obra", "interno") else (mundo if gestor else mundo_ses))
    if gestor:
        mundo_q = mundo if mundo in ("obra", "interno") else None
    else:
        mundo_q = mundo_ses
    sql_t = "SELECT id, tipo, nombre, puntos, giro, tipo_elem, COALESCE(mundo,'obra') AS mundo, COALESCE(auto,0) AS auto FROM plano_trazos"
    args_t = []
    if mundo_q:
        sql_t += " WHERE COALESCE(mundo,'obra')=?"; args_t.append(mundo_q)
        if proyecto_actual() == "torre_a":   # contornos de bloques y areas: se ven en Obra y en Interno
            sql_t += " OR tipo_elem IN ('Bloque','Área')"
    sql_t += " ORDER BY id"
    trazos = db.execute(sql_t, args_t).fetchall()
    ubs = {}
    for u in db.execute("SELECT trazo_id, bloque, area FROM plano_trazo_ub ORDER BY bloque, area"):
        ubs.setdefault(u["trazo_id"], []).append({"bloque": u["bloque"], "area": u["area"]})
    manual = {}
    for r in db.execute(
            f"SELECT ta.trazo_id, {_SQL_ACT_TRAZO} FROM plano_trazo_act ta JOIN actividades a ON a.id=ta.actividad_id "
            "WHERE (a.eliminada IS NULL OR a.eliminada=0) ORDER BY a.giro, a.id"):
        manual.setdefault(r["trazo_id"], []).append(r)

    sal = []
    for tr in trazos:
        try:
            pts = json.loads(tr["puntos"])
        except ValueError:
            pts = []
        lista_ub = ubs.get(tr["id"], [])
        filas = list(manual.get(tr["id"], []))
        ids = {f["actividad_id"] for f in filas}
        auto_ids = set()
        if tr["auto"] and lista_ub:
            cond = ["(a.eliminada IS NULL OR a.eliminada=0)", "(a.aplica IS NULL OR a.aplica<>'NO')",
                    "(a.estado_val IS NULL OR a.estado_val<>'rechazada')"]
            args = []
            piezas = []
            for u in lista_ub:
                if u["area"]:
                    piezas.append("(a.bloque=? AND a.area=?)"); args += [u["bloque"], u["area"]]
                else:
                    piezas.append("(a.bloque=?)"); args.append(u["bloque"])
            cond.append("(" + " OR ".join(piezas) + ")")
            cond.append("(a.mundo = ? OR (a.mundo IS NULL AND ? = 'obra'))"); args += [tr["mundo"], tr["mundo"]]
            for r in db.execute(f"SELECT {_SQL_ACT_TRAZO} FROM actividades a WHERE {' AND '.join(cond)} ORDER BY a.giro, a.id", args):
                if r["actividad_id"] not in ids:
                    filas.append(r); ids.add(r["actividad_id"]); auto_ids.add(r["actividad_id"])
        acts = []
        for r in filas:
            resp = (r["proveedor"] or "").strip() or (r["departamento"] or "").strip() or "Sin responsable"
            esmundo_int = (r["mundo"] or "obra") == "interno"
            duenio = ((r["departamento"] if esmundo_int else r["proveedor"]) or "").strip()
            propia = gestor or (duenio != "" and duenio == proveedor)
            if not gestor and r["reconocida"] == "RECHAZADA":
                continue
            capa, color = capas.get(resp, (resp, PLANO_COLOR_OTROS))
            giro_a = (r["giro"] or "").strip()
            act = {"id": r["actividad_id"], "partida": r["partida"], "giro": r["giro"], "bloque": r["bloque"],
                   "area": r["area"], "resp": resp, "color": color, "avance": r["avance"] or 0, "propia": bool(propia),
                   "auto": r["actividad_id"] in auto_ids,
                   "color_giro": colores.get(giro_a, GIRO_COLOR_OTRAS) if giro_a else GIRO_COLOR_OTRAS}
            if propia:
                act.update({"codigo": r["codigo"], "avance_decl": r["avance_decl"],
                            "reconocida": r["reconocida"], "estado_val": r["estado_val"]})
            acts.append(act)
        if not gestor:
            if not any(a["propia"] for a in acts):
                continue
            if not ver_otros:
                acts = [a for a in acts if a["propia"]]
        n = len(acts)
        g = {}
        for a in acts:
            k = (a["giro"] or "").strip() or "Sin especialidad"
            e = g.setdefault(k, {"giro": k, "n": 0, "color": colores.get(k, GIRO_COLOR_OTRAS) if k != "Sin especialidad" else GIRO_COLOR_OTRAS})
            e["n"] += 1
        giros = sorted(g.values(), key=lambda e: (-e["n"], e["giro"]))
        # avance separado por responsable (dos proveedores pueden trabajar en el mismo muro)
        rs = {}
        for a in acts:
            e = rs.setdefault(a["resp"], {"resp": a["resp"], "n": 0, "suma": 0, "listas": 0, "color": a["color"], "propia": False})
            e["n"] += 1; e["suma"] += a["avance"] or 0; e["listas"] += 1 if (a["avance"] or 0) >= 100 else 0
            e["propia"] = e["propia"] or a["propia"]
        resps = []
        for e in sorted(rs.values(), key=lambda e: (-e["n"], e["resp"])):
            resps.append({"resp": e["resp"], "n": e["n"], "avance": int(round(e["suma"] / e["n"])), "listas": e["listas"],
                          "color": e["color"], "propia": e["propia"]})
        mias = [a for a in acts if a["propia"]]
        giro_st = (tr["giro"] or "").strip()
        principal = giro_st or (giros[0]["giro"] if giros and giros[0]["giro"] != "Sin especialidad" else "")
        if principal and not giro_st:
            # el trazo aún no tenía especialidad principal: se fija ahora (así su color ya no cambia solo)
            db.execute("UPDATE plano_trazos SET giro=? WHERE id=?", (principal, tr["id"]))
            _asegurar_color_giro(db, principal)
            colores = _mapa_colores_giro(db)
            db.commit()
        sal.append({
            "id": tr["id"], "tipo": tr["tipo"], "nombre": tr["nombre"] or "", "puntos": pts, "mundo": tr["mundo"],
            "tipo_elem": tr["tipo_elem"] or "", "auto": bool(tr["auto"]), "ubicaciones": lista_ub,
            "giro": giro_st, "giro_efectivo": principal,
            "color": colores.get(principal, GIRO_COLOR_OTRAS) if principal else GIRO_COLOR_OTRAS,
            "actividades": acts, "n_act": n, "n_auto": len(auto_ids),
            "n_listas": sum(1 for a in acts if (a["avance"] or 0) >= 100),
            "avance": int(round(sum((a["avance"] or 0) for a in acts) / n)) if n else 0,
            "giros": giros, "resps": resps,
            "mi_n": len(mias) if not gestor else 0,
            "mi_avance": int(round(sum((a["avance"] or 0) for a in mias) / len(mias))) if (mias and not gestor) else 0,
        })
    return sal


@app.route("/api/plano/trazos")
@requiere_login
def api_plano_trazos():
    _, rol, proveedor = usuario_actual()
    return jsonify(_trazos_para(get_db(), rol, proveedor, request.args.get("mundo")))


def _tipo_fijo_ajeno(db, data):
    """Los tipos 'Bloque' y 'Area' son solo de Torre A: en otro proyecto no se aceptan (p. ej. pestana vieja abierta)."""
    t = (data.get("tipo_elem") or "").strip() if isinstance(data.get("tipo_elem"), str) else ""
    return bool(t in TIPOS_FIJOS_TA and proyecto_actual() != "torre_a" and t not in _tipos_elem(db))


def _validar_trazo_comun(data):
    mundo = data.get("mundo")
    return mundo if mundo in ("obra", "interno") else None


@app.route("/api/plano/trazos", methods=["POST"])
@requiere_gestor
def api_plano_trazo_crear():
    db = get_db()
    data = request.get_json(silent=True) or {}
    tipo = data.get("tipo")
    if _tipo_fijo_ajeno(db, data):
        return jsonify({"error": "Los contornos de Bloque/Área son de Torre A. Cambia de proyecto y vuelve a intentarlo."}), 409
    if tipo not in TRAZO_TIPOS:
        return jsonify({"error": "El tipo de trazo debe ser línea o zona."}), 400
    pts, err = _leer_puntos(data, tipo)
    if err:
        return jsonify({"error": err}), 400
    mundo = _validar_trazo_comun(data) or "obra"
    giro = _texto_corto(data.get("giro"))
    ubs = _leer_ubicaciones(data)
    nombre = (data.get("nombre") or "").strip()[:120]
    if not nombre and ubs:                              # el nombre sale de la ubicación ligada
        nombre = " + ".join((u["area"] or u["bloque"]) for u in ubs)[:120]
    cur = db.execute(
        "INSERT INTO plano_trazos (tipo,nombre,puntos,creado,creado_por,giro,tipo_elem,mundo,auto) VALUES (?,?,?,?,?,?,?,?,?)",
        (tipo, nombre, json.dumps(pts), datetime.datetime.now().isoformat(timespec="seconds"), session.get("usuario"),
         giro, _texto_corto(data.get("tipo_elem"), 40), mundo, 1 if data.get("auto") else 0))
    _guardar_ubicaciones(db, cur.lastrowid, ubs)
    if giro:
        _asegurar_color_giro(db, giro)
    _guardar_vinculos(db, cur.lastrowid, _ids_actividades(data))
    db.commit()
    return jsonify({"ok": True, "id": cur.lastrowid})


@app.route("/api/plano/trazos/<int:tid>", methods=["PUT"])
@requiere_gestor
def api_plano_trazo_editar(tid):
    db = get_db()
    data = request.get_json(silent=True) or {}
    t = db.execute("SELECT tipo FROM plano_trazos WHERE id=?", (tid,)).fetchone()
    if not t:
        return jsonify({"error": "El trazo no existe."}), 404
    if _tipo_fijo_ajeno(db, data):
        return jsonify({"error": "Los contornos de Bloque/Área son de Torre A. Cambia de proyecto y vuelve a intentarlo."}), 409
    if "puntos" in data:
        pts, err = _leer_puntos(data, t["tipo"])
        if err:
            return jsonify({"error": err}), 400
        if proyecto_actual() == "torre_a":
            # si cambia el numero de lados, las especialidades por lado ya no corresponden: se limpian
            ant = db.execute("SELECT puntos FROM plano_trazos WHERE id=?", (tid,)).fetchone()
            try:
                n_ant = len(json.loads(ant["puntos"])) if ant else 0
            except ValueError:
                n_ant = 0
            if n_ant != len(pts):
                db.execute("DELETE FROM ta_lados WHERE trazo_id=?", (tid,))
        db.execute("UPDATE plano_trazos SET puntos=? WHERE id=?", (json.dumps(pts), tid))
    if "nombre" in data:
        db.execute("UPDATE plano_trazos SET nombre=? WHERE id=?", ((data.get("nombre") or "").strip()[:120], tid))
    if "ubicaciones" in data:
        _guardar_ubicaciones(db, tid, _leer_ubicaciones(data))
    if "tipo_elem" in data:
        db.execute("UPDATE plano_trazos SET tipo_elem=? WHERE id=?", (_texto_corto(data.get("tipo_elem"), 40), tid))
    if "auto" in data:
        db.execute("UPDATE plano_trazos SET auto=? WHERE id=?", (1 if data.get("auto") else 0, tid))
    if "giro" in data:
        g = _texto_corto(data.get("giro"))
        db.execute("UPDATE plano_trazos SET giro=? WHERE id=?", (g, tid))
        _asegurar_color_giro(db, g)
    if "actividades" in data:
        _guardar_vinculos(db, tid, _ids_actividades(data))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/plano/trazos/<int:tid>", methods=["DELETE"])
@requiere_gestor
def api_plano_trazo_borrar(tid):
    db = get_db()
    db.execute("DELETE FROM plano_trazo_act WHERE trazo_id=?", (tid,))
    db.execute("DELETE FROM plano_trazo_ub WHERE trazo_id=?", (tid,))
    if proyecto_actual() == "torre_a":
        db.execute("DELETE FROM ta_lados WHERE trazo_id=?", (tid,))
        db.execute("DELETE FROM ta_alturas WHERE trazo_id=?", (tid,))
    db.execute("DELETE FROM plano_trazos WHERE id=?", (tid,))
    db.commit()
    return jsonify({"ok": True})


# ----------------------------------------------------------------------------
# Melissa (v1.29): asistente de consulta para el admin. Ver melissa.py y DECISIONS.md.
# ----------------------------------------------------------------------------
import melissa as _melissa

_melissa.registrar(app, {
    "requiere_login": requiere_login,
    "requiere_admin": requiere_admin,
    "get_db": get_db,
    "get_db_usuarios": get_db_usuarios,
    "proyecto_actual": proyecto_actual,
    "hoy_local": lambda: _sh_hoy(),
    "ahora_utc": lambda: datetime.datetime.now(datetime.timezone.utc),
})


if __name__ == "__main__":
    print("=" * 60)
    print("  PLATAFORMA DE CONTROL DE OBRA — HAP")
    print("  Base de datos:", DB_PATH)
    print("  Abre en tu navegador:  http://localhost:5000")
    print("=" * 60)
    puerto = int(os.environ.get("PORT", 5000))
    host = "0.0.0.0" if os.environ.get("EN_INTERNET") else "127.0.0.1"
    app.run(host=host, port=puerto, debug=False)
